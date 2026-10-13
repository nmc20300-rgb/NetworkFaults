from __future__ import annotations

from contextlib import contextmanager
from pathlib import Path
from string import Template

from PySide6.QtCore import QSize, QStringListModel, Qt, QTimer, Signal
from PySide6.QtGui import QColor, QFont, QGuiApplication, QKeySequence, QShortcut
from PySide6.QtWidgets import (
    QAbstractItemView,
    QComboBox,
    QCompleter,
    QDialog,
    QFileDialog,
    QFormLayout,
    QFrame,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QMenu,
    QMessageBox,
    QPushButton,
    QTreeWidget,
    QTreeWidgetItem,
    QVBoxLayout,
    QWidget,
    QWizard,
    QWizardPage,
)

from database import get_connection, write_audit_log
from ui.org_cache import OrgStructureCache

MODULE_NAME = "الهيكل التنظيمي"
UNNAMED = "(بدون اسم)"
UNIT_TYPES = ("فرع", "مبنى", "قسم", "مكتب")
ROOT_ONLY_TYPES = frozenset({"فرع", "مبنى"})
PARENT_CAPABLE_TYPES = frozenset({"فرع", "مبنى", "قسم"})
NO_PARENT_LABEL = "— بدون وحدة أعلى —"

TYPE_HINTS = {
    "فرع": "الفرع وحدة رئيسية، ولا يحتاج إلى اختيار عنصر أعلى.",
    "مبنى": "المبنى وحدة رئيسية، ولا يحتاج إلى اختيار عنصر أعلى.",
    "قسم": "يمكن ربط القسم بفرع أو مبنى أو قسم أعلى.",
    "مكتب": "يمكن ربط المكتب بفرع أو مبنى أو قسم.",
}

TYPE_COLORS = {
    "فرع": ("#eff6ff", "#1e40af"),
    "مبنى": ("#f0fdf4", "#166534"),
    "قسم": ("#fefce8", "#854d0e"),
    "مكتب": ("#fdf4ff", "#7e22ce"),
}
DEFAULT_TYPE_COLORS = ("#f8fafc", "#334155")

COL_NAME, COL_TYPE, COL_COUNT, COL_ACTIONS = range(4)
ROW_HEIGHT = 54
TYPE_COLUMN_WIDTH = 130
COUNT_COLUMN_WIDTH = 150
ACTIONS_COLUMN_WIDTH = 190

PALETTE = {
    "bg": "#f4f7fb",
    "surface": "#ffffff",
    "border": "#e2e8f0",
    "border_strong": "#cbd5e1",
    "text": "#0f2338",
    "body": "#334155",
    "muted": "#64748b",
    "faint": "#94a3b8",
    "subtle": "#f1f5f9",
    "hover": "#e2e8f0",
    "field": "#fbfdff",
    "accent": "#0891b2",
    "accent_dark": "#0e7490",
    "accent_soft": "#e0f2fe",
    "accent_text": "#0369a1",
    "row_alt": "#f8fafc",
    "row_hover": "#f0f9ff",
    "root_bg": "#fff7ed",
    "root_fg": "#9a3412",
    "success": "#166534",
    "success_bg": "#ecfdf5",
    "success_text": "#047857",
    "danger": "#991b1b",
    "danger_bg": "#fef2f2",
    "danger_border": "#fecaca",
    "danger_hover": "#fee2e2",
    "edit_bg": "#eff6ff",
    "edit_fg": "#1e40af",
    "edit_border": "#bfdbfe",
    "edit_hover": "#dbeafe",
    "step_active_border": "#7dd3fc",
}

BASE_STYLE = """
QLineEdit, QComboBox {
    background: $field; border: 1px solid $border_strong;
    border-radius: 8px; padding: 7px 10px; min-height: 22px; color: $text;
}
QLineEdit:focus, QComboBox:focus { border: 1.5px solid $accent; }
QComboBox:disabled { background: $subtle; color: $faint; }
QComboBox QAbstractItemView {
    background: $surface; border: 1px solid $border_strong;
    selection-background-color: $accent_soft; selection-color: $text; outline: none;
}
QToolTip {
    background: $text; color: white; border: none; padding: 6px 9px; border-radius: 6px;
}
QPushButton#primaryButton {
    background: $accent; color: white; border: none;
    border-radius: 9px; padding: 9px 18px; font-weight: 700;
}
QPushButton#primaryButton:hover { background: $accent_dark; }
QPushButton#primaryButton:disabled { background: $border_strong; }
QPushButton#toolButton, QPushButton#clearButton {
    background: $subtle; color: $body; border: 1px solid $border_strong;
    border-radius: 8px; padding: 6px 12px; font-weight: 700;
}
QPushButton#toolButton:hover, QPushButton#clearButton:hover { background: $hover; }
"""

PAGE_STYLE = BASE_STYLE + """
#orgPage { background: $bg; }
#pageTitle { color: $text; font-size: 23px; font-weight: 800; }
#pageSubtitle { color: $muted; font-size: 12px; }
#filterCard, #treeCard, #metricCard {
    background: $surface; border: 1px solid $border; border-radius: 12px;
}
#metricTitle { color: $muted; font-size: 11px; font-weight: 700; }
#metricValue { font-size: 24px; font-weight: 800; }
#fieldCaption { color: $muted; font-size: 11px; font-weight: 700; }
#sectionTitle { color: $text; font-size: 14px; font-weight: 800; }
#statsLabel { color: $muted; font-size: 11px; padding: 3px 6px; }
#emptyTitle { color: $text; font-size: 16px; font-weight: 800; }
#emptyHint { color: $muted; font-size: 12px; }
QWidget#rootActionCell { background: $root_bg; border-radius: 6px; }
QTreeWidget {
    background: $surface; border: none; gridline-color: $subtle;
    selection-background-color: $accent_soft; selection-color: $text;
    alternate-background-color: $row_alt; outline: none;
}
QTreeWidget::item { padding: 3px 7px; border-bottom: 1px solid $subtle; }
QTreeWidget::item:hover { background: $row_hover; }
QHeaderView::section {
    background: $row_alt; color: $text; padding: 10px 8px;
    border: none; border-bottom: 2px solid $accent; font-weight: 700;
}
QLineEdit#quickSearch { min-width: 230px; background: $surface; }
QPushButton#editBtn, QPushButton#deleteBtn {
    border-radius: 7px; padding: 3px 10px; font-size: 11px; font-weight: 700; min-height: 30px;
}
QPushButton#editBtn { background: $edit_bg; color: $edit_fg; border: 1px solid $edit_border; }
QPushButton#editBtn:hover { background: $edit_hover; }
QPushButton#deleteBtn { background: $danger_bg; color: $danger; border: 1px solid $danger_border; }
QPushButton#deleteBtn:hover { background: $danger_hover; }
QMenu {
    background: $surface; border: 1px solid $border_strong; border-radius: 8px; padding: 5px;
}
QMenu::item { padding: 7px 22px; border-radius: 6px; }
QMenu::item:selected { background: $accent_soft; color: $text; }
QMenu::separator { height: 1px; background: $border; margin: 4px 8px; }
"""

DIALOG_STYLE = BASE_STYLE + """
QWizard { background: $bg; }
QWizardPage { background: transparent; }
QFrame#wizardCard {
    background: $surface; border: 1px solid $border; border-radius: 12px;
}
QFrame#stepBar { background: transparent; }
QLabel#stepItem {
    background: $subtle; color: $muted; border-radius: 9px;
    padding: 9px 12px; font-weight: 700;
}
QLabel#stepItem[stepState="active"] {
    background: $accent_soft; color: $accent_text; border: 1px solid $step_active_border;
}
QLabel#stepItem[stepState="done"] { background: $success_bg; color: $success_text; }
QLabel#fieldHint { color: $muted; font-size: 12px; }
QLabel#fieldError { color: $danger; font-size: 12px; font-weight: 700; }
QLabel#reviewValue { color: $text; font-weight: 700; }
QWizard QPushButton {
    background: $subtle; color: $text; border: 1px solid $border_strong;
    border-radius: 8px; padding: 8px 18px; font-weight: 700; min-width: 88px;
}
QWizard QPushButton:hover { background: $hover; }
QWizard QPushButton#primaryButton {
    background: $accent; color: white; border: none;
}
QWizard QPushButton#primaryButton:hover { background: $accent_dark; }
"""


def build_style(template):
    return Template(template).substitute(PALETTE)


def org_unit_audit_data(unit, all_units=None):
    if not unit:
        return {}

    parent_id = unit.get("parent_id")
    parent_name = unit.get("parent_name")
    if parent_id is None:
        parent_name = "— (عنصر رئيسي)"
    elif not parent_name:
        parent = next(
            (item for item in (all_units or []) if item.get("id") == parent_id),
            None,
        )
        parent_name = parent.get("name") if parent else "غير معروف"

    return {
        "id": unit.get("id"),
        "name": unit.get("name"),
        "unit_type": unit.get("unit_type"),
        "parent_id": parent_id,
        "parent_name": parent_name,
        "is_active": bool(unit.get("is_active", True)),
    }


@contextmanager
def db_cursor(commit=False):
    conn = get_connection()
    try:
        yield conn.cursor()
        if commit:
            conn.commit()
    except Exception:
        if commit:
            conn.rollback()
        raise
    finally:
        conn.close()


class OrgUnitRepository:
    @staticmethod
    def fetch_all():
        with db_cursor() as cursor:
            cursor.execute("""
                SELECT
                    o.id, o.name, o.unit_type, o.parent_id, o.is_active,
                    p.name AS parent_name
                FROM dbo.org_units o
                LEFT JOIN dbo.org_units p ON p.id = o.parent_id
                ORDER BY o.sort_order, o.unit_type, o.name
            """)
            return [dict(row.items()) for row in cursor.fetchall()]

    @staticmethod
    def insert(data):
        with db_cursor(commit=True) as cursor:
            cursor.execute(
                """
                INSERT INTO dbo.org_units (name, unit_type, parent_id, is_active, sort_order)
                VALUES (?, ?, ?, 1, 0)
                """,
                (data["name"], data["unit_type"], data["parent_id"]),
            )
            cursor.execute("SELECT CAST(SCOPE_IDENTITY() AS INT)")
            row = cursor.fetchone()
            return row[0] if row else None

    @staticmethod
    def update(unit_id, data):
        with db_cursor(commit=True) as cursor:
            cursor.execute(
                """
                UPDATE dbo.org_units
                SET name = ?, unit_type = ?, parent_id = ?, updated_at = SYSDATETIME()
                WHERE id = ?
                """,
                (data["name"], data["unit_type"], data["parent_id"], unit_id),
            )

    @staticmethod
    def set_active(unit_id, active):
        with db_cursor(commit=True) as cursor:
            cursor.execute(
                """
                UPDATE dbo.org_units
                SET is_active = ?, updated_at = SYSDATETIME()
                WHERE id = ?
                """,
                (1 if active else 0, unit_id),
            )

    @staticmethod
    def delete(unit_id):
        with db_cursor(commit=True) as cursor:
            cursor.execute("DELETE FROM dbo.org_units WHERE id = ?", (unit_id,))

    @staticmethod
    def usage_count(unit_id):
        with db_cursor() as cursor:
            cursor.execute(
                """
                SELECT COUNT(*) FROM dbo.ink_transactions
                WHERE org_unit_id = ? OR office_id = ?
                """,
                (unit_id, unit_id),
            )
            return cursor.fetchone()[0]

    @staticmethod
    def child_count(unit_id):
        with db_cursor() as cursor:
            cursor.execute(
                "SELECT COUNT(*) FROM dbo.org_units WHERE parent_id = ?", (unit_id,)
            )
            return cursor.fetchone()[0]


class OrgIndex:
    def __init__(self, units):
        self.units = list(units)
        self.by_id = {unit["id"]: unit for unit in self.units}
        self.children = {}
        for unit in self.units:
            self.children.setdefault(unit.get("parent_id"), []).append(unit["id"])
        self._counts = {}

    def is_root(self, unit):
        return unit.get("parent_id") not in self.by_id

    def parent_of(self, unit):
        return self.by_id.get(unit.get("parent_id"))

    def path_names(self, unit):
        names, seen, current = [], set(), unit
        while current and current["id"] not in seen:
            seen.add(current["id"])
            names.append(current.get("name") or UNNAMED)
            current = self.by_id.get(current.get("parent_id"))
        return names[::-1]

    def path(self, unit):
        return " / ".join(self.path_names(unit))

    def descendants(self, unit_id):
        found = set()
        stack = list(self.children.get(unit_id, []))
        while stack:
            child_id = stack.pop()
            if child_id in found or child_id == unit_id:
                continue
            found.add(child_id)
            stack.extend(self.children.get(child_id, []))
        return found

    def descendant_count(self, unit_id):
        if unit_id not in self._counts:
            self._counts[unit_id] = len(self.descendants(unit_id))
        return self._counts[unit_id]

    def with_ancestors(self, unit_ids):
        visible = set(unit_ids)
        for unit_id in unit_ids:
            current = self.by_id.get(unit_id)
            seen = set()
            while current:
                parent_id = current.get("parent_id")
                if parent_id not in self.by_id or parent_id in seen:
                    break
                seen.add(parent_id)
                visible.add(parent_id)
                current = self.by_id[parent_id]
        return visible

    def walk(self):
        visited = set()
        roots = [unit for unit in self.units if self.is_root(unit)]
        for start in roots + self.units:
            if start["id"] in visited:
                continue
            stack = [(start, 0)]
            while stack:
                unit, level = stack.pop()
                if unit["id"] in visited:
                    continue
                visited.add(unit["id"])
                yield unit, level
                for child_id in reversed(self.children.get(unit["id"], [])):
                    stack.append((self.by_id[child_id], level + 1))


def confirm_action(parent, title, text, accept_text, icon=QMessageBox.Question):
    box = QMessageBox(icon, title, text, QMessageBox.NoButton, parent)
    box.setLayoutDirection(Qt.RightToLeft)
    accept = box.addButton(accept_text, QMessageBox.AcceptRole)
    cancel = box.addButton("إلغاء", QMessageBox.RejectRole)
    box.setDefaultButton(cancel)
    box.exec()
    return box.clickedButton() is accept


class Toast(QLabel):
    STYLES = {
        "success": ("#166534", "#14532d"),
        "warning": ("#b45309", "#92400e"),
        "error": ("#b91c1c", "#991b1b"),
    }

    def __init__(self, parent):
        super().__init__(parent)
        self.setAlignment(Qt.AlignCenter)
        self.setWordWrap(True)
        self.setAttribute(Qt.WA_TransparentForMouseEvents, True)
        self.setContentsMargins(16, 9, 16, 9)
        self.setMinimumHeight(42)
        self._timer = QTimer(self)
        self._timer.setSingleShot(True)
        self._timer.timeout.connect(self.hide)
        self.hide()

    def show_message(self, message, kind="success", duration_ms=3200):
        background, border = self.STYLES.get(kind, self.STYLES["success"])
        self.setStyleSheet(
            f"QLabel {{ background: {background}; color: white; "
            f"border: 1px solid {border}; border-radius: 11px; "
            "font-size: 12px; font-weight: 700; padding: 9px 16px; }"
        )
        host = self.parentWidget()
        max_width = max(200, min(520, host.width() - 32))
        self.setMinimumWidth(min(240, max_width))
        self.setMaximumWidth(max_width)
        self.setText(message)
        self.adjustSize()
        x = max(16, (host.width() - self.width()) // 2)
        y = max(16, host.height() - self.height() - 24)
        self.move(x, y)
        self.show()
        self.raise_()
        self._timer.start(duration_ms)


class MetricCard(QFrame):
    def __init__(self, title, accent, parent=None):
        super().__init__(parent)
        self.setObjectName("metricCard")
        self.setStyleSheet(
            f"QFrame#metricCard {{ border-right: 4px solid {accent}; }}"
        )
        box = QVBoxLayout(self)
        box.setContentsMargins(15, 11, 15, 11)
        box.setSpacing(3)
        caption = QLabel(title)
        caption.setObjectName("metricTitle")
        self._value = QLabel("0")
        self._value.setObjectName("metricValue")
        self._value.setStyleSheet(f"color: {accent};")
        box.addWidget(caption)
        box.addWidget(self._value)

    def set_value(self, value):
        self._value.setText(str(value))


class TypeBadge(QWidget):
    def __init__(self, unit_type, parent=None):
        super().__init__(parent)
        background, foreground = TYPE_COLORS.get(unit_type, DEFAULT_TYPE_COLORS)
        layout = QHBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setAlignment(Qt.AlignCenter)
        label = QLabel(unit_type or "—")
        label.setAlignment(Qt.AlignCenter)
        label.setFixedSize(84, 28)
        label.setStyleSheet(
            f"background: {background}; color: {foreground}; "
            "border-radius: 14px; font-size: 12px; font-weight: 700;"
        )
        layout.addWidget(label)


class RowActions(QWidget):
    edit_requested = Signal()
    delete_requested = Signal()

    def __init__(self, is_root=False, parent=None):
        super().__init__(parent)
        if is_root:
            self.setObjectName("rootActionCell")
        self.setAttribute(Qt.WA_StyledBackground, True)
        self.setMinimumHeight(ROW_HEIGHT - 6)
        layout = QHBoxLayout(self)
        layout.setContentsMargins(5, 4, 5, 4)
        layout.setSpacing(8)

        edit_button = self._make_button("تعديل", "editBtn", "تعديل الوحدة")
        edit_button.clicked.connect(self.edit_requested.emit)
        delete_button = self._make_button("حذف", "deleteBtn", "حذف الوحدة")
        delete_button.clicked.connect(self.delete_requested.emit)
        layout.addStretch()
        layout.addWidget(edit_button)
        layout.addWidget(delete_button)
        layout.addStretch()

    @staticmethod
    def _make_button(text, object_name, tooltip):
        button = QPushButton(text)
        button.setObjectName(object_name)
        button.setToolTip(tooltip)
        button.setCursor(Qt.PointingHandCursor)
        button.setFixedSize(72, 32)
        return button


class StepIndicator(QFrame):
    def __init__(self, steps, active, parent=None):
        super().__init__(parent)
        self.setObjectName("stepBar")
        row = QHBoxLayout(self)
        row.setContentsMargins(0, 4, 0, 10)
        row.setSpacing(8)
        for index, title in enumerate(steps):
            if index < active:
                state = "done"
            elif index == active:
                state = "active"
            else:
                state = "pending"
            label = QLabel(f"{index + 1}  {title}")
            label.setObjectName("stepItem")
            label.setProperty("stepState", state)
            label.setAlignment(Qt.AlignCenter)
            row.addWidget(label, 1)


class _WizardStep(QWizardPage):
    def __init__(self, index, steps, title, subtitle, validator=None):
        super().__init__()
        self._validator = validator
        self.setTitle(title)
        self.setSubTitle(subtitle)
        self.body = QVBoxLayout(self)
        self.body.setContentsMargins(8, 8, 8, 8)
        self.body.setSpacing(12)
        self.body.addWidget(StepIndicator(steps, index))

    def validatePage(self):
        return self._validator() if self._validator else True


class OrgUnitDialog(QWizard):
    STEP_NAMES = ("نوع الوحدة", "البيانات", "مراجعة")

    def __init__(self, parent=None, unit=None, all_units=None, default_parent_id=None):
        super().__init__(parent)
        self.unit = unit
        self.index = OrgIndex(all_units or [])
        is_edit = unit is not None

        self.setWindowTitle(
            "تعديل عنصر في الهيكل" if is_edit else "إضافة عنصر للهيكل التنظيمي"
        )
        self.setMinimumSize(620, 540)
        self.resize(700, 580)
        self.setLayoutDirection(Qt.RightToLeft)
        self.setWizardStyle(QWizard.ModernStyle)
        self.setOption(QWizard.NoBackButtonOnStartPage, True)
        self.setButtonText(QWizard.BackButton, "السابق")
        self.setButtonText(QWizard.NextButton, "التالي")
        self.setButtonText(QWizard.CancelButton, "إلغاء")
        self.setButtonText(
            QWizard.FinishButton, "حفظ التعديلات" if is_edit else "حفظ العنصر"
        )

        self._build_type_page()
        self._build_details_page()
        self._build_review_page()
        self.setStyleSheet(build_style(DIALOG_STYLE))
        self.button(QWizard.FinishButton).setObjectName("primaryButton")

        self.type_combo.currentTextChanged.connect(self._on_type_changed)
        self.name_input.textChanged.connect(self._on_name_changed)
        self.parent_combo.currentIndexChanged.connect(self._refresh_review)
        self.currentIdChanged.connect(self._refresh_review)

        self._apply_initial_state(default_parent_id)
        if is_edit:
            self.setStartId(self.details_page_id)

    def _build_type_page(self):
        page = _WizardStep(
            0, self.STEP_NAMES, "نوع الوحدة", "ابدأ بتحديد نوع العنصر الذي تريد إضافته."
        )
        card = QFrame()
        card.setObjectName("wizardCard")
        layout = QVBoxLayout(card)
        layout.setContentsMargins(20, 18, 20, 18)
        layout.setSpacing(12)
        layout.addWidget(QLabel("نوع العنصر التنظيمي *"))
        self.type_combo = QComboBox()
        self.type_combo.addItems(UNIT_TYPES)
        self.type_combo.setMinimumHeight(42)
        layout.addWidget(self.type_combo)
        self.type_hint = QLabel()
        self.type_hint.setObjectName("fieldHint")
        self.type_hint.setWordWrap(True)
        layout.addWidget(self.type_hint)
        layout.addStretch()
        page.body.addWidget(card, 1)
        self.type_page_id = self.addPage(page)

    def _build_details_page(self):
        page = _WizardStep(
            1,
            self.STEP_NAMES,
            "بيانات الوحدة",
            "أدخل الاسم وحدد موقع الوحدة ضمن الهيكل.",
            validator=self._validate_details,
        )
        card = QFrame()
        card.setObjectName("wizardCard")
        form = QFormLayout(card)
        form.setContentsMargins(20, 20, 20, 20)
        form.setHorizontalSpacing(16)
        form.setVerticalSpacing(14)
        form.setLabelAlignment(Qt.AlignRight | Qt.AlignVCenter)

        self.name_input = QLineEdit()
        self.name_input.setPlaceholderText("مثال: فرع التنظيم")
        self.name_input.setMinimumHeight(42)
        self.name_input.setClearButtonEnabled(True)
        self.name_error = QLabel()
        self.name_error.setObjectName("fieldError")
        self.name_error.setWordWrap(True)
        self.name_error.hide()

        self.parent_combo = QComboBox()
        self.parent_combo.setMinimumHeight(42)
        self.parent_combo.setMaxVisibleItems(12)

        form.addRow("اسم الوحدة *", self.name_input)
        form.addRow("", self.name_error)
        form.addRow("تابعة إلى", self.parent_combo)
        page.body.addWidget(card)
        page.body.addStretch()
        self.details_page = page
        self.details_page_id = self.addPage(page)

    def _build_review_page(self):
        page = _WizardStep(
            2,
            self.STEP_NAMES,
            "مراجعة وحفظ",
            "تأكد من البيانات قبل حفظها في الهيكل.",
        )
        card = QFrame()
        card.setObjectName("wizardCard")
        form = QFormLayout(card)
        form.setContentsMargins(20, 20, 20, 20)
        form.setHorizontalSpacing(16)
        form.setVerticalSpacing(14)
        self.review_name = QLabel("—")
        self.review_type = QLabel("—")
        self.review_parent = QLabel("—")
        self.review_path = QLabel("—")
        for label in (
            self.review_name,
            self.review_type,
            self.review_parent,
            self.review_path,
        ):
            label.setObjectName("reviewValue")
            label.setWordWrap(True)
        form.addRow("اسم الوحدة", self.review_name)
        form.addRow("النوع", self.review_type)
        form.addRow("تابعة إلى", self.review_parent)
        form.addRow("المسار الناتج", self.review_path)
        page.body.addWidget(card)
        page.body.addStretch()
        self.review_page_id = self.addPage(page)

    def _apply_initial_state(self, default_parent_id):
        if self.unit:
            self.name_input.setText(self.unit.get("name") or "")
            self._select_type(self.unit.get("unit_type"))
            self._populate_parent_combo(self.unit.get("parent_id"))
        elif default_parent_id is not None:
            self._select_type("قسم")
            self._populate_parent_combo(default_parent_id)
        else:
            self._populate_parent_combo()
        self._refresh_review()

    def _select_type(self, unit_type):
        index = self.type_combo.findText(unit_type or "")
        if index >= 0:
            self.type_combo.setCurrentIndex(index)

    def _populate_parent_combo(self, selected_id=None):
        previous = selected_id if selected_id is not None else self.parent_combo.currentData()
        self.parent_combo.blockSignals(True)
        self.parent_combo.clear()
        self.parent_combo.addItem(NO_PARENT_LABEL, None)

        allows_parent = self.type_combo.currentText() not in ROOT_ONLY_TYPES
        self.parent_combo.setEnabled(allows_parent)
        if allows_parent:
            excluded = set()
            if self.unit:
                excluded = {self.unit["id"]} | self.index.descendants(self.unit["id"])
            candidates = sorted(
                (
                    unit
                    for unit in self.index.units
                    if unit["id"] not in excluded
                    and unit.get("unit_type") in PARENT_CAPABLE_TYPES
                ),
                key=self.index.path,
            )
            for candidate in candidates:
                self.parent_combo.addItem(
                    f"[{candidate['unit_type']}]  {self.index.path(candidate)}",
                    candidate["id"],
                )
            position = self.parent_combo.findData(previous)
            if position >= 0:
                self.parent_combo.setCurrentIndex(position)
        self.parent_combo.blockSignals(False)

    def _on_type_changed(self, _unit_type):
        self._populate_parent_combo()
        self._refresh_review()

    def _on_name_changed(self, _text):
        self.name_error.hide()
        self._refresh_review()

    def _selected_parent(self):
        return self.index.by_id.get(self.parent_combo.currentData())

    def _is_duplicate(self, name):
        own_id = self.unit["id"] if self.unit else None
        unit_type = self.type_combo.currentText()
        parent_id = self.parent_combo.currentData()
        folded = name.casefold()
        return any(
            unit["id"] != own_id
            and unit.get("unit_type") == unit_type
            and unit.get("parent_id") == parent_id
            and (unit.get("name") or "").strip().casefold() == folded
            for unit in self.index.units
        )

    def _validate_details(self):
        name = self.name_input.text().strip()
        error = ""
        if not name:
            error = "اكتب اسم الوحدة قبل المتابعة."
        elif self._is_duplicate(name):
            error = "توجد وحدة بنفس الاسم والنوع تحت نفس الوحدة الأعلى."
        if error:
            self.name_error.setText(error)
            self.name_error.show()
            self.name_input.setFocus()
            return False
        return True

    def _resulting_path(self):
        name = self.name_input.text().strip()
        if not name:
            return "—"
        parent = self._selected_parent()
        return f"{self.index.path(parent)} / {name}" if parent else name

    def _refresh_review(self, *_args):
        unit_type = self.type_combo.currentText()
        parent = self._selected_parent()
        self.type_hint.setText(TYPE_HINTS.get(unit_type, ""))
        self.review_name.setText(self.name_input.text().strip() or "—")
        self.review_type.setText(unit_type or "—")
        self.review_parent.setText(
            f"{parent.get('name')} — {parent.get('unit_type')}"
            if parent
            else "بدون وحدة أعلى (عنصر رئيسي)"
        )
        self.review_path.setText(self._resulting_path())

    def get_data(self):
        return {
            "name": self.name_input.text().strip(),
            "unit_type": self.type_combo.currentText(),
            "parent_id": self.parent_combo.currentData(),
        }


EXPORT_HEADERS = (
    "المستوى",
    "الوحدة التنظيمية",
    "النوع",
    "الحالة",
    "عدد التابعين",
    "المسار الكامل",
)
EXPORT_WIDTHS = (10, 38, 16, 14, 14, 64)
EXPORT_RIGHT_ALIGNED = (2, 6)


def export_hierarchy_to_excel(file_path, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = MODULE_NAME
    sheet.sheet_view.rightToLeft = True
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90

    last_column = get_column_letter(len(EXPORT_HEADERS))
    centered = Alignment(horizontal="center", vertical="center")

    sheet.merge_cells(f"A1:{last_column}1")
    title = sheet["A1"]
    title.value = MODULE_NAME
    title.font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="10263D")
    title.alignment = centered

    sheet.merge_cells(f"A2:{last_column}2")
    total = sheet["A2"]
    total.value = f"إجمالي الوحدات: {len(rows)}"
    total.font = Font(name="Arial", size=11, bold=True, color="64748B")
    total.alignment = centered

    header_row = 3
    for column, value in enumerate(EXPORT_HEADERS, 1):
        cell = sheet.cell(header_row, column, value)
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1095B0")
        cell.alignment = centered

    border = Border(bottom=Side(style="thin", color="D9DEE5"))
    root_fill = PatternFill("solid", fgColor="FFF7E8")
    for row_index, data in enumerate(rows, header_row + 1):
        is_root = data["level"] == 0
        name = f"{data['name']} — وحدة رئيسية" if is_root else data["name"]
        values = (
            data["level"],
            name,
            data["unit_type"],
            data["status"],
            data["descendants"],
            data["path"],
        )
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_index, column, value)
            cell.font = Font(
                name="Arial",
                size=10,
                bold=is_root,
                color="9A4D18" if is_root else "334155",
            )
            cell.alignment = Alignment(
                horizontal="right" if column in EXPORT_RIGHT_ALIGNED else "center",
                vertical="center",
                wrap_text=column == len(EXPORT_HEADERS),
                indent=min(data["level"] * 2, 15) if column == 2 else 0,
            )
            cell.border = border
            if is_root:
                cell.fill = root_fill
        sheet.row_dimensions[row_index].height = 24
        sheet.row_dimensions[row_index].outlineLevel = min(data["level"], 7)

    for column, width in enumerate(EXPORT_WIDTHS, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 32
    sheet.row_dimensions[2].height = 24
    sheet.row_dimensions[header_row].height = 28
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{last_column}{header_row + len(rows)}"
    sheet.sheet_properties.outlinePr.summaryBelow = True
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = f"1:{header_row}"

    workbook.save(file_path)
    workbook.close()


class OrgStructurePage(QWidget):
    def __init__(self, current_user=None, parent=None):
        super().__init__(parent)
        self.current_user = current_user or {}
        self._units = []
        self._index = OrgIndex([])
        self._items = {}
        self._quick_label_to_id = {}
        self._quick_label_by_id = {}
        self.setLayoutDirection(Qt.RightToLeft)
        self.setObjectName("orgPage")
        self._build_ui()
        self._bind_shortcuts()
        self.reload()

    def _build_ui(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(26, 22, 26, 22)
        root.setSpacing(14)
        root.addLayout(self._build_header())
        root.addLayout(self._build_metrics())
        root.addWidget(self._build_filters())
        root.addWidget(self._build_tree_card(), 1)
        self.setStyleSheet(build_style(PAGE_STYLE))
        self._toast = Toast(self)

        self._search_timer = QTimer(self)
        self._search_timer.setSingleShot(True)
        self._search_timer.setInterval(220)
        self._search_timer.timeout.connect(self._render)

    def _build_header(self):
        row = QHBoxLayout()
        titles = QVBoxLayout()
        titles.setSpacing(3)
        title = QLabel("إدارة الهيكل التنظيمي")
        title.setObjectName("pageTitle")
        subtitle = QLabel("استكشف الفروع والمباني والأقسام والمكاتب في عرض هرمي واضح")
        subtitle.setObjectName("pageSubtitle")
        titles.addWidget(title)
        titles.addWidget(subtitle)
        row.addLayout(titles)
        row.addStretch()

        self.btn_export = QPushButton("تصدير إلى Excel")
        self.btn_export.setObjectName("toolButton")
        self.btn_export.setToolTip("تصدير الهيكل بالكامل إلى ملف Excel")
        self.btn_export.clicked.connect(self._export_excel)

        self.btn_add = QPushButton("＋  إضافة وحدة")
        self.btn_add.setObjectName("primaryButton")
        self.btn_add.setToolTip("إضافة وحدة جديدة (Ctrl+N)")
        self.btn_add.clicked.connect(lambda: self._add_unit())

        for button in (self.btn_export, self.btn_add):
            button.setMinimumHeight(42)
            button.setCursor(Qt.PointingHandCursor)
            row.addWidget(button)
        return row

    def _build_metrics(self):
        row = QHBoxLayout()
        row.setSpacing(10)
        self.metric_total = MetricCard("إجمالي الوحدات", PALETTE["accent"])
        self.metric_roots = MetricCard("وحدات رئيسية", PALETTE["root_fg"])
        self.metric_active = MetricCard("وحدات نشطة", PALETTE["success"])
        self.metric_inactive = MetricCard("وحدات معطلة", PALETTE["danger"])
        for card in (
            self.metric_total,
            self.metric_roots,
            self.metric_active,
            self.metric_inactive,
        ):
            row.addWidget(card, 1)
        return row

    def _build_filters(self):
        card = QFrame()
        card.setObjectName("filterCard")
        row = QHBoxLayout(card)
        row.setContentsMargins(14, 12, 14, 12)
        row.setSpacing(10)

        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("ابحث باسم الفرع أو الوحدة...")
        self.search_input.setClearButtonEnabled(True)
        self.search_input.setMinimumHeight(38)
        self.search_input.textChanged.connect(lambda _text: self._search_timer.start())

        self.filter_type = QComboBox()
        self.filter_type.addItem("كل الأنواع", "")
        for unit_type in UNIT_TYPES:
            self.filter_type.addItem(unit_type, unit_type)
        self.filter_type.setMinimumHeight(38)
        self.filter_type.currentIndexChanged.connect(lambda _index: self._render())

        self.filter_status = QComboBox()
        self.filter_status.addItem("كل الحالات", "")
        self.filter_status.addItem("نشط فقط", "1")
        self.filter_status.addItem("غير نشط فقط", "0")
        self.filter_status.setCurrentIndex(1)
        self.filter_status.setMinimumHeight(38)
        self.filter_status.currentIndexChanged.connect(lambda _index: self._render())

        clear_button = QPushButton("مسح الفلاتر")
        clear_button.setObjectName("clearButton")
        clear_button.setMinimumHeight(38)
        clear_button.setCursor(Qt.PointingHandCursor)
        clear_button.clicked.connect(lambda: self._set_filters())

        for caption, widget, stretch in (
            ("بحث", self.search_input, 1),
            ("النوع", self.filter_type, 0),
            ("الحالة", self.filter_status, 0),
        ):
            label = QLabel(caption)
            label.setObjectName("fieldCaption")
            row.addWidget(label)
            row.addWidget(widget, stretch)
        row.addWidget(clear_button)
        return card

    def _build_tree_card(self):
        card = QFrame()
        card.setObjectName("treeCard")
        layout = QVBoxLayout(card)
        layout.setContentsMargins(10, 10, 10, 10)
        layout.setSpacing(8)
        layout.addLayout(self._build_tree_toolbar())

        self.tree = QTreeWidget()
        self.tree.setColumnCount(4)
        self.tree.setHeaderLabels(
            ["الوحدة التنظيمية", "النوع", "عدد العناصر التابعة", "الإجراءات"]
        )
        self.tree.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.tree.setSelectionMode(QAbstractItemView.SingleSelection)
        self.tree.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.tree.setAlternatingRowColors(True)
        self.tree.setRootIsDecorated(True)
        self.tree.setIndentation(22)
        self.tree.setAnimated(True)
        self.tree.setAllColumnsShowFocus(True)
        self.tree.setContextMenuPolicy(Qt.CustomContextMenu)
        self.tree.customContextMenuRequested.connect(self._show_context_menu)
        self.tree.itemDoubleClicked.connect(self._on_item_double_clicked)

        header = self.tree.header()
        header.setStretchLastSection(False)
        header.setSectionResizeMode(COL_NAME, QHeaderView.Stretch)
        for column, width in (
            (COL_TYPE, TYPE_COLUMN_WIDTH),
            (COL_COUNT, COUNT_COLUMN_WIDTH),
            (COL_ACTIONS, ACTIONS_COLUMN_WIDTH),
        ):
            header.setSectionResizeMode(column, QHeaderView.Fixed)
            header.resizeSection(column, width)
        header_item = self.tree.headerItem()
        for column in (COL_TYPE, COL_COUNT, COL_ACTIONS):
            header_item.setTextAlignment(column, Qt.AlignCenter)
        layout.addWidget(self.tree, 1)

        self.empty_state = self._build_empty_state()
        layout.addWidget(self.empty_state, 1)

        self.stats_lbl = QLabel("")
        self.stats_lbl.setObjectName("statsLabel")
        layout.addWidget(self.stats_lbl)
        return card

    def _build_tree_toolbar(self):
        row = QHBoxLayout()
        title = QLabel("الوحدات والعلاقات التنظيمية")
        title.setObjectName("sectionTitle")
        row.addWidget(title)
        row.addStretch()

        self.quick_search = QLineEdit()
        self.quick_search.setObjectName("quickSearch")
        self.quick_search.setPlaceholderText("بحث سريع داخل الشجرة…")
        self.quick_search.setClearButtonEnabled(True)
        self.quick_search.setToolTip(
            "اكتب اسم الوحدة أو مسارها، ثم اختر النتيجة للانتقال إليها (Ctrl+F)"
        )
        self.quick_search_model = QStringListModel([], self)
        completer = QCompleter(self.quick_search_model, self)
        completer.setCaseSensitivity(Qt.CaseInsensitive)
        completer.setFilterMode(Qt.MatchContains)
        completer.setCompletionMode(QCompleter.PopupCompletion)
        completer.activated[str].connect(self._on_quick_completion)
        self.quick_search_completer = completer
        self.quick_search.setCompleter(completer)
        self.quick_search.returnPressed.connect(self._quick_find_unit)
        row.addWidget(self.quick_search)

        self.quick_find_button = QPushButton("انتقال")
        self.quick_find_button.setObjectName("toolButton")
        self.quick_find_button.clicked.connect(self._quick_find_unit)
        self.btn_collapse = QPushButton("طي الكل")
        self.btn_collapse.setObjectName("toolButton")
        self.btn_collapse.clicked.connect(lambda: self.tree.collapseAll())
        self.btn_expand = QPushButton("توسيع الكل")
        self.btn_expand.setObjectName("toolButton")
        self.btn_expand.clicked.connect(lambda: self.tree.expandAll())
        for button in (self.quick_find_button, self.btn_collapse, self.btn_expand):
            button.setCursor(Qt.PointingHandCursor)
            row.addWidget(button)
        return row

    def _build_empty_state(self):
        frame = QFrame()
        layout = QVBoxLayout(frame)
        layout.setAlignment(Qt.AlignCenter)
        layout.setSpacing(6)
        self.empty_title = QLabel()
        self.empty_title.setObjectName("emptyTitle")
        self.empty_title.setAlignment(Qt.AlignCenter)
        self.empty_hint = QLabel()
        self.empty_hint.setObjectName("emptyHint")
        self.empty_hint.setAlignment(Qt.AlignCenter)
        self.empty_hint.setWordWrap(True)
        layout.addWidget(self.empty_title)
        layout.addWidget(self.empty_hint)
        frame.hide()
        return frame

    def _bind_shortcuts(self):
        shortcuts = (
            (QKeySequence.New, self, lambda: self._add_unit()),
            (QKeySequence.Find, self, self._focus_quick_search),
            (QKeySequence(Qt.Key_Delete), self.tree, self._delete_selected),
        )
        self._shortcuts = []
        for sequence, widget, handler in shortcuts:
            shortcut = QShortcut(QKeySequence(sequence), widget)
            shortcut.setContext(Qt.WidgetWithChildrenShortcut)
            shortcut.activated.connect(handler)
            self._shortcuts.append(shortcut)

    def _focus_quick_search(self):
        self.quick_search.setFocus()
        self.quick_search.selectAll()

    def load_units(self, *_args):
        self.reload()

    def reload(self):
        try:
            units = OrgUnitRepository.fetch_all()
        except Exception as error:
            QMessageBox.critical(self, "خطأ", f"تعذر تحميل البيانات:\n{error}")
            return
        self._units = units
        self._index = OrgIndex(units)
        self._refresh_quick_search_model()
        self._render()

    def _current_filters(self):
        return (
            self.search_input.text().strip().lower(),
            self.filter_type.currentData() or "",
            self.filter_status.currentData() or "",
        )

    @staticmethod
    def _matches(unit, search, type_filter, status_filter):
        if type_filter and unit.get("unit_type") != type_filter:
            return False
        if status_filter and str(int(bool(unit.get("is_active")))) != status_filter:
            return False
        if search and search not in (unit.get("name") or "").lower():
            return False
        return True

    def _set_filters(self, search="", type_index=0, status_index=1):
        widgets = (self.search_input, self.filter_type, self.filter_status)
        for widget in widgets:
            widget.blockSignals(True)
        self.search_input.setText(search)
        self.filter_type.setCurrentIndex(type_index)
        self.filter_status.setCurrentIndex(status_index)
        for widget in widgets:
            widget.blockSignals(False)
        self._search_timer.stop()
        self._render()

    def _render(self):
        search, type_filter, status_filter = self._current_filters()
        selected_id = self._selected_unit_id()
        matched = {
            unit["id"]
            for unit in self._units
            if self._matches(unit, search, type_filter, status_filter)
        }
        visible = self._index.with_ancestors(matched)
        ordered = [unit for unit in self._units if unit["id"] in visible]

        self.tree.setUpdatesEnabled(False)
        self.tree.clear()
        self._items = {unit["id"]: self._build_item(unit) for unit in ordered}
        for unit in ordered:
            item = self._items[unit["id"]]
            parent_item = self._items.get(unit.get("parent_id"))
            if parent_item is not None:
                parent_item.addChild(item)
            else:
                self.tree.addTopLevelItem(item)
        for unit in ordered:
            self._attach_actions(unit, self._items[unit["id"]])

        is_filtering = bool(search or type_filter or status_filter == "0")
        if is_filtering:
            self.tree.expandAll()
        else:
            self.tree.collapseAll()
            for position in range(self.tree.topLevelItemCount()):
                self.tree.topLevelItem(position).setExpanded(True)
        self.tree.setUpdatesEnabled(True)

        if selected_id in self._items:
            self._reveal(self._items[selected_id])
            self.tree.setCurrentItem(self._items[selected_id], COL_NAME)

        self._update_summary(len(matched), len(visible), is_filtering)

    def _build_item(self, unit):
        is_root = self._index.is_root(unit)
        is_active = bool(unit.get("is_active"))
        name = unit.get("name") or UNNAMED
        item = QTreeWidgetItem(
            [
                f"◆ {name}  — وحدة رئيسية" if is_root else name,
                "",
                str(self._index.descendant_count(unit["id"])),
                "",
            ]
        )
        item.setData(COL_NAME, Qt.UserRole, unit["id"])
        item.setSizeHint(COL_NAME, QSize(0, ROW_HEIGHT))
        for column in (COL_TYPE, COL_COUNT):
            item.setTextAlignment(column, Qt.AlignCenter)

        item.setToolTip(COL_TYPE, unit.get("unit_type") or "")
        item.setForeground(COL_COUNT, QColor(PALETTE["accent_text"]))
        item.setToolTip(COL_COUNT, "عدد جميع العناصر التابعة، بكل المستويات والحالات")

        if not is_active:
            item.setForeground(COL_NAME, QColor(PALETTE["faint"]))
        if is_root:
            font = QFont(item.font(COL_NAME))
            font.setPointSize(11)
            font.setBold(True)
            item.setFont(COL_NAME, font)
            item.setToolTip(COL_NAME, "وحدة رئيسية في الهيكل التنظيمي")
            if is_active:
                item.setForeground(COL_NAME, QColor(PALETTE["root_fg"]))
            for column in (COL_NAME, COL_TYPE, COL_COUNT):
                item.setBackground(column, QColor(PALETTE["root_bg"]))
        return item

    def _attach_actions(self, unit, item):
        actions = RowActions(is_root=self._index.is_root(unit))
        unit_id = unit["id"]
        actions.edit_requested.connect(lambda uid=unit_id: self._edit_unit(uid))
        actions.delete_requested.connect(lambda uid=unit_id: self._delete_unit(uid))
        self.tree.setItemWidget(item, COL_TYPE, TypeBadge(unit.get("unit_type")))
        self.tree.setItemWidget(item, COL_ACTIONS, actions)

    def _update_summary(self, matched_count, visible_count, is_filtering):
        total = len(self._units)
        active = sum(bool(unit.get("is_active")) for unit in self._units)
        roots = sum(self._index.is_root(unit) for unit in self._units)
        self.metric_total.set_value(total)
        self.metric_roots.set_value(roots)
        self.metric_active.set_value(active)
        self.metric_inactive.set_value(total - active)
        self.stats_lbl.setText(
            f"إجمالي الوحدات: {total}  |  مطابقة للفلاتر: {matched_count}  |  "
            f"ظاهرة في المسار: {visible_count}"
        )

        has_rows = bool(self._items)
        self.tree.setVisible(has_rows)
        self.empty_state.setVisible(not has_rows)
        if not has_rows:
            if total == 0:
                self.empty_title.setText("لا توجد وحدات بعد")
                self.empty_hint.setText("ابدأ بإضافة أول فرع أو مبنى من زر «إضافة وحدة».")
            else:
                self.empty_title.setText("لا توجد نتائج مطابقة")
                self.empty_hint.setText("جرّب تغيير كلمات البحث أو الضغط على «مسح الفلاتر».")

    def _selected_unit_id(self):
        item = self.tree.currentItem()
        return item.data(COL_NAME, Qt.UserRole) if item else None

    def _reveal(self, item):
        ancestor = item.parent()
        while ancestor is not None:
            ancestor.setExpanded(True)
            ancestor = ancestor.parent()

    def _refresh_quick_search_model(self):
        self._quick_label_to_id = {}
        self._quick_label_by_id = {}
        labels = []
        for unit in self._units:
            label = f"{self._index.path(unit)}  [{unit.get('unit_type') or 'وحدة'}]"
            if label.casefold() in self._quick_label_to_id:
                label = f"{label}  · {unit['id']}"
            self._quick_label_to_id[label.casefold()] = unit["id"]
            self._quick_label_by_id[unit["id"]] = label
            labels.append(label)
        self.quick_search_model.setStringList(labels)

    def _on_quick_completion(self, label):
        unit_id = self._quick_label_to_id.get(str(label).casefold())
        if unit_id is not None:
            self._navigate_to_unit(unit_id)

    def _find_candidates(self, query):
        names = [(unit, (unit.get("name") or "").casefold()) for unit in self._units]
        for predicate in (
            lambda name: name == query,
            lambda name: name.startswith(query),
            lambda name: query in name,
        ):
            found = [unit for unit, name in names if predicate(name)]
            if found:
                return found
        return []

    def _quick_find_unit(self):
        query = self.quick_search.text().strip().casefold()
        if not query:
            return

        unit_id = self._quick_label_to_id.get(query)
        if unit_id is not None:
            self._navigate_to_unit(unit_id)
            return

        candidates = self._find_candidates(query)
        if not candidates:
            self._toast.show_message("لم يتم العثور على وحدة بهذا الاسم.", "warning")
            return
        if len(candidates) > 1:
            self.quick_search.setFocus()
            self.quick_search_completer.complete()
            self._toast.show_message(
                "وجدت أكثر من وحدة مطابقة؛ اختر المسار الصحيح من القائمة.", "warning"
            )
            return
        self._navigate_to_unit(candidates[0]["id"])

    def _navigate_to_unit(self, unit_id):
        if unit_id not in self._items:
            self._set_filters(status_index=0)
        item = self._items.get(unit_id)
        if item is None:
            self._toast.show_message("تعذر إظهار الوحدة في الشجرة.", "warning")
            return
        self._reveal(item)
        if item.childCount():
            item.setExpanded(True)
        self.tree.setCurrentItem(item, COL_NAME)
        self.tree.scrollToItem(item, QAbstractItemView.PositionAtCenter)
        self.tree.setFocus()
        self.quick_search.setText(self._quick_label_by_id.get(unit_id, ""))

    def _on_item_double_clicked(self, item, _column):
        unit_id = item.data(COL_NAME, Qt.UserRole)
        if unit_id is not None:
            self._edit_unit(unit_id)

    def _delete_selected(self):
        unit_id = self._selected_unit_id()
        if unit_id is not None:
            self._delete_unit(unit_id)

    def _show_context_menu(self, position):
        item = self.tree.itemAt(position)
        if item is None:
            return
        unit = self._index.by_id.get(item.data(COL_NAME, Qt.UserRole))
        if unit is None:
            return
        self.tree.setCurrentItem(item)
        unit_id = unit["id"]
        is_active = bool(unit.get("is_active"))

        menu = QMenu(self)
        menu.setLayoutDirection(Qt.RightToLeft)
        menu.addAction("تعديل", lambda: self._edit_unit(unit_id))
        if unit.get("unit_type") in PARENT_CAPABLE_TYPES:
            menu.addAction("إضافة وحدة تابعة", lambda: self._add_unit(unit_id))
        menu.addAction(
            "تعطيل الوحدة" if is_active else "تفعيل الوحدة",
            lambda: self._set_active(unit_id, not is_active),
        )
        menu.addAction("نسخ المسار", lambda: self._copy_path(unit_id))
        menu.addSeparator()
        menu.addAction("حذف", lambda: self._delete_unit(unit_id))
        menu.exec(self.tree.viewport().mapToGlobal(position))

    def _copy_path(self, unit_id):
        unit = self._index.by_id.get(unit_id)
        if unit is None:
            return
        QGuiApplication.clipboard().setText(self._index.path(unit))
        self._toast.show_message("تم نسخ مسار الوحدة")

    def _audit(self, action, description, record_id=None, old_values=None, new_values=None):
        payload = {
            "user": self.current_user,
            "action": action,
            "module_name": MODULE_NAME,
            "description": description,
        }
        if record_id is not None:
            payload["record_id"] = record_id
        if old_values is not None:
            payload["old_values"] = old_values
        if new_values is not None:
            payload["new_values"] = new_values
        write_audit_log(**payload)

    def _commit_change(self, failure_text, operation, success_message):
        try:
            focus_id = operation()
        except Exception as error:
            QMessageBox.critical(self, "خطأ", f"{failure_text}\n{error}")
            return
        # أبطل الكاش المركزي وأشعر كل الشاشات الأخرى بالتغيير
        OrgStructureCache.instance().invalidate()
        self.reload()
        if focus_id in self._items or focus_id in self._index.by_id:
            self._navigate_to_unit(focus_id)
        self._toast.show_message(success_message)

    def _add_unit(self, parent_id=None):
        dialog = OrgUnitDialog(self, all_units=self._units, default_parent_id=parent_id)
        if dialog.exec() != QDialog.Accepted:
            return
        data = dialog.get_data()

        def operation():
            unit_id = OrgUnitRepository.insert(data)
            self._audit(
                "إضافة",
                f"إضافة عنصر: {data['name']} ({data['unit_type']})",
                record_id=unit_id,
                new_values=org_unit_audit_data(
                    {**data, "id": unit_id, "is_active": True}, self._units
                ),
            )
            return unit_id

        self._commit_change("تعذر الإضافة:", operation, f"تمت إضافة «{data['name']}» بنجاح")

    def _edit_unit(self, unit_id):
        unit = self._index.by_id.get(unit_id)
        if unit is None:
            return
        dialog = OrgUnitDialog(self, unit=unit, all_units=self._units)
        if dialog.exec() != QDialog.Accepted:
            return
        data = dialog.get_data()
        unchanged = (
            data["name"] == (unit.get("name") or "")
            and data["unit_type"] == unit.get("unit_type")
            and data["parent_id"] == unit.get("parent_id")
        )
        if unchanged:
            return

        def operation():
            OrgUnitRepository.update(unit_id, data)
            self._audit(
                "تعديل",
                f"تعديل عنصر ID={unit_id}: {data['name']}",
                record_id=unit_id,
                old_values=org_unit_audit_data(unit, self._units),
                new_values=org_unit_audit_data(
                    {**data, "id": unit_id, "is_active": unit.get("is_active", True)},
                    self._units,
                ),
            )
            return unit_id

        self._commit_change("تعذر التعديل:", operation, "تم حفظ التعديلات بنجاح")

    def _set_active(self, unit_id, active):
        unit = self._index.by_id.get(unit_id)
        if unit is None:
            return
        name = unit.get("name")

        def operation():
            OrgUnitRepository.set_active(unit_id, active)
            self._audit(
                "تفعيل" if active else "تعطيل",
                f"{'تفعيل' if active else 'تعطيل'} عنصر: {name}",
                record_id=unit_id,
                old_values=org_unit_audit_data(unit, self._units),
                new_values=org_unit_audit_data({**unit, "is_active": active}, self._units),
            )
            return unit_id

        self._commit_change(
            "تعذر تغيير الحالة:",
            operation,
            f"تم {'تفعيل' if active else 'تعطيل'} «{name}»",
        )

    def _delete_unit(self, unit_id):
        unit = self._index.by_id.get(unit_id)
        if unit is None:
            return
        name = unit.get("name")

        try:
            usage = OrgUnitRepository.usage_count(unit_id)
            children = OrgUnitRepository.child_count(unit_id)
        except Exception as error:
            QMessageBox.critical(self, "خطأ", f"تعذر الفحص:\n{error}")
            return

        if children > 0:
            QMessageBox.warning(
                self,
                "تعذر الحذف",
                f"العنصر «{name}» لديه {children} عنصر تابع.\n"
                "يجب حذف العناصر التابعة أولاً أو نقلها.",
            )
            return

        if usage > 0:
            self._offer_deactivation(unit, usage)
            return

        confirmed = confirm_action(
            self,
            "تأكيد الحذف",
            f"هل تريد حذف العنصر «{name}» نهائياً؟\n\nلا يمكن التراجع عن هذا الإجراء.",
            "حذف نهائي",
            QMessageBox.Warning,
        )
        if not confirmed:
            return

        def operation():
            OrgUnitRepository.delete(unit_id)
            self._audit(
                "حذف",
                f"حذف عنصر: {name}",
                record_id=unit_id,
                old_values=org_unit_audit_data(unit, self._units),
                new_values={"deleted": True},
            )
            return None

        self._commit_change("تعذر الحذف:", operation, f"تم حذف «{name}»")

    def _offer_deactivation(self, unit, usage):
        name = unit.get("name")
        text = (
            f"العنصر «{name}» مرتبط بـ {usage} سجل(ات) أحبار.\n"
            "لا يمكن حذفه فعليًا حفاظًا على سجل العمليات."
        )
        if not unit.get("is_active"):
            QMessageBox.warning(self, "تعذر الحذف", text)
            return
        box = QMessageBox(QMessageBox.Warning, "تعذر الحذف", text, QMessageBox.NoButton, self)
        box.setLayoutDirection(Qt.RightToLeft)
        deactivate = box.addButton("تعطيل الوحدة بدلاً من الحذف", QMessageBox.AcceptRole)
        box.addButton("إغلاق", QMessageBox.RejectRole)
        box.exec()
        if box.clickedButton() is deactivate:
            self._set_active(unit["id"], False)

    def _export_rows(self):
        rows = []
        for unit, level in self._index.walk():
            parent = self._index.parent_of(unit)
            rows.append(
                {
                    "id": unit["id"],
                    "level": level,
                    "name": unit.get("name") or UNNAMED,
                    "unit_type": unit.get("unit_type") or "",
                    "status": "نشط" if unit.get("is_active") else "معطل",
                    "parent_name": parent.get("name") if parent else "—",
                    "descendants": self._index.descendant_count(unit["id"]),
                    "path": self._index.path(unit),
                }
            )
        return rows

    def _export_excel(self):
        if not self._units:
            QMessageBox.information(self, "لا توجد بيانات", "لا توجد وحدات لتصديرها.")
            return

        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "تصدير الهيكل التنظيمي إلى Excel",
            "الهيكل_التنظيمي.xlsx",
            "Excel Workbook (*.xlsx)",
        )
        if not file_path:
            return
        if not file_path.lower().endswith(".xlsx"):
            file_path += ".xlsx"

        try:
            rows = self._export_rows()
            export_hierarchy_to_excel(file_path, rows)
            self._audit(
                "تصدير",
                f"تصدير الهيكل التنظيمي إلى Excel: {Path(file_path).name}",
                new_values={
                    "file_name": Path(file_path).name,
                    "record_count": len(rows),
                },
            )
            self._toast.show_message("تم تصدير الشجرة بالكامل إلى Excel بنجاح")
        except ImportError:
            QMessageBox.critical(
                self,
                "مكتبة غير متاحة",
                "لتصدير Excel، ثبّت مكتبة openpyxl:\n\npython -m pip install openpyxl",
            )
        except Exception as error:
            QMessageBox.critical(
                self, "خطأ في التصدير", f"تعذر إنشاء ملف Excel:\n{error}"
            )