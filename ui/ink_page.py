"""
ui/ink_page.py
==============
شاشة الأحبار - نظام إدارة حركات الأحبار
"""

import datetime
from pathlib import Path

from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel,
    QPushButton, QTableWidget, QTableWidgetItem,
    QHeaderView, QDialog, QFormLayout, QComboBox,
    QSpinBox, QDateEdit, QPlainTextEdit, QLineEdit,
    QMessageBox, QFrame, QDialogButtonBox,
    QAbstractItemView, QScrollArea, QFileDialog,
)
from PySide6.QtCore import Qt, QDate, QByteArray, QSize
from PySide6.QtGui import QColor, QIcon, QPixmap, QPainter
from PySide6.QtSvg import QSvgRenderer

from database import get_connection, write_audit_log


# ─────────────────────────────────────────────────────────────
# STATUS choices
# ─────────────────────────────────────────────────────────────

STATUS_CHOICES = [
    "ملئ",
    "إصلاح",
    "ملئ و إصلاح",
    "تدبير من القديم",
    "تكهين",
]

STATUS_COLORS = {
    "ملئ":             ("#ecfdf5", "#065f46"),
    "إصلاح":           ("#eff6ff", "#1e40af"),
    "ملئ و إصلاح":    ("#f0fdf4", "#166534"),
    "تدبير من القديم": ("#fefce8", "#854d0e"),
    "تكهين":           ("#fef2f2", "#991b1b"),
}


def ink_audit_data(record):
    """Return persisted transaction fields for an audit-log snapshot."""
    if not record:
        return {}
    keys = (
        "id", "serial_number", "trans_date", "org_unit_id", "unit_name",
        "office_id", "office_name", "quantity", "delivery_date", "status",
        "disposal_reason", "created_by", "created_at", "updated_at",
    )
    snapshot = {}
    for key in keys:
        try:
            value = record[key]
        except (KeyError, IndexError, TypeError):
            continue
        if hasattr(value, "isoformat"):
            value = value.isoformat()
        snapshot[key] = value
    return snapshot


def make_action_icon(kind):
    """إنشاء أيقونة SVG حادة وواضحة بنفس الشكل على كل أنظمة التشغيل."""
    if kind == "edit":
        svg = """<svg xmlns="http://www.w3.org/2000/svg" width="22" height="22" viewBox="0 0 24 24">
          <path d="M4 20h4L19.5 8.5a2.12 2.12 0 0 0-3-3L5 17v3Z" fill="none" stroke="#1e40af" stroke-width="1.8" stroke-linejoin="round"/>
          <path d="m14.5 7.5 2 2" fill="none" stroke="#1e40af" stroke-width="1.8" stroke-linecap="round"/>
        </svg>"""
    else:
        svg = """<svg xmlns="http://www.w3.org/2000/svg" width="22" height="22" viewBox="0 0 24 24">
          <path d="M5 7h14M10 4h4l1 3H9l1-3ZM7 7l1 13h8l1-13M10 10v7M14 10v7" fill="none" stroke="#991b1b" stroke-width="1.8" stroke-linecap="round" stroke-linejoin="round"/>
        </svg>"""

    renderer = QSvgRenderer(QByteArray(svg.encode("utf-8")))
    pixmap = QPixmap(22, 22)
    pixmap.fill(Qt.transparent)
    painter = QPainter(pixmap)
    renderer.render(painter)
    painter.end()
    return QIcon(pixmap)


# ─────────────────────────────────────────────────────────────
# DB helpers
# ─────────────────────────────────────────────────────────────

def load_org_units():
    """جلب كل الوحدات النشطة."""
    conn = get_connection()
    cur  = conn.cursor()
    cur.execute("""
        SELECT id, name, unit_type, parent_id
        FROM dbo.org_units
        WHERE is_active = 1
        ORDER BY sort_order, unit_type, name
    """)
    rows = cur.fetchall()
    conn.close()
    return [dict(r.items()) for r in rows]


def load_offices_for_unit(unit_id):
    """جلب المكاتب التابعة لوحدة معينة سواء مباشرة أو متفرعة."""
    if not unit_id:
        return []
    conn = get_connection()
    cur  = conn.cursor()
    cur.execute("""
        WITH Descendants AS (
            SELECT id, name, unit_type, parent_id, sort_order
            FROM dbo.org_units
            WHERE parent_id = ? AND is_active = 1
            UNION ALL
            SELECT child.id, child.name, child.unit_type, child.parent_id, child.sort_order
            FROM dbo.org_units child
            INNER JOIN Descendants d ON child.parent_id = d.id
            WHERE child.is_active = 1
        )
        SELECT id, name
        FROM Descendants
        WHERE unit_type = ?
        ORDER BY sort_order, name
    """, (unit_id, "مكتب"))
    rows = cur.fetchall()
    conn.close()
    return [dict(r.items()) for r in rows]


def next_serial():
    conn = get_connection()
    cur  = conn.cursor()
    try:
        cur.execute("SELECT NEXT VALUE FOR dbo.ink_serial_seq")
        val = cur.fetchone()[0]
    except Exception:
        cur.execute("SELECT ISNULL(MAX(serial_number), 0) + 1 FROM dbo.ink_transactions")
        val = cur.fetchone()[0]
    conn.close()
    return val


# ─────────────────────────────────────────────────────────────
# Add / Edit Dialog
# ─────────────────────────────────────────────────────────────

class InkTransactionDialog(QDialog):

    def __init__(self, parent=None, record=None, current_user=None, can_edit_date=True):
        super().__init__(parent)
        self.record       = record
        self.current_user = current_user or {}
        self.can_edit_date = can_edit_date
        self._org_units   = load_org_units()
        self._offices     = []

        is_edit = record is not None
        self.setWindowTitle("تعديل سجل أحبار" if is_edit else "إضافة سجل أحبار جديد")
        self.setMinimumWidth(520)
        self.setMinimumHeight(600)
        self.setLayoutDirection(Qt.RightToLeft)

        root = QVBoxLayout(self)
        root.setContentsMargins(20, 20, 20, 16)
        root.setSpacing(14)

        # Header
        header = QFrame()
        header.setObjectName("dialogHeader")
        hl = QHBoxLayout(header)
        hl.setContentsMargins(16, 12, 16, 12)

        icon_lbl = QLabel("🖊️")
        icon_lbl.setObjectName("dialogHeaderIcon")
        icon_lbl.setFixedSize(40, 40)
        icon_lbl.setAlignment(Qt.AlignCenter)
        hl.addWidget(icon_lbl)

        title_lbl = QLabel(
            "تعديل بيانات الحركة" if is_edit else "إضافة حركة أحبار جديدة"
        )
        title_lbl.setObjectName("dialogTitle")
        hl.addWidget(title_lbl)
        hl.addStretch()
        root.addWidget(header)

        # Scrollable form
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.NoFrame)

        form_widget = QWidget()
        form = QFormLayout(form_widget)
        form.setHorizontalSpacing(16)
        form.setVerticalSpacing(12)
        form.setLabelAlignment(Qt.AlignRight | Qt.AlignVCenter)

        # ── التاريخ ──────────────────────────────────────────
        self.date_edit = QDateEdit()
        self.date_edit.setCalendarPopup(True)
        self.date_edit.setDate(QDate.currentDate())
        self.date_edit.setMinimumHeight(36)
        self.date_edit.setDisplayFormat("yyyy-MM-dd")
        self.date_edit.setEnabled(self.can_edit_date)
        form.addRow("التاريخ *:", self.date_edit)

        # ── الفرع / المبنى / القسم ───────────────────────────
        self.unit_combo = QComboBox()
        self.unit_combo.setMinimumHeight(36)
        self.unit_combo.addItem("— اختر الوحدة —", None)

        # نعرض فقط الوحدات من النوع فرع / مبنى / قسم
        for u in self._org_units:
            if u["unit_type"] in ("فرع", "مبنى", "قسم"):
                label = f"[{u['unit_type']}]  {u['name']}"
                self.unit_combo.addItem(label, u["id"])

        self.unit_combo.currentIndexChanged.connect(self._on_unit_changed)
        form.addRow("الفرع / المبنى / القسم *:", self.unit_combo)

        # ── المكتب ───────────────────────────────────────────
        self.office_combo = QComboBox()
        self.office_combo.setMinimumHeight(36)
        self.office_combo.addItem("— اختر المكتب —", None)
        form.addRow("المكتب:", self.office_combo)

        self.office_required_lbl = QLabel(
            "⚠ يجب اختيار مكتب"
        )
        self.office_required_lbl.setStyleSheet(
            "color: #991b1b; font-size: 11px;"
        )
        self.office_required_lbl.setVisible(False)
        form.addRow("", self.office_required_lbl)

        # ── العدد ────────────────────────────────────────────
        self.quantity_spin = QSpinBox()
        self.quantity_spin.setRange(1, 99999)
        self.quantity_spin.setValue(1)
        self.quantity_spin.setMinimumHeight(36)
        form.addRow("العدد *:", self.quantity_spin)

        # ── تاريخ التسليم ────────────────────────────────────
        self.delivery_date = QDateEdit()
        self.delivery_date.setCalendarPopup(True)
        self.delivery_date.setDate(QDate.currentDate())
        self.delivery_date.setMinimumHeight(36)
        self.delivery_date.setDisplayFormat("yyyy-MM-dd")
        self.delivery_date.setSpecialValueText("—")
        form.addRow("تاريخ التسليم:", self.delivery_date)

        # ── الملاحظات / الحالة ───────────────────────────────
        self.status_combo = QComboBox()
        self.status_combo.setMinimumHeight(36)
        for s in STATUS_CHOICES:
            self.status_combo.addItem(s)
        self.status_combo.currentTextChanged.connect(self._on_status_changed)
        form.addRow("الملاحظات *:", self.status_combo)

        # ── سبب التكهين ──────────────────────────────────────
        self.disposal_lbl = QLabel("سبب التكهين *:")
        self.disposal_lbl.setStyleSheet("color: #991b1b; font-weight: 700;")

        self.disposal_input = QPlainTextEdit()
        self.disposal_input.setPlaceholderText("اذكر سبب التكهين …")
        self.disposal_input.setFixedHeight(80)

        form.addRow(self.disposal_lbl, self.disposal_input)

        scroll.setWidget(form_widget)
        root.addWidget(scroll, 1)

        # Buttons
        btns = QDialogButtonBox(
            QDialogButtonBox.Save | QDialogButtonBox.Cancel
        )
        btns.button(QDialogButtonBox.Save).setText("حفظ")
        btns.button(QDialogButtonBox.Cancel).setText("إلغاء")
        btns.accepted.connect(self._validate_and_accept)
        btns.rejected.connect(self.reject)
        root.addWidget(btns)

        # Initial state
        self._on_unit_changed()
        self._on_status_changed(self.status_combo.currentText())

        if is_edit:
            self._load_record()

        self.setStyleSheet("""
            #dialogHeader {
                background: #0f2338;
                border-radius: 10px;
                border-bottom: 3px solid #0891b2;
            }
            #dialogHeaderIcon { font-size: 22px; }
            #dialogTitle {
                color: white;
                font-size: 15px;
                font-weight: 800;
            }
            QLineEdit, QComboBox, QSpinBox, QDateEdit, QPlainTextEdit {
                background: white;
                border: 1px solid #cbd5e1;
                border-radius: 8px;
                padding: 5px 10px;
            }
            QLineEdit:focus, QComboBox:focus, QSpinBox:focus,
            QDateEdit:focus, QPlainTextEdit:focus {
                border: 1.5px solid #0891b2;
            }
            QPushButton {
                border-radius: 8px;
                padding: 8px 20px;
                font-weight: 700;
                font-size: 12px;
            }
        """)

    def _on_unit_changed(self, _idx=None):
        unit_id = self.unit_combo.currentData()
        self.office_combo.clear()
        self.office_required_lbl.setVisible(False)

        if not unit_id:
            self.office_combo.addItem("— اختر الوحدة أولاً —", None)
            self.office_combo.setEnabled(False)
            self._offices = []
            return

        self._offices = load_offices_for_unit(unit_id)
        if self._offices:
            self.office_combo.setEnabled(True)
            self.office_combo.addItem("— اختر المكتب —", None)
            for o in self._offices:
                self.office_combo.addItem(o["name"], o["id"])
        else:
            self.office_combo.addItem("— لا توجد مكاتب تابعة —", None)
            self.office_combo.setEnabled(False)

    def _on_status_changed(self, status):
        show = (status == "تكهين")
        self.disposal_lbl.setVisible(show)
        self.disposal_input.setVisible(show)

    def _load_record(self):
        r = self.record

        # Date
        if r.get("trans_date"):
            d = r["trans_date"]
            if hasattr(d, "year"):
                self.date_edit.setDate(QDate(d.year, d.month, d.day))

        # Unit
        for i in range(self.unit_combo.count()):
            if self.unit_combo.itemData(i) == r.get("org_unit_id"):
                self.unit_combo.setCurrentIndex(i)
                self._on_unit_changed()
                break

        # Office (populated after unit change signal)
        office_id = r.get("office_id")
        if office_id:
            for i in range(self.office_combo.count()):
                if self.office_combo.itemData(i) == office_id:
                    self.office_combo.setCurrentIndex(i)
                    break

        # Quantity
        self.quantity_spin.setValue(int(r.get("quantity") or 1))

        # Delivery date
        if r.get("delivery_date"):
            d = r["delivery_date"]
            if hasattr(d, "year"):
                self.delivery_date.setDate(QDate(d.year, d.month, d.day))

        # Status
        idx = self.status_combo.findText(r.get("status") or "")
        if idx >= 0:
            self.status_combo.setCurrentIndex(idx)

        # Disposal reason
        self.disposal_input.setPlainText(r.get("disposal_reason") or "")

    def _validate_and_accept(self):
        if self.unit_combo.currentData() is None:
            QMessageBox.warning(
                self, "تحقق",
                "يجب اختيار الفرع / المبنى / القسم."
            )
            return

        has_offices = self.office_combo.count() > 1
        if has_offices and self.office_combo.currentData() is None:
            self.office_required_lbl.setVisible(True)
            QMessageBox.warning(
                self, "تحقق",
                "يجب اختيار المكتب."
            )
            return
        self.office_required_lbl.setVisible(False)

        if self.status_combo.currentText() == "تكهين":
            reason = self.disposal_input.toPlainText().strip()
            if not reason:
                QMessageBox.warning(
                    self, "تحقق",
                    "يجب إدخال سبب التكهين."
                )
                self.disposal_input.setFocus()
                return

        self.accept()

    def get_data(self):
        status  = self.status_combo.currentText()
        d_date  = self.delivery_date.date()
        t_date  = self.date_edit.date()

        return {
            "trans_date":     f"{t_date.year()}-{t_date.month():02d}-{t_date.day():02d}",
            "org_unit_id":    self.unit_combo.currentData(),
            "office_id":      self.office_combo.currentData(),
            "quantity":       self.quantity_spin.value(),
            "delivery_date":  f"{d_date.year()}-{d_date.month():02d}-{d_date.day():02d}",
            "status":         status,
            "disposal_reason": (
                self.disposal_input.toPlainText().strip()
                if status == "تكهين" else None
            ),
        }


# ─────────────────────────────────────────────────────────────
# View Dialog (read-only)
# ─────────────────────────────────────────────────────────────

class InkViewDialog(QDialog):

    def __init__(self, parent=None, record=None):
        super().__init__(parent)
        self.setWindowTitle("تفاصيل سجل الأحبار")
        self.setMinimumWidth(480)
        self.setLayoutDirection(Qt.RightToLeft)

        root = QVBoxLayout(self)
        root.setContentsMargins(20, 20, 20, 16)
        root.setSpacing(12)

        header = QFrame()
        header.setObjectName("dialogHeader")
        hl = QHBoxLayout(header)
        hl.setContentsMargins(16, 12, 16, 12)
        icon_lbl = QLabel("🖊️")
        icon_lbl.setObjectName("dialogHeaderIcon")
        icon_lbl.setFixedSize(40, 40)
        icon_lbl.setAlignment(Qt.AlignCenter)
        hl.addWidget(icon_lbl)
        title_lbl = QLabel(f"سجل أحبار رقم {record['serial_number']}")
        title_lbl.setObjectName("dialogTitle")
        hl.addWidget(title_lbl)
        hl.addStretch()
        root.addWidget(header)

        if record:
            fields = [
                ("م",                       str(record.get("serial_number", ""))),
                ("التاريخ",                 str(record.get("trans_date", ""))),
                ("الوحدة",                  str(record.get("unit_name", ""))),
                ("المكتب",                  str(record.get("office_name") or "—")),
                ("العدد",                   str(record.get("quantity", ""))),
                ("تاريخ التسليم",           str(record.get("delivery_date") or "—")),
                ("الملاحظات",               str(record.get("status", ""))),
                ("سبب التكهين",             str(record.get("disposal_reason") or "—")),
                ("بواسطة",                  str(record.get("created_by_name") or "—")),
                ("تاريخ الإنشاء",           str(record.get("created_at", ""))[:19]),
            ]

            form = QFormLayout()
            form.setHorizontalSpacing(20)
            form.setVerticalSpacing(10)
            form.setLabelAlignment(Qt.AlignRight | Qt.AlignVCenter)

            for label, value in fields:
                lbl = QLabel(f"{label}:")
                lbl.setStyleSheet("font-weight: 700; color: #334155;")
                val = QLabel(value)
                val.setStyleSheet("color: #0f2338;")
                val.setWordWrap(True)
                form.addRow(lbl, val)

            root.addLayout(form)

        btn_close = QPushButton("إغلاق")
        btn_close.setObjectName("primaryButton")
        btn_close.clicked.connect(self.accept)
        root.addWidget(btn_close)

        self.setStyleSheet("""
            #dialogHeader {
                background: #0f2338;
                border-radius: 10px;
                border-bottom: 3px solid #0891b2;
            }
            #dialogHeaderIcon { font-size: 22px; }
            #dialogTitle { color: white; font-size: 15px; font-weight: 800; }
            QPushButton#primaryButton {
                background: #0891b2; color: white;
                border: none; border-radius: 9px;
                padding: 8px 18px; font-weight: 700;
            }
            QPushButton#primaryButton:hover { background: #0e7490; }
        """)


# ─────────────────────────────────────────────────────────────
# Main Page
# ─────────────────────────────────────────────────────────────

class InkPage(QWidget):

    PAGE_SIZE = 50

    def __init__(self, current_user=None, parent=None):
        super().__init__(parent)
        self.current_user = current_user or {}
        self.setLayoutDirection(Qt.RightToLeft)
        self.setObjectName("inkPage")
        self._page   = 0
        self._total  = 0
        self._org_units_all = []
        self._build_ui()
        self.load_transactions()

    # ── Build UI ─────────────────────────────────────────────

    def _build_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(28, 24, 28, 24)
        layout.setSpacing(16)

        # Title row
        title_row = QHBoxLayout()
        title_col = QVBoxLayout()
        title_lbl = QLabel("الأحبار")
        title_lbl.setObjectName("pageTitle")
        sub_lbl   = QLabel("إدارة ومتابعة حركات الأحبار")
        sub_lbl.setObjectName("pageSubtitle")
        title_col.addWidget(title_lbl)
        title_col.addWidget(sub_lbl)
        title_row.addLayout(title_col)
        title_row.addStretch()

        self.btn_export = QPushButton("تصدير Excel")
        self.btn_export.setObjectName("secondaryButton")
        self.btn_export.setMinimumHeight(38)
        self.btn_export.setToolTip(
            "تصدير جميع الحركات المطابقة للفلاتر الحالية، وليس الصفحة الحالية فقط"
        )
        self.btn_export.clicked.connect(self._export_excel)
        title_row.addWidget(self.btn_export)

        self.btn_add = QPushButton("＋  إضافة حركة")
        self.btn_add.setObjectName("primaryButton")
        self.btn_add.setMinimumHeight(38)
        self.btn_add.clicked.connect(self._add_record)
        title_row.addWidget(self.btn_add)
        layout.addLayout(title_row)

        # ── Filters ──────────────────────────────────────────
        filter_card = QFrame()
        filter_card.setObjectName("filterCard")
        fl = QVBoxLayout(filter_card)
        fl.setContentsMargins(16, 12, 16, 12)
        fl.setSpacing(10)

        filter_lbl = QLabel("🔍  البحث والتصفية")
        filter_lbl.setStyleSheet("font-weight: 800; color: #0f2338; font-size: 13px;")
        fl.addWidget(filter_lbl)

        row1 = QHBoxLayout()
        row1.setSpacing(10)

        self.f_date_from = QDateEdit()
        self.f_date_from.setCalendarPopup(True)
        self.f_date_from.setDate(QDate.currentDate().addDays(-30))
        self.f_date_from.setDisplayFormat("yyyy-MM-dd")
        self.f_date_from.setMinimumHeight(32)
        self.f_date_from.setSpecialValueText("—")

        self.f_date_to = QDateEdit()
        self.f_date_to.setCalendarPopup(True)
        self.f_date_to.setDate(QDate.currentDate())
        self.f_date_to.setDisplayFormat("yyyy-MM-dd")
        self.f_date_to.setMinimumHeight(32)

        self.f_status = QComboBox()
        self.f_status.setMinimumHeight(32)
        self.f_status.addItem("كل الملاحظات", "")
        for s in STATUS_CHOICES:
            self.f_status.addItem(s, s)

        row1.addWidget(QLabel("من:"))
        row1.addWidget(self.f_date_from, 1)
        row1.addWidget(QLabel("إلى:"))
        row1.addWidget(self.f_date_to, 1)
        row1.addWidget(QLabel("الملاحظة:"))
        row1.addWidget(self.f_status, 1)
        fl.addLayout(row1)

        row2 = QHBoxLayout()
        row2.setSpacing(10)

        self.f_unit = QComboBox()
        self.f_unit.setMinimumHeight(32)
        self.f_unit.addItem("كل الوحدات", None)

        self.f_search = QLineEdit()
        self.f_search.setPlaceholderText("بحث بالوحدة أو المكتب …")
        self.f_search.setMinimumHeight(32)

        btn_search = QPushButton("بحث")
        btn_search.setObjectName("primaryButton")
        btn_search.setMinimumHeight(32)
        btn_search.clicked.connect(self._do_search)

        btn_reset = QPushButton("إعادة تعيين")
        btn_reset.setObjectName("secondaryButton")
        btn_reset.setMinimumHeight(32)
        btn_reset.clicked.connect(self._reset_filters)

        row2.addWidget(QLabel("الوحدة:"))
        row2.addWidget(self.f_unit, 1)
        row2.addWidget(self.f_search, 2)
        row2.addWidget(btn_search)
        row2.addWidget(btn_reset)
        fl.addLayout(row2)

        layout.addWidget(filter_card)

        # ── Table ─────────────────────────────────────────────
        self.table = QTableWidget()
        self.table.setColumnCount(9)
        self.table.setHorizontalHeaderLabels([
            "م", "التاريخ", "الفرع / المبنى / القسم",
            "المكتب", "العدد", "تاريخ التسليم",
            "الملاحظات", "عرض", "أدوات",
        ])
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.verticalHeader().setVisible(False)
        self.table.setAlternatingRowColors(True)

        h = self.table.horizontalHeader()
        h.setSectionResizeMode(0, QHeaderView.ResizeToContents)
        h.setSectionResizeMode(1, QHeaderView.ResizeToContents)
        h.setSectionResizeMode(2, QHeaderView.Stretch)
        h.setSectionResizeMode(3, QHeaderView.ResizeToContents)
        h.setSectionResizeMode(4, QHeaderView.ResizeToContents)
        h.setSectionResizeMode(5, QHeaderView.ResizeToContents)
        h.setSectionResizeMode(6, QHeaderView.ResizeToContents)
        h.setSectionResizeMode(7, QHeaderView.ResizeToContents)
        # عمود أيقونات الإجراءات بعرض ثابت ومضغوط.
        h.setSectionResizeMode(8, QHeaderView.Fixed)
        h.resizeSection(8, 96)

        layout.addWidget(self.table)

        # ── Pagination ───────────────────────────────────────
        pag_row = QHBoxLayout()

        self.stats_lbl = QLabel("")
        self.stats_lbl.setStyleSheet("color: #64748b; font-size: 11px;")
        pag_row.addWidget(self.stats_lbl)
        pag_row.addStretch()

        self.btn_prev = QPushButton("◀  السابق")
        self.btn_prev.setObjectName("secondaryButton")
        self.btn_prev.setMinimumHeight(32)
        self.btn_prev.clicked.connect(self._prev_page)

        self.page_lbl = QLabel("صفحة 1")
        self.page_lbl.setAlignment(Qt.AlignCenter)
        self.page_lbl.setStyleSheet("font-weight: 700; min-width: 80px;")

        self.btn_next = QPushButton("التالي  ▶")
        self.btn_next.setObjectName("secondaryButton")
        self.btn_next.setMinimumHeight(32)
        self.btn_next.clicked.connect(self._next_page)

        pag_row.addWidget(self.btn_prev)
        pag_row.addWidget(self.page_lbl)
        pag_row.addWidget(self.btn_next)

        layout.addLayout(pag_row)

        self._apply_style()
        self._load_filter_units()

    def _apply_style(self):
        self.setStyleSheet("""
            #inkPage { background: #f4f7fb; }
            #pageTitle  { color: #0f2338; font-size: 22px; font-weight: 800; }
            #pageSubtitle { color: #64748b; font-size: 12px; }
            #filterCard {
                background: white;
                border: 1px solid #e2e8f0;
                border-radius: 12px;
            }
            QTableWidget {
                background: white;
                border: 1px solid #e2e8f0;
                border-radius: 12px;
                gridline-color: #f1f5f9;
                selection-background-color: #e0f2fe;
                selection-color: #0f2338;
                alternate-background-color: #f8fafc;
            }
            QTableWidget::item { padding: 6px 8px; }
            QHeaderView::section {
                background: #f8fafc;
                color: #0f2338;
                padding: 9px 8px;
                border: none;
                border-bottom: 2px solid #0891b2;
                font-weight: 700;
            }
            QLineEdit, QComboBox, QDateEdit {
                background: white;
                border: 1px solid #cbd5e1;
                border-radius: 8px;
                padding: 4px 10px;
            }
            QLineEdit:focus, QComboBox:focus, QDateEdit:focus {
                border: 1.5px solid #0891b2;
            }
            QPushButton#primaryButton {
                background: #0891b2; color: white;
                border: none; border-radius: 9px;
                padding: 7px 16px; font-weight: 700;
            }
            QPushButton#primaryButton:hover { background: #0e7490; }
            QPushButton#secondaryButton {
                background: #f1f5f9; color: #0f2338;
                border: 1px solid #e2e8f0; border-radius: 9px;
                padding: 7px 16px; font-weight: 700;
            }
            QPushButton#secondaryButton:hover { background: #e2e8f0; }
            QPushButton#viewBtn {
                background: #f0fdf4; color: #166534;
                border: 1px solid #bbf7d0; border-radius: 7px;
                padding: 3px 10px; font-size: 11px; font-weight: 700;
            }
            QPushButton#viewBtn:hover { background: #dcfce7; }
            QPushButton#editBtn {
                background: #eff6ff;
                border: 1px solid #bfdbfe;
                border-radius: 8px;
                padding: 0px;
            }
            QPushButton#editBtn:hover {
                background: #dbeafe;
                border-color: #93c5fd;
            }
            QPushButton#deleteBtn {
                background: #fef2f2;
                border: 1px solid #fecaca;
                border-radius: 8px;
                padding: 0px;
            }
            QPushButton#deleteBtn:hover {
                background: #fee2e2;
                border-color: #fca5a5;
            }
            QPushButton#editBtn:pressed { background: #dbeafe; }
            QPushButton#deleteBtn:pressed { background: #fecaca; }
        """)

    # ── Load filter units ────────────────────────────────────

    def _load_filter_units(self):
        try:
            self._org_units_all = load_org_units()
        except Exception:
            self._org_units_all = []

        self.f_unit.clear()
        self.f_unit.addItem("كل الوحدات", None)
        for u in self._org_units_all:
            if u["unit_type"] in ("فرع", "مبنى", "قسم"):
                self.f_unit.addItem(f"[{u['unit_type']}] {u['name']}", u["id"])

    # ── Load Transactions ────────────────────────────────────

    def load_transactions(self):
        self._load_filter_units()
        if hasattr(self, "btn_add"):
            self.btn_add.setVisible(self._has_perm("ink.add"))
        self._do_search()

    def _do_search(self):
        self._page = 0
        self._fetch_and_display()

    def _reset_filters(self):
        self.f_date_from.setDate(QDate.currentDate().addDays(-30))
        self.f_date_to.setDate(QDate.currentDate())
        self.f_status.setCurrentIndex(0)
        self.f_unit.setCurrentIndex(0)
        self.f_search.clear()
        self._page = 0
        self._fetch_and_display()

    def _prev_page(self):
        if self._page > 0:
            self._page -= 1
            self._fetch_and_display()

    def _next_page(self):
        max_page = max(0, (self._total - 1) // self.PAGE_SIZE)
        if self._page < max_page:
            self._page += 1
            self._fetch_and_display()

    def _build_filtered_query(self):
        date_from = self.f_date_from.date()
        date_to = self.f_date_to.date()
        status = self.f_status.currentData()
        unit_id = self.f_unit.currentData()
        search = self.f_search.text().strip()

        where_clauses = ["1=1"]
        params = []
        where_clauses.append("it.trans_date >= ?")
        params.append(
            f"{date_from.year()}-{date_from.month():02d}-{date_from.day():02d}"
        )
        where_clauses.append("it.trans_date <= ?")
        params.append(
            f"{date_to.year()}-{date_to.month():02d}-{date_to.day():02d}"
        )

        if status:
            where_clauses.append("it.status = ?")
            params.append(status)
        if unit_id:
            where_clauses.append("it.org_unit_id = ?")
            params.append(unit_id)
        if search:
            where_clauses.append("(ou.name LIKE ? OR ofc.name LIKE ?)")
            like = f"%{search}%"
            params.extend([like, like])

        where_sql = " AND ".join(where_clauses)
        base_query = f"""
            SELECT
                it.id, it.serial_number, it.trans_date,
                it.org_unit_id, ou.name AS unit_name,
                it.office_id, ofc.name AS office_name,
                it.quantity, it.delivery_date,
                it.status, it.disposal_reason,
                it.created_by, u.full_name AS created_by_name,
                it.created_at, it.updated_at
            FROM dbo.ink_transactions it
            LEFT JOIN dbo.org_units ou ON ou.id = it.org_unit_id
            LEFT JOIN dbo.org_units ofc ON ofc.id = it.office_id
            LEFT JOIN dbo.users u ON u.id = it.created_by
            WHERE {where_sql}
        """
        return base_query, params

    def _fetch_and_display(self):
        base_query, params = self._build_filtered_query()
        try:
            conn = get_connection()
            cursor = conn.cursor()
            cursor.execute(f"SELECT COUNT(*) FROM ({base_query}) AS cnt", params)
            self._total = cursor.fetchone()[0]

            offset = self._page * self.PAGE_SIZE
            cursor.execute(
                f"{base_query} ORDER BY it.serial_number DESC "
                f"OFFSET ? ROWS FETCH NEXT ? ROWS ONLY",
                params + [offset, self.PAGE_SIZE],
            )
            rows = cursor.fetchall()
            conn.close()
        except Exception as e:
            QMessageBox.critical(self, "خطأ", f"تعذر تحميل البيانات:\n{e}")
            return

        self._display_rows(rows)

    def _export_excel(self):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        except ImportError:
            QMessageBox.critical(
                self,
                "مكتبة غير متاحة",
                "لتصدير Excel، ثبّت مكتبة openpyxl:\n\npython -m pip install openpyxl",
            )
            return

        conn = None
        try:
            base_query, params = self._build_filtered_query()
            conn = get_connection()
            cursor = conn.cursor()
            cursor.execute(
                f"{base_query} ORDER BY it.serial_number DESC",
                params,
            )
            rows = cursor.fetchall()
        except Exception as error:
            QMessageBox.critical(
                self, "خطأ", f"تعذر تجهيز السجلات للتصدير:\n{error}"
            )
            return
        finally:
            if conn is not None:
                conn.close()

        if not rows:
            QMessageBox.information(
                self, "لا توجد بيانات", "لا توجد حركات مطابقة للفلاتر الحالية للتصدير."
            )
            return

        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "تصدير حركات الأحبار إلى Excel",
            "حركات_الأحبار.xlsx",
            "Excel Workbook (*.xlsx)",
        )
        if not file_path:
            return
        if not file_path.lower().endswith(".xlsx"):
            file_path += ".xlsx"

        try:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "حركات الأحبار"
            sheet.sheet_view.rightToLeft = True
            sheet.sheet_view.showGridLines = False
            sheet.sheet_view.zoomScale = 90

            sheet.merge_cells("A1:G1")
            title = sheet["A1"]
            title.value = "سجل حركات الأحبار"
            title.font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
            title.fill = PatternFill("solid", fgColor="10263D")
            title.alignment = Alignment(horizontal="center", vertical="center")

            sheet.merge_cells("A2:G2")
            summary = sheet["A2"]
            summary.value = f"إجمالي الحركات المطابقة للفلاتر: {len(rows)}"
            summary.font = Font(name="Arial", size=11, bold=True, color="64748B")
            summary.alignment = Alignment(horizontal="center", vertical="center")

            headers = [
                "م", "التاريخ", "الفرع / المبنى / القسم", "المكتب",
                "العدد", "تاريخ التسليم", "الملاحظات",
            ]
            header_row = 3
            for column, value in enumerate(headers, 1):
                cell = sheet.cell(header_row, column, value)
                cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1095B0")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            thin_gray = Side(style="thin", color="D9DEE5")
            for row_index, record in enumerate(rows, header_row + 1):
                trans_date = record["trans_date"]
                delivery_date = record["delivery_date"]
                values = [
                    record["serial_number"],
                    trans_date.strftime("%Y-%m-%d") if hasattr(trans_date, "strftime") else str(trans_date)[:10],
                    record["unit_name"] or "",
                    record["office_name"] or "—",
                    record["quantity"],
                    delivery_date.strftime("%Y-%m-%d") if delivery_date and hasattr(delivery_date, "strftime") else (str(delivery_date)[:10] if delivery_date else "—"),
                    record["status"] or "",
                ]
                for column, value in enumerate(values, 1):
                    cell = sheet.cell(row_index, column, value)
                    cell.font = Font(name="Arial", size=10, color="334155")
                    cell.alignment = Alignment(
                        horizontal="right" if column in (3, 4, 7) else "center",
                        vertical="center",
                        wrap_text=(column in (3, 4)),
                    )
                    cell.border = Border(bottom=thin_gray)
                    if column == 7:
                        bg, fg = STATUS_COLORS.get(
                            record["status"] or "", ("#f8fafc", "#334155")
                        )
                        cell.fill = PatternFill("solid", fgColor=bg.lstrip("#"))
                        cell.font = Font(name="Arial", size=10, bold=True, color=fg.lstrip("#"))
                sheet.row_dimensions[row_index].height = 23

            for column, width in enumerate((10, 15, 32, 28, 12, 17, 24), 1):
                sheet.column_dimensions[chr(64 + column)].width = width
            sheet.row_dimensions[1].height = 32
            sheet.row_dimensions[2].height = 24
            sheet.row_dimensions[header_row].height = 30
            sheet.freeze_panes = "A4"
            sheet.auto_filter.ref = f"A{header_row}:G{header_row + len(rows)}"
            sheet.page_setup.orientation = "landscape"
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.print_title_rows = "1:3"

            workbook.save(file_path)
            workbook.close()

            export_filters = {
                "date_from": self.f_date_from.date().toString("yyyy-MM-dd"),
                "date_to": self.f_date_to.date().toString("yyyy-MM-dd"),
                "status": self.f_status.currentData() or "",
                "unit_id": self.f_unit.currentData() or "",
                "search": self.f_search.text().strip(),
            }
            write_audit_log(
                user=self.current_user,
                action="تصدير",
                module_name="الأحبار",
                description=(
                    f"تصدير {len(rows)} حركة أحبار إلى Excel "
                    f"({Path(file_path).name})"
                ),
                new_values={
                    "file": Path(file_path).name,
                    "record_count": len(rows),
                    "filters": export_filters,
                },
            )
            QMessageBox.information(
                self,
                "تم التصدير",
                f"تم تصدير {len(rows)} حركة مطابقة للفلاتر إلى Excel بنجاح.",
            )
        except Exception as error:
            QMessageBox.critical(
                self, "خطأ في التصدير", f"تعذر إنشاء ملف Excel:\n{error}"
            )

    def _display_rows(self, rows):
        self.table.setRowCount(0)

        for i, r in enumerate(rows):
            self.table.insertRow(i)

            row_dict = dict(r.items())

            # م
            sn = QTableWidgetItem(str(r["serial_number"]))
            sn.setTextAlignment(Qt.AlignCenter)
            self.table.setItem(i, 0, sn)

            # التاريخ
            self.table.setItem(
                i, 1,
                QTableWidgetItem(str(r["trans_date"])[:10])
            )

            # الوحدة
            self.table.setItem(i, 2, QTableWidgetItem(r["unit_name"] or ""))

            # المكتب
            self.table.setItem(
                i, 3,
                QTableWidgetItem(r["office_name"] or "—")
            )

            # العدد
            qty = QTableWidgetItem(str(r["quantity"]))
            qty.setTextAlignment(Qt.AlignCenter)
            self.table.setItem(i, 4, qty)

            # تاريخ التسليم
            self.table.setItem(
                i, 5,
                QTableWidgetItem(
                    str(r["delivery_date"])[:10]
                    if r["delivery_date"] else "—"
                )
            )

            # الملاحظات (badge)
            status      = r["status"] or ""
            status_item = QTableWidgetItem(status)
            status_item.setTextAlignment(Qt.AlignCenter)
            bg, fg = STATUS_COLORS.get(status, ("#f8fafc", "#334155"))
            status_item.setBackground(QColor(bg))
            status_item.setForeground(QColor(fg))
            self.table.setItem(i, 6, status_item)

            # زر عرض
            btn_view = QPushButton("عرض")
            btn_view.setObjectName("viewBtn")
            btn_view.setCursor(Qt.PointingHandCursor)
            btn_view.clicked.connect(
                lambda checked=False, rd=row_dict: self._view_record(rd)
            )
            self.table.setCellWidget(i, 7, self._wrap(btn_view))

            # أدوات: أيقونات فقط لتفادي قص النصوص داخل الجدول.
            tools_widget = QWidget()
            tl = QHBoxLayout(tools_widget)
            # مساحة رأسية إضافية تمنع قص أسفل الأيقونات داخل خلية الجدول.
            tools_widget.setMinimumHeight(52)
            tl.setContentsMargins(4, 7, 4, 7)
            tl.setSpacing(6)
            tl.setAlignment(Qt.AlignCenter)

            btn_edit = QPushButton()
            btn_edit.setObjectName("editBtn")
            btn_edit.setIcon(make_action_icon("edit"))
            btn_edit.setIconSize(QSize(20, 20))
            btn_edit.setFixedSize(34, 34)
            btn_edit.setToolTip("تعديل السجل")
            btn_edit.setAccessibleName("تعديل السجل")
            btn_edit.setCursor(Qt.PointingHandCursor)
            btn_edit.clicked.connect(
                lambda checked=False, rd=row_dict: self._edit_record(rd)
            )
            btn_edit.setVisible(self._has_perm("ink.edit"))
            tl.addWidget(btn_edit)

            btn_del = QPushButton()
            btn_del.setObjectName("deleteBtn")
            btn_del.setIcon(make_action_icon("delete"))
            btn_del.setIconSize(QSize(20, 20))
            btn_del.setFixedSize(34, 34)
            btn_del.setToolTip("حذف السجل")
            btn_del.setAccessibleName("حذف السجل")
            btn_del.setCursor(Qt.PointingHandCursor)
            btn_del.clicked.connect(
                lambda checked=False, rd=row_dict: self._delete_record(rd)
            )
            btn_del.setVisible(self._has_perm("ink.delete"))
            tl.addWidget(btn_del)

            self.table.setCellWidget(i, 8, tools_widget)
            self.table.setRowHeight(i, 58)

        # Pagination labels
        total_pages = max(1, (self._total + self.PAGE_SIZE - 1) // self.PAGE_SIZE)
        self.page_lbl.setText(f"صفحة {self._page + 1} / {total_pages}")
        self.btn_prev.setEnabled(self._page > 0)
        self.btn_next.setEnabled(self._page < total_pages - 1)
        self.stats_lbl.setText(
            f"إجمالي السجلات: {self._total}  |  "
            f"معروض: {len(rows)}"
        )

    def _wrap(self, widget):
        w = QWidget()
        l = QHBoxLayout(w)
        l.setContentsMargins(4, 2, 4, 2)
        l.addWidget(widget)
        return w

    # ── Permission helper ─────────────────────────────────────

    def _has_perm(self, code):
        role = str(
            self.current_user.get("role_code", "")
        ).strip().upper()
        if role == "SUPER_ADMIN":
            return True
        return bool(
            self.current_user
            .get("permissions", {})
            .get(code, False)
        )

    # ── Actions ───────────────────────────────────────────────

    def _add_record(self):
        if not self._has_perm("ink.add"):
            QMessageBox.warning(self, "صلاحيات", "ليس لديك صلاحية إضافة سجلات أحبار.")
            return

        role = str(self.current_user.get("role_code", "")).strip().upper()
        can_edit_date = role in ("SUPER_ADMIN", "ADMIN")

        dlg = InkTransactionDialog(
            self,
            current_user=self.current_user,
            can_edit_date=can_edit_date,
        )
        if dlg.exec() != QDialog.Accepted:
            return

        data   = dlg.get_data()
        serial = next_serial()

        try:
            conn   = get_connection()
            cursor = conn.cursor()
            cursor.execute(
                """
                INSERT INTO dbo.ink_transactions
                (serial_number, trans_date, org_unit_id, office_id,
                 quantity, delivery_date, status, disposal_reason,
                 created_by, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, SYSDATETIME(), SYSDATETIME())
                """,
                (
                    serial,
                    data["trans_date"],
                    data["org_unit_id"],
                    data["office_id"],
                    data["quantity"],
                    data["delivery_date"],
                    data["status"],
                    data["disposal_reason"],
                    self.current_user.get("id"),
                )
            )
            cursor.execute(
                """
                SELECT TOP 1 id
                FROM dbo.ink_transactions
                WHERE serial_number = ?
                ORDER BY id DESC
                """,
                (serial,),
            )
            inserted = cursor.fetchone()
            record_id = inserted["id"] if inserted else serial
            conn.commit()
            conn.close()

            audit_new_values = dict(data)
            audit_new_values.update({"id": record_id, "serial_number": serial})
            write_audit_log(
                user=self.current_user,
                action="إضافة",
                module_name="الأحبار",
                record_id=record_id,
                description=f"إضافة سجل أحبار رقم {serial}",
                new_values=audit_new_values,
            )

            self.load_transactions()

        except Exception as e:
            QMessageBox.critical(self, "خطأ", f"تعذر الحفظ:\n{e}")

    def _view_record(self, record):
        dlg = InkViewDialog(self, record=record)
        dlg.exec()

    def _edit_record(self, record):
        if not self._has_perm("ink.edit"):
            QMessageBox.warning(self, "صلاحيات", "ليس لديك صلاحية تعديل سجلات الأحبار.")
            return

        role = str(self.current_user.get("role_code", "")).strip().upper()
        can_edit_date = role in ("SUPER_ADMIN", "ADMIN")

        dlg = InkTransactionDialog(
            self,
            record=record,
            current_user=self.current_user,
            can_edit_date=can_edit_date,
        )
        if dlg.exec() != QDialog.Accepted:
            return

        data = dlg.get_data()
        old_values = ink_audit_data(record)

        try:
            conn   = get_connection()
            cursor = conn.cursor()
            cursor.execute(
                """
                UPDATE dbo.ink_transactions
                SET trans_date     = ?,
                    org_unit_id    = ?,
                    office_id      = ?,
                    quantity       = ?,
                    delivery_date  = ?,
                    status         = ?,
                    disposal_reason = ?,
                    updated_at     = SYSDATETIME()
                WHERE id = ?
                """,
                (
                    data["trans_date"],
                    data["org_unit_id"],
                    data["office_id"],
                    data["quantity"],
                    data["delivery_date"],
                    data["status"],
                    data["disposal_reason"],
                    record["id"],
                )
            )
            conn.commit()
            conn.close()

            audit_new_values = dict(data)
            audit_new_values.update({
                "id": record["id"],
                "serial_number": record["serial_number"],
            })
            write_audit_log(
                user=self.current_user,
                action="تعديل",
                module_name="الأحبار",
                record_id=record["id"],
                description=f"تعديل سجل أحبار رقم {record['serial_number']}",
                old_values=old_values,
                new_values=audit_new_values,
            )

            self.load_transactions()

        except Exception as e:
            QMessageBox.critical(self, "خطأ", f"تعذر التعديل:\n{e}")

    def _delete_record(self, record):
        if not self._has_perm("ink.delete"):
            QMessageBox.warning(self, "صلاحيات", "ليس لديك صلاحية حذف سجلات الأحبار.")
            return

        reply = QMessageBox.question(
            self,
            "تأكيد الحذف",
            f"هل تريد حذف سجل الأحبار رقم {record['serial_number']} نهائياً؟",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No,
        )
        if reply != QMessageBox.Yes:
            return

        try:
            conn   = get_connection()
            cursor = conn.cursor()
            cursor.execute(
                "DELETE FROM dbo.ink_transactions WHERE id = ?",
                (record["id"],)
            )
            conn.commit()
            conn.close()

            write_audit_log(
                user=self.current_user,
                action="حذف",
                module_name="الأحبار",
                record_id=record["id"],
                description=f"حذف سجل أحبار رقم {record['serial_number']}",
                old_values=ink_audit_data(record),
                new_values={"deleted": True},
            )

            self.load_transactions()

        except Exception as e:
            QMessageBox.critical(self, "خطأ", f"تعذر الحذف:\n{e}")
