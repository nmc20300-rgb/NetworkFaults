"""
إدارة الكاميرات والسويتشات وبلاغات الأعطال — داخل نظام «الشبكة».
نسخة مبسطة:
  - لوحة مؤشرات سريعة (إجمالي الكاميرات / تعمل / متوقفة / السويتشات / بلاغات مفتوحة).
  - تبويب الكاميرات مع البحث والفلاتر.
  - تبويب السويتشات مع ربط الكاميرات وأرقام المنافذ.
  - تبويب بلاغات الأعطال مع صرف قطع الغيار.
  - متوافق مع قاعدة البيانات والصلاحيات وسجل العمليات.
"""

from datetime import datetime, date
import html
import ipaddress
from pathlib import Path
import tempfile
import webbrowser

from shiboken6 import isValid as qt_is_valid
from PySide6.QtCore import Qt, QDate, QTimer, QRegularExpression
from PySide6.QtGui import QColor, QRegularExpressionValidator
from PySide6.QtWidgets import (
    QWidget,
    QVBoxLayout,
    QHBoxLayout,
    QFormLayout,
    QLabel,
    QPushButton,
    QLineEdit,
    QTableWidget,
    QTableWidgetItem,
    QMessageBox,
    QDialog,
    QComboBox,
    QTextEdit,
    QDialogButtonBox,
    QHeaderView,
    QFrame,
    QDateEdit,
    QSpinBox,
    QTabWidget,
    QStackedWidget,
    QCheckBox,
    QFileDialog,
    QMenu,
    QScrollArea,
)

from database import get_connection, write_audit_log
from ui.org_cache import OrgStructureCache
from ui.theme import Theme
from ui.faults_page import (
    parse_datetime,
    format_datetime,
    is_unset,
    clear_datetime,
    keep_calendar_on_today,
)
from ui.switches_seed import seed_network_switches, link_cameras_to_switches


# =============================================================
# ثوابت وإعدادات النظام
# =============================================================

CAMERA_STATUSES = ["تعمل", "تحت الصيانة", "متوقفة", "احتياطي", "مستبعدة"]

PLACEMENTS = ["داخلية", "خارجية"]

CAMERA_STATUS_COLORS = {
    "تعمل": ("#065f46", "#ecfdf5", "#a7f3d0"),
    "تحت الصيانة": ("#92400e", "#fffbeb", "#fde68a"),
    "متوقفة": ("#991b1b", "#fef2f2", "#fecaca"),
    "احتياطي": ("#075985", "#f0f9ff", "#bae6fd"),
    "مستبعدة": ("#334155", "#f1f5f9", "#cbd5e1"),
}

CAMERA_TYPES = [
    "ثابتة", "متحركة (PTZ)", "Dome", "Bullet", "Fisheye", "Box", "أخرى",
]

# ثوابت السويتشات
SWITCH_ROLES = [
    "كاميرات مراقبة",
    "شبكة",
    "شبكة وكاميرات",
    "شبكة عسكرية",
    "ربط لاسلكي",
    "أخرى",
]

SWITCH_ROLE_COLORS = {
    "كاميرات مراقبة": ("#065f46", "#ecfdf5", "#a7f3d0"),
    "شبكة وكاميرات": ("#0e7490", "#ecfeff", "#a5f3fc"),
    "شبكة": ("#1e40af", "#eff6ff", "#bfdbfe"),
    "شبكة عسكرية": ("#854d0e", "#fefce8", "#fef08a"),
    "ربط لاسلكي": ("#6b21a8", "#faf5ff", "#e9d5ff"),
    "أخرى": ("#334155", "#f1f5f9", "#cbd5e1"),
}

SWITCH_STATUSES = ["يعمل", "تحت الصيانة", "متوقف"]

STATUS_NEW = "جديد"
STATUS_REPAIRED = "تم الإصلاح"
STATUS_CLOSED = "مغلق"
CLOSED_STATUSES = (STATUS_REPAIRED, STATUS_CLOSED)

FAULT_STATUSES = [
    STATUS_NEW, "قيد الفحص", "قيد الإصلاح",
    "بانتظار قطعة غيار", STATUS_REPAIRED, STATUS_CLOSED,
]

FAULT_STATUS_COLORS = {
    STATUS_NEW: ("#1e40af", "#eff6ff", "#bfdbfe"),
    "قيد الفحص": ("#5b21b6", "#f5f3ff", "#ddd6fe"),
    "قيد الإصلاح": ("#9a3412", "#fff7ed", "#ffedd5"),
    "بانتظار قطعة غيار": ("#86198f", "#fdf4ff", "#f5d0fe"),
    STATUS_REPAIRED: ("#065f46", "#ecfdf5", "#a7f3d0"),
    STATUS_CLOSED: ("#334155", "#f1f5f9", "#cbd5e1"),
}

PRIORITIES = ["عادي", "متوسط", "عاجل", "حرج"]

PRIORITY_COLORS = {
    "عادي": ("#475569", "#f1f5f9", "#cbd5e1"),
    "متوسط": ("#92400e", "#fffbeb", "#fde68a"),
    "عاجل": ("#9a3412", "#fff7ed", "#fed7aa"),
    "حرج": ("#991b1b", "#fef2f2", "#fecaca"),
}

# حد التنبيه عند تسجيل بلاغ جديد لكاميرا كثيرة الأعطال
FREQUENT_FAULTS_MIN = 3
FREQUENT_FAULTS_DAYS = 90


# =============================================================
# تصميم الواجهة وأنماط CSS
# =============================================================

CAMERAS_THEME_STYLE = """
* {
    font-family: "IBM Plex Sans Arabic", "Segoe UI", -apple-system, Tahoma, sans-serif;
}

QWidget#camerasRoot {
    background-color: #f4f7fb;
    color: #334155;
    font-size: 12px;
}

#headerTitle {
    color: #0f2338;
    font-size: 24px;
    font-weight: 800;
    letter-spacing: -0.3px;
}

#headerSubtitle {
    color: #64748b;
    font-size: 12px;
    font-weight: 500;
}

#liveTag {
    background-color: #ecfdf5;
    color: #065f46;
    border: 1px solid #a7f3d0;
    border-radius: 12px;
    padding: 3px 10px;
    font-size: 11px;
    font-weight: 700;
}

#primaryGradientBtn {
    background: qlineargradient(x1:0, y1:0, x2:1, y2:0, stop:0 #0891b2, stop:1 #06b6d4);
    color: white;
    border: none;
    border-radius: 9px;
    padding: 0 16px;
    min-height: 38px;
    font-size: 13px;
    font-weight: 700;
}

#primaryGradientBtn:hover {
    background: qlineargradient(x1:0, y1:0, x2:1, y2:0, stop:0 #0e7490, stop:1 #0891b2);
}

#primaryGradientBtn:pressed {
    background: #155e75;
}

#accentWarningBtn {
    background-color: #fff7ed;
    color: #c2410c;
    border: 1px solid #fed7aa;
    border-radius: 9px;
    padding: 0 15px;
    min-height: 38px;
    font-size: 13px;
    font-weight: 700;
}

#accentWarningBtn:hover {
    background-color: #ffedd5;
    border-color: #fb923c;
    color: #9a3412;
}

#accentWarningBtn:pressed {
    background-color: #fed7aa;
}

#moreOptionsBtn {
    background-color: #ffffff;
    color: #475569;
    border: 1px solid #cbd5e1;
    border-radius: 9px;
    padding: 0 14px;
    min-height: 38px;
    font-size: 13px;
    font-weight: 700;
}

#toolbarBtn, #toolbarDangerBtn {
    background-color: #ffffff;
    border: 1px solid #cbd5e1;
    border-radius: 8px;
    padding: 0 12px;
    min-height: 28px;
    font-size: 12px;
    font-weight: 700;
    color: #334155;
}

#toolbarBtn:hover {
    background-color: #f0f9ff;
    border-color: #0891b2;
    color: #0e7490;
}

#toolbarDangerBtn {
    color: #b91c1c;
    border-color: #fecaca;
}

#toolbarDangerBtn:hover {
    background-color: #fef2f2;
    border-color: #ef4444;
}

#toolbarBtn:disabled, #toolbarDangerBtn:disabled {
    background-color: #f8fafc;
    color: #cbd5e1;
    border-color: #e2e8f0;
}

#moreOptionsBtn:hover {
    background-color: #f8fafc;
    border-color: #94a3b8;
    color: #0f172a;
}

#modernToolbar {
    background-color: #ffffff;
    border: 1px solid #e2e8f0;
    border-radius: 12px;
}

#searchBox {
    background-color: #f8fafc;
    border: 1px solid #cbd5e1;
    border-radius: 8px;
    padding: 6px 12px;
    font-size: 12px;
    min-height: 24px;
    color: #0f2338;
}

#searchBox:focus {
    border: 1.5px solid #0891b2;
    background-color: #ffffff;
}

#filterCombo {
    background-color: #f8fafc;
    border: 1px solid #cbd5e1;
    border-radius: 8px;
    padding: 5px 10px;
    font-size: 12px;
    min-height: 24px;
    color: #334155;
}

#filterCombo:focus {
    border: 1.5px solid #0891b2;
    background-color: #ffffff;
}

#filterCombo::drop-down {
    border: none;
    width: 20px;
}

QTabWidget#mainTabs::pane {
    background: #ffffff;
    border: 1px solid #e2e8f0;
    border-radius: 12px;
    top: -1px;
}

QTabWidget#mainTabs > QTabBar::tab {
    background: #f1f5f9;
    color: #64748b;
    border: 1px solid #e2e8f0;
    border-bottom: none;
    border-top-left-radius: 10px;
    border-top-right-radius: 10px;
    padding: 10px 20px;
    margin-left: 4px;
    font-weight: 700;
    font-size: 13px;
}

QTabWidget#mainTabs > QTabBar::tab:selected {
    background: #ffffff;
    color: #0891b2;
    border-color: #e2e8f0;
    border-bottom: 2px solid #ffffff;
}

QTabWidget#mainTabs > QTabBar::tab:hover:!selected {
    background: #e2e8f0;
    color: #0f2338;
}

#proTableCard {
    background-color: #ffffff;
    border: 1px solid #e2e8f0;
    border-radius: 12px;
}

QTableWidget {
    background-color: #ffffff;
    border: none;
    gridline-color: #f1f5f9;
    selection-background-color: #e0f2fe;
    selection-color: #0f2338;
    alternate-background-color: #fafcff;
    font-size: 12px;
}

QHeaderView::section {
    background-color: #f8fafc;
    color: #475569;
    border: none;
    border-bottom: 2px solid #e2e8f0;
    padding: 9px 8px;
    font-weight: 700;
    font-size: 12px;
}

QScrollBar:vertical {
    border: none;
    background: #f1f5f9;
    width: 8px;
    margin: 0;
    border-radius: 4px;
}

QScrollBar::handle:vertical {
    background: #cbd5e1;
    min-height: 25px;
    border-radius: 4px;
}

QScrollBar::handle:vertical:hover {
    background: #0891b2;
}

QScrollBar:horizontal {
    border: none;
    background: #f1f5f9;
    height: 8px;
    margin: 0;
    border-radius: 4px;
}

QScrollBar::handle:horizontal {
    background: #cbd5e1;
    min-width: 25px;
    border-radius: 4px;
}

QScrollBar::handle:horizontal:hover {
    background: #0891b2;
}

#proNotice {
    background-color: #ecfdf5;
    color: #065f46;
    border: 1px solid #a7f3d0;
    border-radius: 8px;
    padding: 8px 14px;
    font-weight: 700;
    font-size: 12px;
}
"""


# =============================================================
# إنشاء الجداول (لو غير موجودة)
# =============================================================

DDL_STATEMENTS = [
    """
    IF OBJECT_ID('dbo.network_switches', 'U') IS NULL
    CREATE TABLE dbo.network_switches (
        id INT IDENTITY(1,1) PRIMARY KEY,
        switch_code NVARCHAR(30) NOT NULL,
        name NVARCHAR(200) NOT NULL,
        building_id INT NULL,
        location_detail NVARCHAR(300) NULL,
        model NVARCHAR(100) NULL,
        ports_count INT NULL,
        ip_address NVARCHAR(60) NULL,
        role NVARCHAR(100) NULL,
        status NVARCHAR(30) NOT NULL DEFAULT N'يعمل',
        notes NVARCHAR(MAX) NULL,
        is_archived BIT NOT NULL DEFAULT 0,
        archived_at DATETIME NULL,
        created_at DATETIME NOT NULL DEFAULT GETDATE(),
        updated_at DATETIME NOT NULL DEFAULT GETDATE(),
        CONSTRAINT FK_network_switches_building FOREIGN KEY (building_id) REFERENCES dbo.buildings(id)
    )
    """,
    """
    IF NOT EXISTS (SELECT 1 FROM sys.indexes WHERE name = 'UX_network_switches_code')
    CREATE UNIQUE INDEX UX_network_switches_code ON dbo.network_switches (switch_code)
    """,
    """
    IF OBJECT_ID('dbo.cameras', 'U') IS NULL
    CREATE TABLE dbo.cameras (
        id INT IDENTITY(1,1) PRIMARY KEY,
        camera_code NVARCHAR(30) NOT NULL,
        name NVARCHAR(200) NOT NULL,
        building_id INT NULL,
        location_detail NVARCHAR(300) NULL,
        placement NVARCHAR(20) NULL,
        camera_type NVARCHAR(60) NULL,
        manufacturer NVARCHAR(100) NULL,
        model NVARCHAR(100) NULL,
        serial_number NVARCHAR(100) NULL,
        ip_address NVARCHAR(60) NULL,
        install_date DATE NULL,
        warranty_end DATE NULL,
        status NVARCHAR(30) NOT NULL DEFAULT N'تعمل',
        maintenance_interval_days INT NULL,
        last_maintenance_date DATE NULL,
        notes NVARCHAR(MAX) NULL,
        is_archived BIT NOT NULL DEFAULT 0,
        archived_at DATETIME NULL,
        created_at DATETIME NOT NULL DEFAULT GETDATE(),
        updated_at DATETIME NOT NULL DEFAULT GETDATE()
    )
    """,
    """
    IF COL_LENGTH('dbo.cameras', 'placement') IS NULL
        ALTER TABLE dbo.cameras ADD placement NVARCHAR(20) NULL
    """,
    """
    IF COL_LENGTH('dbo.cameras', 'switch_id') IS NULL
        ALTER TABLE dbo.cameras ADD switch_id INT NULL
    """,
    """
    IF COL_LENGTH('dbo.cameras', 'switch_port') IS NULL
        ALTER TABLE dbo.cameras ADD switch_port NVARCHAR(30) NULL
    """,
    """
    IF NOT EXISTS (SELECT 1 FROM sys.foreign_keys WHERE name = 'FK_cameras_switch')
        ALTER TABLE dbo.cameras ADD CONSTRAINT FK_cameras_switch
        FOREIGN KEY (switch_id) REFERENCES dbo.network_switches(id)
    """,
    """
    IF NOT EXISTS (SELECT 1 FROM sys.indexes WHERE name = 'UX_cameras_code')
    CREATE UNIQUE INDEX UX_cameras_code ON dbo.cameras (camera_code)
    """,
    """
    IF NOT EXISTS (SELECT 1 FROM sys.indexes WHERE name = 'UX_cameras_switch_port')
    CREATE UNIQUE INDEX UX_cameras_switch_port ON dbo.cameras (switch_id, switch_port)
    WHERE switch_id IS NOT NULL AND switch_port IS NOT NULL AND is_archived = 0
    """,
    """
    IF OBJECT_ID('dbo.camera_faults', 'U') IS NULL
    CREATE TABLE dbo.camera_faults (
        id INT IDENTITY(1,1) PRIMARY KEY,
        fault_number NVARCHAR(30) NOT NULL,
        camera_id INT NOT NULL,
        entry_date DATETIME NOT NULL,
        fault_description NVARCHAR(MAX) NOT NULL,
        repair_actions NVARCHAR(MAX) NULL,
        spare_part_id INT NULL,
        spare_part_qty INT NULL,
        spare_parts NVARCHAR(300) NULL,
        technician NVARCHAR(150) NULL,
        priority NVARCHAR(20) NOT NULL DEFAULT N'عادي',
        status NVARCHAR(30) NOT NULL DEFAULT N'جديد',
        repair_date DATETIME NULL,
        notes NVARCHAR(MAX) NULL,
        is_archived BIT NOT NULL DEFAULT 0,
        archived_at DATETIME NULL,
        created_at DATETIME NOT NULL DEFAULT GETDATE(),
        updated_at DATETIME NOT NULL DEFAULT GETDATE(),
        CONSTRAINT FK_camera_faults_camera
            FOREIGN KEY (camera_id) REFERENCES dbo.cameras(id)
    )
    """,
    """
    IF NOT EXISTS (SELECT 1 FROM sys.indexes WHERE name = 'UX_camera_faults_number')
    CREATE UNIQUE INDEX UX_camera_faults_number
        ON dbo.camera_faults (fault_number)
    """,
    """
    IF NOT EXISTS (SELECT 1 FROM sys.indexes WHERE name = 'IX_camera_faults_camera')
    CREATE INDEX IX_camera_faults_camera ON dbo.camera_faults (camera_id)
    """,
    # جدول الصيانة القديم يبقى كما هو (لم يعد له واجهة) حتى لا تضيع بياناته
    # ولأن فحص الحذف النهائي يعتمد عليه.
    """
    IF OBJECT_ID('dbo.camera_maintenance', 'U') IS NULL
    CREATE TABLE dbo.camera_maintenance (
        id INT IDENTITY(1,1) PRIMARY KEY,
        camera_id INT NOT NULL,
        maintenance_date DATE NOT NULL,
        maintenance_type NVARCHAR(100) NULL,
        technician NVARCHAR(150) NULL,
        notes NVARCHAR(MAX) NULL,
        created_at DATETIME NOT NULL DEFAULT GETDATE(),
        CONSTRAINT FK_camera_maintenance_camera
            FOREIGN KEY (camera_id) REFERENCES dbo.cameras(id)
    )
    """,
]


def ensure_camera_tables():
    """يُنشئ جداول الكاميرات والسويتشات لو غير موجودة. يرجع (نجح؟، رسالة الخطأ)."""
    connection = get_connection()
    try:
        cursor = connection.cursor()
        for statement in DDL_STATEMENTS:
            cursor.execute(statement)
        connection.commit()
        return True, ""
    except Exception as exc:
        try:
            connection.rollback()
        except Exception:
            pass
        return False, str(exc)
    finally:
        connection.close()


# =============================================================
# أدوات مساعدة وقواعد بيانات
# =============================================================

def fetch_all(sql, params=()):
    connection = get_connection()
    try:
        cursor = connection.cursor()
        cursor.execute(sql, params)
        return cursor.fetchall()
    finally:
        connection.close()


def fetch_one(sql, params=()):
    rows = fetch_all(sql, params)
    return rows[0] if rows else None


def to_date(value):
    """يحوّل أي قيمة تاريخ إلى date أو None."""
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    parsed = parse_datetime(value)
    return parsed.date() if parsed else None


def fmt_date(value):
    parsed = to_date(value)
    return parsed.strftime("%Y-%m-%d") if parsed else "-"


def fmt_dt(value):
    return format_datetime(value) if value else "-"


def make_item(text, colors=None, bold=False, tooltip=None, muted=False, align_right=False):
    """عنصر جدول منسق."""
    item = QTableWidgetItem(str(text if text not in (None, "") else "-"))
    if align_right:
        item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter | Qt.AlignAbsolute)
    else:
        item.setTextAlignment(Qt.AlignCenter)

    if colors:
        fg, bg = colors[0], colors[1]
        item.setForeground(QColor(fg))
        if bg:
            item.setBackground(QColor(bg))
        bold = True

    if muted or item.text() == "-":
        item.setForeground(QColor("#94a3b8"))

    if bold:
        font = item.font()
        font.setBold(True)
        item.setFont(font)

    if tooltip:
        item.setToolTip(tooltip)

    return item


def make_table(headers, stretch_column=None):
    """إنشاء جدول بتصميم موحد."""
    table = QTableWidget()
    table.setColumnCount(len(headers))
    table.setHorizontalHeaderLabels(headers)
    table.setSelectionBehavior(QTableWidget.SelectRows)
    table.setSelectionMode(QTableWidget.SingleSelection)
    table.setEditTriggers(QTableWidget.NoEditTriggers)
    table.setAlternatingRowColors(True)
    table.setShowGrid(False)
    table.setWordWrap(False)
    table.verticalHeader().setVisible(False)
    table.verticalHeader().setDefaultSectionSize(40)
    table.setContextMenuPolicy(Qt.CustomContextMenu)

    header = table.horizontalHeader()
    header.setSectionResizeMode(QHeaderView.ResizeToContents)
    header.setHighlightSections(False)

    if stretch_column is not None:
        header.setSectionResizeMode(stretch_column, QHeaderView.Stretch)

    return table


def wrap_card(widget):
    """تغليف أي عنصر ببطاقة بحدود خفيفة."""
    card = QFrame()
    card.setObjectName("proTableCard")
    card.setStyleSheet("""
        QFrame#proTableCard {
            background-color: #ffffff;
            border: 1px solid #e2e8f0;
            border-radius: 12px;
        }
    """)
    layout = QVBoxLayout(card)
    layout.setContentsMargins(10, 10, 10, 10)
    layout.addWidget(widget)
    return card


def optional_date_edit():
    edit = QDateEdit()
    edit.setCalendarPopup(True)
    edit.setDisplayFormat("yyyy-MM-dd")
    edit.setSpecialValueText("غير محدد")
    clear_datetime(edit)
    keep_calendar_on_today(edit)
    return edit


def optional_date_text(edit):
    return None if is_unset(edit) else edit.date().toString("yyyy-MM-dd")


def set_optional_date(edit, value):
    parsed = to_date(value)
    if parsed:
        edit.setDate(QDate(parsed.year, parsed.month, parsed.day))


def with_clear_button(edit):
    box = QHBoxLayout()
    box.setContentsMargins(0, 0, 0, 0)
    box.setSpacing(6)

    button = QPushButton("مسح")
    button.setStyleSheet("""
        QPushButton {
            background: #f1f5f9;
            color: #475569;
            border: 1px solid #cbd5e1;
            border-radius: 7px;
            padding: 5px 12px;
            font-weight: 700;
        }
        QPushButton:hover {
            background: #e2e8f0;
            color: #0f172a;
        }
    """)
    button.clicked.connect(lambda: clear_datetime(edit))

    box.addWidget(edit, 1)
    box.addWidget(button)

    container = QWidget()
    container.setLayout(box)
    return container


def building_choices():
    """المباني والفروع النشطة من الكاش المركزي."""
    try:
        return OrgStructureCache.instance().get_active_by_types("مبنى", "فرع")
    except Exception:
        return []


def switches_choices(building_id=None):
    """قائمة السويتشات النشطة، مع إمكانية التصفية بمبنى معين."""
    sql = """
        SELECT s.id, s.switch_code, s.name, s.building_id, b.name as building_name
        FROM network_switches s
        LEFT JOIN buildings b ON s.building_id = b.id
        WHERE s.is_archived = 0
    """
    params = []
    if building_id:
        sql += " AND (s.building_id = ? OR s.building_id IS NULL)"
        params.append(building_id)
    sql += " ORDER BY s.switch_code"
    return fetch_all(sql, params)


def open_print_page(content, prefix):
    path = (
        Path(tempfile.gettempdir())
        / f"{prefix}_{datetime.now():%Y%m%d_%H%M%S}.html"
    )
    path.write_text(content, encoding="utf-8")
    webbrowser.open(path.as_uri())


def build_print_html(title, subtitle, headers, rows, summary=""):
    """تقرير مطبوع (A4 أفقي)."""
    head = "".join(f"<th>{html.escape(h)}</th>" for h in headers)
    body = "".join(
        "<tr>"
        + "".join(
            f"<td>{html.escape(str(v if v not in (None, '') else '-'))}</td>"
            for v in row
        )
        + "</tr>"
        for row in rows
    ) or f"<tr><td colspan='{len(headers)}'>لا توجد بيانات متاحة للعرض</td></tr>"

    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    summary_html = f'<div class="summary">{html.escape(summary)}</div>' if summary else ""

    return f"""<!doctype html>
<html lang="ar" dir="rtl"><head><meta charset="utf-8">
<title>{html.escape(title)}</title>
<style>
@page {{ size: A4 landscape; margin: 12mm; }}
* {{ box-sizing: border-box; }}
body {{ font-family: "Segoe UI", Tahoma, Arial, sans-serif; color:#1e293b; font-size:11px; margin:0; }}
.head {{ border-bottom:3px solid #0891b2; padding-bottom:10px; margin-bottom:12px; display:flex; justify-content:space-between; align-items:flex-end; }}
h1 {{ margin:0; font-size:22px; color:#0f2338; font-weight:800; }}
.sub {{ color:#64748b; margin-top:4px; font-size:12px; }}
.badge {{ background:#ecfeff; color:#0891b2; border:1px solid #a5f3fc; padding:4px 10px; border-radius:12px; font-weight:700; font-size:11px; }}
.summary {{ background:#f8fafc; border:1px solid #e2e8f0; border-radius:8px; padding:10px 14px; margin-bottom:12px; font-weight:700; color:#334155; }}
table {{ width:100%; border-collapse:collapse; margin-top:6px; }}
thead {{ display:table-header-group; }}
tr {{ break-inside:avoid; }}
th, td {{ border:1px solid #cbd5e1; padding:7px 6px; text-align:center; }}
th {{ background:#0f2338; color:#ffffff; font-size:11px; font-weight:700; }}
tbody tr:nth-child(even) td {{ background:#f8fafc; }}
.foot {{ margin-top:14px; border-top:1px solid #e2e8f0; padding-top:8px; display:flex; justify-content:space-between; color:#64748b; font-size:10px; }}
</style></head><body>
<div class="head">
  <div>
    <h1>{html.escape(title)}</h1>
    <div class="sub">{html.escape(subtitle)}</div>
  </div>
  <div class="badge">نظام إدارة الشبكة والكاميرات</div>
</div>
{summary_html}
<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>
<div class="foot">
  <span>نظام الشبكة — الكاميرات والسويتشات</span>
  <span>تاريخ وساعة الطباعة: {now}</span>
</div>
<script>window.onload=()=>window.print()</script>
</body></html>"""


# =============================================================
# بطاقة إحصائيات تفاعلية
# =============================================================

class ProMetricCard(QFrame):
    """بطاقة مؤشر بسيطة؛ تدعم النقر للتصفية المباشرة."""

    def __init__(self, title, value="0", subtitle="", accent="#0891b2",
                 icon="", on_click=None, parent=None):
        super().__init__(parent)
        self.on_click = on_click
        self.accent = accent

        self.setObjectName("proMetricCard")
        self.setCursor(Qt.PointingHandCursor if on_click else Qt.ArrowCursor)

        self.setStyleSheet(f"""
            QFrame#proMetricCard {{
                background-color: #ffffff;
                border: 1px solid #e2e8f0;
                border-radius: 12px;
            }}
            QFrame#proMetricCard:hover {{
                border-color: {accent};
                background-color: #fafcff;
            }}
        """)
        self.setMinimumHeight(86)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(16, 12, 16, 12)
        layout.setSpacing(2)

        # العنوان مع نقطة لونية (بدل الشريط العلوي)
        top_row = QHBoxLayout()
        top_row.setSpacing(7)

        dot = QLabel("●")
        dot.setStyleSheet(f"color: {accent}; font-size: 11px; background: transparent;")
        top_row.addWidget(dot)

        self.title_label = QLabel(title)
        self.title_label.setStyleSheet(
            "color: #64748b; font-size: 12px; font-weight: 700; background: transparent;"
        )
        top_row.addWidget(self.title_label)
        top_row.addStretch()
        layout.addLayout(top_row)

        self.val_label = QLabel(str(value))
        self.val_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        self.val_label.setStyleSheet(
            "color: #0f2338; font-size: 28px; font-weight: 800; background: transparent;"
        )
        layout.addWidget(self.val_label)

        # سطر فرعي يظهر فقط لو فيه نص
        self.sub_label = QLabel(subtitle)
        self.sub_label.setStyleSheet("color: #94a3b8; font-size: 11px; background: transparent;")
        self.sub_label.setVisible(bool(subtitle))
        layout.addWidget(self.sub_label)

        # نسق تفاعلي: نخلي الكارت بلا فراغ زيادة لما مفيش سطر فرعي
        layout.addStretch()

    def set_value(self, val, sub=None):
        self.val_label.setText(str(val))
        if sub is not None:
            self.sub_label.setText(str(sub))
            self.sub_label.setVisible(bool(sub))

    def mousePressEvent(self, event):
        if event.button() == Qt.LeftButton and self.on_click:
            self.on_click()
        super().mousePressEvent(event)


# =============================================================
# نافذة إضافة / تعديل سويتش شبكة
# =============================================================

class SwitchDialog(QDialog):
    """إدخال وتعديل بيانات سويتش الشبكة."""

    def __init__(self, parent=None, switch=None):
        super().__init__(parent)
        self.switch = switch

        self.setWindowTitle("تعديل سويتش شبكة" if switch else "إضافة سويتش شبكة جديد")
        self.resize(580, 560)
        self.setLayoutDirection(Qt.RightToLeft)
        self.setStyleSheet(CAMERAS_THEME_STYLE)

        root = QVBoxLayout(self)
        root.setContentsMargins(20, 18, 20, 16)
        root.setSpacing(14)

        header_box = QHBoxLayout()
        icon = QLabel("")
        icon.setStyleSheet("font-size: 26px; color: #0891b2; padding-left: 6px;")
        header_box.addWidget(icon)

        title_vbox = QVBoxLayout()
        title = QLabel("تعديل بيانات السويتش" if switch else "تسجيل سويتش شبكة جديد")
        title.setStyleSheet("font-size: 18px; font-weight: 800; color: #0f2338;")
        sub = QLabel("تحديد مواصفات السويتش وموقعه ودوره في المنظومة")
        sub.setStyleSheet("font-size: 11px; color: #64748b;")
        title_vbox.addWidget(title)
        title_vbox.addWidget(sub)
        header_box.addLayout(title_vbox)
        header_box.addStretch()
        root.addLayout(header_box)

        form_frame = QFrame()
        form_frame.setStyleSheet("""
            QFrame {
                background-color: #ffffff;
                border: 1px solid #e2e8f0;
                border-radius: 10px;
                padding: 12px;
            }
        """)
        form = QFormLayout(form_frame)
        form.setLabelAlignment(Qt.AlignRight)
        form.setHorizontalSpacing(14)
        form.setVerticalSpacing(10)

        self.code = QLineEdit()
        self.code.setReadOnly(True)
        self.code.setStyleSheet("background: #f1f5f9; color: #475569; font-weight: 700;")

        self.name = QLineEdit()
        self.name.setPlaceholderText("مثال: سويتش الإدارة — الدور الأرضي")

        self.building = QComboBox()
        self.building.addItem("غير محدد", None)
        for unit in building_choices():
            self.building.addItem(f"[{unit['unit_type']}] {unit['name']}", unit["id"])

        self.location_detail = QLineEdit()
        self.location_detail.setPlaceholderText("الدور / الغرفة / الموضع (مثال: الدور الأرضي)")

        self.model = QComboBox()
        self.model.setEditable(True)
        self.model.addItems(["Cisco 24 Port", "Cisco 48 Port", "Cisco 12 Port", "Tp-Link", "Ubiquiti NanoStation", "أخرى"])

        self.ports = QSpinBox()
        self.ports.setRange(1, 96)
        self.ports.setValue(24)
        self.ports.setSuffix(" منفذ")

        self.role = QComboBox()
        self.role.addItems(SWITCH_ROLES)

        self.ip_address = QLineEdit()
        self.ip_address.setPlaceholderText("192.168.x.x (مرجع إداري للسويتش)")

        self.status = QComboBox()
        self.status.addItems(SWITCH_STATUSES)

        self.notes = QTextEdit()
        self.notes.setFixedHeight(65)
        self.notes.setPlaceholderText("ملاحظات إضافية حول التغذية أو التوصيلات...")

        form.addRow("كود السويتش:", self.code)
        form.addRow("اسم السويتش *:", self.name)
        form.addRow("المبنى / الفرع:", self.building)
        form.addRow("الموضع التفصيلي:", self.location_detail)
        form.addRow("الموديل والنوع:", self.model)
        form.addRow("عدد المنافذ:", self.ports)
        form.addRow("الغرض والاستخدام:", self.role)
        form.addRow("عنوان IP:", self.ip_address)
        form.addRow("الحالة التشغيلية:", self.status)
        form.addRow("ملاحظات:", self.notes)

        root.addWidget(form_frame, 1)

        btn_box = QHBoxLayout()
        save_btn = QPushButton("  حفظ السويتش")
        save_btn.setObjectName("primaryGradientBtn")
        save_btn.setCursor(Qt.PointingHandCursor)
        save_btn.clicked.connect(self.validate_and_accept)

        cancel_btn = QPushButton("إلغاء")
        cancel_btn.setObjectName("moreOptionsBtn")
        cancel_btn.setCursor(Qt.PointingHandCursor)
        cancel_btn.clicked.connect(self.reject)

        btn_box.addStretch()
        btn_box.addWidget(cancel_btn)
        btn_box.addWidget(save_btn)
        root.addLayout(btn_box)

        if switch:
            self.load_switch()

    def load_switch(self):
        s = self.switch
        self.code.setText(s["switch_code"])
        self.name.setText(s["name"] or "")

        idx = self.building.findData(s["building_id"])
        if idx >= 0:
            self.building.setCurrentIndex(idx)

        self.location_detail.setText(s["location_detail"] or "")
        self.model.setCurrentText(s["model"] or "")
        self.ports.setValue(int(s["ports_count"] or 24))

        idx = self.role.findText(s["role"] or "")
        if idx >= 0:
            self.role.setCurrentIndex(idx)

        self.ip_address.setText(s["ip_address"] or "")

        idx = self.status.findText(s["status"] or "")
        if idx >= 0:
            self.status.setCurrentIndex(idx)

        self.notes.setPlainText(s["notes"] or "")

    def validate_and_accept(self):
        if not self.name.text().strip():
            QMessageBox.warning(self, "تنبيه", "يرجى كتابة اسم السويتش.")
            self.name.setFocus()
            return
        self.accept()

    def get_data(self):
        return {
            "switch_code": self.code.text(),
            "name": self.name.text().strip(),
            "building_id": self.building.currentData(),
            "location_detail": self.location_detail.text().strip(),
            "model": self.model.currentText().strip(),
            "ports_count": self.ports.value(),
            "role": self.role.currentText(),
            "ip_address": self.ip_address.text().strip(),
            "status": self.status.currentText(),
            "notes": self.notes.toPlainText().strip(),
        }


# =============================================================
# نافذة ربط الكاميرات بالسويتش
# =============================================================

class SwitchLinkCamerasDialog(QDialog):
    """ربط الكاميرات بسويتش معين بنقرة واحدة."""

    def __init__(self, switch_id, parent=None):
        super().__init__(parent)
        self.switch_id = switch_id

        self.setWindowTitle("ربط الكاميرات بالسويتش")
        self.resize(750, 600)
        self.setLayoutDirection(Qt.RightToLeft)
        self.setStyleSheet(CAMERAS_THEME_STYLE)

        root = QVBoxLayout(self)
        root.setContentsMargins(20, 18, 20, 16)
        root.setSpacing(12)

        self.switch_info = fetch_one("""
            SELECT s.*, b.name as building_name
            FROM network_switches s
            LEFT JOIN buildings b ON s.building_id = b.id
            WHERE s.id = ?
        """, (switch_id,))

        title = QLabel(f" ربط الكاميرات بالسويتش: {self.switch_info['name']}")
        title.setStyleSheet("font-size: 16px; font-weight: 800; color: #0891b2;")
        root.addWidget(title)

        sub = QLabel(
            f"المبنى: <b>{html.escape(self.switch_info['building_name'] or 'غير محدد')}</b> | "
            f"الموديل: <b>{html.escape(self.switch_info['model'] or '-')}</b> | "
            f"الغرض: <b>{html.escape(self.switch_info['role'] or '-')}</b>"
        )
        sub.setStyleSheet("font-size: 12px; color: #475569;")
        root.addWidget(sub)

        filter_box = QHBoxLayout()
        self.filter_scope = QCheckBox("عرض كاميرات نفس المبنى فقط")
        self.filter_scope.setChecked(True)
        self.filter_scope.toggled.connect(self.load_cameras)
        filter_box.addWidget(self.filter_scope)
        filter_box.addStretch()
        root.addLayout(filter_box)

        self.table = make_table(
            ["تحديد", "ID", "كود الكاميرا", "اسم الكاميرا", "المبنى", "السويتش الحالي"],
            stretch_column=3,
        )
        self.table.setColumnHidden(1, True)
        root.addWidget(wrap_card(self.table), 1)

        btn_box = QHBoxLayout()
        save_btn = QPushButton("  حفظ تعيينات الكاميرات")
        save_btn.setObjectName("primaryGradientBtn")
        save_btn.setCursor(Qt.PointingHandCursor)
        save_btn.clicked.connect(self.save_links)

        cancel_btn = QPushButton("إلغاء")
        cancel_btn.setObjectName("moreOptionsBtn")
        cancel_btn.setCursor(Qt.PointingHandCursor)
        cancel_btn.clicked.connect(self.reject)

        btn_box.addStretch()
        btn_box.addWidget(cancel_btn)
        btn_box.addWidget(save_btn)
        root.addLayout(btn_box)

        self.load_cameras()

    def load_cameras(self):
        same_building = self.filter_scope.isChecked()
        building_id = self.switch_info["building_id"]

        sql = """
            SELECT c.id, c.camera_code, c.name, c.switch_id,
                   b.name as building_name, s.name as current_switch_name
            FROM cameras c
            LEFT JOIN buildings b ON c.building_id = b.id
            LEFT JOIN network_switches s ON c.switch_id = s.id
            WHERE c.is_archived = 0
        """
        params = []
        if same_building and building_id:
            sql += " AND (c.building_id = ? OR c.switch_id = ?)"
            params.extend([building_id, self.switch_id])
        sql += " ORDER BY c.camera_code"

        rows = fetch_all(sql, params)
        self.table.setRowCount(0)

        for r in rows:
            row = self.table.rowCount()
            self.table.insertRow(row)

            is_linked = (r["switch_id"] == self.switch_id)

            chk_item = QTableWidgetItem()
            chk_item.setFlags(Qt.ItemIsUserCheckable | Qt.ItemIsEnabled)
            chk_item.setCheckState(Qt.Checked if is_linked else Qt.Unchecked)
            self.table.setItem(row, 0, chk_item)

            self.table.setItem(row, 1, make_item(r["id"]))
            self.table.setItem(row, 2, make_item(r["camera_code"], bold=True))
            self.table.setItem(row, 3, make_item(r["name"], align_right=True))
            self.table.setItem(row, 4, make_item(r["building_name"] or "-", align_right=True))

            current_sw = r["current_switch_name"] or "غير مربوط"
            sw_color = ("#065f46", "#ecfdf5") if is_linked else ("#64748b", None)
            self.table.setItem(row, 5, make_item(current_sw, colors=sw_color, align_right=True))

    def save_links(self):
        conn = get_connection()
        try:
            cur = conn.cursor()
            for r in range(self.table.rowCount()):
                chk = self.table.item(r, 0).checkState() == Qt.Checked
                cam_id = int(self.table.item(r, 1).text())

                if chk:
                    # لا نمسح المنفذ لو الكاميرا مربوطة بالفعل بنفس السويتش
                    cur.execute("""
                        UPDATE cameras
                        SET switch_port = CASE WHEN switch_id = ? THEN switch_port ELSE NULL END,
                            switch_id = ?, updated_at = GETDATE()
                        WHERE id = ?
                    """, (self.switch_id, self.switch_id, cam_id))
                else:
                    cur.execute("""
                        UPDATE cameras
                        SET switch_id = NULL, switch_port = NULL, updated_at = GETDATE()
                        WHERE id = ? AND switch_id = ?
                    """, (cam_id, self.switch_id))

            conn.commit()
            self.accept()
        except Exception as exc:
            conn.rollback()
            QMessageBox.critical(self, "خطأ", f"تعذر حفظ الربط:\n{exc}")
        finally:
            conn.close()


# =============================================================
# نافذة إضافة / تعديل كاميرا
# =============================================================

class CameraDialog(QDialog):
    """نافذة إضافة/تعديل كاميرا مع تحقق فوري."""

    def __init__(self, parent=None, camera=None):
        super().__init__(parent)
        self.camera = camera
        self._loading = True
        self._duplicate_serial = False

        self.setWindowTitle("تعديل بيانات الكاميرا" if camera else "إضافة كاميرا جديدة")
        self.resize(760, 700)
        self.setMinimumSize(680, 600)
        self.setLayoutDirection(Qt.RightToLeft)
        self.setStyleSheet(CAMERAS_THEME_STYLE + """
            QLineEdit, QComboBox, QSpinBox, QTextEdit {
                min-height: 30px;
                border: 1px solid #cbd5e1;
                border-radius: 7px;
                padding: 3px 9px;
                background: #ffffff;
            }
            QLineEdit:focus, QComboBox:focus, QSpinBox:focus, QTextEdit:focus {
                border: 2px solid #0891b2;
            }
            QLineEdit[error="true"] { border: 2px solid #dc2626; background: #fff7f7; }
            QLabel.fieldHint { color: #64748b; font-size: 10px; }
            QLabel.fieldError { color: #dc2626; font-size: 10px; font-weight: 700; }
            QTabWidget::pane { border: 1px solid #e2e8f0; border-radius: 10px; background: #fff; }
            QTabBar::tab { padding: 10px 20px; font-weight: 700; color: #64748b; }
            QTabBar::tab:selected { color: #0891b2; border-bottom: 2px solid #0891b2; }
        """)

        root = QVBoxLayout(self)
        root.setContentsMargins(18, 16, 18, 14)
        root.setSpacing(10)

        header = QHBoxLayout()
        icon = QLabel("")
        icon.setStyleSheet("font-size: 25px; padding-left: 7px;")
        header.addWidget(icon)
        title_box = QVBoxLayout()
        title = QLabel("تعديل بيانات الكاميرا" if camera else "إضافة كاميرا جديدة")
        title.setStyleSheet("font-size: 19px; font-weight: 800; color: #0f2338;")
        subtitle = QLabel("أدخل البيانات الأساسية، ثم راجع الشبكة قبل الحفظ")
        subtitle.setStyleSheet("font-size: 11px; color: #64748b;")
        title_box.addWidget(title)
        title_box.addWidget(subtitle)
        header.addLayout(title_box)
        header.addStretch()
        self.step_hint = QLabel("البيانات الأساسية")
        self.step_hint.setStyleSheet("color:#0891b2; font-weight:700;")
        header.addWidget(self.step_hint)
        root.addLayout(header)

        self.step_buttons = []
        stepper = QHBoxLayout()
        for i, label in enumerate(("1  البيانات الأساسية", "2  الشبكة والمواصفات", "3  الحالة والضمان")):
            btn = QPushButton(label)
            btn.setCheckable(True)
            btn.clicked.connect(lambda checked=False, n=i: self._go_to_step(n))
            stepper.addWidget(btn)
            self.step_buttons.append(btn)
        root.addLayout(stepper)

        self.tabs = QStackedWidget()
        self.tabs.currentChanged.connect(self._tab_changed)
        root.addWidget(self.tabs, 1)

        self._build_identity_tab()
        self._build_network_tab()
        self._build_operations_tab()

        footer = QHBoxLayout()
        footer.setSpacing(9)
        self.status_hint = QLabel("الحقول المعلّمة بنجمة مطلوبة")
        self.status_hint.setStyleSheet("color:#64748b; font-size:11px;")
        footer.addWidget(self.status_hint)
        footer.addStretch()

        self.back_button = QPushButton("السابق")
        self.back_button.setObjectName("moreOptionsBtn")
        self.back_button.clicked.connect(
            lambda: self.tabs.setCurrentIndex(max(0, self.tabs.currentIndex() - 1))
        )
        self.next_button = QPushButton("التالي ←")
        self.next_button.setObjectName("primaryGradientBtn")
        self.next_button.clicked.connect(self._step_next)
        cancel = QPushButton("إلغاء")
        cancel.setObjectName("moreOptionsBtn")
        cancel.clicked.connect(self.reject)
        self.save_button = QPushButton("مراجعة وحفظ")
        self.save_button.setObjectName("primaryGradientBtn")
        self.save_button.setDefault(True)
        self.save_button.clicked.connect(self.validate_and_accept)
        footer.addWidget(self.back_button)
        footer.addWidget(self.next_button)
        footer.addWidget(cancel)
        footer.addWidget(self.save_button)
        root.addLayout(footer)

        self._load_switches_into_combo()
        self._connect_live_validation()
        if camera:
            self.load_camera()
        else:
            self._set_new_defaults()
        self._loading = False
        self._tab_changed(0)
        self.name.setFocus()

    def _card(self, title_text):
        card = QFrame()
        card.setStyleSheet("""
            QFrame { background:#fff; border:1px solid #e2e8f0; border-radius:10px; }
            QFrame QLabel.cardTitle { color:#0891b2; font-size:13px; font-weight:800; }
        """)
        layout = QVBoxLayout(card)
        layout.setContentsMargins(14, 11, 14, 14)
        heading = QLabel(title_text)
        heading.setProperty("class", "cardTitle")
        heading.setObjectName("cardTitle")
        layout.addWidget(heading)
        return card

    def _form_card(self, title_text):
        card = self._card(title_text)
        form = QFormLayout()
        form.setLabelAlignment(Qt.AlignRight | Qt.AlignVCenter)
        form.setHorizontalSpacing(16)
        form.setVerticalSpacing(9)
        card.layout().addLayout(form)
        return card, form

    def _build_identity_tab(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        layout.setContentsMargins(5, 10, 5, 10)
        layout.setSpacing(10)

        card, form = self._form_card(" التعريف والموقع")
        self.code = QLineEdit()
        self.code.setReadOnly(True)
        self.code.setPlaceholderText("يُنشأ تلقائيًا عند الحفظ")
        self.code.setStyleSheet("background:#f1f5f9; color:#475569; font-weight:700;")

        self.name = QLineEdit()
        self.name.setPlaceholderText("مثال: كاميرا المدخل الرئيسي")
        self.name.setClearButtonEnabled(True)
        self.building = QComboBox()
        self.building.addItem("غير محدد", None)
        for unit in building_choices():
            self.building.addItem(f"[{unit['unit_type']}] {unit['name']}", unit["id"])
        self.building.currentIndexChanged.connect(self._building_changed)

        self.location_detail = QLineEdit()
        self.location_detail.setPlaceholderText("الدور / الغرفة / الموضع التفصيلي")
        self.placement = QComboBox()
        self.placement.addItem("غير محدد", "")
        for value in PLACEMENTS:
            self.placement.addItem(value, value)
        self.camera_type = QComboBox()
        self.camera_type.setEditable(True)
        self.camera_type.addItems(CAMERA_TYPES)

        form.addRow("كود الكاميرا:", self.code)
        form.addRow("الاسم التعريفي *:", self.name)
        form.addRow("المبنى / الفرع:", self.building)
        form.addRow("الموضع التفصيلي:", self.location_detail)
        form.addRow("التثبيت:", self.placement)
        form.addRow("نوع الكاميرا:", self.camera_type)
        layout.addWidget(card)
        layout.addStretch()
        self.tabs.addWidget(page)

    def _build_network_tab(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        layout.setContentsMargins(5, 10, 5, 10)
        layout.setSpacing(10)
        card, form = self._form_card(" الهوية الفنية والربط الشبكي")

        self.switch_combo = QComboBox()
        self.switch_port = QLineEdit()
        self.switch_port.setPlaceholderText("مثال: 1/12 أو 12")
        self.switch_port.setMaxLength(30)
        self.manufacturer = QLineEdit()
        self.manufacturer.setPlaceholderText("Hikvision / Dahua / Uniview")
        self.model = QLineEdit()
        self.model.setPlaceholderText("رقم الموديل")
        self.serial_number = QLineEdit()
        self.serial_number.setPlaceholderText("Serial Number")
        self.ip_address = QLineEdit()
        self.ip_address.setPlaceholderText("192.168.1.20")
        self.ip_address.setValidator(QRegularExpressionValidator(QRegularExpression(r"[0-9.]{0,15}")))

        form.addRow("السويتش المتصل به:", self.switch_combo)
        form.addRow("منفذ السويتش:", self.switch_port)
        form.addRow("الشركة المصنعة:", self.manufacturer)
        form.addRow("الموديل:", self.model)
        form.addRow("الرقم التسلسلي:", self.serial_number)
        form.addRow("عنوان IP:", self.ip_address)
        layout.addWidget(card)
        hint = QLabel("يفضل تسجيل السويتش والمنفذ لتسهيل التتبع لاحقًا")
        hint.setProperty("class", "fieldHint")
        layout.addWidget(hint)
        layout.addStretch()
        self.tabs.addWidget(page)

    def _build_operations_tab(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        layout.setContentsMargins(5, 10, 5, 10)
        layout.setSpacing(10)
        card, form = self._form_card(" الحالة والضمان")
        self.status = QComboBox()
        self.status.addItems(CAMERA_STATUSES)
        self.install_date = optional_date_edit()
        self.warranty_end = optional_date_edit()
        self.notes = QTextEdit()
        self.notes.setFixedHeight(90)
        self.notes.setPlaceholderText("ملاحظات فنية أو إدارية إضافية...")

        form.addRow("الحالة التشغيلية:", self.status)
        form.addRow("تاريخ التركيب:", with_clear_button(self.install_date))
        form.addRow("انتهاء الضمان:", with_clear_button(self.warranty_end))
        form.addRow("ملاحظات:", self.notes)
        layout.addWidget(card)
        layout.addStretch()
        self.tabs.addWidget(page)

    def _connect_live_validation(self):
        self.name.textChanged.connect(self._validate_name_live)
        self.ip_address.textChanged.connect(self._validate_ip_live)
        self.serial_number.editingFinished.connect(self._check_serial_duplicate)

    def _set_new_defaults(self):
        self.status.setCurrentText("تعمل")
        self.camera_type.setCurrentText("ثابتة")

    def _tab_changed(self, index):
        labels = ["البيانات الأساسية", "الشبكة والمواصفات", "الحالة والضمان"]
        self.step_hint.setText(f"الخطوة {index + 1} من 3 • {labels[index]}")
        for i, btn in enumerate(self.step_buttons):
            btn.setChecked(i == index)
        self.back_button.setEnabled(index > 0)
        self.next_button.setVisible(index < 2)
        self.save_button.setVisible(index == 2)

    def _go_to_step(self, index):
        current = self.tabs.currentIndex()
        if index > current and not self._validate_step(current):
            return
        self.tabs.setCurrentIndex(index)

    def _step_next(self):
        current = self.tabs.currentIndex()
        if self._validate_step(current):
            self.tabs.setCurrentIndex(min(2, current + 1))

    def _validation_warning(self, title, message, widget=None):
        if widget is not None:
            widget.setFocus()
            self._set_error(widget, True)
        self.status_hint.setText(message)
        self.status_hint.setStyleSheet("color:#dc2626; font-size:11px; font-weight:700;")
        QMessageBox.warning(self, title, message)
        return False

    def _validate_step(self, index):
        if index == 0:
            if not self.name.text().strip():
                return self._validation_warning("بيانات ناقصة", "اكتب اسم الكاميرا قبل الانتقال للخطوة التالية.", self.name)
            if not self.code.text().strip():
                return self._validation_warning("كود الكاميرا ناقص", "يجب تحديد كود الكاميرا قبل المتابعة.", self.code)
            if not self._check_code_duplicate(True):
                return self._validation_warning("كود مكرر", "كود الكاميرا مستخدم بالفعل، اختر كودًا آخر.", self.code)
            return True
        if index == 1:
            if not self._validate_ip_live(self.ip_address.text()):
                return self._validation_warning("عنوان IP غير صحيح", "اكتب عنوان IPv4 صحيحًا أو اترك الحقل فارغًا.", self.ip_address)
            self._check_serial_duplicate()
            if self._duplicate_serial:
                return self._validation_warning("رقم تسلسلي مكرر", "هذا الرقم التسلسلي مستخدم في كاميرا أخرى.", self.serial_number)
            if not self._check_port_duplicate(True):
                return self._validation_warning("منفذ مستخدم", "منفذ السويتش مستخدم بالفعل على نفس السويتش.", self.switch_port)
            return True
        return True

    def _validate_name_live(self, value):
        self._set_error(self.name, not bool(value.strip()))
        if value.strip():
            self.status_hint.setText("البيانات الأساسية جاهزة")
            self.status_hint.setStyleSheet("color:#64748b; font-size:11px;")
        else:
            self.status_hint.setText("اكتب اسمًا واضحًا وسهل التمييز")

    def _validate_ip_live(self, value):
        if not value.strip():
            self._set_error(self.ip_address, False)
            return True
        try:
            ipaddress.ip_address(value.strip())
            self._set_error(self.ip_address, False)
            return True
        except ValueError:
            self._set_error(self.ip_address, True)
            return False

    def _set_error(self, widget, state):
        widget.setProperty("error", state)
        widget.style().unpolish(widget)
        widget.style().polish(widget)

    def _check_serial_duplicate(self):
        serial = self.serial_number.text().strip()
        self._duplicate_serial = False
        if not serial:
            return
        exclude_id = self.camera["id"] if self.camera else 0
        duplicate = fetch_one(
            """SELECT TOP 1 camera_code FROM cameras
               WHERE serial_number = ? AND id <> ? AND is_archived = 0""",
            (serial, exclude_id),
        )
        if duplicate:
            self._duplicate_serial = True
            self.status_hint.setText(f"الرقم التسلسلي مستخدم في الكاميرا {duplicate['camera_code']}")
            self.status_hint.setStyleSheet("color:#dc2626; font-size:11px; font-weight:700;")
        else:
            self.status_hint.setText("الرقم التسلسلي متاح")
            self.status_hint.setStyleSheet("color:#64748b; font-size:11px;")

    @staticmethod
    def _row_value(row, key, default=None):
        """قراءة قيمة من dict أو SQLServerRow بنفس الطريقة."""
        try:
            return row[key]
        except (KeyError, IndexError, TypeError, AttributeError):
            try:
                return getattr(row, key)
            except AttributeError:
                return default

    def _load_switches_into_combo(self, selected_building_id=None):
        selected = self.switch_combo.currentData() if hasattr(self, "switch_combo") else None
        self.switch_combo.blockSignals(True)
        self.switch_combo.clear()
        self.switch_combo.addItem("غير مربوط بسويتش", None)
        switches = switches_choices()
        same = [s for s in switches if selected_building_id and s["building_id"] == selected_building_id]
        same_ids = {s["id"] for s in same}
        other = [s for s in switches if s["id"] not in same_ids]
        for s in same:
            self.switch_combo.addItem(f" {s['name']} (نفس المبنى)", s["id"])
        for s in other:
            suffix = f" ({s['building_name']})" if self._row_value(s, "building_name") else ""
            self.switch_combo.addItem(f"{s['name']}{suffix}", s["id"])
        if selected is not None:
            idx = self.switch_combo.findData(selected)
            if idx >= 0:
                self.switch_combo.setCurrentIndex(idx)
        self.switch_combo.blockSignals(False)

    def _building_changed(self):
        if not self._loading:
            self._load_switches_into_combo(self.building.currentData())

    def load_camera(self):
        c = self.camera
        self._loading = True
        self.code.setText(self._row_value(c, "camera_code") or "")
        self.name.setText(self._row_value(c, "name") or "")
        idx = self.building.findData(self._row_value(c, "building_id"))
        if idx >= 0:
            self.building.setCurrentIndex(idx)
        self.location_detail.setText(self._row_value(c, "location_detail") or "")
        self.placement.setCurrentIndex(max(0, self.placement.findData(self._row_value(c, "placement") or "")))
        self.camera_type.setCurrentText(self._row_value(c, "camera_type") or "")
        self.manufacturer.setText(self._row_value(c, "manufacturer") or "")
        self.model.setText(self._row_value(c, "model") or "")
        self.serial_number.setText(self._row_value(c, "serial_number") or "")
        self.ip_address.setText(self._row_value(c, "ip_address") or "")
        self._load_switches_into_combo(self._row_value(c, "building_id"))
        idx = self.switch_combo.findData(self._row_value(c, "switch_id"))
        if idx >= 0:
            self.switch_combo.setCurrentIndex(idx)
        self.switch_port.setText(self._row_value(c, "switch_port") or "")
        set_optional_date(self.install_date, self._row_value(c, "install_date"))
        set_optional_date(self.warranty_end, self._row_value(c, "warranty_end"))
        self.status.setCurrentText(self._row_value(c, "status") or "تعمل")
        self.notes.setPlainText(self._row_value(c, "notes") or "")
        self._loading = False

    def _check_code_duplicate(self, show_message=False):
        code = self.code.text().strip()
        if not code:
            return True
        exclude_id = self._row_value(self.camera, "id", 0) if self.camera else 0
        duplicate = fetch_one(
            "SELECT TOP 1 id FROM cameras WHERE camera_code = ? AND id <> ? AND is_archived = 0",
            (code, exclude_id),
        )
        if duplicate:
            if show_message:
                QMessageBox.warning(self, "كود مكرر", f"كود الكاميرا {code} مستخدم بالفعل.")
            return False
        return True

    def _check_port_duplicate(self, show_message=False):
        switch_id, port = self.switch_combo.currentData(), self.switch_port.text().strip()
        if not switch_id or not port:
            return True
        exclude_id = self._row_value(self.camera, "id", 0) if self.camera else 0
        duplicate = fetch_one(
            "SELECT TOP 1 camera_code FROM cameras "
            "WHERE switch_id = ? AND UPPER(LTRIM(RTRIM(switch_port))) = UPPER(?) "
            "AND id <> ? AND is_archived = 0",
            (switch_id, port, exclude_id),
        )
        if duplicate:
            if show_message:
                QMessageBox.warning(self, "منفذ مستخدم", f"المنفذ {port} مستخدم بالفعل في الكاميرا {duplicate['camera_code']}.")
            return False
        return True

    def validate_and_accept(self):
        if not self._validate_step(0):
            self.tabs.setCurrentIndex(0)
            return
        if not self._validate_step(1):
            self.tabs.setCurrentIndex(1)
            return
        preview = CameraPreviewDialog(
            self, self.get_data(), self.building.currentText(), self.switch_combo.currentText()
        )
        if preview.exec() == QDialog.Accepted:
            self.accept()

    def get_data(self):
        return {
            "camera_code": self.code.text().strip(),
            "name": self.name.text().strip(),
            "building_id": self.building.currentData(),
            "location_detail": self.location_detail.text().strip(),
            "placement": self.placement.currentData() or None,
            "camera_type": self.camera_type.currentText().strip(),
            "manufacturer": self.manufacturer.text().strip(),
            "model": self.model.text().strip(),
            "serial_number": self.serial_number.text().strip(),
            "ip_address": self.ip_address.text().strip(),
            "switch_id": self.switch_combo.currentData(),
            "switch_port": self.switch_port.text().strip() or None,
            "install_date": optional_date_text(self.install_date),
            "warranty_end": optional_date_text(self.warranty_end),
            "status": self.status.currentText(),
            "notes": self.notes.toPlainText().strip(),
        }


class CameraPreviewDialog(QDialog):
    def __init__(self, parent, data, building_text, switch_text):
        super().__init__(parent)
        self.setWindowTitle("مراجعة بيانات الكاميرا قبل الحفظ")
        self.resize(650, 600)
        self.setLayoutDirection(Qt.RightToLeft)
        root = QVBoxLayout(self)
        root.setContentsMargins(20, 18, 20, 16)
        root.setSpacing(12)
        title = QLabel("راجع البيانات قبل الحفظ النهائي")
        title.setStyleSheet("font-size:19px; font-weight:800; color:#0f2338;")
        subtitle = QLabel("إذا كانت البيانات صحيحة اضغط اعتماد وحفظ، أو عدّلها بالضغط على رجوع")
        subtitle.setStyleSheet("color:#64748b; font-size:11px;")
        root.addWidget(title)
        root.addWidget(subtitle)

        def val(value):
            return html.escape(str(value or "غير محدد"))

        rows = [
            ("كود الكاميرا", data.get("camera_code")),
            ("الاسم التعريفي", data.get("name")),
            ("المبنى / الفرع", building_text),
            ("الموضع", data.get("location_detail")),
            ("نوع الكاميرا", data.get("camera_type")),
            ("التثبيت", data.get("placement")),
            ("السويتش", switch_text),
            ("منفذ السويتش", data.get("switch_port")),
            ("عنوان IP", data.get("ip_address")),
            ("الشركة / الموديل", f"{data.get('manufacturer') or 'غير محدد'} / {data.get('model') or 'غير محدد'}"),
            ("الرقم التسلسلي", data.get("serial_number")),
            ("الحالة", data.get("status")),
            ("تاريخ التركيب", data.get("install_date")),
            ("انتهاء الضمان", data.get("warranty_end")),
            ("الملاحظات", data.get("notes")),
        ]
        rows_html = "".join(
            f"<tr><td class='label'>{val(k)}</td><td>{val(v)}</td></tr>" for k, v in rows
        )
        view = QTextEdit()
        view.setReadOnly(True)
        view.setHtml(
            "<style>body { color:#334155; } table { width:100%; } "
            "td { padding:9px; border-bottom:1px solid #e2e8f0; } "
            ".label { color:#64748b; font-weight:700; background:#f8fafc; width:34%; }</style>"
            f"<table>{rows_html}</table>"
        )
        root.addWidget(view, 1)
        buttons = QDialogButtonBox()
        buttons.addButton("رجوع للتعديل", QDialogButtonBox.RejectRole)
        buttons.addButton("اعتماد وحفظ", QDialogButtonBox.AcceptRole)
        buttons.rejected.connect(self.reject)
        buttons.accepted.connect(self.accept)
        root.addWidget(buttons)


# =============================================================
# نافذة إضافة / تعديل بلاغ عطل كاميرا
# =============================================================

class CameraFaultDialog(QDialog):
    """تسجيل بلاغات الأعطال وربطها بقطع الغيار والفنيين."""

    def __init__(self, parent=None, fault=None, camera_id=None):
        super().__init__(parent)
        self.fault = fault

        self.setWindowTitle("تعديل بلاغ العطل" if fault else "تسجيل بلاغ عطل كاميرا")
        self.resize(700, 720)
        self.setLayoutDirection(Qt.RightToLeft)
        self.setStyleSheet(CAMERAS_THEME_STYLE)

        root = QVBoxLayout(self)
        root.setContentsMargins(20, 18, 20, 16)
        root.setSpacing(14)

        header_box = QHBoxLayout()
        icon = QLabel("")
        icon.setStyleSheet("font-size: 26px; color: #c2410c; padding-left: 6px;")
        header_box.addWidget(icon)

        title_vbox = QVBoxLayout()
        title = QLabel("تعديل تفاصيل البلاغ" if fault else "تسجيل بلاغ عطل فني بكاميرا")
        title.setStyleSheet("font-size: 18px; font-weight: 800; color: #0f2338;")
        sub = QLabel("متابعة العطل وصرف قطع الغيار تلقائيًا من المخزن")
        sub.setStyleSheet("font-size: 11px; color: #64748b;")
        title_vbox.addWidget(title)
        title_vbox.addWidget(sub)
        header_box.addLayout(title_vbox)
        header_box.addStretch()

        root.addLayout(header_box)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.NoFrame)

        container = QWidget()
        form = QFormLayout(container)
        form.setLabelAlignment(Qt.AlignRight)
        form.setHorizontalSpacing(14)
        form.setVerticalSpacing(10)

        self.fault_number = QLineEdit()
        self.fault_number.setReadOnly(True)
        self.fault_number.setStyleSheet("background:#f1f5f9; color:#475569; font-weight:700;")

        self.entry_date = QDateEdit()
        self.entry_date.setCalendarPopup(True)
        self.entry_date.setDisplayFormat("yyyy-MM-dd")
        self.entry_date.setDate(QDate.currentDate())

        self.camera = QComboBox()
        self.camera.setEditable(True)
        self.camera.setInsertPolicy(QComboBox.NoInsert)
        self.camera.lineEdit().setPlaceholderText("اختر الكاميرا أو ابحث بالكود والاسم...")

        self.description = QTextEdit()
        self.description.setFixedHeight(75)
        self.description.setPlaceholderText("شرح دقيق للعطل ومظاهره (مثال: توقف الصورة، مشكلة في التغذية)...")

        self.priority = QComboBox()
        self.priority.addItems(PRIORITIES)

        self.status = QComboBox()
        self.status.addItems(FAULT_STATUSES)
        self.status.activated.connect(self.status_chosen)

        self.technician = QComboBox()
        self.technician.setEditable(True)
        self.technician.setInsertPolicy(QComboBox.NoInsert)

        self.repair_actions = QTextEdit()
        self.repair_actions.setFixedHeight(75)
        self.repair_actions.setPlaceholderText("الإجراءات المتخذة للإصلاح...")

        self.spare_part = QComboBox()
        self.spare_part_qty = QSpinBox()
        self.spare_part_qty.setRange(1, 999999)
        self.spare_part_qty.setFixedWidth(110)

        part_box = QHBoxLayout()
        part_box.setContentsMargins(0, 0, 0, 0)
        part_box.addWidget(self.spare_part, 1)
        part_box.addWidget(QLabel("الكمية:"))
        part_box.addWidget(self.spare_part_qty)
        part_container = QWidget()
        part_container.setLayout(part_box)

        self.repair_date = optional_date_edit()

        self.notes = QTextEdit()
        self.notes.setFixedHeight(60)

        rows = [
            ("رقم البلاغ:", self.fault_number),
            ("تاريخ البلاغ:", self.entry_date),
            ("الكاميرا المعطلة *:", self.camera),
            ("وصف العطل *:", self.description),
            ("درجة الأولوية:", self.priority),
            ("حالة البلاغ:", self.status),
            ("الفني المسؤول:", self.technician),
            ("إجراءات الإصلاح:", self.repair_actions),
            ("قطعة الغيار المصروفة:", part_container),
            ("تاريخ إتمام الإصلاح:", with_clear_button(self.repair_date)),
            ("ملاحظات إضافية:", self.notes),
        ]

        for label_text, widget in rows:
            lbl = QLabel(label_text)
            lbl.setStyleSheet("font-weight: 700; color: #475569;")
            form.addRow(lbl, widget)

        scroll.setWidget(container)
        root.addWidget(scroll, 1)

        btn_box = QHBoxLayout()
        save_btn = QPushButton("  حفظ البلاغ")
        save_btn.setObjectName("primaryGradientBtn")
        save_btn.setCursor(Qt.PointingHandCursor)
        save_btn.clicked.connect(self.validate_and_accept)

        cancel_btn = QPushButton("إلغاء")
        cancel_btn.setObjectName("moreOptionsBtn")
        cancel_btn.setCursor(Qt.PointingHandCursor)
        cancel_btn.clicked.connect(self.reject)

        btn_box.addStretch()
        btn_box.addWidget(cancel_btn)
        btn_box.addWidget(save_btn)
        root.addLayout(btn_box)

        self.load_cameras(camera_id or (fault["camera_id"] if fault else None))
        self.load_technicians()
        self.load_spare_parts(fault["spare_part_id"] if fault else None)

        if fault:
            self.load_fault()

    def load_cameras(self, selected_id):
        rows = fetch_all("""
            SELECT id, camera_code, name FROM cameras
            WHERE is_archived = 0 OR id = ?
            ORDER BY camera_code
        """, (selected_id or 0,))

        self.camera.clear()
        self.camera.addItem("اختر الكاميرا...", None)
        for row in rows:
            self.camera.addItem(f"{row['camera_code']} — {row['name']}", row["id"])

        index = self.camera.findData(selected_id) if selected_id else -1
        self.camera.setCurrentIndex(index if index >= 0 else 0)

    def load_technicians(self):
        rows = fetch_all("""
            SELECT DISTINCT LTRIM(RTRIM(technician)) AS technician
            FROM camera_faults
            WHERE technician IS NOT NULL AND LTRIM(RTRIM(technician)) <> ''
            ORDER BY LTRIM(RTRIM(technician))
        """)
        self.technician.clear()
        self.technician.addItem("")
        for row in rows:
            self.technician.addItem(row["technician"])

    def load_spare_parts(self, current_id):
        rows = fetch_all("""
            SELECT id, part_name, part_number, quantity
            FROM spare_parts
            WHERE is_archived = 0 OR id = ?
            ORDER BY part_name
        """, (current_id or 0,))

        self.spare_part.clear()
        self.spare_part.addItem("بدون قطعة غيار", None)
        for part in rows:
            label = f"{part['part_name']} | رصيد متاح: {part['quantity']}"
            if part["part_number"]:
                label += f" | {part['part_number']}"
            self.spare_part.addItem(label, part["id"])

    def load_fault(self):
        f = self.fault
        self.fault_number.setText(f["fault_number"])

        entry = parse_datetime(f["entry_date"])
        if entry:
            self.entry_date.setDate(QDate(entry.year, entry.month, entry.day))

        self.description.setPlainText(f["fault_description"] or "")
        self.repair_actions.setPlainText(f["repair_actions"] or "")
        self.notes.setPlainText(f["notes"] or "")
        self.technician.setCurrentText((f["technician"] or "").strip())

        idx = self.priority.findText(f["priority"] or "")
        if idx >= 0:
            self.priority.setCurrentIndex(idx)

        idx = self.status.findText(f["status"] or "")
        if idx >= 0:
            self.status.setCurrentIndex(idx)

        idx = self.spare_part.findData(f["spare_part_id"])
        if idx >= 0:
            self.spare_part.setCurrentIndex(idx)
            self.spare_part_qty.setValue(max(1, int(f["spare_part_qty"] or 1)))

        set_optional_date(self.repair_date, f["repair_date"])

    def status_chosen(self, *_args):
        if self.status.currentText() in CLOSED_STATUSES and is_unset(self.repair_date):
            self.repair_date.setDate(QDate.currentDate())

    def validate_and_accept(self):
        if not self.camera.currentData():
            QMessageBox.warning(self, "تنبيه", "من فضلك اختر الكاميرا المتعطلة.")
            self.camera.setFocus()
            return
        if not self.description.toPlainText().strip():
            QMessageBox.warning(self, "تنبيه", "من فضلك اكتب وصفًا للعطل.")
            self.description.setFocus()
            return
        self.accept()

    def get_data(self):
        part_id = self.spare_part.currentData()
        return {
            "fault_number": self.fault_number.text(),
            "camera_id": self.camera.currentData(),
            "entry_date": self.entry_date.date().toString("yyyy-MM-dd") + " 00:00:00",
            "fault_description": self.description.toPlainText().strip(),
            "repair_actions": self.repair_actions.toPlainText().strip(),
            "priority": self.priority.currentText(),
            "status": self.status.currentText(),
            "technician": self.technician.currentText().strip(),
            "spare_part_id": part_id,
            "spare_part_qty": self.spare_part_qty.value() if part_id else 0,
            "spare_parts": (
                f"{self.spare_part.currentText().split(' | رصيد متاح:')[0]} × {self.spare_part_qty.value()}"
                if part_id
                else ""
            ),
            "repair_date": (
                None
                if is_unset(self.repair_date)
                else self.repair_date.date().toString("yyyy-MM-dd") + " 00:00:00"
            ),
            "notes": self.notes.toPlainText().strip(),
        }


# =============================================================
# نافذة السجل التاريخي للكاميرا (مواصفات + بلاغات)
# =============================================================

class CameraHistoryDialog(QDialog):
    """ملف الكاميرا: بطاقة مواصفات وسجل البلاغات."""

    def __init__(self, camera_id, parent=None, current_user=None):
        super().__init__(parent)
        self.camera_id = camera_id
        self.current_user = current_user or {}

        self.setWindowTitle("السجل التاريخي والملف الفني للكاميرا")
        self.resize(1080, 680)
        self.setLayoutDirection(Qt.RightToLeft)
        self.setStyleSheet(CAMERAS_THEME_STYLE)

        root = QVBoxLayout(self)
        root.setContentsMargins(18, 16, 18, 16)
        root.setSpacing(12)

        self.hero = QFrame()
        self.hero.setStyleSheet("""
            QFrame {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0, stop:0 #0b1d33, stop:1 #0e4a63);
                border-radius: 12px;
                padding: 16px;
            }
        """)
        hero_layout = QHBoxLayout(self.hero)
        hero_layout.setContentsMargins(16, 12, 16, 12)

        self.hero_info = QLabel("")
        self.hero_info.setTextFormat(Qt.RichText)
        hero_layout.addWidget(self.hero_info, 1)

        root.addWidget(self.hero)

        self.specs_ribbon = QLabel("")
        self.specs_ribbon.setWordWrap(True)
        self.specs_ribbon.setStyleSheet("""
            background: #ffffff;
            border: 1px solid #e2e8f0;
            border-radius: 10px;
            padding: 12px 16px;
            color: #334155;
            font-size: 12px;
        """)
        root.addWidget(self.specs_ribbon)

        caption = QLabel("سجل الأعطال والإصلاحات")
        caption.setStyleSheet("font-weight: 800; color: #0f2338; font-size: 14px;")
        root.addWidget(caption)

        self.faults_table = make_table(
            ["رقم البلاغ", "التاريخ", "وصف العطل", "الأولوية", "الحالة",
             "الفني", "قطع الغيار", "تاريخ الإصلاح"],
            stretch_column=2,
        )
        root.addWidget(wrap_card(self.faults_table), 1)

        footer = QHBoxLayout()

        print_single_btn = QPushButton("  طباعة تقرير الكاميرا")
        print_single_btn.setObjectName("moreOptionsBtn")
        print_single_btn.setCursor(Qt.PointingHandCursor)
        print_single_btn.clicked.connect(self.print_single_camera_report)
        footer.addWidget(print_single_btn)

        footer.addStretch()

        close_btn = QPushButton("إغلاق")
        close_btn.setObjectName("moreOptionsBtn")
        close_btn.setMinimumWidth(100)
        close_btn.clicked.connect(self.accept)
        footer.addWidget(close_btn)

        root.addLayout(footer)

        self.load()

    def load(self):
        camera = fetch_one("""
            SELECT c.*, b.name AS building_name, s.name as switch_name, s.switch_code
            FROM cameras c
            LEFT JOIN buildings b ON c.building_id = b.id
            LEFT JOIN network_switches s ON c.switch_id = s.id
            WHERE c.id = ?
        """, (self.camera_id,))

        if not camera:
            return

        self.camera_data = camera
        self.camera_name = camera["name"]

        status = camera["status"] or "تعمل"
        status_colors = CAMERA_STATUS_COLORS.get(status, ("#065f46", "#ecfdf5", "#a7f3d0"))

        switch_display = "-"
        if camera["switch_name"]:
            switch_display = camera["switch_name"]
            if camera["switch_port"]:
                switch_display += f" — منفذ {camera['switch_port']}"

        archived_badge = (
            " &nbsp;<span style='background:#64748b; color:#fff; border-radius:8px; padding:2px 6px;'>مؤرشفة</span>"
            if camera["is_archived"] else ""
        )

        self.hero_info.setText(
            f"<div style='font-size:20px; font-weight:800; color:#ffffff;'>"
            f" {html.escape(camera['name'])}"
            f"</div>"
            f"<div style='margin-top:6px; font-size:12px; color:#9be7f5;'>"
            f"كود: <b>{html.escape(camera['camera_code'])}</b> &nbsp;|&nbsp; "
            f"المبنى: <b>{html.escape(camera['building_name'] or 'بدون مبنى')}</b> &nbsp;|&nbsp; "
            f"<span style='background-color:{status_colors[1]}; color:{status_colors[0]}; border-radius:10px; padding:2px 8px; font-weight:700;'>"
            f"● {html.escape(status)}</span>"
            f"{archived_badge}"
            f"</div>"
        )

        def v(k):
            return html.escape(str(camera[k] if camera[k] not in (None, "") else "-"))

        self.specs_ribbon.setText(
            f"<b>الموضع:</b> {v('location_detail')} &nbsp;|&nbsp; "
            f"<b>النوع والتثبيت:</b> {v('placement')} {v('camera_type')} &nbsp;|&nbsp; "
            f"<b>الشركة والموديل:</b> {v('manufacturer')} / {v('model')}<br>"
            f"<b>الرقم التسلسلي:</b> {v('serial_number')} &nbsp;|&nbsp; "
            f"<b>عنوان IP:</b> {v('ip_address')} &nbsp;|&nbsp; "
            f"<b>السويتش المتصل:</b> <span style='color:#0891b2; font-weight:700;'>{html.escape(switch_display)}</span><br>"
            f"<b>تاريخ التركيب:</b> {fmt_date(camera['install_date'])} &nbsp;|&nbsp; "
            f"<b>انتهاء الضمان:</b> {fmt_date(camera['warranty_end'])}"
        )

        faults = fetch_all("""
            SELECT fault_number, entry_date, fault_description, priority,
                   status, technician, spare_parts, repair_date, is_archived
            FROM camera_faults
            WHERE camera_id = ?
            ORDER BY entry_date DESC, id DESC
        """, (self.camera_id,))

        self.faults_table.setRowCount(0)
        for f in faults:
            row = self.faults_table.rowCount()
            self.faults_table.insertRow(row)

            desc = (f["fault_description"] or "").replace("\n", " ")
            values = [
                make_item(f["fault_number"], bold=True),
                make_item(fmt_date(f["entry_date"])),
                make_item(desc, tooltip=desc, align_right=True),
                make_item(f["priority"], PRIORITY_COLORS.get(f["priority"])),
                make_item(f["status"], FAULT_STATUS_COLORS.get(f["status"])),
                make_item((f["technician"] or "").strip()),
                make_item(f["spare_parts"]),
                make_item(fmt_date(f["repair_date"])),
            ]
            for col, it in enumerate(values):
                self.faults_table.setItem(row, col, it)

    def print_single_camera_report(self):
        c = getattr(self, "camera_data", None)
        if not c:
            return

        fault_rows = []
        for r in range(self.faults_table.rowCount()):
            fault_rows.append([
                self.faults_table.item(r, col).text()
                for col in range(self.faults_table.columnCount())
            ])

        headers = ["رقم البلاغ", "التاريخ", "وصف العطل", "الأولوية", "الحالة", "الفني", "قطع الغيار", "تاريخ الإصلاح"]
        summary = (
            f"الكاميرا: {c['camera_code']} — {c['name']} | "
            f"المبنى: {c['building_name'] or 'غير محدد'} | "
            f"السويتش: {c['switch_name'] or 'غير مربوط'} | "
            f"الحالة: {c['status']} | إجمالي البلاغات المسجلة: {len(fault_rows)}"
        )

        open_print_page(
            build_print_html(f"السجل الفني للكاميرا {c['camera_code']}", c["name"], headers, fault_rows, summary),
            f"camera_{c['camera_code']}_history",
        )


# =============================================================
# الصفحة الرئيسية (CamerasPage)
# =============================================================

class CamerasPage(QWidget):
    """الصفحة الرئيسية: الكاميرات، السويتشات، وبلاغات الأعطال."""

    def __init__(self, current_user=None):
        super().__init__()
        self.current_user = current_user or {}

        self._camera_export = []
        self._switch_export = []
        self._fault_export = []

        ok, error = ensure_camera_tables()

        self.setup_ui()
        self.apply_permissions()

        self._org_cache = OrgStructureCache.instance()
        self._org_cache.data_changed.connect(self.load_building_filter)
        self._org_cache_connection_active = True

        if not ok:
            QMessageBox.warning(
                self,
                "جداول الكاميرات والسويتشات",
                "تعذر إنشاء جداول الكاميرات والسويتشات تلقائيًا:\n\n" + error,
            )
            return

        # استيراد أولي تلقائي للسويتشات إن كان الجدول فارغًا
        try:
            count = fetch_one("SELECT COUNT(*) as n FROM network_switches")["n"]
            if count == 0:
                seed_network_switches(auto_link=True)
        except Exception:
            pass

        self.load_building_filter()
        self.refresh_all(notify=False)

    def dispose(self):
        """فصل إشارة الكاش قبل حذف الصفحة لتجنب استدعاء عناصر Qt المحذوفة."""
        if not getattr(self, "_org_cache_connection_active", False):
            return
        try:
            self._org_cache.data_changed.disconnect(self.load_building_filter)
        except (RuntimeError, TypeError):
            pass
        self._org_cache_connection_active = False

    # =========================================================
    # الصلاحيات
    # =========================================================

    def has_permission(self, code):
        if not self.current_user:
            return False
        if self.current_user.get("role_code") == "SUPER_ADMIN":
            return True
        return bool(self.current_user.get("permissions", {}).get(code, False))

    def apply_permissions(self):
        can_add = self.has_permission("cameras.add")
        can_reports = self.has_permission("cameras.reports")

        self.add_camera_button.setVisible(can_add)
        self.edit_camera_button.setVisible(self.has_permission("cameras.edit"))
        self.delete_camera_button.setVisible(self.has_permission("cameras.delete"))
        self.add_fault_button.setVisible(can_add)
        self.more_button.setVisible(can_reports or can_add)
        self.act_import.setVisible(can_add)
        self.act_autolink.setVisible(can_add)

        for action in self.report_actions:
            action.setVisible(can_reports)

        self.setEnabled(self.has_permission("cameras.view"))

    # =========================================================
    # بناء الواجهة
    # =========================================================

    def setup_ui(self):
        self.setObjectName("camerasRoot")
        self.setStyleSheet(CAMERAS_THEME_STYLE)
        self.setLayoutDirection(Qt.RightToLeft)

        main = QVBoxLayout(self)
        main.setContentsMargins(24, 20, 24, 20)
        main.setSpacing(14)

        # 1. الرأس
        header = QHBoxLayout()
        header.setSpacing(10)

        title_box = QVBoxLayout()
        title_box.setSpacing(3)

        title_row = QHBoxLayout()
        title_row.setSpacing(8)

        title = QLabel("الكاميرات والسويتشات")
        title.setObjectName("headerTitle")
        title_row.addWidget(title)
        title_row.addStretch()
        title_box.addLayout(title_row)

        subtitle = QLabel("جرد الكاميرات، ربطها بالسويتشات والمنافذ، ومتابعة الأعطال")
        subtitle.setObjectName("headerSubtitle")
        title_box.addWidget(subtitle)

        header.addLayout(title_box)
        header.addStretch()

        self.add_camera_button = QPushButton("＋  إضافة كاميرا")
        self.add_camera_button.setObjectName("primaryGradientBtn")
        self.add_camera_button.setCursor(Qt.PointingHandCursor)
        self.add_camera_button.clicked.connect(self.add_camera)
        header.addWidget(self.add_camera_button)

        # أزرار التعديل والحذف تتحط في شريط أدوات الكاميرات (تتفعل عند اختيار صف)
        self.edit_camera_button = QPushButton("✎  تعديل")
        self.edit_camera_button.setObjectName("toolbarBtn")
        self.edit_camera_button.setCursor(Qt.PointingHandCursor)
        self.edit_camera_button.clicked.connect(self.edit_camera)

        self.delete_camera_button = QPushButton("حذف")
        self.delete_camera_button.setObjectName("toolbarDangerBtn")
        self.delete_camera_button.setCursor(Qt.PointingHandCursor)
        self.delete_camera_button.clicked.connect(self.delete_camera)

        self.add_fault_button = QPushButton("  تسجيل بلاغ عطل")
        self.add_fault_button.setObjectName("accentWarningBtn")
        self.add_fault_button.setCursor(Qt.PointingHandCursor)
        self.add_fault_button.clicked.connect(lambda: self.add_fault())
        header.addWidget(self.add_fault_button)

        self.more_button = QPushButton("خيارات متقدمة ▾")
        self.more_button.setObjectName("moreOptionsBtn")
        self.more_button.setCursor(Qt.PointingHandCursor)

        menu = QMenu(self.more_button)
        self.act_refresh = menu.addAction("  ⟳ تحديث البيانات")
        self.act_refresh.triggered.connect(self.refresh_all)
        menu.addSeparator()

        self.act_autolink = menu.addAction("   إعادة ربط الكاميرات بالسويتشات تلقائيًا...")
        self.act_autolink.triggered.connect(self.auto_link_all_switches)

        self.act_import = menu.addAction("   استيراد توزيع الكاميرات الأولي (35 موقعًا)...")
        self.act_import.triggered.connect(self.import_initial_distribution)

        self.report_actions = []
        menu.addSeparator()

        for text, handler in (
            ("   تصدير Excel (الكاميرات + السويتشات + البلاغات)", self.export_excel),
            ("   طباعة جرد الكاميرات", self.print_inventory),
            ("   طباعة تقرير السويتشات", self.print_switches),
            ("   طباعة بلاغات الأعطال", self.print_faults),
        ):
            action = menu.addAction(text)
            action.triggered.connect(handler)
            self.report_actions.append(action)

        self.more_button.setMenu(menu)
        header.addWidget(self.more_button)

        main.addLayout(header)

        # 2. بطاقات الإحصائيات
        stats_layout = QHBoxLayout()
        stats_layout.setSpacing(12)

        self.stat_cards = {}
        deck_configs = [
            ("total", "إجمالي الكاميرات", "#0891b2", lambda: self._drilldown_status("الكل")),
            ("working", "تعمل", "#10b981", lambda: self._drilldown_status("تعمل")),
            ("stopped", "متوقفة", "#ef4444", lambda: self._drilldown_status("متوقفة")),
            ("switches", "السويتشات", "#0e7490", lambda: self.tabs.setCurrentIndex(1)),
            ("open_faults", "بلاغات مفتوحة", "#6366f1", self._drilldown_open_faults),
        ]

        for key, lbl, color, fn in deck_configs:
            card = ProMetricCard(lbl, "0", accent=color, on_click=fn)
            stats_layout.addWidget(card, 1)
            self.stat_cards[key] = card

        main.addLayout(stats_layout)

        # 3. التبويبات
        self.tabs = QTabWidget()
        self.tabs.setObjectName("mainTabs")
        self.tabs.addTab(self._build_cameras_tab(), "الكاميرات")
        self.tabs.addTab(self._build_switches_tab(), "السويتشات والشبكة")
        self.tabs.addTab(self._build_faults_tab(), "بلاغات الأعطال")
        main.addWidget(self.tabs, 1)

        # 4. شريط الإشعارات المؤقت
        self.notice = QLabel("")
        self.notice.setObjectName("proNotice")
        self.notice.setVisible(False)
        main.addWidget(self.notice)

        self._notice_timer = QTimer(self)
        self._notice_timer.setSingleShot(True)
        self._notice_timer.setInterval(4000)
        self._notice_timer.timeout.connect(lambda: self.notice.setVisible(False))

        self._search_timer = QTimer(self)
        self._search_timer.setSingleShot(True)
        self._search_timer.setInterval(250)
        self._search_timer.timeout.connect(self._reload_lists)

    def show_notice(self, message):
        self.notice.setText(f"  {message}")
        self.notice.setVisible(True)
        self._notice_timer.start()

    # =========================================================
    # تفاعلات النقر على البطاقات
    # =========================================================

    def _drilldown_status(self, status_text):
        self.tabs.setCurrentIndex(0)
        if status_text == "الكل":
            self.cam_status_filter.setCurrentIndex(0)
        else:
            idx = self.cam_status_filter.findText(status_text)
            if idx >= 0:
                self.cam_status_filter.setCurrentIndex(idx)

    def _drilldown_open_faults(self):
        self.tabs.setCurrentIndex(2)
        self.flt_status_filter.setCurrentIndex(1)  # "المفتوحة فقط"

    # =========================================================
    # التبويب الأول: الكاميرات
    # =========================================================

    def _build_cameras_tab(self):
        tab = QWidget()
        layout = QVBoxLayout(tab)
        layout.setContentsMargins(12, 14, 12, 12)
        layout.setSpacing(10)

        bar = QFrame()
        bar.setObjectName("modernToolbar")
        b_layout = QHBoxLayout(bar)
        b_layout.setContentsMargins(12, 8, 12, 8)
        b_layout.setSpacing(10)

        self.cam_search = QLineEdit()
        self.cam_search.setObjectName("searchBox")
        self.cam_search.setClearButtonEnabled(True)
        self.cam_search.setPlaceholderText("  بحث بالكود، الاسم، الموديل، السيريال، IP، أو السويتش...")
        self.cam_search.textChanged.connect(lambda *_: self._search_timer.start())
        b_layout.addWidget(self.cam_search, 2)

        self.cam_building_filter = QComboBox()
        self.cam_building_filter.setObjectName("filterCombo")
        self.cam_building_filter.currentIndexChanged.connect(self.load_cameras)
        b_layout.addWidget(self.cam_building_filter, 1)

        self.cam_status_filter = QComboBox()
        self.cam_status_filter.setObjectName("filterCombo")
        self.cam_status_filter.addItem("كل الحالات")
        self.cam_status_filter.addItems(CAMERA_STATUSES)
        self.cam_status_filter.currentIndexChanged.connect(self.load_cameras)
        b_layout.addWidget(self.cam_status_filter)

        self.cam_placement_filter = QComboBox()
        self.cam_placement_filter.setObjectName("filterCombo")
        self.cam_placement_filter.addItem("كل التثبيتات")
        self.cam_placement_filter.addItems(PLACEMENTS)
        self.cam_placement_filter.currentIndexChanged.connect(self.load_cameras)
        b_layout.addWidget(self.cam_placement_filter)

        self.show_archived = QCheckBox("المؤرشفة")
        self.show_archived.toggled.connect(self.load_cameras)
        b_layout.addWidget(self.show_archived)

        # فاصل خفيف ثم أزرار الإجراء على الصف المحدد
        sep = QFrame()
        sep.setFrameShape(QFrame.VLine)
        sep.setStyleSheet("color: #e2e8f0;")
        b_layout.addWidget(sep)
        b_layout.addWidget(self.edit_camera_button)
        b_layout.addWidget(self.delete_camera_button)

        layout.addWidget(bar)

        # جدول مبسط: باقي التفاصيل (الموضع، النوع، الموديل، الضمان) في تلميح الاسم
        # وفي ملف الكاميرا (دبل كليك) وفي ملف Excel.
        self.cam_master_table = make_table([
            "الكود", "اسم الكاميرا", "المبنى / الفرع", "السويتش", "IP", "الحالة", "بلاغات",
        ], stretch_column=1)
        self.cam_master_table.doubleClicked.connect(lambda *_: self.show_history())
        self.cam_master_table.customContextMenuRequested.connect(self.camera_menu)
        self.cam_master_table.itemSelectionChanged.connect(self._update_camera_buttons)
        self._update_camera_buttons()

        layout.addWidget(wrap_card(self.cam_master_table), 1)

        self.cam_count = QLabel("")
        self.cam_count.setStyleSheet("color: #64748b; font-size: 11px; font-weight: 600; padding: 2px 4px;")
        layout.addWidget(self.cam_count)

        return tab

    # =========================================================
    # التبويب الثاني: السويتشات والشبكة
    # =========================================================

    def _build_switches_tab(self):
        tab = QWidget()
        layout = QVBoxLayout(tab)
        layout.setContentsMargins(12, 14, 12, 12)
        layout.setSpacing(10)

        bar = QFrame()
        bar.setObjectName("modernToolbar")
        b_layout = QHBoxLayout(bar)
        b_layout.setContentsMargins(12, 8, 12, 8)
        b_layout.setSpacing(10)

        self.sw_search = QLineEdit()
        self.sw_search.setObjectName("searchBox")
        self.sw_search.setClearButtonEnabled(True)
        self.sw_search.setPlaceholderText("  بحث باسم السويتش، الكود، الموديل، المبنى، أو عنوان IP...")
        self.sw_search.textChanged.connect(lambda *_: self._search_timer.start())
        b_layout.addWidget(self.sw_search, 2)

        self.sw_building_filter = QComboBox()
        self.sw_building_filter.setObjectName("filterCombo")
        self.sw_building_filter.currentIndexChanged.connect(self.load_switches)
        b_layout.addWidget(self.sw_building_filter, 1)

        self.sw_role_filter = QComboBox()
        self.sw_role_filter.setObjectName("filterCombo")
        self.sw_role_filter.addItem("كل الأغراض")
        self.sw_role_filter.addItems(SWITCH_ROLES)
        self.sw_role_filter.currentIndexChanged.connect(self.load_switches)
        b_layout.addWidget(self.sw_role_filter)

        self.sw_show_archived = QCheckBox("المؤرشفة")
        self.sw_show_archived.toggled.connect(self.load_switches)
        b_layout.addWidget(self.sw_show_archived)

        self.add_switch_button = QPushButton("＋ إضافة سويتش")
        self.add_switch_button.setObjectName("primaryGradientBtn")
        self.add_switch_button.setCursor(Qt.PointingHandCursor)
        self.add_switch_button.clicked.connect(self.add_switch)
        b_layout.addWidget(self.add_switch_button)

        layout.addWidget(bar)

        self.sw_table = make_table(
            ["ID", "كود السويتش", "اسم السويتش", "المبنى / الموقع", "الموضع التفصيلي",
             "الموديل", "المنافذ", "الغرض والاستخدام", "الكاميرات المتصلة", "عنوان IP", "الحالة التشغيلية"],
            stretch_column=2,
        )
        self.sw_table.setColumnHidden(0, True)
        self.sw_table.doubleClicked.connect(lambda *_: self.edit_switch())
        self.sw_table.customContextMenuRequested.connect(self.switch_menu)

        layout.addWidget(wrap_card(self.sw_table), 1)

        self.sw_count = QLabel("")
        self.sw_count.setStyleSheet("color: #64748b; font-size: 11px; font-weight: 600; padding: 2px 4px;")
        layout.addWidget(self.sw_count)

        return tab

    # =========================================================
    # التبويب الثالث: بلاغات الأعطال
    # =========================================================

    def _build_faults_tab(self):
        tab = QWidget()
        layout = QVBoxLayout(tab)
        layout.setContentsMargins(12, 14, 12, 12)
        layout.setSpacing(10)

        bar = QFrame()
        bar.setObjectName("modernToolbar")
        b_layout = QHBoxLayout(bar)
        b_layout.setContentsMargins(12, 8, 12, 8)
        b_layout.setSpacing(10)

        self.flt_search = QLineEdit()
        self.flt_search.setObjectName("searchBox")
        self.flt_search.setClearButtonEnabled(True)
        self.flt_search.setPlaceholderText("  بحث برقم البلاغ، الكاميرا، وصف العطل، أو الفني...")
        self.flt_search.textChanged.connect(lambda *_: self._search_timer.start())
        b_layout.addWidget(self.flt_search, 2)

        self.flt_status_filter = QComboBox()
        self.flt_status_filter.setObjectName("filterCombo")
        self.flt_status_filter.addItem("كل الحالات")
        self.flt_status_filter.addItem("المفتوحة فقط")
        self.flt_status_filter.addItems(FAULT_STATUSES)
        self.flt_status_filter.currentIndexChanged.connect(self.load_faults)
        b_layout.addWidget(self.flt_status_filter)

        self.flt_priority_filter = QComboBox()
        self.flt_priority_filter.setObjectName("filterCombo")
        self.flt_priority_filter.addItem("كل الأولويات")
        self.flt_priority_filter.addItems(PRIORITIES)
        self.flt_priority_filter.currentIndexChanged.connect(self.load_faults)
        b_layout.addWidget(self.flt_priority_filter)

        self.flt_show_archived = QCheckBox("المؤرشفة")
        self.flt_show_archived.toggled.connect(self.load_faults)
        b_layout.addWidget(self.flt_show_archived)

        layout.addWidget(bar)

        self.flt_table = make_table(
            ["ID", "رقم البلاغ", "تاريخ البلاغ", "الكاميرا", "المبنى", "وصف العطل",
             "الأولوية", "الحالة", "الفني", "قطع الغيار المصروفة", "تاريخ الإصلاح"],
            stretch_column=5,
        )
        self.flt_table.setColumnHidden(0, True)
        self.flt_table.doubleClicked.connect(lambda *_: self.edit_fault())
        self.flt_table.customContextMenuRequested.connect(self.fault_menu)

        layout.addWidget(wrap_card(self.flt_table), 1)

        self.flt_count = QLabel("")
        self.flt_count.setStyleSheet("color: #64748b; font-size: 11px; font-weight: 600; padding: 2px 4px;")
        layout.addWidget(self.flt_count)

        return tab

    # =========================================================
    # تحميل الفلاتر والبيانات
    # =========================================================

    def load_building_filter(self):
        if not (
            qt_is_valid(getattr(self, "cam_building_filter", None))
            and qt_is_valid(getattr(self, "sw_building_filter", None))
        ):
            return

        current = self.cam_building_filter.currentData()
        current_sw = self.sw_building_filter.currentData()

        self.cam_building_filter.blockSignals(True)
        self.cam_building_filter.clear()
        self.cam_building_filter.addItem("كل المباني والفروع", None)

        self.sw_building_filter.blockSignals(True)
        self.sw_building_filter.clear()
        self.sw_building_filter.addItem("كل المباني والفروع", None)

        for unit in building_choices():
            label = f"[{unit['unit_type']}] {unit['name']}"
            cam_index = self.cam_building_filter.count()
            self.cam_building_filter.addItem(label, unit["id"])
            self.cam_building_filter.setItemData(cam_index, unit["name"], Qt.UserRole + 1)

            self.sw_building_filter.addItem(label, unit["id"])

        index = self.cam_building_filter.findData(current) if current else -1
        self.cam_building_filter.setCurrentIndex(max(index, 0))
        index = self.sw_building_filter.findData(current_sw) if current_sw else -1
        self.sw_building_filter.setCurrentIndex(max(index, 0))
        self.cam_building_filter.blockSignals(False)
        self.sw_building_filter.blockSignals(False)

    def _reload_lists(self):
        self.load_cameras()
        self.load_switches()
        self.load_faults()

    def refresh_all(self, notify=True):
        self.load_cameras()
        self.load_switches()
        self.load_faults()
        self.load_stats()
        if notify:
            self.show_notice("تم تحديث البيانات")

    # =========================================================
    # تحميل الكاميرات
    # =========================================================

    def load_cameras(self, *_args):
        clauses = ["c.is_archived = ?"]
        params = [
            CLOSED_STATUSES[0], CLOSED_STATUSES[1],
            1 if self.show_archived.isChecked() else 0,
        ]

        search = self.cam_search.text().strip()
        if search:
            clauses.append(
                "(c.camera_code LIKE ? OR c.name LIKE ? OR c.model LIKE ? "
                "OR c.serial_number LIKE ? OR c.ip_address LIKE ? "
                "OR c.location_detail LIKE ? OR s.name LIKE ? OR s.switch_code LIKE ?)"
            )
            params.extend([f"%{search}%"] * 8)

        if self.cam_status_filter.currentIndex() > 0:
            clauses.append("c.status = ?")
            params.append(self.cam_status_filter.currentText())

        building_id = self.cam_building_filter.currentData()
        if building_id:
            # بعض السجلات القديمة قد تحتوي على معرّف تنظيمي مختلف رغم
            # ارتباطها بنفس اسم المبنى؛ لذلك نطابق بالمعرّف أو الاسم.
            try:
                building_id = int(building_id)
            except (TypeError, ValueError):
                pass
            building_name = self.cam_building_filter.currentData(Qt.UserRole + 1)
            clauses.append("(c.building_id = ? OR b.id = ? OR b.name = ?)")
            params.extend([building_id, building_id, building_name or ""])

        if self.cam_placement_filter.currentIndex() > 0:
            clauses.append("c.placement = ?")
            params.append(self.cam_placement_filter.currentText())

        rows = fetch_all(f"""
            SELECT TOP 3000 c.*, b.name AS building_name,
                s.name as switch_name, s.switch_code,
                (SELECT COUNT(*) FROM camera_faults f
                 WHERE f.camera_id = c.id AND f.is_archived = 0) AS faults_count,
                (SELECT COUNT(*) FROM camera_faults f
                 WHERE f.camera_id = c.id AND f.is_archived = 0
                   AND f.status NOT IN (?, ?)) AS open_count
            FROM cameras c
            LEFT JOIN buildings b ON c.building_id = b.id
            LEFT JOIN network_switches s ON c.switch_id = s.id
            WHERE {' AND '.join(clauses)}
            ORDER BY c.camera_code
        """, params)

        self._camera_export = []
        self.cam_master_table.setUpdatesEnabled(False)
        self.cam_master_table.setRowCount(0)

        for c in rows:
            switch_str = c["switch_name"] or "-"
            if c["switch_name"] and c["switch_port"]:
                switch_str = f"{c['switch_name']} — {c['switch_port']}"

            self._camera_export.append([
                c["camera_code"], c["name"], c["building_name"] or "",
                c["location_detail"] or "", switch_str, c["placement"] or "",
                c["camera_type"] or "", c["manufacturer"] or "", c["model"] or "",
                c["serial_number"] or "", c["ip_address"] or "",
                c["status"] or "", fmt_date(c["install_date"]),
                fmt_date(c["warranty_end"]), c["faults_count"],
            ])

            t_row = self.cam_master_table.rowCount()
            self.cam_master_table.insertRow(t_row)

            st = c["status"] or "تعمل"
            c_colors = CAMERA_STATUS_COLORS.get(st, ("#334155", "#f1f5f9", "#cbd5e1"))

            details = [
                c["location_detail"] or "",
                f"{c['placement'] or ''} {c['camera_type'] or ''}".strip(),
                " / ".join(dict.fromkeys(
                    p for p in (c["manufacturer"], c["model"]) if p
                )),
                f"ضمان حتى {fmt_date(c['warranty_end'])}" if c["warranty_end"] else "",
            ]
            name_tip = "\n".join(d for d in details if d)

            m_values = [
                make_item(c["camera_code"], bold=True),
                make_item(c["name"], bold=True, align_right=True, tooltip=name_tip or None),
                make_item(c["building_name"] or "-", align_right=True),
                make_item(switch_str, align_right=True,
                          tooltip=f"كود: {c['switch_code']}" if c["switch_code"] else "غير مربوط بسويتش"),
                make_item(c["ip_address"] or "-"),
                make_item(f"● {st}", (c_colors[0], None)),  # لون النص فقط بدون خلفية
                make_item(f"{c['open_count']} مفتوح" if c["open_count"] else "-",
                          colors=("#c2410c", None) if c["open_count"] else None),
            ]
            for col_idx, item in enumerate(m_values):
                item.setData(Qt.UserRole, int(c["id"]))
                self.cam_master_table.setItem(t_row, col_idx, item)

        self.cam_master_table.setUpdatesEnabled(True)

        working_total = sum(1 for c in rows if c["status"] == "تعمل")
        total_cameras = len(rows)

        self.cam_count.setText(f"إجمالي النتائج: {total_cameras} كاميرا   •   تعمل: {working_total}")
        self.tabs.setTabText(0, f"الكاميرات ({total_cameras})")

    # =========================================================
    # تحميل السويتشات
    # =========================================================

    def load_switches(self, *_args):
        clauses = ["s.is_archived = ?"]
        params = [1 if self.sw_show_archived.isChecked() else 0]

        search = self.sw_search.text().strip()
        if search:
            clauses.append(
                "(s.switch_code LIKE ? OR s.name LIKE ? OR s.model LIKE ? "
                "OR s.location_detail LIKE ? OR s.ip_address LIKE ? OR b.name LIKE ?)"
            )
            params.extend([f"%{search}%"] * 6)

        b_id = self.sw_building_filter.currentData()
        if b_id:
            clauses.append("s.building_id = ?")
            params.append(b_id)

        if self.sw_role_filter.currentIndex() > 0:
            clauses.append("s.role = ?")
            params.append(self.sw_role_filter.currentText())

        rows = fetch_all(f"""
            SELECT s.*, b.name as building_name,
                (SELECT COUNT(*) FROM cameras c WHERE c.switch_id = s.id AND c.is_archived = 0) as cameras_count
            FROM network_switches s
            LEFT JOIN buildings b ON s.building_id = b.id
            WHERE {' AND '.join(clauses)}
            ORDER BY s.switch_code
        """, params)

        self._switch_export = []
        self.sw_table.setUpdatesEnabled(False)
        self.sw_table.setRowCount(0)

        for s in rows:
            row = self.sw_table.rowCount()
            self.sw_table.insertRow(row)

            role = s["role"] or "شبكة"
            role_colors = SWITCH_ROLE_COLORS.get(role, ("#334155", "#f1f5f9", "#cbd5e1"))

            status = s["status"] or "يعمل"
            if status == "يعمل":
                status_colors = ("#065f46", "#ecfdf5")
            elif status == "تحت الصيانة":
                status_colors = ("#92400e", "#fffbeb")
            else:
                status_colors = ("#991b1b", "#fef2f2")

            cams_count = s["cameras_count"] or 0
            cams_chip = f"{cams_count} كاميرا" if cams_count else "لا توجد كاميرات"
            cams_color = ("#0891b2", "#ecfeff") if cams_count else ("#94a3b8", None)

            values = [
                make_item(s["id"]),
                make_item(s["switch_code"], bold=True),
                make_item(s["name"], bold=True, align_right=True),
                make_item(s["building_name"] or "-", align_right=True),
                make_item(s["location_detail"] or "-", align_right=True),
                make_item(s["model"] or "-"),
                make_item(f"{s['ports_count'] or '-'} Port"),
                make_item(role, role_colors),
                make_item(cams_chip, colors=cams_color, bold=bool(cams_count)),
                make_item(s["ip_address"] or "-"),
                make_item(f"● {status}", colors=status_colors),
            ]

            for col, it in enumerate(values):
                it.setData(Qt.UserRole, int(s["id"]))
                self.sw_table.setItem(row, col, it)

            self._switch_export.append([
                s["switch_code"], s["name"], s["building_name"] or "",
                s["location_detail"] or "", s["model"] or "", s["ports_count"] or "",
                role, cams_count, s["ip_address"] or "", status,
            ])

        self.sw_table.setUpdatesEnabled(True)

        total_sw = len(rows)
        cam_sw = sum(1 for s in rows if "كاميرا" in (s["role"] or ""))
        total_linked = sum(s["cameras_count"] or 0 for s in rows)

        self.sw_count.setText(
            f"إجمالي السويتشات: {total_sw}   •   "
            f"سويتشات الكاميرات: {cam_sw}   •   "
            f"الكاميرات المربوطة: {total_linked}"
        )
        self.tabs.setTabText(1, f"السويتشات والشبكة ({total_sw})")

    def _selected_switch_id(self):
        rows = self.sw_table.selectionModel().selectedRows() if self.sw_table.selectionModel() else []
        if not rows:
            return None
        item = self.sw_table.item(rows[0].row(), 0)
        return int(item.text()) if item else None

    def switch_menu(self, pos):
        sw_id = self._selected_switch_id()
        if not sw_id:
            return

        archived = self.sw_show_archived.isChecked()
        menu = QMenu(self)

        if self.has_permission("cameras.edit") and not archived:
            menu.addAction(" ربط وتعيين كاميرات لهذا السويتش").triggered.connect(
                lambda: self.open_switch_link_dialog(sw_id)
            )
            menu.addAction(" تعديل بيانات السويتش").triggered.connect(self.edit_switch)

        if self.has_permission("cameras.delete"):
            menu.addSeparator()
            menu.addAction(
                " استرجاع من الأرشيف" if archived else " أرشفة السويتش"
            ).triggered.connect(self.toggle_archive_switch)

        if not menu.isEmpty():
            menu.exec(self.sw_table.viewport().mapToGlobal(pos))

    def _next_switch_code(self):
        row = fetch_one("""
            SELECT ISNULL(MAX(CAST(SUBSTRING(switch_code, 4, 10) AS INT)), 0) AS n
            FROM network_switches
            WHERE switch_code LIKE 'SW-%'
              AND ISNUMERIC(SUBSTRING(switch_code, 4, 10)) = 1
        """)
        return f"SW-{(row['n'] or 0) + 1:03d}"

    def add_switch(self):
        if not self.has_permission("cameras.add"):
            QMessageBox.warning(self, "غير مسموح", "ليس لديك صلاحية إضافة سويتشات.")
            return

        dialog = SwitchDialog(self)
        dialog.code.setText(self._next_switch_code())

        if dialog.exec() != QDialog.Accepted:
            return

        data = dialog.get_data()
        data["switch_code"] = self._next_switch_code()

        conn = get_connection()
        try:
            cur = conn.cursor()
            cur.execute("""
                INSERT INTO network_switches (
                    switch_code, name, building_id, location_detail,
                    model, ports_count, role, ip_address, status, notes
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                data["switch_code"], data["name"], data["building_id"],
                data["location_detail"], data["model"], data["ports_count"],
                data["role"], data["ip_address"], data["status"], data["notes"],
            ))
            cur.execute("SELECT CAST(SCOPE_IDENTITY() AS INT) AS id")
            new_id = cur.fetchone()["id"]
            conn.commit()
        except Exception as exc:
            conn.rollback()
            QMessageBox.critical(self, "خطأ", f"تعذر حفظ السويتش:\n{exc}")
            return
        finally:
            conn.close()

        write_audit_log(
            user=self.current_user, action="إضافة", module_name="السويتشات",
            record_id=new_id, description=f"إضافة سويتش {data['switch_code']} - {data['name']}",
            new_values=data,
        )

        self.load_switches()
        self.load_stats()
        self.show_notice(f"تمت إضافة السويتش {data['switch_code']}")

    def edit_switch(self):
        if not self.has_permission("cameras.edit"):
            QMessageBox.warning(self, "غير مسموح", "ليس لديك صلاحية تعديل السويتشات.")
            return

        sw_id = self._selected_switch_id()
        if not sw_id:
            QMessageBox.information(self, "تنبيه", "يرجى تحديد سويتش أولاً.")
            return

        old = fetch_one("SELECT * FROM network_switches WHERE id = ?", (sw_id,))
        if not old:
            return

        dialog = SwitchDialog(self, switch=old)
        if dialog.exec() != QDialog.Accepted:
            return

        data = dialog.get_data()
        conn = get_connection()
        try:
            cur = conn.cursor()
            cur.execute("""
                UPDATE network_switches SET
                    name = ?, building_id = ?, location_detail = ?,
                    model = ?, ports_count = ?, role = ?, ip_address = ?,
                    status = ?, notes = ?, updated_at = GETDATE()
                WHERE id = ?
            """, (
                data["name"], data["building_id"], data["location_detail"],
                data["model"], data["ports_count"], data["role"],
                data["ip_address"], data["status"], data["notes"], sw_id,
            ))
            conn.commit()
        except Exception as exc:
            conn.rollback()
            QMessageBox.critical(self, "خطأ", f"تعذر تعديل السويتش:\n{exc}")
            return
        finally:
            conn.close()

        write_audit_log(
            user=self.current_user, action="تعديل", module_name="السويتشات",
            record_id=sw_id, description=f"تعديل السويتش {old['switch_code']}",
            old_values={k: str(old[k]) for k in old.keys()}, new_values=data,
        )

        self.load_switches()
        self.load_cameras()
        self.show_notice(f"تم تعديل السويتش {old['switch_code']}")

    def toggle_archive_switch(self):
        if not self.has_permission("cameras.delete"):
            QMessageBox.warning(self, "غير مسموح", "ليس لديك صلاحية الأرشفة.")
            return

        sw_id = self._selected_switch_id()
        if not sw_id:
            return

        sw = fetch_one("SELECT * FROM network_switches WHERE id = ?", (sw_id,))
        if not sw:
            return

        restoring = bool(sw["is_archived"])
        conn = get_connection()
        try:
            cur = conn.cursor()
            cur.execute("""
                UPDATE network_switches SET is_archived = ?, archived_at = ?, updated_at = GETDATE()
                WHERE id = ?
            """, (0 if restoring else 1, None if restoring else datetime.now(), sw_id))
            conn.commit()
        except Exception as exc:
            conn.rollback()
            QMessageBox.critical(self, "خطأ", str(exc))
            return
        finally:
            conn.close()

        write_audit_log(
            user=self.current_user,
            action="استرجاع" if restoring else "أرشفة",
            module_name="السويتشات", record_id=sw_id,
            description=f"{'استرجاع' if restoring else 'أرشفة'} السويتش {sw['switch_code']}",
        )

        self.load_switches()
        self.load_cameras()
        self.load_stats()
        self.show_notice(f"تم {'استرجاع' if restoring else 'أرشفة'} {sw['switch_code']}")

    def open_switch_link_dialog(self, switch_id):
        dialog = SwitchLinkCamerasDialog(switch_id, self)
        if dialog.exec() == QDialog.Accepted:
            write_audit_log(
                user=self.current_user, action="تعديل", module_name="السويتشات",
                record_id=switch_id, description="تحديث ربط الكاميرات بالسويتش",
            )
            self.load_switches()
            self.load_cameras()
            self.show_notice("تم تحديث ربط الكاميرات بالسويتش")

    def auto_link_all_switches(self):
        ans = QMessageBox.question(
            self, "ربط الكاميرات بالسويتشات تلقائيًا",
            "سيقوم النظام بفحص جميع الكاميرات غير المربوطة ومطابقتها تلقائيًا بالسويتشات المخصصة لها في نفس المبنى والدور.\n\nهل تريد المتابعة؟",
            QMessageBox.Yes | QMessageBox.No, QMessageBox.Yes,
        )
        if ans != QMessageBox.Yes:
            return

        conn = get_connection()
        try:
            cur = conn.cursor()
            linked = link_cameras_to_switches(cur)
            conn.commit()
            QMessageBox.information(self, "اكتمل الربط التلقائي", f"تم ربط {linked} كاميرا بنجاح بالسويتشات المناسبة لها.")
            self.load_switches()
            self.load_cameras()
        except Exception as exc:
            conn.rollback()
            QMessageBox.critical(self, "خطأ", f"تعذر الربط التلقائي:\n{exc}")
        finally:
            conn.close()

    # =========================================================
    # تحميل بلاغات الأعطال
    # =========================================================

    def load_faults(self, *_args):
        clauses = ["f.is_archived = ?"]
        params = [1 if self.flt_show_archived.isChecked() else 0]

        search = self.flt_search.text().strip()
        if search:
            clauses.append(
                "(f.fault_number LIKE ? OR c.name LIKE ? OR c.camera_code LIKE ? "
                "OR f.fault_description LIKE ? OR f.technician LIKE ?)"
            )
            params.extend([f"%{search}%"] * 5)

        idx = self.flt_status_filter.currentIndex()
        if idx == 1:
            clauses.append("f.status NOT IN (?, ?)")
            params.extend(CLOSED_STATUSES)
        elif idx > 1:
            clauses.append("f.status = ?")
            params.append(self.flt_status_filter.currentText())

        if self.flt_priority_filter.currentIndex() > 0:
            clauses.append("f.priority = ?")
            params.append(self.flt_priority_filter.currentText())

        rows = fetch_all(f"""
            SELECT TOP 3000 f.*, c.name AS camera_name,
                   c.camera_code, b.name AS building_name
            FROM camera_faults f
            JOIN cameras c ON f.camera_id = c.id
            LEFT JOIN buildings b ON c.building_id = b.id
            WHERE {' AND '.join(clauses)}
            ORDER BY f.id DESC
        """, params)

        self._fault_export = []
        self.flt_table.setUpdatesEnabled(False)
        self.flt_table.setRowCount(0)

        for f in rows:
            row = self.flt_table.rowCount()
            self.flt_table.insertRow(row)

            desc = (f["fault_description"] or "").replace("\n", " ").strip()
            short_desc = desc if len(desc) <= 75 else desc[:75] + "…"

            values = [
                make_item(f["id"]),
                make_item(f["fault_number"], bold=True),
                make_item(fmt_date(f["entry_date"])),
                make_item(f"{f['camera_code']} — {f['camera_name']}", align_right=True),
                make_item(f["building_name"] or "-", align_right=True),
                make_item(short_desc, tooltip=desc, align_right=True),
                make_item(f["priority"], PRIORITY_COLORS.get(f["priority"])),
                make_item(f["status"], FAULT_STATUS_COLORS.get(f["status"])),
                make_item((f["technician"] or "").strip()),
                make_item(f["spare_parts"] or "-"),
                make_item(fmt_date(f["repair_date"])),
            ]

            for col, it in enumerate(values):
                self.flt_table.setItem(row, col, it)

            self._fault_export.append([
                f["fault_number"], fmt_date(f["entry_date"]),
                f["camera_code"], f["camera_name"], f["building_name"] or "",
                desc, f["priority"] or "", f["status"] or "",
                (f["technician"] or "").strip(), f["repair_actions"] or "",
                f["spare_parts"] or "", fmt_date(f["repair_date"]),
            ])

        self.flt_table.setUpdatesEnabled(True)

        open_in_list = sum(1 for f in rows if f["status"] not in CLOSED_STATUSES)
        self.flt_count.setText(f"إجمالي المعروض: {len(rows)} بلاغ (منها {open_in_list} مفتوح)")
        self.tabs.setTabText(2, f"بلاغات الأعطال ({open_in_list} مفتوح)" if open_in_list else "بلاغات الأعطال")

    # =========================================================
    # الإحصائيات
    # =========================================================

    def load_stats(self):
        stats = fetch_one("""
            SELECT COUNT(*) AS total,
                   SUM(CASE WHEN status = N'تعمل' THEN 1 ELSE 0 END) AS working,
                   SUM(CASE WHEN status = N'متوقفة' THEN 1 ELSE 0 END) AS stopped
            FROM cameras WHERE is_archived = 0
        """)
        open_faults = fetch_one("""
            SELECT COUNT(*) AS n FROM camera_faults
            WHERE is_archived = 0 AND status NOT IN (?, ?)
        """, CLOSED_STATUSES)["n"]
        switches = fetch_one("""
            SELECT COUNT(*) AS n FROM network_switches WHERE is_archived = 0
        """)["n"]

        self.stat_cards["total"].set_value(stats["total"] or 0, "")
        self.stat_cards["working"].set_value(stats["working"] or 0, "")
        self.stat_cards["stopped"].set_value(stats["stopped"] or 0, "")
        self.stat_cards["switches"].set_value(switches or 0, "")
        self.stat_cards["open_faults"].set_value(open_faults or 0, "")

    # =========================================================
    # اختيار الصفوف والقوائم المنبثقة
    # =========================================================

    def _update_camera_buttons(self):
        """تعديل/حذف يشتغلوا فقط لما يكون فيه كاميرا مختارة وغير مؤرشفة."""
        has_row = self._selected_camera_id() is not None
        enabled = has_row and not self.show_archived.isChecked()
        self.edit_camera_button.setEnabled(enabled)
        self.delete_camera_button.setEnabled(enabled)

    def _selected_camera_id(self):
        rows = self.cam_master_table.selectionModel().selectedRows() if self.cam_master_table.selectionModel() else []
        if not rows:
            return None
        item = self.cam_master_table.item(rows[0].row(), 0)
        val = item.data(Qt.UserRole) if item else None
        return int(val) if val else None

    def _selected_fault_id(self):
        rows = self.flt_table.selectionModel().selectedRows() if self.flt_table.selectionModel() else []
        if not rows:
            return None
        item = self.flt_table.item(rows[0].row(), 0)
        return int(item.text()) if item else None

    def camera_menu(self, pos):
        camera_id = self._selected_camera_id()
        if not camera_id:
            return

        archived = self.show_archived.isChecked()
        menu = QMenu(self)

        menu.addAction(" السجل التاريخي والملف الفني").triggered.connect(self.show_history)

        if self.has_permission("cameras.add") and not archived:
            menu.addAction(" تسجيل بلاغ عطل لهذه الكاميرا").triggered.connect(
                lambda: self.add_fault(camera_id)
            )

        if self.has_permission("cameras.edit") and not archived:
            menu.addAction(" تعديل بيانات الكاميرا").triggered.connect(self.edit_camera)
        if self.has_permission("cameras.delete") and not archived:
            menu.addAction("حذف نهائي").triggered.connect(self.delete_camera)

        if self.has_permission("cameras.delete"):
            menu.addSeparator()
            menu.addAction(
                " استرجاع من الأرشيف" if archived else " أرشفة الكاميرا"
            ).triggered.connect(self.toggle_archive_camera)

        menu.exec(self.cam_master_table.viewport().mapToGlobal(pos))

    def fault_menu(self, pos):
        item = self.flt_table.itemAt(pos)
        if not item:
            return

        self.flt_table.selectRow(item.row())
        archived = self.flt_show_archived.isChecked()

        menu = QMenu(self)
        if self.has_permission("cameras.edit") and not archived:
            quick = menu.addMenu(" تغيير حالة البلاغ بسرعة")
            for st in ("قيد الفحص", "قيد الإصلاح", "بانتظار قطعة غيار", STATUS_REPAIRED, STATUS_CLOSED):
                quick.addAction(st).triggered.connect(
                    lambda checked=False, s=st: self.quick_status(s)
                )

            menu.addAction(" تعديل البلاغ الكامل").triggered.connect(self.edit_fault)

        if self.has_permission("cameras.delete"):
            menu.addSeparator()
            menu.addAction(
                " استرجاع من الأرشيف" if archived else " أرشفة البلاغ"
            ).triggered.connect(self.toggle_archive_fault)

        if not menu.isEmpty():
            menu.exec(self.flt_table.viewport().mapToGlobal(pos))

    # =========================================================
    # إجراءات الكاميرات (إضافة / تعديل / حذف / أرشفة / سجل)
    # =========================================================

    def _next_camera_code(self):
        row = fetch_one("""
            SELECT ISNULL(MAX(CAST(SUBSTRING(camera_code, 5, 10) AS INT)), 0) AS n
            FROM cameras
            WHERE camera_code LIKE 'CAM-%'
              AND ISNUMERIC(SUBSTRING(camera_code, 5, 10)) = 1
        """)
        return f"CAM-{(row['n'] or 0) + 1:04d}"

    def add_camera(self):
        if not self.has_permission("cameras.add"):
            QMessageBox.warning(self, "غير مسموح", "ليس لديك صلاحية إضافة الكاميرات.")
            return

        dialog = CameraDialog(self)
        dialog.code.setText(self._next_camera_code())

        if dialog.exec() != QDialog.Accepted:
            return

        data = dialog.get_data()
        data["camera_code"] = self._next_camera_code()

        connection = get_connection()
        try:
            cursor = connection.cursor()
            cursor.execute("""
                INSERT INTO cameras (
                    camera_code, name, building_id, location_detail,
                    placement, camera_type, manufacturer, model,
                    serial_number, ip_address, switch_id, switch_port,
                    install_date, warranty_end, status, notes
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                data["camera_code"], data["name"], data["building_id"],
                data["location_detail"], data["placement"], data["camera_type"],
                data["manufacturer"], data["model"], data["serial_number"],
                data["ip_address"], data["switch_id"], data["switch_port"],
                data["install_date"], data["warranty_end"], data["status"],
                data["notes"],
            ))
            cursor.execute("SELECT CAST(SCOPE_IDENTITY() AS INT) AS id")
            new_id = cursor.fetchone()["id"]
            connection.commit()
        except Exception as exc:
            connection.rollback()
            QMessageBox.critical(self, "خطأ", f"تعذر حفظ الكاميرا:\n{exc}")
            return
        finally:
            connection.close()

        write_audit_log(
            user=self.current_user, action="إضافة", module_name="الكاميرات",
            record_id=new_id,
            description=f"تمت إضافة الكاميرا {data['camera_code']} - {data['name']}",
            new_values=data,
        )

        self.refresh_all(notify=False)
        self.show_notice(f"تمت إضافة الكاميرا {data['camera_code']} بنجاح")

    def edit_camera(self):
        if not self.has_permission("cameras.edit"):
            QMessageBox.warning(self, "غير مسموح", "ليس لديك صلاحية تعديل الكاميرات.")
            return

        camera_id = self._selected_camera_id()
        if not camera_id:
            QMessageBox.information(self, "تنبيه", "من فضلك حدد كاميرا أولاً.")
            return

        old = fetch_one("SELECT * FROM cameras WHERE id = ?", (camera_id,))
        if not old:
            return

        dialog = CameraDialog(self, camera=old)
        if dialog.exec() != QDialog.Accepted:
            return

        data = dialog.get_data()
        connection = get_connection()
        try:
            cursor = connection.cursor()
            cursor.execute("""
                UPDATE cameras SET
                    name = ?, building_id = ?, location_detail = ?,
                    placement = ?, camera_type = ?, manufacturer = ?, model = ?,
                    serial_number = ?, ip_address = ?, switch_id = ?, switch_port = ?,
                    install_date = ?, warranty_end = ?, status = ?, notes = ?,
                    updated_at = GETDATE()
                WHERE id = ?
            """, (
                data["name"], data["building_id"], data["location_detail"],
                data["placement"], data["camera_type"], data["manufacturer"],
                data["model"], data["serial_number"], data["ip_address"],
                data["switch_id"], data["switch_port"],
                data["install_date"], data["warranty_end"], data["status"],
                data["notes"], camera_id,
            ))
            connection.commit()
        except Exception as exc:
            connection.rollback()
            QMessageBox.critical(self, "خطأ", f"تعذر تعديل الكاميرا:\n{exc}")
            return
        finally:
            connection.close()

        write_audit_log(
            user=self.current_user, action="تعديل", module_name="الكاميرات",
            record_id=camera_id,
            description=f"تم تعديل الكاميرا {old['camera_code']}",
            old_values={k: str(old[k]) for k in old.keys()},
            new_values=data,
        )

        self.refresh_all(notify=False)
        self.show_notice(f"تم تعديل بيانات الكاميرا {old['camera_code']}")

    def delete_camera(self):
        """حذف نهائي آمن؛ السجلات المرتبطة تمنع الحذف وتبقى الأرشفة متاحة."""
        if not self.has_permission("cameras.delete"):
            QMessageBox.warning(self, "غير مسموح", "ليس لديك صلاحية حذف الكاميرات.")
            return
        camera_id = self._selected_camera_id()
        if not camera_id:
            QMessageBox.information(self, "تنبيه", "اختر كاميرا أولًا.")
            return
        camera = fetch_one("SELECT * FROM cameras WHERE id = ?", (camera_id,))
        if not camera:
            return
        counts = fetch_one(
            "SELECT (SELECT COUNT(*) FROM camera_faults WHERE camera_id = ?) AS faults_count, "
            "(SELECT COUNT(*) FROM camera_maintenance WHERE camera_id = ?) AS maintenance_count",
            (camera_id, camera_id),
        )
        if counts["faults_count"] or counts["maintenance_count"]:
            QMessageBox.warning(
                self, "لا يمكن الحذف النهائي",
                "الكاميرا مرتبطة ببلاغات أو سجلات صيانة. استخدم أرشفة الكاميرا للحفاظ على التاريخ.",
            )
            return
        answer = QMessageBox.warning(
            self, "تأكيد الحذف النهائي",
            f"سيتم حذف الكاميرا {camera['camera_code']} نهائيًا ولا يمكن استرجاعها.\nهل تريد المتابعة؟",
            QMessageBox.Yes | QMessageBox.No, QMessageBox.No,
        )
        if answer != QMessageBox.Yes:
            return
        connection = get_connection()
        try:
            connection.cursor().execute("DELETE FROM cameras WHERE id = ?", (camera_id,))
            connection.commit()
        except Exception as exc:
            connection.rollback()
            QMessageBox.critical(self, "خطأ", f"تعذر حذف الكاميرا:\n{exc}")
            return
        finally:
            connection.close()
        write_audit_log(
            user=self.current_user, action="حذف نهائي", module_name="الكاميرات",
            record_id=camera_id,
            description=f"تم حذف الكاميرا نهائيًا {camera['camera_code']}",
            old_values={k: str(camera[k]) for k in camera.keys()},
        )
        self.refresh_all(notify=False)
        self.show_notice(f"تم حذف الكاميرا {camera['camera_code']} نهائيًا")

    def toggle_archive_camera(self):
        if not self.has_permission("cameras.delete"):
            QMessageBox.warning(self, "غير مسموح", "ليس لديك صلاحية الأرشفة.")
            return

        camera_id = self._selected_camera_id()
        if not camera_id:
            QMessageBox.information(self, "تنبيه", "اختر كاميرا أولاً.")
            return

        camera = fetch_one("SELECT * FROM cameras WHERE id = ?", (camera_id,))
        if not camera:
            return

        restoring = bool(camera["is_archived"])
        if not restoring:
            open_count = fetch_one("""
                SELECT COUNT(*) AS n FROM camera_faults
                WHERE camera_id = ? AND is_archived = 0 AND status NOT IN (?, ?)
            """, (camera_id, *CLOSED_STATUSES))["n"]

            if open_count:
                QMessageBox.warning(
                    self, "لا يمكن الأرشفة",
                    f"للكاميرا {open_count} بلاغ مفتوح. يُرجى إغلاق البلاغات أولاً.",
                )
                return

            ans = QMessageBox.question(
                self, "تأكيد الأرشفة",
                f"سيتم نقل الكاميرا {camera['camera_code']} إلى الأرشيف مع الاحتفاظ بكافة بياناتها وسجلاتها.\nهل تريد المتابعة؟",
                QMessageBox.Yes | QMessageBox.No, QMessageBox.No,
            )
            if ans != QMessageBox.Yes:
                return

        connection = get_connection()
        try:
            connection.cursor().execute("""
                UPDATE cameras SET is_archived = ?, archived_at = ?, updated_at = GETDATE()
                WHERE id = ?
            """, (0 if restoring else 1, None if restoring else datetime.now(), camera_id))
            connection.commit()
        except Exception as exc:
            connection.rollback()
            QMessageBox.critical(self, "خطأ", str(exc))
            return
        finally:
            connection.close()

        write_audit_log(
            user=self.current_user,
            action="استرجاع" if restoring else "أرشفة",
            module_name="الكاميرات", record_id=camera_id,
            description=f"{'استرجاع' if restoring else 'أرشفة'} الكاميرا {camera['camera_code']}",
        )

        self.refresh_all(notify=False)
        self.show_notice(f"تم {'استرجاع' if restoring else 'أرشفة'} {camera['camera_code']}")

    def show_history(self):
        camera_id = self._selected_camera_id()
        if camera_id is not None:
            self._open_history(camera_id)
        else:
            QMessageBox.information(self, "تنبيه", "يرجى تحديد كاميرا من الجدول أولاً.")

    def _open_history(self, camera_id):
        CameraHistoryDialog(camera_id, self, current_user=self.current_user).exec()
        self.refresh_all(notify=False)

    # =========================================================
    # إجراءات بلاغات الأعطال
    # =========================================================

    def _next_fault_number(self):
        year = datetime.now().year
        prefix = f"CFLT-{year}-"
        row = fetch_one("""
            SELECT ISNULL(MAX(CAST(SUBSTRING(fault_number, ?, 10) AS INT)), 0) AS n
            FROM camera_faults
            WHERE fault_number LIKE ?
              AND ISNUMERIC(SUBSTRING(fault_number, ?, 10)) = 1
        """, (len(prefix) + 1, prefix + "%", len(prefix) + 1))
        return f"{prefix}{(row['n'] or 0) + 1:04d}"

    @staticmethod
    def _sync_spare_part(cursor, fault_number, new_part_id, new_qty):
        note = f"صرف تلقائي لعطل كاميرا {fault_number}"
        cursor.execute("""
            SELECT TOP 1 id, spare_part_id, quantity
            FROM spare_part_movements
            WHERE movement_type = ? AND notes = ?
            ORDER BY id DESC
        """, ("صرف", note))
        old = cursor.fetchone()

        if old:
            cursor.execute("""
                UPDATE spare_parts
                SET quantity = quantity + ?, updated_at = CURRENT_TIMESTAMP
                WHERE id = ?
            """, (old["quantity"], old["spare_part_id"]))

            if cursor.rowcount != 1:
                raise ValueError("تعذر إرجاع قطعة الغيار السابقة للمخزن.")

            cursor.execute("DELETE FROM spare_part_movements WHERE id = ?", (old["id"],))

        if new_part_id and new_qty > 0:
            cursor.execute("""
                SELECT part_name, quantity FROM spare_parts WHERE id = ?
            """, (new_part_id,))
            part = cursor.fetchone()

            if not part:
                raise ValueError("قطعة الغيار المختارة غير موجودة بالمخزون.")

            if int(part["quantity"]) < new_qty:
                raise ValueError(
                    f"الرصيد المتاح للقطعة «{part['part_name']}» ({part['quantity']}) لا يكفي للكمية المطلوبة ({new_qty})."
                )

            cursor.execute("""
                UPDATE spare_parts
                SET quantity = quantity - ?, updated_at = CURRENT_TIMESTAMP
                WHERE id = ? AND quantity >= ?
            """, (new_qty, new_part_id, new_qty))

            if cursor.rowcount != 1:
                raise ValueError("تعذر خصم قطعة الغيار من المخزون.")

            cursor.execute("""
                INSERT INTO spare_part_movements (spare_part_id, movement_type, quantity, notes)
                VALUES (?, ?, ?, ?)
            """, (new_part_id, "صرف", new_qty, note))

        return old

    @staticmethod
    def _refresh_camera_status(cursor, camera_id):
        cursor.execute("""
            SELECT COUNT(*) AS n FROM camera_faults
            WHERE camera_id = ? AND is_archived = 0 AND status NOT IN (?, ?)
        """, (camera_id, *CLOSED_STATUSES))
        open_count = cursor.fetchone()["n"]

        if open_count:
            cursor.execute("""
                UPDATE cameras SET status = ?, updated_at = GETDATE()
                WHERE id = ? AND status = ?
            """, ("تحت الصيانة", camera_id, "تعمل"))
        else:
            cursor.execute("""
                UPDATE cameras SET status = ?, updated_at = GETDATE()
                WHERE id = ? AND status = ?
            """, ("تعمل", camera_id, "تحت الصيانة"))

    def add_fault(self, camera_id=None):
        if not self.has_permission("cameras.add"):
            QMessageBox.warning(self, "غير مسموح", "ليس لديك صلاحية تسجيل البلاغات.")
            return

        if isinstance(camera_id, bool):
            camera_id = None

        dialog = CameraFaultDialog(self, camera_id=camera_id)
        dialog.fault_number.setText(self._next_fault_number())

        if dialog.exec() != QDialog.Accepted:
            return

        data = dialog.get_data()
        data["fault_number"] = self._next_fault_number()

        prev_faults = fetch_one("""
            SELECT COUNT(*) AS n FROM camera_faults
            WHERE camera_id = ? AND is_archived = 0 AND entry_date >= DATEADD(day, ?, GETDATE())
        """, (data["camera_id"], -FREQUENT_FAULTS_DAYS))["n"]

        if prev_faults >= FREQUENT_FAULTS_MIN - 1:
            ans = QMessageBox.question(
                self, "كاميرا متكررة الأعطال",
                f"تنبيه: لهذه الكاميرا {prev_faults} أعطال خلال آخر {FREQUENT_FAULTS_DAYS} يوم.\nهل ترغب في تسجيل هذا البلاغ الإضافي؟",
                QMessageBox.Yes | QMessageBox.No, QMessageBox.Yes,
            )
            if ans != QMessageBox.Yes:
                return

        connection = get_connection()
        try:
            cursor = connection.cursor()
            self._sync_spare_part(
                cursor, data["fault_number"],
                data["spare_part_id"], data["spare_part_qty"],
            )

            cursor.execute("""
                INSERT INTO camera_faults (
                    fault_number, camera_id, entry_date, fault_description,
                    repair_actions, spare_part_id, spare_part_qty,
                    spare_parts, technician, priority, status,
                    repair_date, notes
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                data["fault_number"], data["camera_id"], data["entry_date"],
                data["fault_description"], data["repair_actions"],
                data["spare_part_id"], data["spare_part_qty"],
                data["spare_parts"], data["technician"], data["priority"],
                data["status"], data["repair_date"], data["notes"],
            ))
            cursor.execute("SELECT CAST(SCOPE_IDENTITY() AS INT) AS id")
            new_id = cursor.fetchone()["id"]

            self._refresh_camera_status(cursor, data["camera_id"])
            connection.commit()
        except ValueError as exc:
            connection.rollback()
            QMessageBox.warning(self, "تنبيه المخزون", str(exc))
            return
        except Exception as exc:
            connection.rollback()
            QMessageBox.critical(self, "خطأ", f"تعذر حفظ البلاغ:\n{exc}")
            return
        finally:
            connection.close()

        write_audit_log(
            user=self.current_user, action="إضافة",
            module_name="أعطال الكاميرات", record_id=new_id,
            description=f"تم تسجيل بلاغ العطل {data['fault_number']}",
            new_values=data,
        )

        self.refresh_all(notify=False)
        self.show_notice(f"تم تسجيل البلاغ {data['fault_number']} بنجاح")

    def edit_fault(self):
        if not self.has_permission("cameras.edit"):
            QMessageBox.warning(self, "غير مسموح", "ليس لديك صلاحية تعديل البلاغات.")
            return

        fault_id = self._selected_fault_id()
        if not fault_id:
            QMessageBox.information(self, "تنبيه", "يرجى تحديد بلاغ أولاً.")
            return

        old = fetch_one("SELECT * FROM camera_faults WHERE id = ?", (fault_id,))
        if not old:
            return

        if old["is_archived"]:
            QMessageBox.information(self, "بلاغ مؤرشف", "استرجع البلاغ من الأرشيف قبل تعديله.")
            return

        dialog = CameraFaultDialog(self, fault=old)
        if dialog.exec() != QDialog.Accepted:
            return

        data = dialog.get_data()
        connection = get_connection()
        try:
            cursor = connection.cursor()
            self._sync_spare_part(
                cursor, old["fault_number"],
                data["spare_part_id"], data["spare_part_qty"],
            )

            cursor.execute("""
                UPDATE camera_faults SET
                    camera_id = ?, entry_date = ?, fault_description = ?,
                    repair_actions = ?, spare_part_id = ?, spare_part_qty = ?,
                    spare_parts = ?, technician = ?, priority = ?, status = ?,
                    repair_date = ?, notes = ?, updated_at = GETDATE()
                WHERE id = ?
            """, (
                data["camera_id"], data["entry_date"],
                data["fault_description"], data["repair_actions"],
                data["spare_part_id"], data["spare_part_qty"],
                data["spare_parts"], data["technician"], data["priority"],
                data["status"], data["repair_date"], data["notes"], fault_id,
            ))

            self._refresh_camera_status(cursor, data["camera_id"])
            if old["camera_id"] != data["camera_id"]:
                self._refresh_camera_status(cursor, old["camera_id"])

            connection.commit()
        except ValueError as exc:
            connection.rollback()
            QMessageBox.warning(self, "تنبيه المخزون", str(exc))
            return
        except Exception as exc:
            connection.rollback()
            QMessageBox.critical(self, "خطأ", f"تعذر تعديل البلاغ:\n{exc}")
            return
        finally:
            connection.close()

        write_audit_log(
            user=self.current_user, action="تعديل",
            module_name="أعطال الكاميرات", record_id=fault_id,
            description=f"تم تعديل البلاغ {old['fault_number']}",
            old_values={k: str(old[k]) for k in old.keys()},
            new_values=data,
        )

        self.refresh_all(notify=False)
        self.show_notice(f"تم تعديل البلاغ {old['fault_number']}")

    def quick_status(self, status):
        if not self.has_permission("cameras.edit"):
            return

        fault_id = self._selected_fault_id()
        if not fault_id:
            return

        old = fetch_one("SELECT * FROM camera_faults WHERE id = ?", (fault_id,))
        if not old:
            return

        connection = get_connection()
        try:
            cursor = connection.cursor()
            if status in CLOSED_STATUSES:
                cursor.execute("""
                    UPDATE camera_faults
                    SET status = ?, repair_date = COALESCE(repair_date, GETDATE()),
                        updated_at = GETDATE()
                    WHERE id = ?
                """, (status, fault_id))
            else:
                cursor.execute("""
                    UPDATE camera_faults SET status = ?, updated_at = GETDATE()
                    WHERE id = ?
                """, (status, fault_id))

            self._refresh_camera_status(cursor, old["camera_id"])
            connection.commit()
        except Exception as exc:
            connection.rollback()
            QMessageBox.critical(self, "خطأ", str(exc))
            return
        finally:
            connection.close()

        write_audit_log(
            user=self.current_user, action="تعديل",
            module_name="أعطال الكاميرات", record_id=fault_id,
            description=f"تغيير حالة البلاغ {old['fault_number']} من {old['status']} إلى {status}",
            old_values={"status": old["status"]}, new_values={"status": status},
        )

        self.refresh_all(notify=False)
        self.show_notice(f"تم تحديث حالة البلاغ إلى: {status}")

    def toggle_archive_fault(self):
        if not self.has_permission("cameras.delete"):
            QMessageBox.warning(self, "غير مسموح", "ليس لديك صلاحية الأرشفة.")
            return

        fault_id = self._selected_fault_id()
        if not fault_id:
            QMessageBox.information(self, "تنبيه", "اختر بلاغًا أولاً.")
            return

        fault = fetch_one("SELECT * FROM camera_faults WHERE id = ?", (fault_id,))
        if not fault:
            return

        restoring = bool(fault["is_archived"])
        if not restoring:
            ans = QMessageBox.question(
                self, "تأكيد الأرشفة",
                f"سيتم أرشفة البلاغ {fault['fault_number']} والاحتفاظ به في السجل التاريخي.\nهل تريد المتابعة؟",
                QMessageBox.Yes | QMessageBox.No, QMessageBox.No,
            )
            if ans != QMessageBox.Yes:
                return

        connection = get_connection()
        try:
            cursor = connection.cursor()
            cursor.execute("""
                UPDATE camera_faults SET is_archived = ?, archived_at = ?, updated_at = GETDATE()
                WHERE id = ?
            """, (0 if restoring else 1, None if restoring else datetime.now(), fault_id))

            self._refresh_camera_status(cursor, fault["camera_id"])
            connection.commit()
        except Exception as exc:
            connection.rollback()
            QMessageBox.critical(self, "خطأ", str(exc))
            return
        finally:
            connection.close()

        write_audit_log(
            user=self.current_user,
            action="استرجاع" if restoring else "أرشفة",
            module_name="أعطال الكاميرات", record_id=fault_id,
            description=f"{'استرجاع' if restoring else 'أرشفة'} البلاغ {fault['fault_number']}",
        )

        self.refresh_all(notify=False)
        self.show_notice(f"تم {'استرجاع' if restoring else 'أرشفة'} {fault['fault_number']}")

    # =========================================================
    # استيراد التوزيع الأولي
    # =========================================================

    @staticmethod
    def _norm(text):
        return " ".join(str(text or "").split())

    def import_initial_distribution(self):
        if not self.has_permission("cameras.add"):
            QMessageBox.warning(self, "غير مسموح", "ليس لديك صلاحية إضافة الكاميرات.")
            return

        try:
            from ui.camera_seed import expand_tree, building_name_for
        except ImportError:
            QMessageBox.critical(
                self, "ملف ناقص", "ملف camera_seed.py غير موجود بجوار cameras.py داخل مجلد ui."
            )
            return

        planned = expand_tree()
        existing = {self._norm(r["name"]) for r in fetch_all("SELECT name FROM cameras")}
        new_rows = [r for r in planned if self._norm(r["name"]) not in existing]

        if not new_rows:
            QMessageBox.information(self, "تم مسبقًا", "كل كاميرات التوزيع الأولي مسجلة بالفعل بالمنظومة.")
            return

        buildings = {
            self._norm(b["name"]): b["id"]
            for b in fetch_all("SELECT id, name FROM buildings")
        }
        locations = sorted({building_name_for(r["location"]) for r in new_rows})
        missing = [l for l in locations if self._norm(l) not in buildings]

        msg = (
            f"سيتم إضافة {len(new_rows)} كاميرا من {len(locations)} موقعًا "
            f"(المتخطى لتسجيله مسبقًا: {len(planned) - len(new_rows)}).\n\n"
        )

        create_missing = False
        if missing:
            msg += (
                f"{len(missing)} موقع غير مسجل بنفس الاسم في المباني:\n"
                + "، ".join(missing[:10])
                + ("، ..." if len(missing) > 10 else "")
                + "\n\nنعم = إنشاء هذه المواقع في المباني وربط الكاميرات بها\n"
                "لا = إضافة الكاميرات بدون ربط المبنى\n"
                "إلغاء = إيقاف الاستيراد"
            )
            box = QMessageBox.question(
                self, "استيراد التوزيع الأولي", msg,
                QMessageBox.Yes | QMessageBox.No | QMessageBox.Cancel, QMessageBox.No,
            )
            if box == QMessageBox.Cancel:
                return
            create_missing = (box == QMessageBox.Yes)
        else:
            msg += "كل المواقع متطابقة ومسجلة بالمباني. هل تريد المتابعة؟"
            if QMessageBox.question(
                self, "استيراد التوزيع الأولي", msg,
                QMessageBox.Yes | QMessageBox.No, QMessageBox.Yes,
            ) != QMessageBox.Yes:
                return

        connection = get_connection()
        try:
            cursor = connection.cursor()
            if create_missing:
                for location in missing:
                    cursor.execute("INSERT INTO buildings (name) VALUES (?)", (location,))
                    cursor.execute("SELECT CAST(SCOPE_IDENTITY() AS INT) AS id")
                    buildings[self._norm(location)] = cursor.fetchone()["id"]

            cursor.execute("""
                SELECT ISNULL(MAX(CAST(SUBSTRING(camera_code, 5, 10) AS INT)), 0) AS n
                FROM cameras
                WHERE camera_code LIKE 'CAM-%'
                  AND ISNUMERIC(SUBSTRING(camera_code, 5, 10)) = 1
            """)
            counter = cursor.fetchone()["n"] or 0

            for r in new_rows:
                counter += 1
                cursor.execute("""
                    INSERT INTO cameras (
                        camera_code, name, building_id, location_detail,
                        placement, camera_type, status, manufacturer, model
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    f"CAM-{counter:04d}", r["name"],
                    buildings.get(self._norm(building_name_for(r["location"]))),
                    " — ".join(p for p in (r["location"], r["detail"]) if p),
                    r["placement"], r["camera_type"], r["status"],
                    r.get("model") or "Hikvision", r.get("model") or "Hikvision",
                ))

            connection.commit()
        except Exception as exc:
            connection.rollback()
            QMessageBox.critical(self, "خطأ", f"تعذر الاستيراد، وتم التراجع عن كافة العمليات:\n{exc}")
            return
        finally:
            connection.close()

        write_audit_log(
            user=self.current_user, action="استيراد", module_name="الكاميرات",
            description=f"استيراد التوزيع الأولي للكاميرات: {len(new_rows)} كاميرا",
            new_values={
                "imported": len(new_rows),
                "locations": len(locations),
                "created_buildings": len(missing) if create_missing else 0,
            },
        )

        self.load_building_filter()
        self.refresh_all(notify=False)
        self.show_notice(f"تم بنجاح استيراد {len(new_rows)} كاميرا")

    # =========================================================
    # التصدير والطباعة
    # =========================================================

    CAMERA_HEADERS = [
        "الكود", "اسم الكاميرا", "المبنى / الفرع", "الموضع التفصيلي", "السويتش / المنفذ", "التثبيت", "النوع",
        "الشركة المصنعة", "الموديل", "الرقم التسلسلي", "IP", "الحالة",
        "تاريخ التركيب", "انتهاء الضمان", "عدد الأعطال",
    ]

    SWITCH_HEADERS = [
        "كود السويتش", "اسم السويتش", "المبنى / الفرع", "الموضع التفصيلي",
        "الموديل", "عدد المنافذ", "الغرض والاستخدام", "عدد الكاميرات المتصلة", "عنوان IP", "الحالة التشغيلية",
    ]

    FAULT_HEADERS = [
        "رقم البلاغ", "تاريخ البلاغ", "كود الكاميرا", "الكاميرا", "المبنى",
        "وصف العطل", "الأولوية", "الحالة", "الفني", "إجراءات الإصلاح",
        "قطع الغيار", "تاريخ الإصلاح",
    ]

    def _check_reports_permission(self):
        if self.has_permission("cameras.reports"):
            return True
        QMessageBox.warning(self, "غير مسموح", "ليس لديك صلاحية استخراج التقارير والتصدير.")
        return False

    def export_excel(self):
        if not self._check_reports_permission():
            return

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            QMessageBox.critical(
                self, "مكتبة ناقصة", "مكتبة openpyxl غير مثبتة.\nشغّل:\npip install openpyxl"
            )
            return

        path, _ = QFileDialog.getSaveFileName(
            self, "تصدير بيانات الكاميرات والسويتشات إلى Excel",
            f"cameras_and_switches_{datetime.now():%Y%m%d_%H%M}.xlsx",
            "Excel Files (*.xlsx)",
        )
        if not path:
            return

        workbook = Workbook()

        def fill_sheet(sheet, headers, rows):
            sheet.sheet_view.rightToLeft = True
            sheet.append(headers)

            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF", size=11)
                cell.fill = PatternFill("solid", fgColor="0F2338")
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for r_idx, row in enumerate(rows, 2):
                sheet.append(row)
                if r_idx % 2 == 0:
                    for cell in sheet[r_idx]:
                        cell.fill = PatternFill("solid", fgColor="F8FAFC")

            for col in sheet.columns:
                max_len = max(len(str(c.value or "")) for c in col)
                sheet.column_dimensions[col[0].column_letter].width = min(max(max_len + 3, 12), 48)

            sheet.freeze_panes = "A2"

        first = workbook.active
        first.title = "سجل الكاميرات"
        fill_sheet(first, self.CAMERA_HEADERS, self._camera_export)

        second = workbook.create_sheet("سجل السويتشات")
        fill_sheet(second, self.SWITCH_HEADERS, self._switch_export)

        third = workbook.create_sheet("سجل بلاغات الأعطال")
        fill_sheet(third, self.FAULT_HEADERS, self._fault_export)

        try:
            workbook.save(path)
        except Exception as exc:
            QMessageBox.critical(self, "خطأ", f"تعذر حفظ الملف:\n{exc}")
            return

        write_audit_log(
            user=self.current_user, action="تصدير", module_name="الكاميرات",
            description=f"تصدير بيانات الكاميرات والسويتشات والأعطال إلى {Path(path).name}",
            new_values={
                "cameras": len(self._camera_export),
                "switches": len(self._switch_export),
                "faults": len(self._fault_export),
            },
        )
        self.show_notice(f"تم التصدير بنجاح: {Path(path).name}")

    def print_inventory(self):
        if not self._check_reports_permission():
            return
        open_print_page(
            build_print_html(
                "تقرير جرد الكاميرات",
                "بيانات الكاميرات ومواقعها وحالتها والسويتش المغذي وفق الفلاتر الحالية",
                self.CAMERA_HEADERS, self._camera_export,
                summary=f"إجمالي عدد الكاميرات: {len(self._camera_export)}",
            ),
            "cameras_inventory",
        )

    def print_switches(self):
        if not self._check_reports_permission():
            return
        open_print_page(
            build_print_html(
                "تقرير السويتشات",
                "توزيع السويتشات وموديلاتها وعدد المنافذ والكاميرات المتصلة بها",
                self.SWITCH_HEADERS, self._switch_export,
                summary=f"إجمالي عدد السويتشات: {len(self._switch_export)}",
            ),
            "network_switches_report",
        )

    def print_faults(self):
        if not self._check_reports_permission():
            return
        open_print_page(
            build_print_html(
                "تقرير بلاغات أعطال الكاميرات",
                "سجل الأعطال وحالات الإصلاح وقطع الغيار المصروفة",
                self.FAULT_HEADERS, self._fault_export,
                summary=f"إجمالي عدد البلاغات: {len(self._fault_export)}",
            ),
            "camera_faults",
        )