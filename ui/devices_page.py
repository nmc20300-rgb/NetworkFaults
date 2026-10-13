import os

from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QWidget,
    QVBoxLayout,
    QHBoxLayout,
    QLabel,
    QPushButton,
    QLineEdit,
    QTableWidget,
    QTableWidgetItem,
    QMessageBox,
    QDialog,
    QFormLayout,
    QComboBox,
    QDialogButtonBox,
    QHeaderView,
    QFrame,
    QFileDialog,
    QProgressDialog,
    QScrollArea,
    QMenu,
)
from PySide6.QtGui import QKeySequence, QShortcut

from database import (
    get_connection,
    write_audit_log,
)


# ============================================================
# STYLE
# ============================================================

STYLE = """
QWidget {
    font-family: "IBM Plex Sans Arabic", "Segoe UI", "Tahoma";
    font-size: 12px;
    color: #334155;
}

#root {
    background: #f4f7fb;
}

#title {
    color: #0f2338;
    font-size: 26px;
    font-weight: 800;
}

#subtitle {
    color: #64748b;
    font-size: 12px;
}

#card {
    background: white;
    border: 1px solid #e2e8f0;
    border-radius: 12px;
}

#primary {
    background: #0891b2;
    color: white;
    border: none;
    border-radius: 9px;
    padding: 9px 16px;
    font-weight: 700;
}

#primary:hover {
    background: #0e7490;
}

#primary:pressed {
    background: #155e75;
}

#secondary {
    background: #f1f5f9;
    color: #0f2338;
    border: 1px solid #cbd5e1;
    border-radius: 9px;
    padding: 9px 16px;
    font-weight: 700;
}

#secondary:hover {
    background: #e2e8f0;
}

#danger {
    background: #fef2f2;
    color: #991b1b;
    border: 1px solid #fecaca;
    border-radius: 9px;
    padding: 9px 16px;
    font-weight: 700;
}

#danger:hover {
    background: #fee2e2;
}

#devicesAddBtn {
    background: qlineargradient(x1:0, y1:0, x2:1, y2:0, stop:0 #0891b2, stop:1 #06b6d4);
    color: white;
    border: none;
    border-radius: 9px;
    padding: 0 16px;
    min-height: 38px;
    max-height: 38px;
    font-size: 13px;
    font-weight: 700;
}

#devicesAddBtn:hover {
    background: qlineargradient(x1:0, y1:0, x2:1, y2:0, stop:0 #0e7490, stop:1 #0891b2);
}

#devicesAddBtn:pressed {
    background: #155e75;
}

#devicesImportBtn {
    background: #ecfeff;
    color: #0e7490;
    border: 1px solid #a5f3fc;
    border-radius: 9px;
    padding: 0 15px;
    min-height: 38px;
    max-height: 38px;
    font-size: 13px;
    font-weight: 700;
}

#devicesImportBtn:hover {
    background: #cffafe;
    border-color: #67e8f9;
    color: #0891b2;
}

#devicesImportBtn:pressed {
    background: #a5f3fc;
}

#devicesMoreBtn {
    background: white;
    color: #475569;
    border: 1px solid #cbd5e1;
    border-radius: 9px;
    padding: 0 14px;
    min-height: 38px;
    max-height: 38px;
    font-size: 13px;
    font-weight: 700;
}

#devicesMoreBtn:hover {
    background: #f8fafc;
    border-color: #94a3b8;
    color: #0f172a;
}

#devicesMoreBtn:pressed {
    background: #e2e8f0;
}

#selectionBar {
    background: #f0fdf4;
    border: 1px solid #bbf7d0;
    border-radius: 9px;
}

#selectionBadge {
    color: #166534;
    font-weight: 800;
    font-size: 12px;
    padding: 0 4px;
}

#devicesEditBtn {
    background: white;
    color: #334155;
    border: 1px solid #cbd5e1;
    border-radius: 7px;
    padding: 0 13px;
    min-height: 30px;
    max-height: 30px;
    font-size: 12px;
    font-weight: 700;
}

#devicesEditBtn:hover {
    background: #f8fafc;
    border-color: #94a3b8;
    color: #0f172a;
}

#devicesEditBtn:pressed {
    background: #e2e8f0;
}

#devicesDeleteBtn {
    background: #fff1f2;
    color: #e11d48;
    border: 1px solid #fecdd3;
    border-radius: 7px;
    padding: 0 13px;
    min-height: 30px;
    max-height: 30px;
    font-size: 12px;
    font-weight: 700;
}

#devicesDeleteBtn:hover {
    background: #ffe4e6;
    border-color: #fda4af;
    color: #be123c;
}

#devicesDeleteBtn:pressed {
    background: #fecdd3;
}

#selectionDeselectBtn {
    background: transparent;
    color: #64748b;
    border: none;
    border-radius: 6px;
    padding: 2px 8px;
    min-height: 30px;
    font-size: 12px;
    font-weight: 700;
}

#selectionDeselectBtn:hover {
    background: #dcfce7;
    color: #15803d;
}

QMenu {
    background: white;
    border: 1px solid #cbd5e1;
    border-radius: 9px;
    padding: 5px;
    font-size: 12px;
}

QMenu::item {
    padding: 7px 20px;
    border-radius: 6px;
    color: #334155;
    font-weight: 600;
}

QMenu::item:selected {
    background: #f1f5f9;
    color: #0f2338;
}

QMenu::separator {
    height: 1px;
    background: #e2e8f0;
    margin: 4px 6px;
}


#search {
    background: #f8fafc;
    border: 1px solid #cbd5e1;
    border-radius: 9px;
    padding: 9px 12px;
}

#search:focus {
    border: 1.5px solid #0891b2;
    background: white;
}

QTableWidget {
    background: white;
    border: none;
    gridline-color: #f1f5f9;
    selection-background-color: #e0f2fe;
    selection-color: #0f2338;
    alternate-background-color: #f8fafc;
}

QHeaderView::section {
    background: #f8fafc;
    color: #0f2338;
    border: none;
    border-bottom: 2px solid #0891b2;
    padding: 10px;
    font-weight: 700;
}

QLineEdit,
QComboBox {
    background: #fbfdff;
    border: 1px solid #cbd5e1;
    border-radius: 8px;
    padding: 8px 10px;
    min-height: 20px;
}

QLineEdit:focus,
QComboBox:focus {
    border: 1.5px solid #0891b2;
    background: white;
}
"""


# ============================================================
# Excel Columns
# ============================================================

EXTRA_COLUMNS = {
    "location": "TEXT",
    "user_name_en": "TEXT",
    "lock_number": "TEXT",
    "network_connection": "TEXT",
    "motherboard": "TEXT",
    "graphics_card": "TEXT",
    "hard_disk_serial": "TEXT",
    "ms_office": "TEXT",
    "scanner": "TEXT",
    "ip_address": "TEXT",
}


def ensure_device_columns():
    """
    أعمدة الأجهزة موجودة بالفعل في SQL Server.
    لا نحتاج إلى ALTER TABLE أثناء التشغيل.
    """
    return


# ============================================================
# Helpers
# ============================================================

def clean(value):
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def norm(value):
    return " ".join(
        clean(value)
        .replace("\n", " ")
        .replace("\r", " ")
        .split()
    )


def device_audit_data(device):
    """
    تحويل بيانات الجهاز إلى Dictionary مناسب لسجل العمليات.
    """

    if not device:
        return {}

    keys = [
        "id",
        "device_number",
        "device_name",
        "user_name",
        "processor",
        "ram",
        "hard_disk",
        "monitor",
        "operating_system",
        "printer",
        "status",
        "building_id",
        "location",
        "user_name_en",
        "lock_number",
        "network_connection",
        "motherboard",
        "graphics_card",
        "hard_disk_serial",
        "ms_office",
        "scanner",
        "ip_address",
        "is_archived",
    ]

    result = {}

    for key in keys:
        try:
            result[key] = clean(device[key])
        except Exception:
            pass

    try:
        result["building_name"] = clean(
            device["building_name"]
        )
    except Exception:
        pass

    return result


# ============================================================
# Excel Structure
# ============================================================

EXCEL_COLUMNS = [
    "م",
    "المبنى",
    "مكان التواجد",
    "اسم الجهاز",
    "اسم المستخدم",
    "User name",
    "رقم القفل",
    "التوصيل بالشبكة",
    "Mother Board",
    "كارت الشاشة الخارجى",
    "الهارد ديسك",
    "الرقم التسلسلى للهارد",
    "المعالج المركزى",
    "-",
    "طراز الشاشة",
    "نظام التشغيل",
    "MS Office",
    "طراز الطابعة",
    "الماسح الضوئى",
    "IP",
    "الحالة الفنية",
]


HEADER_ALIASES = {

    "م": [
        "م",
        "م.",
        "الرقم",
        "رقم",
    ],

    "المبنى": [
        "المبنى",
        "المبنى/الفرع",
        "الفرع",
    ],

    "مكان التواجد": [
        "مكان التواجد",
        "مكان التواجد ",
    ],

    "اسم الجهاز": [
        "اسم الجهاز",
        "الجهاز",
    ],

    "اسم المستخدم": [
        "اسم المستخدم",
        "المستخدم",
    ],

    "User name": [
        "User name",
        "User Name",
        "Username",
        "USER NAME",
    ],

    "رقم القفل": [
        "رقم القفل",
        "رقم القفل ",
    ],

    "التوصيل بالشبكة": [
        "التوصيل بالشبكة",
        "التوصيل بالشبكه",
    ],

    "Mother Board": [
        "Mother Board",
        "Motherboard",
        "MOTHER BOARD",
    ],

    "كارت الشاشة الخارجى": [
        "كارت الشاشة الخارجى",
        "كارت الشاشة الخارجي",
        "كارت الشاشة الخارجى ",
    ],

    "الهارد ديسك": [
        "الهارد ديسك",
        "الهارد",
        "Hard Disk",
    ],

    "الرقم التسلسلى للهارد": [
        "الرقم التسلسلى للهارد",
        "الرقم التسلسلي للهارد",
        "Hard Disk Serial",
    ],

    "المعالج المركزى": [
        "المعالج المركزى",
        "المعالج المركزي",
        "Processor",
    ],

    "-": [
        "-",
        " - ",
        "RAM",
        "الرام",
        "RAM ",
    ],

    "طراز الشاشة": [
        "طراز الشاشة",
        "الشاشة",
        "Monitor",
    ],

    "نظام التشغيل": [
        "نظام التشغيل",
        "Operating System",
        "OS",
    ],

    "MS Office": [
        "MS Office",
        "MS OFFICE",
        "Office",
    ],

    "طراز الطابعة": [
        "طراز الطابعة",
        "الطابعة",
        "Printer",
    ],

    "الماسح الضوئى": [
        "الماسح الضوئى",
        "الماسح الضوئي",
        "Scanner",
        "الماسح",
    ],

    "IP": [
        "IP",
        "Ip",
        "ip",
        "IP Address",
    ],

    "الحالة الفنية": [
        "الحالة الفنية",
        "الحالة",
        "Status",
    ],
}


def find_header_indexes(headers):
    """
    مطابقة عناوين Excel مع الأعمدة الفعلية.

    يتم الاعتماد أولاً على اسم العمود،
    ثم يتم استخدام ترتيب الأعمدة كـ fallback.
    """

    indexes = {}

    normalized = [
        norm(header)
        for header in headers
    ]

    for wanted in EXCEL_COLUMNS:

        aliases = HEADER_ALIASES.get(
            wanted,
            [wanted],
        )

        aliases = {
            norm(alias)
            for alias in aliases
        }

        for index, header in enumerate(
            normalized
        ):

            if header in aliases:

                indexes[wanted] = index

                break

    # Fallback حسب ترتيب الملف
    for index, wanted in enumerate(
        EXCEL_COLUMNS
    ):

        if (
            wanted not in indexes
            and index < len(headers)
        ):

            indexes[wanted] = index

    return indexes


# ============================================================
# Device Dialog
# ============================================================

class DeviceDialog(QDialog):

    def __init__(
        self,
        parent=None,
        device=None,
    ):
        super().__init__(parent)

        self.device = device

        self.setWindowTitle(
            "تعديل جهاز"
            if device
            else "إضافة جهاز"
        )

        self.resize(
            720,
            760,
        )

        self.setStyleSheet(
            STYLE
        )

        root = QVBoxLayout(self)

        root.setContentsMargins(
            20,
            20,
            20,
            18,
        )

        title = QLabel(
            "بيانات الجهاز كاملة"
        )

        title.setStyleSheet(
            "background:#0f2338;"
            "color:white;"
            "font-size:19px;"
            "font-weight:800;"
            "padding:14px;"
            "border-radius:10px;"
            "border-bottom:3px solid #0891b2;"
        )

        root.addWidget(title)

        scroll = QScrollArea()

        scroll.setWidgetResizable(
            True
        )

        scroll.setFrameShape(
            QFrame.NoFrame
        )

        body = QWidget()

        form = QFormLayout(body)

        form.setVerticalSpacing(
            8
        )

        form.setHorizontalSpacing(
            18
        )

        self.fields = {}

        field_list = [
            (
                "device_number",
                "م / رقم الجهاز",
            ),
            (
                "device_name",
                "اسم الجهاز",
            ),
            (
                "building",
                "المبنى",
            ),
            (
                "location",
                "مكان التواجد",
            ),
            (
                "user_name",
                "اسم المستخدم",
            ),
            (
                "user_name_en",
                "User name",
            ),
            (
                "lock_number",
                "رقم القفل",
            ),
            (
                "network_connection",
                "التوصيل بالشبكة",
            ),
            (
                "motherboard",
                "Mother Board",
            ),
            (
                "graphics_card",
                "كارت الشاشة الخارجى",
            ),
            (
                "hard_disk",
                "الهارد ديسك",
            ),
            (
                "hard_disk_serial",
                "الرقم التسلسلى للهارد",
            ),
            (
                "processor",
                "المعالج المركزى",
            ),
            (
                "ram",
                "RAM",
            ),
            (
                "monitor",
                "طراز الشاشة",
            ),
            (
                "operating_system",
                "نظام التشغيل",
            ),
            (
                "ms_office",
                "MS Office",
            ),
            (
                "printer",
                "طراز الطابعة",
            ),
            (
                "scanner",
                "الماسح الضوئى",
            ),
            (
                "ip_address",
                "IP",
            ),
        ]

        for key, label in field_list:

            widget = QLineEdit()

            self.fields[key] = widget

            form.addRow(
                label + ":",
                widget,
            )

        self.status = QComboBox()

        self.status.addItems(
            [
                "صالح",
                "يحتاج فحص",
                "عطل",
                "خارج الخدمة",
            ]
        )

        form.addRow(
            "الحالة الفنية:",
            self.status,
        )

        scroll.setWidget(
            body
        )

        root.addWidget(
            scroll,
            1,
        )

        buttons = QDialogButtonBox(
            QDialogButtonBox.Save
            | QDialogButtonBox.Cancel
        )

        buttons.accepted.connect(
            self.validate
        )

        buttons.rejected.connect(
            self.reject
        )

        root.addWidget(
            buttons
        )

        if device:
            self.load_device()

    def load_device(self):

        for key, widget in self.fields.items():

            if key == "building":

                value = ""

                try:
                    value = clean(
                        self.device[
                            "building_name"
                        ]
                    )
                except Exception:
                    pass

            else:

                try:
                    value = clean(
                        self.device[key]
                    )
                except Exception:
                    value = ""

            widget.setText(
                value
            )

        try:

            status = (
                clean(
                    self.device[
                        "status"
                    ]
                )
                or "صالح"
            )

        except Exception:

            status = "صالح"

        index = self.status.findText(
            status
        )

        if index >= 0:

            self.status.setCurrentIndex(
                index
            )

    def validate(self):

        if not self.fields[
            "device_name"
        ].text().strip():

            QMessageBox.warning(
                self,
                "تنبيه",
                "اسم الجهاز مطلوب.",
            )

            return

        self.accept()

    def data(self):

        result = {
            key: widget.text().strip()
            for key, widget in self.fields.items()
        }

        result["status"] = (
            self.status.currentText()
        )

        return result


# ============================================================
# Devices Page
# ============================================================

class DevicesPage(QWidget):

    PAGE_SIZE = 50

    def __init__(
        self,
        current_user=None,
    ):
        super().__init__()

        self.current_user = (
            current_user or {}
        )

        self._page = 0
        self._total = 0

        ensure_device_columns()

        self.setup_ui()

        self.apply_permissions()

        self.load_devices()

    # ========================================================
    # Permission System
    # ========================================================

    def has_permission(self, code):

        role = str(
            self.current_user.get(
                "role_code",
                "",
            )
        ).strip().upper()

        if role == "SUPER_ADMIN":
            return True

        permissions = (
            self.current_user.get(
                "permissions",
                {},
            )
        )

        if not isinstance(
            permissions,
            dict,
        ):
            return False

        return bool(
            permissions.get(
                code,
                False,
            )
        )

    def apply_permissions(self):

        can_add = self.has_permission(
            "devices.add"
        )

        can_edit = self.has_permission(
            "devices.edit"
        )

        can_delete = self.has_permission(
            "devices.delete"
        )

        if hasattr(self, "import_button"):
            self.import_button.setVisible(can_add)

        if hasattr(self, "add_button"):
            self.add_button.setVisible(can_add)

        if hasattr(self, "act_clear_all"):
            self.act_clear_all.setVisible(can_delete)

    # ========================================================
    # UI
    # ========================================================

    def setup_ui(self):

        self.setObjectName(
            "root"
        )

        self.setStyleSheet(
            STYLE
        )

        root = QVBoxLayout(self)

        root.setContentsMargins(
            25,
            22,
            25,
            22,
        )

        root.setSpacing(
            13
        )

        # ----------------------------------------------------
        # Header
        # ----------------------------------------------------

        head = QHBoxLayout()
        head.setSpacing(8)
        head.setAlignment(Qt.AlignVCenter)

        titles = QVBoxLayout()

        title = QLabel(
            "إدارة الأجهزة"
        )

        title.setObjectName(
            "title"
        )

        titles.addWidget(
            title
        )

        subtitle = QLabel(
            "استيراد وعرض وتعديل جميع بيانات الأجهزة كما هي في ملف Excel"
        )

        subtitle.setObjectName(
            "subtitle"
        )

        titles.addWidget(
            subtitle
        )

        head.addLayout(
            titles
        )

        head.addStretch()

        # ----------------------------------------------------
        # Action Buttons Toolbar
        # ----------------------------------------------------

        # Contextual Selection Bar (Appears ONLY when rows are selected)
        self.selection_bar = QFrame()
        self.selection_bar.setObjectName("selectionBar")
        sel_layout = QHBoxLayout(self.selection_bar)
        sel_layout.setContentsMargins(8, 2, 8, 2)
        sel_layout.setSpacing(6)

        self.device_selection_label = QLabel("")
        self.device_selection_label.setObjectName("selectionBadge")
        sel_layout.addWidget(self.device_selection_label)

        self.edit_button = QPushButton("  تعديل")
        self.edit_button.setObjectName("devicesEditBtn")
        self.edit_button.setCursor(Qt.PointingHandCursor)
        self.edit_button.setToolTip("تعديل بيانات الجهاز المحدد")
        self.edit_button.clicked.connect(self.edit_device)
        sel_layout.addWidget(self.edit_button)

        self.device_delete_button = QPushButton("  حذف")
        self.device_delete_button.setObjectName("devicesDeleteBtn")
        self.device_delete_button.setCursor(Qt.PointingHandCursor)
        self.device_delete_button.setToolTip("حذف الجهاز أو الأجهزة المحددة")
        self.device_delete_button.clicked.connect(self.delete_device)
        sel_layout.addWidget(self.device_delete_button)

        self.deselect_button = QPushButton("✕  إلغاء")
        self.deselect_button.setObjectName("selectionDeselectBtn")
        self.deselect_button.setCursor(Qt.PointingHandCursor)
        self.deselect_button.setToolTip("إلغاء التحديد")
        self.deselect_button.clicked.connect(
            lambda: self.table.clearSelection() if hasattr(self, "table") else None
        )
        sel_layout.addWidget(self.deselect_button)

        self.selection_bar.setVisible(False)
        head.addWidget(self.selection_bar)

        # 1. Primary Action: إضافة جهاز
        self.add_button = QPushButton("  إضافة جهاز")
        self.add_button.setObjectName("devicesAddBtn")
        self.add_button.setCursor(Qt.PointingHandCursor)
        self.add_button.setToolTip("إضافة جهاز كمبيوتر أو ملحق جديد")
        self.add_button.clicked.connect(self.add_device)
        head.addWidget(self.add_button)

        # 2. Companion Action: استيراد Excel
        self.import_button = QPushButton("  استيراد Excel")
        self.import_button.setObjectName("devicesImportBtn")
        self.import_button.setCursor(Qt.PointingHandCursor)
        self.import_button.setToolTip("استيراد بيانات الأجهزة دفعة واحدة من ملف Excel")
        self.import_button.clicked.connect(self.import_excel)
        head.addWidget(self.import_button)

        # 3. More Options Menu: خيارات ▾ (تحديد الكل، مسح الكل)
        self.more_button = QPushButton("خيارات ▾")
        self.more_button.setObjectName("devicesMoreBtn")
        self.more_button.setCursor(Qt.PointingHandCursor)
        self.more_button.setToolTip("خيارات إضافية للجدول والبيانات")

        more_menu = QMenu(self.more_button)
        self.act_select_all = more_menu.addAction("  تحديد الكل")
        self.act_select_all.triggered.connect(self.toggle_select_all)

        more_menu.addSeparator()

        self.act_clear_all = more_menu.addAction("  مسح جميع الأجهزة...")
        self.act_clear_all.triggered.connect(self.clear_all_devices)

        self.more_button.setMenu(more_menu)
        head.addWidget(self.more_button)

        # Backward compatibility references
        self.select_all_button = self.more_button
        self.clear_all_button = QPushButton()
        self.clear_all_button.setVisible(False)

        root.addLayout(
            head
        )

        # ----------------------------------------------------
        # Search / Filters
        # ----------------------------------------------------

        search_card = QFrame()

        search_card.setObjectName(
            "card"
        )

        search_layout = QHBoxLayout(
            search_card
        )

        search_layout.setContentsMargins(
            12,
            10,
            12,
            10,
        )

        search_layout.addWidget(
            QLabel("🔎")
        )

        self.search = QLineEdit()

        self.search.setObjectName(
            "search"
        )

        self.search.setPlaceholderText(
            "بحث برقم الجهاز أو المبنى أو المكان أو المستخدم أو IP أو Serial..."
        )

        self.search.textChanged.connect(
            self._reset_page_and_load_devices
        )

        search_layout.addWidget(
            self.search,
            1,
        )

        self.building_filter = QComboBox()

        self.building_filter.setMinimumWidth(
            170
        )

        self.building_filter.addItem(
            "كل المباني",
            "",
        )

        self.building_filter.currentIndexChanged.connect(
            self._reset_page_and_load_devices
        )

        search_layout.addWidget(
            self.building_filter
        )

        self.status_filter = QComboBox()

        self.status_filter.setMinimumWidth(
            150
        )

        self.status_filter.addItem(
            "كل الحالات",
            "",
        )

        self.status_filter.currentIndexChanged.connect(
            self._reset_page_and_load_devices
        )

        search_layout.addWidget(
            self.status_filter
        )

        root.addWidget(
            search_card
        )

        # ----------------------------------------------------
        # Table
        # ----------------------------------------------------

        table_card = QFrame()

        table_card.setObjectName(
            "card"
        )

        table_layout = QVBoxLayout(
            table_card
        )

        table_layout.setContentsMargins(
            7,
            7,
            7,
            7,
        )

        self.table = QTableWidget()

        self.table.setColumnCount(
            22
        )

        self.table.setHorizontalHeaderLabels(
            [
                "ID",
                "م",
                "المبنى",
                "مكان التواجد",
                "اسم الجهاز",
                "اسم المستخدم",
                "User name",
                "رقم القفل",
                "التوصيل بالشبكة",
                "Mother Board",
                "كارت الشاشة الخارجى",
                "الهارد ديسك",
                "الرقم التسلسلى للهارد",
                "المعالج المركزى",
                "RAM",
                "طراز الشاشة",
                "نظام التشغيل",
                "MS Office",
                "طراز الطابعة",
                "الماسح الضوئى",
                "IP",
                "الحالة الفنية",
            ]
        )

        self.table.setColumnHidden(
            0,
            True,
        )

        self.table.setSelectionBehavior(
            QTableWidget.SelectRows
        )

        self.table.setSelectionMode(
            QTableWidget.MultiSelection
        )

        self.table.itemSelectionChanged.connect(
            self._on_device_selection_changed
        )

        self.table.setContextMenuPolicy(
            Qt.CustomContextMenu
        )
        self.table.customContextMenuRequested.connect(
            self._show_context_menu
        )

        QShortcut(
            QKeySequence(Qt.Key_Escape),
            self.table
        ).activated.connect(
            self.table.clearSelection
        )

        self.table.setEditTriggers(
            QTableWidget.NoEditTriggers
        )

        self.table.setAlternatingRowColors(
            True
        )

        self.table.verticalHeader().setVisible(
            False
        )

        header = self.table.horizontalHeader()

        header.setSectionResizeMode(
            QHeaderView.ResizeToContents
        )

        header.setStretchLastSection(
            True
        )

        self.table.doubleClicked.connect(
            self.edit_device
        )

        table_layout.addWidget(
            self.table
        )

        # ----------------------------------------------------
        # Pagination
        # ----------------------------------------------------

        pagination_row = QHBoxLayout()

        self.stats_label = QLabel("")

        self.stats_label.setStyleSheet(
            "color: #64748b; font-size: 11px;"
        )

        pagination_row.addWidget(
            self.stats_label
        )

        pagination_row.addStretch()

        self.previous_button = QPushButton(
            "◀  السابق"
        )

        self.previous_button.setObjectName(
            "secondary"
        )

        self.previous_button.setMinimumHeight(
            32
        )

        self.previous_button.clicked.connect(
            self._previous_page
        )

        self.page_label = QLabel(
            "صفحة 1 / 1"
        )

        self.page_label.setAlignment(
            Qt.AlignCenter
        )

        self.page_label.setStyleSheet(
            "font-weight: 700; min-width: 80px;"
        )

        self.next_button = QPushButton(
            "التالي  ▶"
        )

        self.next_button.setObjectName(
            "secondary"
        )

        self.next_button.setMinimumHeight(
            32
        )

        self.next_button.clicked.connect(
            self._next_page
        )

        pagination_row.addWidget(
            self.previous_button
        )

        pagination_row.addWidget(
            self.page_label
        )

        pagination_row.addWidget(
            self.next_button
        )

        table_layout.addLayout(
            pagination_row
        )

        root.addWidget(
            table_card,
            1,
        )

    # ========================================================
    # Load Devices
    # ========================================================

    def _reset_page_and_load_devices(
        self,
        *_args,
    ):

        self._page = 0

        self.load_devices()

    def _previous_page(
        self,
        _checked=False,
    ):

        if self._page > 0:

            self._page -= 1

            self.load_devices()

    def _next_page(
        self,
        _checked=False,
    ):

        total_pages = max(
            1,
            (
                self._total
                + self.PAGE_SIZE
                - 1
            )
            // self.PAGE_SIZE,
        )

        if self._page < total_pages - 1:

            self._page += 1

            self.load_devices()

    def load_devices(self):

        search = self.search.text().strip()

        selected_building = (
            self.building_filter.currentData()
            if hasattr(
                self,
                "building_filter",
            )
            else ""
        )

        selected_status = (
            self.status_filter.currentData()
            if hasattr(
                self,
                "status_filter",
            )
            else ""
        )

        con = get_connection()

        cur = con.cursor()

        try:

            # ------------------------------------------------
            # Buildings
            # ------------------------------------------------

            cur.execute(
                """
                SELECT DISTINCT
                    b.id,
                    b.name
                FROM buildings b
                INNER JOIN devices d
                    ON d.building_id = b.id
                WHERE d.is_archived = 0
                  AND b.name IS NOT NULL
                  AND LTRIM(RTRIM(b.name)) <> ''
                ORDER BY b.name
                """
            )

            buildings = cur.fetchall()

            current_building = (
                selected_building
            )

            self.building_filter.blockSignals(
                True
            )

            self.building_filter.clear()

            self.building_filter.addItem(
                "كل المباني",
                "",
            )

            for building in buildings:

                self.building_filter.addItem(
                    clean(
                        building["name"]
                    ),
                    str(
                        building["id"]
                    ),
                )

            index = (
                self.building_filter.findData(
                    current_building
                )
            )

            if index >= 0:

                self.building_filter.setCurrentIndex(
                    index
                )

                selected_building = (
                    current_building
                )

            else:

                self.building_filter.setCurrentIndex(
                    0
                )

                selected_building = ""

            self.building_filter.blockSignals(
                False
            )

            # ------------------------------------------------
            # Statuses
            # ------------------------------------------------

            cur.execute(
                """
                SELECT DISTINCT
                    status
                FROM devices
                WHERE is_archived = 0
                  AND status IS NOT NULL
                  AND LTRIM(RTRIM(status)) <> ''
                ORDER BY status
                """
            )

            statuses = cur.fetchall()

            current_status = (
                selected_status
            )

            self.status_filter.blockSignals(
                True
            )

            self.status_filter.clear()

            self.status_filter.addItem(
                "كل الحالات",
                "",
            )

            for status in statuses:

                self.status_filter.addItem(
                    clean(
                        status["status"]
                    ),
                    clean(
                        status["status"]
                    ),
                )

            index = (
                self.status_filter.findData(
                    current_status
                )
            )

            if index >= 0:

                self.status_filter.setCurrentIndex(
                    index
                )

                selected_status = (
                    current_status
                )

            else:

                self.status_filter.setCurrentIndex(
                    0
                )

                selected_status = ""

            self.status_filter.blockSignals(
                False
            )

            # ------------------------------------------------
            # Devices Query
            # ------------------------------------------------

            query = """
                SELECT
                    devices.*,
                    buildings.name AS building_name
                FROM devices
                LEFT JOIN buildings
                    ON devices.building_id = buildings.id
                WHERE devices.is_archived = 0
            """

            params = []

            if search:

                query += """
                    AND (
                        devices.device_number LIKE ?
                        OR devices.device_name LIKE ?
                        OR devices.user_name LIKE ?
                        OR devices.serial_number LIKE ?
                        OR devices.ip_address LIKE ?
                        OR devices.location LIKE ?
                        OR buildings.name LIKE ?
                    )
                """

                key = f"%{search}%"

                params.extend(
                    [key] * 7
                )

            if selected_building:

                query += """
                    AND devices.building_id = ?
                """

                params.append(
                    int(
                        selected_building
                    )
                )

            if selected_status:

                query += """
                    AND devices.status = ?
                """

                params.append(
                    selected_status
                )

            query += """
                ORDER BY devices.id ASC
            """

            cur.execute(
                query,
                params,
            )

            rows = cur.fetchall()

        finally:

            con.close()

        # ----------------------------------------------------
        # Apply pagination
        # ----------------------------------------------------

        self._total = len(rows)

        total_pages = max(
            1,
            (
                self._total
                + self.PAGE_SIZE
                - 1
            )
            // self.PAGE_SIZE,
        )

        self._page = min(
            self._page,
            total_pages - 1,
        )

        start = (
            self._page
            * self.PAGE_SIZE
        )

        visible_rows = rows[
            start:
            start + self.PAGE_SIZE
        ]

        # ----------------------------------------------------
        # Fill Table
        # ----------------------------------------------------

        self.table.setRowCount(
            0
        )

        for device in visible_rows:

            row = self.table.rowCount()

            self.table.insertRow(
                row
            )

            values = [
                device["id"],
                device["device_number"],
                device["building_name"] or "",
                device["location"],
                device["device_name"],
                device["user_name"],
                device["user_name_en"],
                device["lock_number"],
                device["network_connection"],
                device["motherboard"],
                device["graphics_card"],
                device["hard_disk"],
                device["hard_disk_serial"],
                device["processor"],
                device["ram"],
                device["monitor"],
                device["operating_system"],
                device["ms_office"],
                device["printer"],
                device["scanner"],
                device["ip_address"],
                device["status"],
            ]

            for column, value in enumerate(
                values
            ):

                item = QTableWidgetItem(
                    clean(value)
                )

                item.setTextAlignment(
                    Qt.AlignCenter
                )

                self.table.setItem(
                    row,
                    column,
                    item,
                )

        self.page_label.setText(
            f"صفحة {self._page + 1} / {total_pages}"
        )

        self.previous_button.setEnabled(
            self._page > 0
        )

        self.next_button.setEnabled(
            self._page < total_pages - 1
        )

        self.stats_label.setText(
            f"إجمالي الأجهزة: {self._total}  |  "
            f"المعروض في الصفحة: {len(visible_rows)}"
        )

    # ========================================================
    # Selected Device
    # ========================================================

    def get_selected_device(self):

        selected = (
            self.table.selectedItems()
        )

        if not selected:

            QMessageBox.information(
                self,
                "تنبيه",
                "من فضلك اختر جهازاً أولاً.",
            )

            return None

        row = selected[0].row()

        device_id = (
            self.table.item(
                row,
                0,
            ).text()
        )

        con = get_connection()

        cur = con.cursor()

        try:

            cur.execute(
                """
                SELECT
                    devices.*,
                    buildings.name AS building_name
                FROM devices
                LEFT JOIN buildings
                    ON devices.building_id = buildings.id
                WHERE devices.id = ?
                """,
                (
                    device_id,
                ),
            )

            device = cur.fetchone()

        finally:

            con.close()

        return device

    # ========================================================
    # Building
    # ========================================================

    def get_building_id(
        self,
        building_name,
        con,
    ):

        name = clean(
            building_name
        )

        if not name:
            return None

        cur = con.cursor()

        cur.execute(
            """
            SELECT TOP 1
                id
            FROM buildings
            WHERE name = ?
            """,
            (
                name,
            ),
        )

        row = cur.fetchone()

        if row:
            return row["id"]

        cur.execute(
            """
            INSERT INTO buildings (
                name,
                location,
                notes
            )
            VALUES (?, ?, ?)
            """,
            (
                name,
                "",
                "تم إنشاؤه تلقائياً من استيراد Excel",
            ),
        )

        cur.execute(
            """
            SELECT CAST(
                SCOPE_IDENTITY()
                AS INT
            ) AS id
            """
        )

        row = cur.fetchone()

        return (
            row["id"]
            if row
            else None
        )

    # ========================================================
    # Add Device
    # ========================================================

    def add_device(self):

        if not self.has_permission(
            "devices.add"
        ):

            QMessageBox.warning(
                self,
                "غير مسموح",
                "ليس لديك صلاحية إضافة الأجهزة.",
            )

            return

        dialog = DeviceDialog(
            self
        )

        if dialog.exec() != QDialog.Accepted:
            return

        data = dialog.data()

        con = get_connection()

        cur = con.cursor()

        try:

            building_id = (
                self.get_building_id(
                    data["building"],
                    con,
                )
            )

            cur.execute(
                """
                INSERT INTO devices (
                    device_number,
                    device_name,
                    user_name,
                    department,
                    device_type,
                    serial_number,
                    processor,
                    ram,
                    hard_disk,
                    monitor,
                    operating_system,
                    printer,
                    production_year,
                    status,
                    building_id,
                    location,
                    user_name_en,
                    lock_number,
                    network_connection,
                    motherboard,
                    graphics_card,
                    hard_disk_serial,
                    ms_office,
                    scanner,
                    ip_address
                )
                VALUES (
                    ?,?,?,?,?,?,?,?,?,?,
                    ?,?,?,?,?,?,?,?,?,?,
                    ?,?,?,?,?
                )
                """,
                (
                    data["device_number"],
                    data["device_name"],
                    data["user_name"],
                    "",
                    "",
                    "",
                    data["processor"],
                    data["ram"],
                    data["hard_disk"],
                    data["monitor"],
                    data["operating_system"],
                    data["printer"],
                    "",
                    data["status"],
                    building_id,
                    data["location"],
                    data["user_name_en"],
                    data["lock_number"],
                    data["network_connection"],
                    data["motherboard"],
                    data["graphics_card"],
                    data["hard_disk_serial"],
                    data["ms_office"],
                    data["scanner"],
                    data["ip_address"],
                ),
            )

            new_device_id = None

            try:

                cur.execute(
                    """
                    SELECT CAST(
                        SCOPE_IDENTITY()
                        AS INT
                    ) AS id
                    """
                )

                identity_row = cur.fetchone()

                if identity_row:
                    new_device_id = identity_row["id"]

            except Exception:
                pass

            con.commit()

        except Exception as error:

            con.rollback()

            QMessageBox.critical(
                self,
                "خطأ",
                f"حدث خطأ أثناء إضافة الجهاز:\n\n{error}",
            )

            return

        finally:

            con.close()

        audit_new_values = dict(
            data
        )

        audit_new_values[
            "building_id"
        ] = building_id

        write_audit_log(
            user=self.current_user,
            action="إضافة",
            module_name="الأجهزة",
            record_id=new_device_id,
            description=(
                f"تمت إضافة الجهاز: "
                f"{data['device_name']}"
            ),
            new_values=audit_new_values,
        )

        self.load_devices()

        QMessageBox.information(
            self,
            "تم",
            "تمت إضافة الجهاز بنجاح.",
        )

    # ========================================================
    # Edit Device
    # ========================================================

    def edit_device(self):

        if not self.has_permission(
            "devices.edit"
        ):

            QMessageBox.warning(
                self,
                "غير مسموح",
                "ليس لديك صلاحية تعديل الأجهزة.",
            )

            return

        device = self.get_selected_device()

        if not device:
            return

        old_values = device_audit_data(
            device
        )

        dialog = DeviceDialog(
            self,
            device,
        )

        if dialog.exec() != QDialog.Accepted:
            return

        data = dialog.data()

        con = get_connection()

        cur = con.cursor()

        try:

            building_id = (
                self.get_building_id(
                    data["building"],
                    con,
                )
            )

            cur.execute(
                """
                UPDATE devices
                SET
                    device_number=?,
                    device_name=?,
                    user_name=?,
                    processor=?,
                    ram=?,
                    hard_disk=?,
                    monitor=?,
                    operating_system=?,
                    printer=?,
                    status=?,
                    building_id=?,
                    location=?,
                    user_name_en=?,
                    lock_number=?,
                    network_connection=?,
                    motherboard=?,
                    graphics_card=?,
                    hard_disk_serial=?,
                    ms_office=?,
                    scanner=?,
                    ip_address=?,
                    updated_at=CURRENT_TIMESTAMP
                WHERE id=?
                """,
                (
                    data["device_number"],
                    data["device_name"],
                    data["user_name"],
                    data["processor"],
                    data["ram"],
                    data["hard_disk"],
                    data["monitor"],
                    data["operating_system"],
                    data["printer"],
                    data["status"],
                    building_id,
                    data["location"],
                    data["user_name_en"],
                    data["lock_number"],
                    data["network_connection"],
                    data["motherboard"],
                    data["graphics_card"],
                    data["hard_disk_serial"],
                    data["ms_office"],
                    data["scanner"],
                    data["ip_address"],
                    device["id"],
                ),
            )

            con.commit()

        except Exception as error:

            con.rollback()

            QMessageBox.critical(
                self,
                "خطأ",
                f"حدث خطأ أثناء تعديل الجهاز:\n\n{error}",
            )

            return

        finally:

            con.close()

        audit_new_values = dict(
            data
        )

        audit_new_values[
            "building_id"
        ] = building_id

        write_audit_log(
            user=self.current_user,
            action="تعديل",
            module_name="الأجهزة",
            record_id=device["id"],
            description=(
                f"تم تعديل بيانات الجهاز: "
                f"{data['device_name']}"
            ),
            old_values=old_values,
            new_values=audit_new_values,
        )

        self.load_devices()

        QMessageBox.information(
            self,
            "تم",
            "تم تعديل الجهاز بنجاح.",
        )

    # ========================================================
    # Delete All Devices
    # ========================================================

    def clear_all_devices(self):

        if not self.has_permission(
            "devices.delete"
        ):

            QMessageBox.warning(
                self,
                "غير مسموح",
                "ليس لديك صلاحية حذف الأجهزة.",
            )

            return

        answer = QMessageBox.question(
            self,
            "تأكيد مسح جميع الأجهزة",
            "⚠ هل أنت متأكد من مسح جميع الأجهزة نهائيًا؟\n\n"
            "سيتم حذف الأجهزة النشطة والمؤرشفة من قاعدة البيانات، "
            "ولا يمكن التراجع عن العملية.",
            QMessageBox.Yes
            | QMessageBox.No,
            QMessageBox.No,
        )

        if answer != QMessageBox.Yes:
            return

        con = get_connection()

        cur = con.cursor()

        try:

            # حماية سجل الأعطال
            cur.execute(
                """
                SELECT COUNT(*) AS total
                FROM faults
                WHERE device_id IS NOT NULL
                """
            )

            fault_count = int(
                cur.fetchone()["total"] or 0
            )

            if fault_count > 0:

                QMessageBox.warning(
                    self,
                    "لا يمكن المسح",
                    "يوجد أعطال مرتبطة بأجهزة مسجلة.\n\n"
                    "لم يتم حذف أي جهاز حفاظًا على سجل الأعطال.\n"
                    "احذف/عالج سجلات الأعطال المرتبطة أولًا "
                    "ثم استخدم مسح الكل.",
                )

                return

            # معرفة العدد قبل الحذف
            cur.execute(
                """
                SELECT COUNT(*) AS total
                FROM devices
                """
            )

            total_before = int(
                cur.fetchone()["total"] or 0
            )

            cur.execute(
                "DELETE FROM devices"
            )

            deleted = cur.rowcount

            con.commit()

        except Exception as error:

            con.rollback()

            QMessageBox.critical(
                self,
                "خطأ",
                f"حدث خطأ أثناء مسح الأجهزة:\n{error}",
            )

            return

        finally:

            con.close()

        write_audit_log(
            user=self.current_user,
            action="حذف",
            module_name="الأجهزة",
            description=(
                "تم مسح جميع الأجهزة نهائيًا "
                f"وعدد السجلات المحذوفة: {deleted}"
            ),
            old_values={
                "عدد الأجهزة قبل المسح": total_before,
                "عدد الأجهزة المحذوفة": deleted,
            },
        )

        QMessageBox.information(
            self,
            "تم",
            f"تم مسح جميع الأجهزة بنجاح.\n\n"
            f"عدد الأجهزة المحذوفة: {deleted}",
        )

        self.load_devices()

    # ========================================================
    # Helper: Selection
    # ========================================================

    def toggle_select_all(self):
        """تحديد أو إلغاء تحديد كل العناصر بضغطة زر واحدة ذكية."""
        selected_count = len(
            set(idx.row() for idx in self.table.selectedIndexes())
        )
        total_rows = self.table.rowCount()
        if total_rows == 0:
            return

        if selected_count >= total_rows:
            self.table.clearSelection()
        else:
            self.table.selectAll()

    def _on_device_selection_changed(self):
        """تحديث عداد التحديد والزر الذكي وزر الحذف عند تغيير الصفوف المختارة."""
        selected_rows = len(
            set(
                idx.row()
                for idx in
                self.table.selectedIndexes()
            )
        )
        total_rows = self.table.rowCount()

        if hasattr(self, "act_select_all"):
            if total_rows > 0 and selected_rows >= total_rows:
                self.act_select_all.setText("  إلغاء التحديد")
            else:
                self.act_select_all.setText("  تحديد الكل")

        can_edit = self.has_permission("devices.edit")
        can_delete = self.has_permission("devices.delete")

        if hasattr(self, "selection_bar"):
            if selected_rows > 0:
                self.selection_bar.setVisible(True)
                self.device_selection_label.setText(f"✔  {selected_rows} محدد")
                self.edit_button.setVisible(selected_rows == 1 and can_edit)
                self.device_delete_button.setVisible(can_delete)
                if selected_rows > 1:
                    self.device_delete_button.setText(f"🗑  حذف ({selected_rows})")
                else:
                    self.device_delete_button.setText("🗑  حذف")
            else:
                self.selection_bar.setVisible(False)

    def _show_context_menu(self, pos):
        """قائمة كليك يمين أنيقة للجدول بدون زحمة أزرار في الواجهة."""
        menu = QMenu(self)
        menu.setStyleSheet("""
            QMenu {
                background: white;
                border: 1px solid #cbd5e1;
                border-radius: 8px;
                padding: 4px;
                font-size: 12px;
            }
            QMenu::item {
                padding: 6px 20px;
                border-radius: 4px;
            }
            QMenu::item:selected {
                background: #e0f2fe;
                color: #0369a1;
            }
        """)

        act_select_all = menu.addAction("☑  تحديد الكل")
        act_select_all.triggered.connect(self.table.selectAll)

        act_deselect = menu.addAction("☐  إلغاء التحديد")
        act_deselect.triggered.connect(self.table.clearSelection)

        selected_count = len(
            set(idx.row() for idx in self.table.selectedIndexes())
        )
        if selected_count > 0:
            menu.addSeparator()
            act_delete = menu.addAction(f"🗑  حذف المحدد ({selected_count})")
            act_delete.triggered.connect(self.delete_device)

        menu.exec(self.table.viewport().mapToGlobal(pos))

    def _get_selected_device_ids(self):
        """إرجاع قائمة بمعرّفات (IDs) جميع الصفوف المحددة."""
        rows = set(
            idx.row()
            for idx in
            self.table.selectedIndexes()
        )

        ids = []

        for row in sorted(rows):
            item = self.table.item(row, 0)
            if item:
                ids.append(item.text())

        return ids

    # ========================================================
    # Delete Device
    # ========================================================

    def delete_device(self):

        if not self.has_permission(
            "devices.delete"
        ):

            QMessageBox.warning(
                self,
                "غير مسموح",
                "ليس لديك صلاحية حذف الأجهزة.",
            )

            return

        device_ids = self._get_selected_device_ids()

        if not device_ids:

            QMessageBox.information(
                self,
                "تنبيه",
                "من فضلك اختر جهازًا أو أكثر أولاً.",
            )

            return

        count = len(device_ids)

        answer = QMessageBox.question(
            self,
            "تأكيد الحذف",
            f"هل تريد حذف {count} جهاز نهائيًا؟\n\n"
            "الأجهزة المرتبطة بسجلات أعطال لن تُحذف.\n"
            "لا يمكن التراجع عن هذا الحذف.",
            QMessageBox.Yes
            | QMessageBox.No,
            QMessageBox.No,
        )

        if answer != QMessageBox.Yes:
            return

        con = get_connection()
        cur = con.cursor()

        deleted_count = 0
        skipped = []
        errors = []

        for device_id in device_ids:

            try:

                cur.execute(
                    """
                    SELECT
                        devices.*,
                        buildings.name AS building_name
                    FROM devices
                    LEFT JOIN buildings
                        ON devices.building_id = buildings.id
                    WHERE devices.id = ?
                    """,
                    (device_id,),
                )

                device = cur.fetchone()

                if not device:
                    continue

                old_values = device_audit_data(device)

                cur.execute(
                    """
                    SELECT COUNT(*) AS total
                    FROM faults
                    WHERE device_id = ?
                    """,
                    (device_id,),
                )

                fault_count = int(
                    cur.fetchone()["total"] or 0
                )

                if fault_count:
                    skipped.append(
                        f"{device['device_name']} – مرتبط بسجلات أعطال"
                    )
                    continue

                cur.execute(
                    "DELETE FROM devices WHERE id = ?",
                    (device_id,),
                )

                write_audit_log(
                    user=self.current_user,
                    action="حذف",
                    module_name="الأجهزة",
                    record_id=device["id"],
                    description=(
                        f"تم حذف الجهاز نهائيًا: "
                        f"{device['device_name']}"
                    ),
                    old_values=old_values,
                    new_values={"deleted": True},
                )

                deleted_count += 1

            except Exception as error:

                errors.append(
                    f"{device_id}: {error}"
                )

        try:
            con.commit()
        except Exception as exc:
            con.rollback()
            QMessageBox.critical(
                self,
                "خطأ",
                f"حدث خطأ أثناء حفظ التغييرات:\n\n{exc}",
            )
            return
        finally:
            con.close()

        self.load_devices()

        summary = f"تم حذف {deleted_count} جهاز بنجاح."

        if skipped:
            summary += (
                "\n\nتم تخطي الأجهزة التالية "
                "(مرتبطة بسجلات أعطال):\n"
                + "\n".join(skipped)
            )

        if errors:
            summary += (
                "\n\nأخطاء:\n" + "\n".join(errors)
            )

        if skipped or errors:
            QMessageBox.warning(
                self,
                "اكتمل مع تحذيرات",
                summary,
            )
        else:
            QMessageBox.information(
                self,
                "تم",
                summary,
            )

    # ========================================================
    # Import Excel
    # ========================================================

    def import_excel(self):

        if not self.has_permission(
            "devices.add"
        ):

            QMessageBox.warning(
                self,
                "غير مسموح",
                "ليس لديك صلاحية استيراد وإضافة الأجهزة.",
            )

            return

        try:

            from openpyxl import load_workbook

        except ImportError:

            QMessageBox.critical(
                self,
                "مكتبة ناقصة",
                "ثبت openpyxl أولاً:\n\n"
                "python -m pip install openpyxl",
            )

            return

        file_path, _ = (
            QFileDialog.getOpenFileName(
                self,
                "اختيار ملف Excel",
                "",
                "Excel (*.xlsx *.xlsm)",
            )
        )

        if not file_path:
            return

        file_name = os.path.basename(
            file_path
        )

        wb = None
        con = None
        progress = None

        try:

            # ==================================================
            # 1) قراءة Excel
            # ==================================================

            wb = load_workbook(
                file_path,
                read_only=True,
                data_only=True,
            )

            ws = wb.active

            iterator = ws.iter_rows(
                values_only=True
            )

            header = next(
                iterator,
                None,
            )

            if not header:

                wb.close()

                wb = None

                QMessageBox.warning(
                    self,
                    "تنبيه",
                    "ملف Excel فارغ.",
                )

                return

            indexes = find_header_indexes(
                header
            )

            # ==================================================
            # Required Columns
            # ==================================================

            if (
                "اسم الجهاز" not in indexes
                or "المبنى" not in indexes
            ):

                wb.close()

                wb = None

                QMessageBox.warning(
                    self,
                    "خطأ في الملف",
                    "لم أستطع التعرف على أعمدة "
                    "المبنى واسم الجهاز.",
                )

                return

            def cell(
                row,
                name,
            ):

                index = indexes.get(
                    name
                )

                if (
                    index is None
                    or index >= len(row)
                ):

                    return ""

                return clean(
                    row[index]
                )

            imported_rows = []

            # ==================================================
            # 2) قراءة الصفوف
            # ==================================================

            for row in iterator:

                if (
                    not row
                    or not any(
                        clean(value)
                        for value in row
                    )
                ):
                    continue

                device_name = cell(
                    row,
                    "اسم الجهاز",
                )

                if not device_name:
                    continue

                imported_rows.append(
                    {
                        "device_number": cell(
                            row,
                            "م",
                        ),

                        "building": cell(
                            row,
                            "المبنى",
                        ),

                        "location": cell(
                            row,
                            "مكان التواجد",
                        ),

                        "device_name": device_name,

                        "user_name": cell(
                            row,
                            "اسم المستخدم",
                        ),

                        "user_name_en": cell(
                            row,
                            "User name",
                        ),

                        "lock_number": cell(
                            row,
                            "رقم القفل",
                        ),

                        "network_connection": cell(
                            row,
                            "التوصيل بالشبكة",
                        ),

                        "motherboard": cell(
                            row,
                            "Mother Board",
                        ),

                        "graphics_card": cell(
                            row,
                            "كارت الشاشة الخارجى",
                        ),

                        "hard_disk": cell(
                            row,
                            "الهارد ديسك",
                        ),

                        "hard_disk_serial": cell(
                            row,
                            "الرقم التسلسلى للهارد",
                        ),

                        "processor": cell(
                            row,
                            "المعالج المركزى",
                        ),

                        "ram": cell(
                            row,
                            "-",
                        ),

                        "monitor": cell(
                            row,
                            "طراز الشاشة",
                        ),

                        "operating_system": cell(
                            row,
                            "نظام التشغيل",
                        ),

                        "ms_office": cell(
                            row,
                            "MS Office",
                        ),

                        "printer": cell(
                            row,
                            "طراز الطابعة",
                        ),

                        "scanner": cell(
                            row,
                            "الماسح الضوئى",
                        ),

                        "ip_address": cell(
                            row,
                            "IP",
                        ),

                        "status": (
                            cell(
                                row,
                                "الحالة الفنية",
                            )
                            or "صالح"
                        ),
                    }
                )

            wb.close()

            wb = None

            if not imported_rows:

                QMessageBox.warning(
                    self,
                    "تنبيه",
                    "لم يتم العثور على صفوف أجهزة.",
                )

                return

            # ==================================================
            # 3) اكتشاف الأسماء المكررة داخل Excel
            # ==================================================

            duplicate_names = {}

            for data in imported_rows:

                name_key = norm(
                    data["device_name"]
                )

                if not name_key:
                    continue

                duplicate_names.setdefault(
                    name_key,
                    0,
                )

                duplicate_names[
                    name_key
                ] += 1

            duplicate_names = {
                name: count
                for name, count
                in duplicate_names.items()
                if count > 1
            }

            duplicate_note = ""

            if duplicate_names:

                duplicate_note = (
                    "\n\n⚠ تم اكتشاف "
                    f"{len(duplicate_names)} "
                    "اسم جهاز مكرر داخل Excel."
                    "\nلن يتم اعتبار اسم الجهاز وحده "
                    "مفتاحًا للمطابقة."
                )

            # ==================================================
            # 4) Preview
            # ==================================================

            preview = "\n".join(
                f"{index + 1}. "
                f"{data['device_number']} | "
                f"{data['device_name']} | "
                f"{data['building']} | "
                f"{data['status']}"
                for index, data
                in enumerate(
                    imported_rows[:12]
                )
            )

            more = ""

            if len(imported_rows) > 12:

                more = (
                    f"\n... و "
                    f"{len(imported_rows) - 12}"
                    f" جهاز آخر"
                )

            answer = QMessageBox.question(
                self,
                "تأكيد الاستيراد",
                f"تم العثور على "
                f"{len(imported_rows)} جهاز.\n\n"
                f"{preview}"
                f"{more}"
                f"{duplicate_note}"
                "\n\n"
                "سيتم نقل كل الأعمدة الـ21 إلى النظام.\n"
                "المباني غير الموجودة سيتم إنشاؤها تلقائياً.\n\n"
                "طريقة مطابقة الجهاز:\n"
                "1) رقم الجهاز\n"
                "2) Serial الهارد\n"
                "3) اسم الجهاز + رقم القفل\n"
                "4) اسم الجهاز + User name\n"
                "5) إذا لم توجد مطابقة سيتم إنشاء جهاز جديد.\n\n"
                "هل تريد بدء الاستيراد؟",
                QMessageBox.Yes
                | QMessageBox.No,
                QMessageBox.No,
            )

            if answer != QMessageBox.Yes:
                return

            # ==================================================
            # 5) Progress
            # ==================================================

            progress = QProgressDialog(
                "جاري استيراد جميع بيانات الأجهزة...",
                "إلغاء",
                0,
                len(imported_rows),
                self,
            )

            progress.setWindowTitle(
                "استيراد Excel"
            )

            progress.setWindowModality(
                Qt.WindowModal
            )

            progress.setMinimumDuration(
                0
            )

            progress.show()

            con = get_connection()

            cur = con.cursor()

            new_count = 0
            update_count = 0
            cancel = False

            # ==================================================
            # مهم جدًا:
            #
            # يمنع استخدام نفس سجل SQL Server مرتين
            # أثناء نفس عملية الاستيراد.
            #
            # هذا هو الجزء الذي يحل مشكلة الأسماء المكررة.
            # ==================================================

            matched_device_ids = set()

            # لتسجيل كيفية مطابقة كل جهاز
            match_by_number = 0
            match_by_serial = 0
            match_by_name_lock = 0
            match_by_name_username = 0
            inserted_new = 0

            # ==================================================
            # 6) Import Loop
            # ==================================================

            for number, data in enumerate(
                imported_rows,
                1,
            ):

                progress.setValue(
                    number
                )

                if progress.wasCanceled():

                    cancel = True

                    break

                # ------------------------------------------------
                # Building
                # ------------------------------------------------

                building_id = (
                    self.get_building_id(
                        data["building"],
                        con,
                    )
                )

                existing = None
                match_method = ""

                # ==================================================
                # MATCH #1
                # Device Number
                #
                # لا نستخدم السجل إذا كان قد تم استخدامه
                # بالفعل في صف Excel آخر.
                # ==================================================

                device_number = norm(
                    data["device_number"]
                )

                if device_number:

                    cur.execute(
                        """
                        SELECT TOP 1
                            id
                        FROM devices
                        WHERE LTRIM(RTRIM(device_number)) = ?
                          AND is_archived = 0
                        ORDER BY id ASC
                        """,
                        (
                            device_number,
                        ),
                    )

                    candidates = cur.fetchall()

                    for candidate in candidates:

                        candidate_id = int(
                            candidate["id"]
                        )

                        if (
                            candidate_id
                            not in matched_device_ids
                        ):

                            existing = candidate

                            match_method = (
                                "device_number"
                            )

                            break

                # ==================================================
                # MATCH #2
                # Hard Disk Serial
                # ==================================================

                if not existing:

                    hard_disk_serial = norm(
                        data[
                            "hard_disk_serial"
                        ]
                    )

                    if hard_disk_serial:

                        cur.execute(
                            """
                            SELECT TOP 1
                                id
                            FROM devices
                            WHERE LTRIM(RTRIM(hard_disk_serial)) = ?
                              AND is_archived = 0
                            ORDER BY id ASC
                            """,
                            (
                                hard_disk_serial,
                            ),
                        )

                        candidates = cur.fetchall()

                        for candidate in candidates:

                            candidate_id = int(
                                candidate["id"]
                            )

                            if (
                                candidate_id
                                not in matched_device_ids
                            ):

                                existing = candidate

                                match_method = (
                                    "hard_disk_serial"
                                )

                                break

                # ==================================================
                # MATCH #3
                # Device Name + Lock Number
                #
                # مهم للأجهزة التي لها نفس اسم الجهاز.
                # ==================================================

                if not existing:

                    device_name = norm(
                        data[
                            "device_name"
                        ]
                    )

                    lock_number = norm(
                        data[
                            "lock_number"
                        ]
                    )

                    if (
                        device_name
                        and lock_number
                    ):

                        cur.execute(
                            """
                            SELECT TOP 1
                                id
                            FROM devices
                            WHERE LTRIM(RTRIM(device_name)) = ?
                              AND LTRIM(RTRIM(lock_number)) = ?
                              AND is_archived = 0
                            ORDER BY id ASC
                            """,
                            (
                                device_name,
                                lock_number,
                            ),
                        )

                        candidates = cur.fetchall()

                        for candidate in candidates:

                            candidate_id = int(
                                candidate["id"]
                            )

                            if (
                                candidate_id
                                not in matched_device_ids
                            ):

                                existing = candidate

                                match_method = (
                                    "name_lock"
                                )

                                break

                # ==================================================
                # MATCH #4
                # Device Name + User Name EN
                # ==================================================

                if not existing:

                    device_name = norm(
                        data[
                            "device_name"
                        ]
                    )

                    user_name_en = norm(
                        data[
                            "user_name_en"
                        ]
                    )

                    if (
                        device_name
                        and user_name_en
                    ):

                        cur.execute(
                            """
                            SELECT TOP 1
                                id
                            FROM devices
                            WHERE LTRIM(RTRIM(device_name)) = ?
                              AND LTRIM(RTRIM(user_name_en)) = ?
                              AND is_archived = 0
                            ORDER BY id ASC
                            """,
                            (
                                device_name,
                                user_name_en,
                            ),
                        )

                        candidates = cur.fetchall()

                        for candidate in candidates:

                            candidate_id = int(
                                candidate["id"]
                            )

                            if (
                                candidate_id
                                not in matched_device_ids
                            ):

                                existing = candidate

                                match_method = (
                                    "name_username"
                                )

                                break

                # ==================================================
                # Values
                # ==================================================

                values = (
                    data["device_number"],
                    data["device_name"],
                    data["user_name"],
                    "",
                    "",
                    "",
                    data["processor"],
                    data["ram"],
                    data["hard_disk"],
                    data["monitor"],
                    data["operating_system"],
                    data["printer"],
                    "",
                    data["status"],
                    building_id,
                    data["location"],
                    data["user_name_en"],
                    data["lock_number"],
                    data["network_connection"],
                    data["motherboard"],
                    data["graphics_card"],
                    data["hard_disk_serial"],
                    data["ms_office"],
                    data["scanner"],
                    data["ip_address"],
                )

                # ==================================================
                # UPDATE EXISTING
                # ==================================================

                if existing:

                    existing_id = int(
                        existing["id"]
                    )

                    cur.execute(
                        """
                        UPDATE devices
                        SET
                            device_number=?,
                            device_name=?,
                            user_name=?,
                            department=?,
                            device_type=?,
                            serial_number=?,
                            processor=?,
                            ram=?,
                            hard_disk=?,
                            monitor=?,
                            operating_system=?,
                            printer=?,
                            production_year=?,
                            status=?,
                            building_id=?,
                            location=?,
                            user_name_en=?,
                            lock_number=?,
                            network_connection=?,
                            motherboard=?,
                            graphics_card=?,
                            hard_disk_serial=?,
                            ms_office=?,
                            scanner=?,
                            ip_address=?,
                            updated_at=CURRENT_TIMESTAMP
                        WHERE id=?
                        """,
                        values
                        + (
                            existing_id,
                        ),
                    )

                    matched_device_ids.add(
                        existing_id
                    )

                    update_count += 1

                    if (
                        match_method
                        == "device_number"
                    ):
                        match_by_number += 1

                    elif (
                        match_method
                        == "hard_disk_serial"
                    ):
                        match_by_serial += 1

                    elif (
                        match_method
                        == "name_lock"
                    ):
                        match_by_name_lock += 1

                    elif (
                        match_method
                        == "name_username"
                    ):
                        match_by_name_username += 1

                # ==================================================
                # INSERT NEW
                # ==================================================

                else:

                    cur.execute(
                        """
                        INSERT INTO devices (
                            device_number,
                            device_name,
                            user_name,
                            department,
                            device_type,
                            serial_number,
                            processor,
                            ram,
                            hard_disk,
                            monitor,
                            operating_system,
                            printer,
                            production_year,
                            status,
                            building_id,
                            location,
                            user_name_en,
                            lock_number,
                            network_connection,
                            motherboard,
                            graphics_card,
                            hard_disk_serial,
                            ms_office,
                            scanner,
                            ip_address
                        )
                        VALUES (
                            ?,?,?,?,?,?,?,?,?,?,
                            ?,?,?,?,?,?,?,?,?,?,
                            ?,?,?,?,?
                        )
                        """,
                        values,
                    )

                    # الحصول على ID الجديد
                    new_id = None

                    try:

                        cur.execute(
                            """
                            SELECT CAST(
                                SCOPE_IDENTITY()
                                AS INT
                            ) AS id
                            """
                        )

                        identity_row = (
                            cur.fetchone()
                        )

                        if identity_row:

                            new_id = (
                                identity_row["id"]
                            )

                    except Exception:

                        new_id = None

                    if new_id is not None:

                        matched_device_ids.add(
                            int(new_id)
                        )

                    new_count += 1
                    inserted_new += 1

            # ==================================================
            # 7) Save
            # ==================================================

            con.commit()

            con.close()

            con = None

            progress.close()

            progress = None

            self.load_devices()

            # ==================================================
            # 8) Audit Log
            # ==================================================

            write_audit_log(
                user=self.current_user,
                action="استيراد",
                module_name="الأجهزة",
                description=(
                    "تم استيراد ملف Excel للأجهزة"
                    f"\nالملف: {file_name}"
                    f"\nإجمالي الصفوف: "
                    f"{len(imported_rows)}"
                    f"\nأجهزة جديدة: "
                    f"{new_count}"
                    f"\nأجهزة محدثة: "
                    f"{update_count}"
                    f"\nمطابقة برقم الجهاز: "
                    f"{match_by_number}"
                    f"\nمطابقة بـ Serial الهارد: "
                    f"{match_by_serial}"
                    f"\nمطابقة باسم الجهاز + القفل: "
                    f"{match_by_name_lock}"
                    f"\nمطابقة باسم الجهاز + User name: "
                    f"{match_by_name_username}"
                    f"\nتم الإلغاء: "
                    f"{'نعم' if cancel else 'لا'}"
                ),
                new_values={
                    "file": file_name,
                    "total_rows": len(
                        imported_rows
                    ),
                    "new_devices": new_count,
                    "updated_devices": update_count,
                    "match_by_device_number":
                        match_by_number,
                    "match_by_hard_disk_serial":
                        match_by_serial,
                    "match_by_name_lock":
                        match_by_name_lock,
                    "match_by_name_username":
                        match_by_name_username,
                    "cancelled": cancel,
                },
            )

            # ==================================================
            # 9) Result
            # ==================================================

            if cancel:

                QMessageBox.information(
                    self,
                    "تم إيقاف الاستيراد",
                    "تم حفظ ما تم استيراده حتى لحظة الإلغاء.\n\n"
                    f"أجهزة جديدة: {new_count}\n"
                    f"أجهزة محدثة: {update_count}\n\n"
                    f"إجمالي الصفوف المعالجة: "
                    f"{new_count + update_count}",
                )

            else:

                QMessageBox.information(
                    self,
                    "تم الاستيراد بالكامل",
                    "تم استيراد الملف بالكامل.\n\n"
                    f"إجمالي صفوف Excel: "
                    f"{len(imported_rows)}\n\n"
                    f"أجهزة جديدة: {new_count}\n"
                    f"أجهزة محدثة: {update_count}\n\n"
                    "تم التعامل مع أسماء الأجهزة المكررة "
                    "بدون استبدال سجل بجهاز آخر.\n\n"
                    "افتح صفحة الأجهزة وتأكد أن العدد أصبح مطابقًا للملف.",
                )

        except Exception as error:

            if wb is not None:

                try:
                    wb.close()
                except Exception:
                    pass

            if con is not None:

                try:
                    con.rollback()
                except Exception:
                    pass

                try:
                    con.close()
                except Exception:
                    pass

            if progress is not None:

                try:
                    progress.close()
                except Exception:
                    pass

            QMessageBox.critical(
                self,
                "خطأ في الاستيراد",
                "حدث خطأ أثناء الاستيراد:\n\n"
                + str(error),
            )