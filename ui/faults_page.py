from datetime import datetime
import html
import re
import tempfile
import webbrowser
from pathlib import Path

from PySide6.QtCore import Qt, QDate, QDateTime, QEvent, QObject, QTimer, Signal
from PySide6.QtGui import QColor, QKeySequence, QShortcut
from PySide6.QtWidgets import (
    QApplication,
    QWidget,
    QVBoxLayout,
    QHBoxLayout,
    QGridLayout,
    QLabel,
    QPushButton,
    QLineEdit,
    QTableWidget,
    QTableWidgetItem,
    QMessageBox,
    QDialog,
    QComboBox,
    QTextEdit,
    QDialogButtonBox,
    QInputDialog,
    QHeaderView,
    QFrame,
    QStyle,
    QDateTimeEdit,
    QDateEdit,
    QSpinBox,
    QCalendarWidget,
    QScrollArea,
    QSizePolicy,
    QCompleter,
    QFileDialog,
    QProgressDialog,
    QMenu,
    QSplitter,
)

from database import get_connection, write_audit_log
from ui.org_cache import OrgStructureCache


# =============================================================
# ثوابت
# =============================================================

STATUS_REPAIRED = "تم الإصلاح"
STATUS_DELIVERED = "تم التسليم"

# الحالة: (لون النص، لون الخلفية)
STATUS_COLORS = {
    "جديد": ("#0369a1", "#e0f2fe"),
    "قيد الفحص": ("#6d28d9", "#ede9fe"),
    "قيد الإصلاح": ("#b45309", "#fef3c7"),
    "بانتظار قطعة غيار": ("#c2410c", "#ffedd5"),
    STATUS_REPAIRED: ("#15803d", "#dcfce7"),
    STATUS_DELIVERED: ("#0f766e", "#ccfbf1"),
}

PRIORITY_COLORS = {
    "عادي": ("#475569", None),
    "متوسط": ("#b45309", "#fef3c7"),
    "عاجل": ("#c2410c", "#ffedd5"),
    "حرج": ("#b91c1c", "#fee2e2"),
}


def fault_audit_data(fault):
    """Return a JSON-friendly snapshot of the persisted fault fields."""
    if not fault:
        return {}

    keys = (
        "id", "fault_number", "entry_date", "building_id", "device_id",
        "user_name", "receiver_name", "fault_description", "repair_actions",
        "spare_parts", "technician", "priority", "status", "repair_date",
        "delivery_date", "notes", "created_at", "updated_at",
    )

    snapshot = {}

    for key in keys:
        try:
            value = fault[key]
        except (KeyError, IndexError, TypeError):
            continue

        if hasattr(value, "isoformat"):
            value = value.isoformat()

        snapshot[key] = value

    return snapshot


# =============================================================
# مساعدات التاريخ
# =============================================================

def is_unset(edit):
    """الحقل في وضع (غير محدد / كل التواريخ)؟"""
    return edit.dateTime() <= edit.minimumDateTime()


def clear_datetime(edit):
    edit.setDateTime(edit.minimumDateTime())


class _CalendarTodayFilter(QObject):
    """يخلي تقويم الاختيار يفتح على الشهر الحالي لما الحقل يكون فاضي.

    من غير ده التقويم بيفتح على سنة 1752 (أقل تاريخ في Qt).
    """

    def __init__(self, edit):
        super().__init__(edit)
        self._edit = edit
        edit.calendarWidget().installEventFilter(self)

    def eventFilter(self, obj, event):
        if event.type() == QEvent.Show and is_unset(self._edit):
            today = QDate.currentDate()
            obj.setCurrentPage(today.year(), today.month())
        return False


def keep_calendar_on_today(edit):
    return _CalendarTodayFilter(edit)


def parse_datetime(value):
    """يحوّل القيمة المخزنة (نص أو datetime) إلى datetime أو None."""
    if not value:
        return None

    if hasattr(value, "strftime"):
        return value

    raw = str(value).strip()

    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            continue

    return None


def format_datetime(value):
    parsed = parse_datetime(value)

    if parsed:
        return parsed.strftime("%Y-%m-%d %H:%M")

    return str(value or "")


APP_STYLE = """
QWidget {
    font-family: "IBM Plex Sans Arabic", "Segoe UI", "Tahoma";
    font-size: 12px;
    color: #334155;
}

#pageRoot {
    background: #f4f7fb;
}

#pageTitle {
    color: #0f2338;
    font-size: 26px;
    font-weight: 800;
}

#pageSubtitle {
    color: #64748b;
    font-size: 12px;
}

#toolBar {
    background: white;
    border: 1px solid #e2e8f0;
    border-radius: 12px;
}

#filterCard {
    background: white;
    border: 1px solid #e2e8f0;
    border-radius: 12px;
}

#filterLabel {
    color: #64748b;
    font-weight: 700;
}

#statTitle {
    color: #475569;
    font-size: 12px;
    font-weight: 700;
}

#statValue {
    color: #0f2338;
    font-size: 28px;
    font-weight: 800;
}

#noticeLabel {
    color: #15803d;
    font-weight: 700;
}

#emptyLabel {
    color: #94a3b8;
    font-size: 14px;
    font-weight: 600;
    padding: 40px;
}

#searchBox {
    background: #f8fafc;
    border: 1px solid #cbd5e1;
    border-radius: 9px;
    padding: 9px 12px;
    min-height: 22px;
}

#searchBox:focus {
    border: 1.5px solid #0891b2;
    background: white;
}

#primaryButton {
    background: #0891b2;
    color: white;
    border: none;
    border-radius: 9px;
    padding: 9px 16px;
    font-weight: 700;
}

#primaryButton:hover {
    background: #0e7490;
}

#primaryButton:pressed {
    background: #155e75;
}

#primaryButton:disabled {
    background: #cbd5e1;
}

#secondaryButton {
    background: #f1f5f9;
    color: #0f2338;
    border: 1px solid #cbd5e1;
    border-radius: 9px;
    padding: 9px 16px;
    font-weight: 700;
}

#secondaryButton:hover {
    background: #e2e8f0;
}

#secondaryButton:disabled {
    color: #94a3b8;
    background: #f8fafc;
}

#dangerButton {
    background: #fef2f2;
    color: #991b1b;
    border: 1px solid #fecaca;
    border-radius: 9px;
    padding: 9px 16px;
    font-weight: 700;
}

#dangerButton:hover {
    background: #fee2e2;
}

#dangerButton:disabled {
    color: #cbd5e1;
    background: #f8fafc;
    border: 1px solid #e2e8f0;
}

#clearFilterButton {
    background: #fff7ed;
    color: #9a3412;
    border: 1px solid #fed7aa;
    border-radius: 9px;
    padding: 8px 14px;
    font-weight: 700;
}

#clearFilterButton:hover {
    background: #ffedd5;
}

#clearFilterButton:disabled {
    color: #cbd5e1;
    background: #f8fafc;
    border: 1px solid #e2e8f0;
}

QComboBox {
    background: #fbfdff;
    border: 1px solid #cbd5e1;
    border-radius: 8px;
    padding: 8px 10px;
    min-width: 145px;
}

QComboBox:focus {
    border: 1.5px solid #0891b2;
    background: white;
}

QDateEdit {
    background: #fbfdff;
    border: 1px solid #cbd5e1;
    border-radius: 8px;
    padding: 8px 10px;
    min-width: 0;
}

QDateEdit:focus {
    border: 1.5px solid #0891b2;
    background: white;
}

#tableCard {
    background: white;
    border: 1px solid #e2e8f0;
    border-radius: 14px;
}

QTableWidget {
    background: white;
    border: none;
    gridline-color: #f1f5f9;
    color: #334155;
    selection-background-color: #e0f2fe;
    selection-color: #0f2338;
    alternate-background-color: #f8fafc;
}

QHeaderView::section {
    background: #f8fafc;
    color: #0f2338;
    border: none;
    border-bottom: 2px solid #0891b2;
    padding: 10px;
    font-weight: 700;
}

#dialogRoot {
    background: #f5f8fb;
}

#dialogHeader {
    background: #0f2338;
    color: white;
    border-radius: 12px;
    border-bottom: 3px solid #0891b2;
}

#dialogHeaderIcon {
    color: #26c6da;
    font-size: 28px;
    font-weight: 900;
}

#dialogTitle {
    color: white;
    font-size: 19px;
    font-weight: 800;
}

#dialogHint {
    color: #94a3b8;
    font-size: 11px;
}

#formLabel {
    color: #4b6074;
    font-weight: 700;
    padding: 2px 3px;
}

QLineEdit, QComboBox, QDateTimeEdit, QTextEdit {
    background: #fbfdff;
    border: 1px solid #d5dfe9;
    border-radius: 8px;
    padding: 8px 10px;
}

QLineEdit:focus, QComboBox:focus, QDateTimeEdit:focus, QTextEdit:focus {
    border: 1.5px solid #0891b2;
    background: white;
}

QTextEdit {
    min-height: 55px;
}

QDialogButtonBox QPushButton {
    min-width: 95px;
    padding: 8px 14px;
    border-radius: 8px;
    border: 1px solid #d5dfe9;
    background: white;
}

QDialogButtonBox QPushButton:hover {
    background: #eef3f7;
}

/* ---------- أزرار شريط الأدوات الجديد (مثل شاشة الأجهزة) ---------- */

#faultsAddBtn {
    background: qlineargradient(x1:0, y1:0, x2:1, y2:0, stop:0 #0891b2, stop:1 #06b6d4);
    color: white;
    border: none;
    border-radius: 9px;
    padding: 0 16px;
    min-height: 38px;
    max-height: 38px;
    font-size: 13px;
    font-weight: 700;
}

#faultsAddBtn:hover {
    background: qlineargradient(x1:0, y1:0, x2:1, y2:0, stop:0 #0e7490, stop:1 #0891b2);
}

#faultsAddBtn:pressed {
    background: #155e75;
}

#faultsImportBtn {
    background: #ecfeff;
    color: #0e7490;
    border: 1px solid #a5f3fc;
    border-radius: 9px;
    padding: 0 15px;
    min-height: 38px;
    max-height: 38px;
    font-size: 13px;
    font-weight: 700;
}

#faultsImportBtn:hover {
    background: #cffafe;
    border-color: #67e8f9;
    color: #0891b2;
}

#faultsImportBtn:pressed {
    background: #a5f3fc;
}

#faultsMoreBtn {
    background: white;
    color: #475569;
    border: 1px solid #cbd5e1;
    border-radius: 9px;
    padding: 0 14px;
    min-height: 38px;
    max-height: 38px;
    font-size: 13px;
    font-weight: 700;
}

#faultsMoreBtn:hover {
    background: #f8fafc;
    border-color: #94a3b8;
    color: #0f172a;
}

#faultsMoreBtn:pressed {
    background: #e2e8f0;
}

#faultsMoreBtn::menu-indicator {
    image: none;
    width: 0;
}

#selectionBar {
    background: #f0fdf4;
    border: 1px solid #bbf7d0;
    border-radius: 9px;
}

#selectionBadge {
    color: #166534;
    font-weight: 800;
    font-size: 12px;
    padding: 0 4px;
}

#faultsEditBtn {
    background: white;
    color: #334155;
    border: 1px solid #cbd5e1;
    border-radius: 7px;
    padding: 0 13px;
    min-height: 30px;
    max-height: 30px;
    font-size: 12px;
    font-weight: 700;
}

#faultsEditBtn:hover {
    background: #f8fafc;
    border-color: #94a3b8;
    color: #0f172a;
}

#faultsEditBtn:pressed {
    background: #e2e8f0;
}

#faultsDeleteBtn {
    background: #fff1f2;
    color: #e11d48;
    border: 1px solid #fecdd3;
    border-radius: 7px;
    padding: 0 13px;
    min-height: 30px;
    max-height: 30px;
    font-size: 12px;
    font-weight: 700;
}

#faultsDeleteBtn:hover {
    background: #ffe4e6;
    border-color: #fda4af;
    color: #be123c;
}

#faultsDeleteBtn:pressed {
    background: #fecdd3;
}

#selectionDeselectBtn {
    background: transparent;
    color: #64748b;
    border: none;
    border-radius: 6px;
    padding: 2px 8px;
    min-height: 30px;
    font-size: 12px;
    font-weight: 700;
}

#selectionDeselectBtn:hover {
    background: #dcfce7;
    color: #15803d;
}

QMenu {
    background: white;
    border: 1px solid #cbd5e1;
    border-radius: 9px;
    padding: 5px;
    font-size: 12px;
}

QMenu::item {
    padding: 7px 20px;
    border-radius: 6px;
    color: #334155;
    font-weight: 600;
}

QMenu::item:selected {
    background: #f1f5f9;
    color: #0f2338;
}

QMenu::item:disabled {
    color: #cbd5e1;
}

QMenu::separator {
    height: 1px;
    background: #e2e8f0;
    margin: 4px 6px;
}
"""


# =============================================================
# كارت إحصائية
# =============================================================

class FaultStatCard(QFrame):
    """كارت إحصائية بنقطة لون ورقم كبير. بعض الكروت قابلة للضغط للفلترة."""

    clicked = Signal()

    def __init__(self, title, accent="#0891b2", clickable=False):
        super().__init__()

        self.accent = accent
        self._clickable = clickable
        self._active = False

        self.setObjectName("faultStatCard")
        self.setMinimumHeight(92)

        if clickable:
            self.setCursor(Qt.PointingHandCursor)
            self.setToolTip("اضغط للفلترة")

        layout = QVBoxLayout(self)
        layout.setContentsMargins(16, 12, 16, 12)
        layout.setSpacing(4)

        top = QHBoxLayout()
        top.setSpacing(8)

        dot = QLabel()
        dot.setFixedSize(10, 10)
        dot.setStyleSheet(
            f"background: {accent}; border-radius: 5px;"
        )
        top.addWidget(dot)

        title_label = QLabel(title)
        title_label.setObjectName("statTitle")
        top.addWidget(title_label)
        top.addStretch()

        layout.addLayout(top)

        self.value_label = QLabel("0")
        self.value_label.setObjectName("statValue")
        layout.addWidget(self.value_label)

        self._apply_style()

    def _apply_style(self):
        bg = "#f8fafc" if self._active else "white"
        border = self.accent if self._active else "#e2e8f0"

        hover = ""

        if self._clickable:
            hover = f"""
            #faultStatCard:hover {{
                border: 1px solid {self.accent};
                border-right: 4px solid {self.accent};
            }}
            """

        self.setStyleSheet(f"""
            #faultStatCard {{
                background: {bg};
                border: 1px solid {border};
                border-right: 4px solid {self.accent};
                border-radius: 12px;
            }}
            {hover}
        """)

    def set_value(self, value):
        self.value_label.setText(str(value))

    def set_active(self, active):
        if active != self._active:
            self._active = active
            self._apply_style()

    def mousePressEvent(self, event):
        if self._clickable and event.button() == Qt.LeftButton:
            self.clicked.emit()

        super().mousePressEvent(event)


# =============================================================
# نافذة سجل أعطال جهاز
# =============================================================

HISTORY_STYLE = """
#histHeader {
    background: qlineargradient(x1:0, y1:0, x2:1, y2:0,
                                stop:0 #0b1d33, stop:1 #0e4a63);
    border-radius: 16px;
    border-right: 5px solid #22d3ee;
}

#histEyebrow {
    color: #9be7f5;
    font-size: 13px;
    font-weight: 700;
}

#histDeviceName {
    color: white;
    font-size: 26px;
    font-weight: 800;
}

#histHeaderChip {
    background: rgba(255, 255, 255, 0.10);
    color: #e2f4f8;
    border: 1px solid rgba(255, 255, 255, 0.22);
    border-radius: 12px;
    padding: 5px 14px;
    font-size: 11px;
    font-weight: 700;
}

#histHeaderMeta {
    background: rgba(255, 255, 255, 0.08);
    color: #c4ecf5;
    border: 1px solid rgba(255, 255, 255, 0.14);
    border-radius: 10px;
    padding: 3px 11px;
    font-size: 11px;
    font-weight: 600;
}

#histTableCard {
    background: white;
    border: 1px solid #e2e8f0;
    border-radius: 14px;
}

#histTable {
    background: white;
    border: none;
    alternate-background-color: #f8fafc;
    selection-background-color: #e0f2fe;
    selection-color: #0f2338;
}

#histTable::item {
    padding: 0 8px;
    border-bottom: 1px solid #f1f5f9;
}

#histTable QHeaderView::section {
    background: white;
    color: #64748b;
    border: none;
    border-bottom: 1px solid #e2e8f0;
    padding: 10px 8px;
    font-size: 11px;
    font-weight: 800;
}

#histPanel {
    background: white;
    border: 1px solid #e2e8f0;
    border-radius: 14px;
}

#histFaultNo {
    color: #0f2338;
    font-size: 20px;
    font-weight: 800;
}

#histMetaChip {
    background: #f1f5f9;
    color: #475569;
    border-radius: 11px;
    padding: 4px 12px;
    font-size: 11px;
    font-weight: 700;
}

#histStepTitle {
    color: #64748b;
    font-size: 10px;
    font-weight: 700;
}

#histSectionTitle {
    font-size: 11px;
    font-weight: 800;
}

#histValue {
    color: #0f2338;
    font-size: 12px;
    font-weight: 700;
}

#histText {
    background: transparent;
    border: none;
    padding: 0;
    color: #334155;
    font-size: 12px;
    min-height: 44px;
}

#histHint {
    color: #94a3b8;
    font-size: 11px;
}

#histEmpty {
    color: #64748b;
    font-size: 15px;
    font-weight: 700;
    background: white;
    border: 1px dashed #cbd5e1;
    border-radius: 14px;
    padding: 50px;
}

QSplitter::handle {
    background: transparent;
}
"""


class DeviceFaultHistoryDialog(QDialog):
    """كل الأعطال السابقة لجهاز واحد: امتى اتصلح وايه اللي اتعمل.

    القائمة على اليمين وتفاصيل العطل المحدد على الشمال في نفس النافذة.
    """

    HEADERS = [
        "رقم العطل", "تاريخ الدخول", "وصف العطل",
        "الحالة", "القائم بالإصلاح", "تاريخ الإصلاح",
    ]

    # (عنوان الخطوة، لون التظليل، لون الإطار، لون التمييز)
    STEP_STYLES = [
        ("تاريخ الدخول", "#ecfeff", "#a5f3fc", "#0891b2"),
        ("تاريخ الإصلاح", "#f0fdf4", "#bbf7d0", "#16a34a"),
        ("تاريخ التسليم", "#f0fdfa", "#99f6e4", "#0f766e"),
    ]

    def __init__(self, device_id, device_name="", parent=None,
                 exclude_fault_id=None):
        super().__init__(parent)

        self.device_id = device_id
        self.device_name = device_name or "-"
        self.exclude_fault_id = exclude_fault_id
        self._all_rows = []
        self._rows = []

        self.setWindowTitle(f"سجل أعطال الجهاز: {self.device_name}")
        self.resize(1320, 860)
        self.setMinimumSize(1100, 700)
        self.setLayoutDirection(Qt.RightToLeft)
        self.setObjectName("dialogRoot")
        self.setStyleSheet(APP_STYLE + HISTORY_STYLE)

        root = QVBoxLayout(self)
        root.setContentsMargins(20, 18, 20, 14)
        root.setSpacing(14)

        root.addWidget(self._build_header())
        root.addLayout(self._build_stats())

        # قائمة الأعطال + لوحة التفاصيل في نفس الشاشة
        self.content = QWidget()
        content_layout = QVBoxLayout(self.content)
        content_layout.setContentsMargins(0, 0, 0, 0)
        content_layout.setSpacing(10)
        content_layout.addWidget(self._build_toolbar())

        self.splitter = QSplitter(Qt.Horizontal)
        self.splitter.setChildrenCollapsible(False)
        self.splitter.setHandleWidth(12)
        self.splitter.addWidget(self._build_table_card())
        self.splitter.addWidget(self._build_details_panel())
        self.splitter.setStretchFactor(0, 5)
        self.splitter.setStretchFactor(1, 6)
        self.splitter.setSizes([520, 700])
        content_layout.addWidget(self.splitter, 1)

        root.addWidget(self.content, 1)

        # ---------- حالة عدم وجود أعطال ----------
        self.empty_label = QLabel(
            "✓\n\nلا توجد أعطال سابقة لهذا الجهاز\n"
            "سجل نظيف حتى الآن"
        )
        self.empty_label.setObjectName("histEmpty")
        self.empty_label.setAlignment(Qt.AlignCenter)
        self.empty_label.setVisible(False)
        root.addWidget(self.empty_label, 1)

        # ---------- الشريط السفلي ----------
        footer = QHBoxLayout()
        hint = QLabel("اختر أي عطل من القائمة لعرض تفاصيله ومسار الإصلاح")
        hint.setObjectName("histHint")
        footer.addWidget(hint)
        footer.addStretch()

        close_button = QPushButton("إغلاق")
        close_button.setObjectName("primaryButton")
        close_button.setCursor(Qt.PointingHandCursor)
        close_button.setMinimumWidth(120)
        close_button.clicked.connect(self.accept)
        footer.addWidget(close_button)

        root.addLayout(footer)

        self.load_history()

    # =========================================================
    # بناء الواجهة
    # =========================================================

    def _load_device_meta(self):
        """بيانات مختصرة عن الجهاز تظهر تحت اسمه في الهيدر."""
        row = None
        connection = get_connection()

        try:
            cursor = connection.cursor()
            cursor.execute("""
                SELECT
                    devices.user_name,
                    devices.serial_number,
                    buildings.name AS building_name
                FROM devices
                LEFT JOIN buildings
                    ON devices.building_id = buildings.id
                WHERE devices.id = ?
            """, (self.device_id,))
            row = cursor.fetchone()
        except Exception:
            row = None
        finally:
            connection.close()

        meta = []

        if row:
            for label, key in (
                ("المبنى", "building_name"),
                ("المستخدم", "user_name"),
                ("الرقم التسلسلي", "serial_number"),
            ):
                value = str(row[key] or "").strip()

                if value:
                    meta.append(f"{label}: {value}")

        return meta

    def _build_header(self):
        header = QFrame()
        header.setObjectName("histHeader")
        header.setLayoutDirection(Qt.RightToLeft)

        layout = QHBoxLayout(header)
        layout.setContentsMargins(26, 18, 26, 18)
        layout.setSpacing(14)

        # محاذاة فعلية لليمين بغض النظر عن اتجاه الواجهة
        right = Qt.AlignAbsolute | Qt.AlignRight | Qt.AlignVCenter

        text_box = QVBoxLayout()
        text_box.setSpacing(5)

        eyebrow = QLabel("الأعطال والإصلاحات السابقة")
        eyebrow.setObjectName("histEyebrow")
        eyebrow.setAlignment(right)

        name = QLabel(self.device_name)
        name.setObjectName("histDeviceName")
        name.setWordWrap(True)
        name.setAlignment(right)

        text_box.addWidget(eyebrow)
        text_box.addWidget(name)

        meta = self._load_device_meta()

        if meta:
            meta_row = QHBoxLayout()
            meta_row.setSpacing(8)
            meta_row.setContentsMargins(0, 4, 0, 0)

            for text in meta:
                chip = QLabel(text)
                chip.setObjectName("histHeaderMeta")
                meta_row.addWidget(chip)

            meta_row.addStretch()
            text_box.addLayout(meta_row)

        layout.addLayout(text_box, 1)

        self.header_count_chip = QLabel("")
        self.header_count_chip.setObjectName("histHeaderChip")
        layout.addWidget(self.header_count_chip, 0, Qt.AlignVCenter)

        if self.exclude_fault_id:
            note = QLabel("لا يشمل العطل الحالي")
            note.setObjectName("histHeaderChip")
            layout.addWidget(note, 0, Qt.AlignVCenter)

        return header

    def _build_stats(self):
        row = QHBoxLayout()
        row.setSpacing(12)

        self.stat_cards = {}

        items = [
            ("total", "إجمالي الأعطال", "#0891b2"),
            ("last_fault", "آخر عطل", "#7c3aed"),
            ("last_repair", "آخر إصلاح", "#16a34a"),
            ("avg_time", "متوسط مدة الإصلاح", "#d97706"),
            ("technician", "الأكثر إصلاحًا للجهاز", "#0f766e"),
        ]

        for key, title, accent in items:
            card = FaultStatCard(title, accent, False)
            card.setMinimumHeight(88)
            card.value_label.setStyleSheet(
                "font-size:18px;font-weight:800;color:#0f2338;"
            )
            row.addWidget(card, 1)
            self.stat_cards[key] = card

        return row

    def _build_toolbar(self):
        bar = QFrame()
        bar.setObjectName("toolBar")

        layout = QHBoxLayout(bar)
        layout.setContentsMargins(12, 8, 12, 8)
        layout.setSpacing(10)

        self.search = QLineEdit()
        self.search.setObjectName("searchBox")
        self.search.setClearButtonEnabled(True)
        self.search.setPlaceholderText(
            "ابحث في وصف العطل أو الإجراء أو قطع الغيار أو القائم بالإصلاح..."
        )
        self.search.textChanged.connect(self._apply_filter)
        layout.addWidget(self.search, 1)

        self.status_filter = QComboBox()
        self.status_filter.addItem("كل الحالات")
        self.status_filter.currentIndexChanged.connect(self._apply_filter)
        layout.addWidget(self.status_filter)

        self.result_label = QLabel("")
        self.result_label.setObjectName("histHint")
        layout.addWidget(self.result_label)

        return bar

    def _build_table_card(self):
        card = QFrame()
        card.setObjectName("histTableCard")

        layout = QVBoxLayout(card)
        layout.setContentsMargins(6, 6, 6, 6)

        self.table = QTableWidget()
        self.table.setObjectName("histTable")
        self.table.setColumnCount(len(self.HEADERS))
        self.table.setHorizontalHeaderLabels(self.HEADERS)
        self.table.setSelectionBehavior(QTableWidget.SelectRows)
        self.table.setSelectionMode(QTableWidget.SingleSelection)
        self.table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.table.setAlternatingRowColors(True)
        self.table.setShowGrid(False)
        self.table.setWordWrap(False)
        self.table.setFrameShape(QFrame.NoFrame)
        self.table.verticalHeader().setVisible(False)
        self.table.verticalHeader().setDefaultSectionSize(46)

        header = self.table.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.ResizeToContents)
        header.setSectionResizeMode(2, QHeaderView.Stretch)
        header.setHighlightSections(False)

        self.table.itemSelectionChanged.connect(self._show_details)

        # القائمة مختصرة، والباقي يظهر في لوحة التفاصيل
        self.table.setColumnHidden(4, True)
        self.table.setColumnHidden(5, True)

        layout.addWidget(self.table)

        return card

    def _make_step(self, title, tint, border, accent):
        frame = QFrame()
        frame.setObjectName("histStep")

        layout = QHBoxLayout(frame)
        layout.setContentsMargins(16, 10, 16, 10)
        layout.setSpacing(12)

        dot = QLabel()
        dot.setFixedSize(12, 12)
        layout.addWidget(dot, 0, Qt.AlignVCenter)

        text = QVBoxLayout()
        text.setSpacing(2)

        title_label = QLabel(title)
        title_label.setObjectName("histStepTitle")

        value_label = QLabel("-")
        value_label.setObjectName("histValue")

        text.addWidget(title_label)
        text.addWidget(value_label)
        layout.addLayout(text, 1)

        step = {
            "frame": frame,
            "dot": dot,
            "value": value_label,
            "tint": tint,
            "border": border,
            "accent": accent,
        }

        self._set_step(step, "-")

        return step

    @staticmethod
    def _set_step(step, text):
        done = text != "-"

        step["frame"].setStyleSheet(
            "#histStep {"
            f"background:{step['tint'] if done else '#f8fafc'};"
            f"border:1px {'solid' if done else 'dashed'} "
            f"{step['border'] if done else '#cbd5e1'};"
            "border-radius:12px;}"
        )

        step["dot"].setStyleSheet(
            f"background:{step['accent'] if done else '#cbd5e1'};"
            "border-radius:6px;"
        )

        step["value"].setText(text)
        step["value"].setStyleSheet(
            "color:#0f2338;font-weight:800;font-size:12px;"
            if done
            else "color:#94a3b8;font-weight:600;font-size:12px;"
        )

    def _make_section(self, title, body, accent):
        section = QFrame()
        section.setObjectName("histSection")
        section.setStyleSheet(
            "#histSection {"
            "background:#fbfdff;"
            "border:1px solid #e2e8f0;"
            f"border-right:4px solid {accent};"
            "border-radius:12px;}"
        )

        layout = QVBoxLayout(section)
        layout.setContentsMargins(16, 12, 16, 12)
        layout.setSpacing(6)

        title_label = QLabel(title)
        title_label.setObjectName("histSectionTitle")
        title_label.setStyleSheet(f"color:{accent};")

        layout.addWidget(title_label)
        layout.addWidget(body)

        return section

    def _make_text_box(self):
        box = QTextEdit()
        box.setObjectName("histText")
        box.setReadOnly(True)
        box.setFrameShape(QFrame.NoFrame)
        box.setMinimumHeight(60)
        box.setMaximumHeight(130)
        return box

    def _build_details_panel(self):
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.NoFrame)
        scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        scroll.setStyleSheet("QScrollArea{background:transparent;}")

        panel = QFrame()
        panel.setObjectName("histPanel")

        layout = QVBoxLayout(panel)
        layout.setContentsMargins(20, 16, 20, 16)
        layout.setSpacing(14)

        # --- رقم العطل + الحالة + الأولوية ---
        top = QHBoxLayout()
        top.setSpacing(8)

        self.detail_number = QLabel("-")
        self.detail_number.setObjectName("histFaultNo")
        top.addWidget(self.detail_number)

        self.detail_status = QLabel("")
        self.detail_status.setVisible(False)
        top.addWidget(self.detail_status)

        self.detail_priority = QLabel("")
        self.detail_priority.setVisible(False)
        top.addWidget(self.detail_priority)

        top.addStretch()
        layout.addLayout(top)

        # --- شريحتا المدة والعمر ---
        chips = QHBoxLayout()
        chips.setSpacing(8)

        self.detail_duration = QLabel("")
        self.detail_duration.setObjectName("histMetaChip")
        self.detail_duration.setVisible(False)
        chips.addWidget(self.detail_duration)

        self.detail_age = QLabel("")
        self.detail_age.setObjectName("histMetaChip")
        self.detail_age.setVisible(False)
        chips.addWidget(self.detail_age)

        chips.addStretch()
        layout.addLayout(chips)

        # --- الخط الزمني: دخول ← إصلاح ← تسليم ---
        timeline = QHBoxLayout()
        timeline.setSpacing(0)

        self.steps = [self._make_step(*style) for style in self.STEP_STYLES]
        self.connectors = []

        for index, step in enumerate(self.steps):
            timeline.addWidget(step["frame"], 1)

            if index < len(self.steps) - 1:
                line = QFrame()
                line.setFixedHeight(3)
                line.setMinimumWidth(24)
                line.setStyleSheet("background:#e2e8f0;border-radius:1px;")
                timeline.addWidget(line, 0, Qt.AlignVCenter)
                self.connectors.append(line)

        layout.addLayout(timeline)

        # --- الوصف + الإجراء ---
        self.detail_description = self._make_text_box()
        self.detail_actions = self._make_text_box()

        layout.addWidget(
            self._make_section("وصف العطل", self.detail_description, "#dc2626")
        )
        layout.addWidget(
            self._make_section("الإجراء المتخذ", self.detail_actions, "#16a34a")
        )

        # --- قطع الغيار + القائم بالإصلاح ---
        row1 = QHBoxLayout()
        row1.setSpacing(12)

        self.detail_parts = QLabel("-")
        self.detail_parts.setObjectName("histValue")
        self.detail_parts.setWordWrap(True)

        self.detail_technician = QLabel("-")
        self.detail_technician.setObjectName("histValue")
        self.detail_technician.setWordWrap(True)

        row1.addWidget(
            self._make_section(
                "قطع الغيار المستخدمة", self.detail_parts, "#d97706"
            ), 1
        )
        row1.addWidget(
            self._make_section(
                "القائم بالإصلاح", self.detail_technician, "#0891b2"
            ), 1
        )
        layout.addLayout(row1)

        # --- المستخدم + المستلم ---
        row2 = QHBoxLayout()
        row2.setSpacing(12)

        self.detail_user = QLabel("-")
        self.detail_user.setObjectName("histValue")
        self.detail_user.setWordWrap(True)

        self.detail_receiver = QLabel("-")
        self.detail_receiver.setObjectName("histValue")
        self.detail_receiver.setWordWrap(True)

        row2.addWidget(
            self._make_section("اسم المستخدم", self.detail_user, "#7c3aed"), 1
        )
        row2.addWidget(
            self._make_section("اسم المستلم", self.detail_receiver, "#0f766e"), 1
        )
        layout.addLayout(row2)

        # --- ملاحظات ---
        self.detail_notes = self._make_text_box()
        self.detail_notes.setMinimumHeight(44)
        layout.addWidget(
            self._make_section("ملاحظات", self.detail_notes, "#64748b")
        )

        layout.addStretch()

        scroll.setWidget(panel)
        return scroll

    # =========================================================
    # مساعدات
    # =========================================================

    @staticmethod
    def _valid_dt(value):
        """يرجع datetime صالح أو None (ويتجاهل القيمة الافتراضية القديمة 2000-01-01)."""
        parsed = parse_datetime(value)

        if not parsed:
            return None

        if not isinstance(parsed, datetime):
            try:
                parsed = datetime(parsed.year, parsed.month, parsed.day)
            except AttributeError:
                return None

        if parsed == datetime(2000, 1, 1):
            return None

        return parsed

    @classmethod
    def _fmt_optional(cls, value, with_time=True):
        parsed = cls._valid_dt(value)

        if not parsed:
            return "-"

        return parsed.strftime("%Y-%m-%d %H:%M" if with_time else "%Y-%m-%d")

    @staticmethod
    def _badge_style(colors):
        foreground, background = colors
        return (
            f"color:{foreground};"
            f"background:{background or '#f1f5f9'};"
            "border-radius:12px;padding:5px 15px;"
            "font-weight:800;font-size:11px;"
        )

    @staticmethod
    def _format_days(days):
        if days < 1:
            return "أقل من يوم"
        whole = int(round(days))
        if whole == 1:
            return "يوم واحد"
        if whole == 2:
            return "يومان"
        return f"{whole} يوم"

    # =========================================================
    # تحميل البيانات
    # =========================================================

    def load_history(self):
        sql = """
            SELECT id, fault_number, entry_date, fault_description,
                   repair_actions, spare_parts, technician, priority,
                   status, repair_date, delivery_date, user_name,
                   receiver_name, notes
            FROM faults
            WHERE device_id = ?
        """
        params = [self.device_id]

        if self.exclude_fault_id:
            sql += " AND id <> ?"
            params.append(self.exclude_fault_id)

        sql += " ORDER BY entry_date DESC, id DESC"

        connection = get_connection()
        try:
            cursor = connection.cursor()
            cursor.execute(sql, params)
            self._all_rows = cursor.fetchall()
        finally:
            connection.close()

        total = len(self._all_rows)

        self.header_count_chip.setText(f"{total} عطل مسجل")

        self._fill_stats()

        # قائمة الحالات الموجودة فعلًا في سجل هذا الجهاز
        present_statuses = []
        for fault in self._all_rows:
            status = fault["status"] or ""
            if status and status not in present_statuses:
                present_statuses.append(status)

        self.status_filter.blockSignals(True)
        self.status_filter.clear()
        self.status_filter.addItem("كل الحالات")
        for status in present_statuses:
            self.status_filter.addItem(status)
        self.status_filter.blockSignals(False)

        has_rows = total > 0
        self.content.setVisible(has_rows)
        self.empty_label.setVisible(not has_rows)

        if has_rows:
            self._apply_filter()

    def _fill_stats(self):
        from collections import Counter

        rows = self._all_rows
        cards = self.stat_cards

        cards["total"].set_value(len(rows))

        if not rows:
            for key in ("last_fault", "last_repair", "avg_time", "technician"):
                cards[key].set_value("-")
            return

        cards["last_fault"].set_value(
            self._fmt_optional(rows[0]["entry_date"], with_time=False)
        )

        repair_dates = [
            d for d in (self._valid_dt(f["repair_date"]) for f in rows) if d
        ]
        cards["last_repair"].set_value(
            max(repair_dates).strftime("%Y-%m-%d") if repair_dates else "-"
        )

        durations = []
        for fault in rows:
            entry = self._valid_dt(fault["entry_date"])
            repair = self._valid_dt(fault["repair_date"])

            if entry and repair and repair >= entry:
                durations.append((repair - entry).total_seconds() / 86400)

        cards["avg_time"].set_value(
            self._format_days(sum(durations) / len(durations))
            if durations
            else "-"
        )

        technicians = Counter(
            (f["technician"] or "").strip()
            for f in rows
            if (f["technician"] or "").strip()
        )

        if technicians:
            name, count = technicians.most_common(1)[0]
            cards["technician"].set_value(f"{name} ({count})")
        else:
            cards["technician"].set_value("-")

    # =========================================================
    # الفلترة وملء الجدول
    # =========================================================

    def _apply_filter(self, *_args):
        text = self.search.text().strip().lower()
        status = (
            self.status_filter.currentText()
            if self.status_filter.currentIndex() > 0
            else None
        )

        def matches(fault):
            if status and (fault["status"] or "") != status:
                return False

            if not text:
                return True

            haystack = " ".join(
                str(fault[key] or "")
                for key in (
                    "fault_number", "fault_description", "repair_actions",
                    "spare_parts", "technician",
                )
            ).lower()

            return text in haystack

        self._rows = [f for f in self._all_rows if matches(f)]
        self._fill_table()

    def _fill_table(self):
        self.table.setUpdatesEnabled(False)
        self.table.setRowCount(0)

        def one_line(value):
            return (value or "").replace("\n", " ").strip()

        for fault in self._rows:
            row = self.table.rowCount()
            self.table.insertRow(row)

            description = one_line(fault["fault_description"])

            values = [
                fault["fault_number"],
                format_datetime(fault["entry_date"]),
                description,
                fault["status"] or "",
                fault["technician"] or "",
                self._fmt_optional(fault["repair_date"]),
            ]

            for column, value in enumerate(values):
                item = QTableWidgetItem(str(value))
                item.setTextAlignment(Qt.AlignCenter)

                if column == 0:
                    font = item.font()
                    font.setBold(True)
                    item.setFont(font)

                if column == 2:
                    item.setToolTip(description)

                if column == 3 and value in STATUS_COLORS:
                    FaultsPage._paint_badge(item, STATUS_COLORS[value])

                if column in (4, 5) and value in ("", "-"):
                    item.setText("-")
                    item.setForeground(QColor("#94a3b8"))

                self.table.setItem(row, column, item)

        self.table.setUpdatesEnabled(True)

        shown = len(self._rows)
        total = len(self._all_rows)

        self.result_label.setText(
            f"{shown} من {total}" if shown != total else f"{total} عطل"
        )

        if self._rows:
            self.table.selectRow(0)
        else:
            self._clear_details()

    # =========================================================
    # التفاصيل
    # =========================================================

    def _update_connectors(self):
        for index, line in enumerate(self.connectors):
            next_step = self.steps[index + 1]
            done = next_step["value"].text() != "-"

            line.setStyleSheet(
                f"background:{next_step['accent'] if done else '#e2e8f0'};"
                "border-radius:1px;"
            )

    def _clear_details(self):
        self.detail_number.setText("لا توجد نتائج مطابقة")
        self.detail_status.setVisible(False)
        self.detail_priority.setVisible(False)
        self.detail_duration.setVisible(False)
        self.detail_age.setVisible(False)

        for step in self.steps:
            self._set_step(step, "-")

        self._update_connectors()

        self.detail_description.clear()
        self.detail_actions.clear()
        self.detail_notes.clear()
        self.detail_parts.setText("-")
        self.detail_technician.setText("-")
        self.detail_user.setText("-")
        self.detail_receiver.setText("-")

    def _show_details(self):
        selected = self.table.selectedItems()

        if not selected:
            return

        index = selected[0].row()

        if index >= len(self._rows):
            return

        fault = self._rows[index]

        self.detail_number.setText(f"عطل رقم {fault['fault_number']}")

        status = fault["status"] or ""
        self.detail_status.setVisible(bool(status))
        self.detail_status.setText(status)
        self.detail_status.setStyleSheet(
            self._badge_style(
                STATUS_COLORS.get(status, ("#475569", "#f1f5f9"))
            )
        )

        priority = fault["priority"] or ""
        self.detail_priority.setVisible(bool(priority))
        self.detail_priority.setText(f"الأولوية: {priority}")
        self.detail_priority.setStyleSheet(
            self._badge_style(
                PRIORITY_COLORS.get(priority, ("#475569", "#f1f5f9"))
            )
        )

        entry = self._valid_dt(fault["entry_date"])
        repair = self._valid_dt(fault["repair_date"])

        # مدة الإصلاح
        if entry and repair and repair >= entry:
            self.detail_duration.setText(
                "مدة الإصلاح: "
                + self._format_days(
                    (repair - entry).total_seconds() / 86400
                )
            )
            self.detail_duration.setVisible(True)
        else:
            self.detail_duration.setVisible(False)

        # عمر العطل
        if entry:
            days = (datetime.now() - entry).days

            if days <= 0:
                age_text = "سُجّل اليوم"
            elif days == 1:
                age_text = "سُجّل أمس"
            else:
                age_text = f"سُجّل منذ {days} يوم"

            self.detail_age.setText(age_text)
            self.detail_age.setVisible(True)
        else:
            self.detail_age.setVisible(False)

        self._set_step(
            self.steps[0], self._fmt_optional(fault["entry_date"])
        )
        self._set_step(
            self.steps[1], self._fmt_optional(fault["repair_date"])
        )
        self._set_step(
            self.steps[2], self._fmt_optional(fault["delivery_date"])
        )
        self._update_connectors()

        self.detail_description.setPlainText(
            fault["fault_description"] or "-"
        )
        self.detail_actions.setPlainText(fault["repair_actions"] or "-")
        self.detail_notes.setPlainText(fault["notes"] or "-")
        self.detail_parts.setText(fault["spare_parts"] or "-")
        self.detail_technician.setText(
            (fault["technician"] or "").strip() or "-"
        )
        self.detail_user.setText((fault["user_name"] or "").strip() or "-")
        self.detail_receiver.setText(
            (fault["receiver_name"] or "").strip() or "-"
        )


# =============================================================
# نافذة إضافة / تعديل عطل
# =============================================================

class FaultDialog(QDialog):

    def __init__(self, parent=None, fault=None):
        super().__init__(parent)

        self.fault = fault

        self.setWindowTitle("تعديل عطل" if fault else "إضافة عطل")

        screen = self.screen() or QApplication.primaryScreen()
        available = screen.availableGeometry() if screen else None

        if available:
            width = max(320, min(840, int(available.width() * 0.86)))
            height = max(380, min(640, int(available.height() * 0.82)))
        else:
            width, height = 800, 620

        self.resize(width, height)
        self.setMinimumSize(min(620, width), min(480, height))

        if available:
            self.move(available.center() - self.rect().center())

        self.setObjectName("dialogRoot")
        self.setStyleSheet(APP_STYLE)

        root = QVBoxLayout(self)
        root.setContentsMargins(14, 14, 14, 12)
        root.setSpacing(12)

        header = QFrame()
        header.setObjectName("dialogHeader")

        hl = QHBoxLayout(header)
        hl.setContentsMargins(16, 13, 16, 13)

        icon = QLabel("")
        icon.setObjectName("dialogHeaderIcon")
        icon.setAlignment(Qt.AlignCenter)
        icon.setFixedSize(42, 42)
        hl.addWidget(icon)

        text = QVBoxLayout()

        title = QLabel(
            "تسجيل عطل جديد" if not fault else "تعديل بيانات العطل"
        )
        title.setObjectName("dialogTitle")

        hint = QLabel("سجل بيانات العطل ومتابعة الإصلاح من مكان واحد")
        hint.setObjectName("dialogHint")

        text.addWidget(title)
        text.addWidget(hint)

        hl.addLayout(text)
        hl.addStretch()

        root.addWidget(header)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.NoFrame)
        scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)

        form_widget = QWidget()

        form = QGridLayout(form_widget)
        form.setContentsMargins(6, 8, 6, 8)
        form.setHorizontalSpacing(12)
        form.setVerticalSpacing(8)
        form.setColumnMinimumWidth(0, 86)
        form.setColumnMinimumWidth(2, 86)
        form.setColumnStretch(1, 1)
        form.setColumnStretch(3, 1)

        self.fault_number = QLineEdit()
        self.fault_number.setReadOnly(True)

        self.entry_date = QDateTimeEdit()
        self.entry_date.setCalendarPopup(True)
        self.entry_date.setDateTime(datetime.now())
        self.entry_date.setDisplayFormat("yyyy-MM-dd HH:mm")

        self.device = QComboBox()
        self.device.setEditable(True)
        self.device.setInsertPolicy(QComboBox.NoInsert)
        self.device.setPlaceholderText(
            "اكتب اسم الجهاز للبحث أو اختر من القائمة..."
        )

        self.device.setCompleter(QCompleter())
        self.device.completer().setCaseSensitivity(Qt.CaseInsensitive)
        self.device.completer().setFilterMode(Qt.MatchContains)
        self.device.completer().setCompletionMode(
            QCompleter.PopupCompletion
        )

        self.device.lineEdit().editingFinished.connect(
            self.select_typed_device
        )
        self.device.currentIndexChanged.connect(self.device_changed)

        self.device_specs = QTextEdit()
        self.device_specs.setReadOnly(True)
        self.device_specs.setFixedHeight(72)
        self.device_specs.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        self.device_specs.setHorizontalScrollBarPolicy(
            Qt.ScrollBarAlwaysOff
        )

        self.building = QLineEdit()
        self.building.setReadOnly(True)

        self.user_name = QLineEdit()
        self.receiver_name = QLineEdit()

        self.fault_description = QTextEdit()
        self.repair_actions = QTextEdit()

        self.fault_description.setFixedHeight(74)
        self.repair_actions.setFixedHeight(68)

        self.spare_part = QComboBox()
        self.spare_part.addItem("بدون قطعة غيار", None)

        self.spare_part_qty = QSpinBox()
        self.spare_part_qty.setRange(1, 999999)
        self.spare_part_qty.setValue(1)
        self.spare_part_qty.setFixedWidth(110)

        spare_part_box = QHBoxLayout()
        spare_part_box.setContentsMargins(0, 0, 0, 0)
        spare_part_box.setSpacing(8)
        spare_part_box.addWidget(self.spare_part, 1)
        spare_part_box.addWidget(QLabel("الكمية:"))
        spare_part_box.addWidget(self.spare_part_qty)

        self.spare_part_container = QWidget()
        self.spare_part_container.setLayout(spare_part_box)

        self.technician = QComboBox()
        self.technician.setInsertPolicy(QComboBox.NoInsert)
        self.technician.addItem("اختر القائم بالإصلاح...", None)
        self.technician.addItem("＋ إضافة اسم جديد...", "__add_new__")
        self.technician.activated.connect(self._technician_choice)
        self.load_technicians()

        self.priority = QComboBox()
        self.priority.addItems(["عادي", "متوسط", "عاجل", "حرج"])

        self.status = QComboBox()
        self.status.addItems([
            "جديد",
            "قيد الفحص",
            "قيد الإصلاح",
            "بانتظار قطعة غيار",
            STATUS_REPAIRED,
            STATUS_DELIVERED,
        ])

        # activated = تغيير يدوي من المستخدم فقط (مش أثناء تحميل بيانات قديمة)
        self.status.activated.connect(self.status_chosen)

        # تاريخ الإصلاح / التسليم: الافتراضي "غير محدد" (مش 2000-01-01)
        self.repair_date = self._optional_datetime_edit()
        self.delivery_date = self._optional_datetime_edit()

        # زرار سجل أعطال الجهاز جنب قائمة الأجهزة
        self.history_button = QPushButton("سجل الأعطال (0)")
        self.history_button.setObjectName("secondaryButton")
        self.history_button.setStyleSheet("padding: 6px 12px;")
        self.history_button.setToolTip("عرض الأعطال السابقة لهذا الجهاز")
        self.history_button.setEnabled(False)
        self.history_button.clicked.connect(self.show_device_history)

        device_box = QHBoxLayout()
        device_box.setContentsMargins(0, 0, 0, 0)
        device_box.setSpacing(8)
        device_box.addWidget(self.device, 1)
        device_box.addWidget(self.history_button)

        self.device_container = QWidget()
        self.device_container.setLayout(device_box)

        def add_pair(row, right_label, right_widget, left_label, left_widget):
            for column, label_text, widget in (
                (0, right_label, right_widget),
                (2, left_label, left_widget),
            ):
                label = QLabel(label_text + ":")
                label.setObjectName("formLabel")
                label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)

                form.addWidget(label, row, column)
                form.addWidget(widget, row, column + 1)

        def add_full(row, label_text, widget):
            label = QLabel(label_text + ":")
            label.setObjectName("formLabel")
            label.setAlignment(Qt.AlignRight | Qt.AlignTop)

            form.addWidget(label, row, 0)
            form.addWidget(widget, row, 1, 1, 3)

        add_pair(0, "رقم العطل", self.fault_number,
                 "تاريخ الدخول", self.entry_date)
        add_pair(1, "الجهاز *", self.device_container,
                 "المبنى / الفرع", self.building)
        add_full(2, "مواصفات الجهاز", self.device_specs)
        add_pair(3, "اسم المستخدم", self.user_name,
                 "اسم المستلم", self.receiver_name)
        add_full(4, "توصيف العطل *", self.fault_description)
        add_full(5, "إجراءات الإصلاح", self.repair_actions)
        add_full(6, "قطع الغيار المستخدمة", self.spare_part_container)
        add_pair(7, "القائم بالإصلاح", self.technician,
                 "الأولوية", self.priority)
        add_pair(8, "حالة العطل", self.status,
                 "تاريخ الإصلاح", self._with_clear_button(self.repair_date))
        add_full(9, "تاريخ التسليم",
                 self._with_clear_button(self.delivery_date))

        scroll.setWidget(form_widget)
        root.addWidget(scroll, 1)

        buttons = QDialogButtonBox(
            QDialogButtonBox.Save | QDialogButtonBox.Cancel
        )
        buttons.accepted.connect(self.validate_and_accept)
        buttons.rejected.connect(self.reject)
        root.addWidget(buttons)

        self.load_devices()
        self.load_spare_parts()

        if fault:
            self.load_fault()
        else:
            self.spare_part_qty.setEnabled(True)

    # ---------- مساعدات الواجهة ----------

    def _optional_datetime_edit(self):
        edit = QDateTimeEdit()
        edit.setCalendarPopup(True)
        edit.setDisplayFormat("yyyy-MM-dd HH:mm")
        edit.setSpecialValueText("غير محدد")
        clear_datetime(edit)
        keep_calendar_on_today(edit)
        return edit

    def _with_clear_button(self, edit):
        box = QHBoxLayout()
        box.setContentsMargins(0, 0, 0, 0)
        box.setSpacing(8)

        clear_button = QPushButton("مسح")
        clear_button.setObjectName("secondaryButton")
        clear_button.setStyleSheet("padding: 6px 12px;")
        clear_button.setToolTip("إرجاع التاريخ إلى (غير محدد)")
        clear_button.clicked.connect(lambda: clear_datetime(edit))

        box.addWidget(edit, 1)
        box.addWidget(clear_button)

        container = QWidget()
        container.setLayout(box)
        return container

    def status_chosen(self, *_args):
        """لما المستخدم يختار (تم الإصلاح / تم التسليم) نملأ التاريخ بالوقت الحالي لو فاضي."""
        status = self.status.currentText()
        now = QDateTime.currentDateTime()

        if status == STATUS_REPAIRED and is_unset(self.repair_date):
            self.repair_date.setDateTime(now)

        elif status == STATUS_DELIVERED and is_unset(self.delivery_date):
            self.delivery_date.setDateTime(now)

    def select_typed_device(self):
        typed = self.device.currentText().strip()

        if not typed:
            return

        for i in range(self.device.count()):
            if self.device.itemText(i).strip() == typed:
                self.device.setCurrentIndex(i)
                return

    # ---------- تحميل البيانات ----------

    def load_devices(self):
        connection = get_connection()
        cursor = connection.cursor()

        cursor.execute("""
            SELECT
                devices.id,
                devices.device_name,
                devices.device_number,
                devices.user_name,
                devices.processor,
                devices.ram,
                devices.hard_disk,
                devices.serial_number,
                devices.monitor,
                devices.operating_system,
                devices.printer,
                buildings.name AS building_name
            FROM devices
            LEFT JOIN buildings
                ON devices.building_id = buildings.id
            WHERE devices.is_archived = 0
            ORDER BY devices.device_name
        """)

        devices = cursor.fetchall()
        connection.close()

        self.device.clear()
        self.device.addItem("اختر الجهاز...", None)

        for device in devices:
            self.device.addItem(device["device_name"], device["id"])

        self.device.completer().setModel(self.device.model())

    def load_technicians(self):
        """قائمة موحدة بالأسماء المسجلة، مع خيار لإضافة اسم جديد مرة واحدة."""
        current = self.fault["technician"] if self.fault else ""
        connection = get_connection()
        try:
            cursor = connection.cursor()
            cursor.execute("""
                SELECT DISTINCT LTRIM(RTRIM(technician)) AS technician
                FROM faults
                WHERE technician IS NOT NULL AND LTRIM(RTRIM(technician)) <> ''
                ORDER BY LTRIM(RTRIM(technician))
            """)
            names = [row["technician"] for row in cursor.fetchall()]
        finally:
            connection.close()
        self.technician.blockSignals(True)
        self.technician.clear()
        self.technician.addItem("اختر القائم بالإصلاح...", None)
        for name in names:
            self.technician.addItem(name, name)
        self.technician.addItem("＋ إضافة اسم جديد...", "__add_new__")
        if current:
            index = self.technician.findData(current.strip())
            if index < 0:
                self.technician.insertItem(self.technician.count() - 1, current.strip(), current.strip())
                index = self.technician.findData(current.strip())
            self.technician.setCurrentIndex(index)
        self.technician.blockSignals(False)

    def _technician_choice(self, index):
        if self.technician.itemData(index) != "__add_new__":
            return
        name, accepted = QInputDialog.getText(
            self, "إضافة قائم بالإصلاح", "أدخل اسم القائم بالإصلاح الجديد:"
        )
        name = name.strip()
        if not accepted or not name:
            self.technician.setCurrentIndex(0)
            return
        existing = self.technician.findText(name, Qt.MatchFixedString)
        if existing >= 0:
            self.technician.setCurrentIndex(existing)
            return
        self.technician.insertItem(self.technician.count() - 1, name, name)
        self.technician.setCurrentIndex(self.technician.findData(name))

    def load_spare_parts(self):
        connection = get_connection()
        cursor = connection.cursor()

        cursor.execute("""
            SELECT
                id,
                part_name,
                part_number,
                quantity
            FROM spare_parts
            WHERE is_archived = 0
            ORDER BY part_name
        """)

        parts = cursor.fetchall()
        connection.close()

        self.spare_part.clear()
        self.spare_part.addItem("بدون قطعة غيار", None)

        for part in parts:
            label = f"{part['part_name']} | المتاح: {part['quantity']}"

            if part["part_number"]:
                label += f" | {part['part_number']}"

            self.spare_part.addItem(label, part["id"])

    def _previous_faults_count(self, device_id):
        connection = get_connection()
        try:
            cursor = connection.cursor()
            sql = "SELECT COUNT(*) FROM faults WHERE device_id = ?"
            params = [device_id]
            if self.fault:
                sql += " AND id <> ?"
                params.append(self.fault["id"])
            cursor.execute(sql, params)
            return cursor.fetchone()[0] or 0
        finally:
            connection.close()

    def show_device_history(self):
        device_id = self.device.currentData()
        if not device_id:
            return

        DeviceFaultHistoryDialog(
            device_id,
            self.device.currentText(),
            self,
            exclude_fault_id=self.fault["id"] if self.fault else None,
        ).exec()

    def device_changed(self):
        device_id = self.device.currentData()

        if not device_id:
            self.building.clear()
            self.user_name.clear()
            self.device_specs.clear()
            self.history_button.setText("سجل الأعطال (0)")
            self.history_button.setEnabled(False)
            return

        connection = get_connection()
        cursor = connection.cursor()

        cursor.execute("""
            SELECT
                devices.*,
                buildings.name AS building_name
            FROM devices
            LEFT JOIN buildings
                ON devices.building_id = buildings.id
            WHERE devices.id = ?
        """, (device_id,))

        device = cursor.fetchone()
        connection.close()

        if not device:
            return

        self.building.setText(device["building_name"] or "")
        self.user_name.setText(device["user_name"] or "")

        specs = (
            f"النوع: {device['device_type'] or '-'}\n"
            f"المعالج: {device['processor'] or '-'}\n"
            f"RAM: {device['ram'] or '-'}\n"
            f"الهارد: {device['hard_disk'] or '-'}\n"
            f"Serial: {device['serial_number'] or '-'}\n"
            f"الشاشة: {device['monitor'] or '-'}\n"
            f"النظام: {device['operating_system'] or '-'}\n"
            f"الطابعة: {device['printer'] or '-'}"
        )

        self.device_specs.setPlainText(specs)

        count = self._previous_faults_count(device_id)
        self.history_button.setText(f"سجل الأعطال ({count})")
        self.history_button.setEnabled(True)

    def generate_fault_number(self):
        year = datetime.now().year

        connection = get_connection()
        cursor = connection.cursor()

        cursor.execute(
            "SELECT COUNT(*) FROM faults WHERE fault_number LIKE ?",
            (f"FLT-{year}-%",)
        )

        count = cursor.fetchone()[0]
        connection.close()

        return f"FLT-{year}-{count + 1:04d}"

    def validate_and_accept(self):
        if not self.device.currentData():
            QMessageBox.warning(self, "تنبيه", "من فضلك اختر الجهاز.")
            self.device.setFocus()
            return

        if not self.fault_description.toPlainText().strip():
            QMessageBox.warning(self, "تنبيه", "من فضلك اكتب توصيف العطل.")
            self.fault_description.setFocus()
            return

        self.accept()

    @staticmethod
    def _optional_text(edit):
        if is_unset(edit):
            return None

        return edit.dateTime().toString("yyyy-MM-dd HH:mm:ss")

    def get_data(self):
        building_id = None
        device_id = self.device.currentData()

        if device_id:
            connection = get_connection()
            cursor = connection.cursor()

            cursor.execute(
                "SELECT building_id FROM devices WHERE id = ?",
                (device_id,)
            )

            result = cursor.fetchone()
            connection.close()

            if result:
                building_id = result["building_id"]

        return {
            "fault_number": self.fault_number.text(),
            "entry_date": self.entry_date.dateTime().toString(
                "yyyy-MM-dd HH:mm:ss"
            ),
            "building_id": building_id,
            "device_id": device_id,
            "user_name": self.user_name.text().strip(),
            "receiver_name": self.receiver_name.text().strip(),
            "fault_description": self.fault_description.toPlainText().strip(),
            "repair_actions": self.repair_actions.toPlainText().strip(),
            "spare_part_id": self.spare_part.currentData(),
            "spare_part_qty": (
                self.spare_part_qty.value()
                if self.spare_part.currentData()
                else 0
            ),
            "spare_parts": (
                f"{self.spare_part.currentText().split(' | المتاح:')[0]} "
                f"× {self.spare_part_qty.value()}"
                if self.spare_part.currentData()
                else ""
            ),
            "technician": (self.technician.currentData() or "").strip(),
            "priority": self.priority.currentText(),
            "status": self.status.currentText(),
            "repair_date": self._optional_text(self.repair_date),
            "delivery_date": self._optional_text(self.delivery_date),
            "notes": (
                self.fault["notes"]
                if self.fault and "notes" in self.fault.keys()
                else ""
            ),
        }

    def load_fault(self):
        self.fault_number.setText(self.fault["fault_number"])

        entry = parse_datetime(self.fault["entry_date"])
        self.entry_date.setDateTime(entry or datetime.now())

        index = self.device.findData(self.fault["device_id"])

        if index >= 0:
            self.device.setCurrentIndex(index)

        self.user_name.setText(self.fault["user_name"] or "")
        self.receiver_name.setText(self.fault["receiver_name"] or "")
        self.fault_description.setPlainText(
            self.fault["fault_description"] or ""
        )
        self.repair_actions.setPlainText(self.fault["repair_actions"] or "")

        saved_parts = self.fault["spare_parts"] or ""

        self.spare_part.setEnabled(True)
        self.spare_part_qty.setEnabled(True)

        if saved_parts:
            match = re.search(r"(.+?)\s*[×xX]\s*(\d+)\s*$", saved_parts)

            if match:
                part_name = match.group(1).strip()
                qty = int(match.group(2))

                self.spare_part_qty.setValue(max(1, qty))

                for i in range(self.spare_part.count()):
                    label = self.spare_part.itemText(i)

                    if label.split(" | المتاح:")[0].strip() == part_name:
                        self.spare_part.setCurrentIndex(i)
                        break

            else:
                self.spare_part.setCurrentIndex(0)

        saved_technician = (self.fault["technician"] or "").strip()
        if saved_technician:
            index = self.technician.findData(saved_technician)
            if index < 0:
                self.technician.insertItem(self.technician.count() - 1, saved_technician, saved_technician)
                index = self.technician.findData(saved_technician)
            self.technician.setCurrentIndex(index)

        priority_index = self.priority.findText(self.fault["priority"])

        if priority_index >= 0:
            self.priority.setCurrentIndex(priority_index)

        status_index = self.status.findText(self.fault["status"])

        if status_index >= 0:
            self.status.setCurrentIndex(status_index)

        # القيمة 2000-01-01 كانت بتتسجل بالغلط كـ "تاريخ افتراضي" في بيانات قديمة
        legacy_default = datetime(2000, 1, 1, 0, 0, 0)

        for edit, key in (
            (self.repair_date, "repair_date"),
            (self.delivery_date, "delivery_date"),
        ):
            parsed = parse_datetime(self.fault[key])

            if parsed and parsed != legacy_default:
                edit.setDateTime(parsed)


# =============================================================
# نافذة تقارير موقف الأعطال
# =============================================================

class FaultReportsDialog(QDialog):

    def __init__(self, parent=None):
        super().__init__(parent)

        self.setWindowTitle("تقرير الأعطال حسب فترة")
        self.resize(1100, 700)
        self.setLayoutDirection(Qt.RightToLeft)

        root = QVBoxLayout(self)
        root.setContentsMargins(18, 18, 18, 18)
        root.setSpacing(12)

        title = QLabel("تقرير الأعطال حسب فترة")
        title.setStyleSheet("font-size:22px;font-weight:800;color:#0f2338;")
        root.addWidget(title)

        subtitle = QLabel(
            "حدد تاريخ البداية والنهاية لعرض الأعطال المسجلة خلال الفترة ثم معاينة التقرير وطباعته."
        )
        subtitle.setStyleSheet("color:#718096;")
        root.addWidget(subtitle)

        controls = QHBoxLayout()
        controls.setSpacing(10)

        controls.addWidget(QLabel("من تاريخ:"))
        self.date_from = QDateEdit()
        self.date_from.setCalendarPopup(True)
        self.date_from.setDisplayFormat("dd/MM/yyyy")
        self.date_from.setDate(QDate.currentDate())
        controls.addWidget(self.date_from)

        controls.addWidget(QLabel("إلى تاريخ:"))
        self.date_to = QDateEdit()
        self.date_to.setCalendarPopup(True)
        self.date_to.setDisplayFormat("dd/MM/yyyy")
        self.date_to.setDate(QDate.currentDate())
        controls.addWidget(self.date_to)
        self.date_from.dateChanged.connect(self.load_report)
        self.date_to.dateChanged.connect(self.load_report)

        root.addLayout(controls)

        self.period_label = QLabel()
        self.period_label.setStyleSheet(
            "font-weight:700;color:#334155;padding:5px;"
        )
        root.addWidget(self.period_label)

        self.summary = QLabel()
        self.summary.setStyleSheet(
            "background:#f4f7fb;"
            "border:1px solid #e3eaf2;"
            "border-radius:8px;"
            "padding:10px;"
        )
        root.addWidget(self.summary)

        self.table = QTableWidget()
        self.table.setColumnCount(8)
        self.table.setHorizontalHeaderLabels([
            "رقم العطل", "التاريخ", "الجهاز", "المستخدم",
            "الحالة", "الأولوية", "القائم بالإصلاح", "وصف العطل",
        ])
        self.table.setSelectionBehavior(QTableWidget.SelectRows)
        self.table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.table.verticalHeader().setVisible(False)
        self.table.horizontalHeader().setSectionResizeMode(
            QHeaderView.Stretch
        )
        self.table.setAlternatingRowColors(True)
        root.addWidget(self.table, 1)

        buttons = QHBoxLayout()

        self.refresh_button = QPushButton("تحديث")
        self.print_button = QPushButton("معاينة وطباعة")
        self.close_button = QPushButton("إغلاق")

        self.refresh_button.clicked.connect(self.load_report)
        self.print_button.clicked.connect(self.print_report)
        self.close_button.clicked.connect(self.accept)

        buttons.addStretch()
        buttons.addWidget(self.refresh_button)
        buttons.addWidget(self.print_button)
        buttons.addWidget(self.close_button)

        root.addLayout(buttons)

        self.load_report()

    def date_range(self):
        start = self.date_from.date()
        end = self.date_to.date()
        if start > end:
            start, end = end, start
        return start.toString("yyyy-MM-dd"), end.toString("yyyy-MM-dd")

    def load_report(self):
        start, end = self.date_range()

        self.period_label.setText(
            f"الفترة المحددة: {start}"
            if start == end
            else f"الفترة المحددة: من {start} إلى {end}"
        )

        con = get_connection()
        cur = con.cursor()

        cur.execute("""
            SELECT
                faults.fault_number,
                faults.entry_date,
                devices.device_name,
                faults.user_name,
                faults.status,
                faults.priority,
                faults.technician,
                faults.fault_description
            FROM faults
            LEFT JOIN devices
                ON faults.device_id = devices.id
            WHERE TRY_CONVERT(date, faults.entry_date) BETWEEN ? AND ?
            ORDER BY
                faults.entry_date DESC,
                faults.id DESC
        """, (start, end))

        rows = cur.fetchall()
        con.close()

        self.table.setRowCount(0)

        for fault in rows:
            row = self.table.rowCount()
            self.table.insertRow(row)

            values = [
                fault["fault_number"],
                fault["entry_date"],
                fault["device_name"] or "",
                fault["user_name"] or "",
                fault["status"] or "",
                fault["priority"] or "",
                fault["technician"] or "",
                fault["fault_description"] or "",
            ]

            for col, value in enumerate(values):
                item = QTableWidgetItem(str(value))
                item.setTextAlignment(Qt.AlignCenter)
                self.table.setItem(row, col, item)

        counts = {}

        for fault in rows:
            status = fault["status"] or "غير محدد"
            counts[status] = counts.get(status, 0) + 1

        details = "  |  ".join(f"{k}: {v}" for k, v in counts.items())

        self.summary.setText(
            f"إجمالي الأعطال: {len(rows)}"
            + (
                f"  |  {details}"
                if details
                else "  |  لا توجد أعطال خلال الفترة المحددة."
            )
        )

    def print_report(self):
        try:
            report_html = self.build_print_html()

            reports_dir = Path(tempfile.gettempdir()) / "NetworkFaultsReports"
            reports_dir.mkdir(parents=True, exist_ok=True)

            file_name = (
                "fault_report_"
                f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.html"
            )

            file_path = reports_dir / file_name
            file_path.write_text(report_html, encoding="utf-8")

            webbrowser.open(file_path.resolve().as_uri())

        except Exception as exc:
            QMessageBox.critical(
                self,
                "خطأ في التقرير",
                "حدث خطأ أثناء تجهيز التقرير للطباعة:\n\n"
                f"{exc}"
            )

    def build_print_html(self):
        title = html.escape("تقرير موقف الأعطال")
        period = html.escape(self.period_label.text())
        summary = html.escape(self.summary.text())

        rows_html = ""

        for row in range(self.table.rowCount()):
            cells = []

            for column in range(self.table.columnCount()):
                item = self.table.item(row, column)
                value = item.text() if item else ""
                cells.append(html.escape(str(value)))

            rows_html += f"""
                <tr>
                    <td>{cells[0]}</td>
                    <td>{cells[1]}</td>
                    <td>{cells[2]}</td>
                    <td>{cells[3]}</td>
                    <td>{cells[4]}</td>
                    <td>{cells[5]}</td>
                    <td>{cells[6]}</td>
                    <td class="description">{cells[7]}</td>
                </tr>
            """

        if not rows_html:
            rows_html = """
                <tr>
                    <td colspan="8" class="empty-row">
                        لا توجد أعطال خلال الفترة المحددة.
                    </td>
                </tr>
            """

        generated_at = html.escape(datetime.now().strftime("%Y-%m-%d %H:%M"))

        return f"""<!DOCTYPE html>
<html lang="ar" dir="rtl">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{title}</title>
<style>
@page {{ size: A4 landscape; margin: 10mm; }}
* {{ box-sizing: border-box; }}
html, body {{ margin: 0; padding: 0; }}
body {{
    direction: rtl;
    background: white;
    color: #1e293b;
    font-family: "IBM Plex Sans Arabic", "Segoe UI", Tahoma, Arial, sans-serif;
    font-size: 11px;
    line-height: 1.5;
}}
.report {{ width: 100%; max-width: 100%; margin: 0 auto; }}
.report-header {{
    border-bottom: 3px solid #0891b2;
    padding-bottom: 10px;
    margin-bottom: 12px;
}}
.report-title {{ font-size: 24px; font-weight: 800; color: #0f2338; margin-bottom: 4px; }}
.report-subtitle {{ color: #64748b; font-size: 11px; }}
.info-grid {{
    display: grid;
    grid-template-columns: 1fr 1fr 1fr;
    gap: 8px;
    margin-bottom: 12px;
}}
.info-card {{
    border: 1px solid #dbe3ec;
    background: #f8fafc;
    border-radius: 6px;
    padding: 8px 10px;
}}
.info-label {{ display: block; color: #64748b; font-size: 9px; margin-bottom: 2px; }}
.info-value {{ color: #0f2338; font-size: 11px; font-weight: 700; }}
table {{
    width: 100%;
    border-collapse: collapse;
    table-layout: fixed;
    margin-top: 8px;
}}
thead {{ display: table-header-group; }}
tfoot {{ display: table-footer-group; }}
tr {{ page-break-inside: avoid; break-inside: avoid; }}
th, td {{
    border: 1px solid #cbd5e1;
    padding: 6px 5px;
    text-align: center;
    vertical-align: middle;
    overflow-wrap: anywhere;
    word-break: break-word;
}}
th {{
    background: #0f2338;
    color: white;
    font-weight: 800;
    font-size: 10px;
    border-bottom: 2px solid #0891b2;
}}
td {{ font-size: 9.5px; background: white; }}
tbody tr:nth-child(even) td {{ background: #f8fafc; }}
.description {{ text-align: right; line-height: 1.45; }}
.empty-row {{ text-align: center; padding: 25px; color: #64748b; font-size: 12px; }}
.footer {{
    margin-top: 12px;
    padding-top: 8px;
    border-top: 1px solid #dbe3ec;
    display: flex;
    justify-content: space-between;
    color: #64748b;
    font-size: 9px;
}}
@media screen {{
    body {{ background: #e5e7eb; padding: 20px; }}
    .report {{
        background: white;
        padding: 18px;
        box-shadow: 0 3px 15px rgba(0, 0, 0, 0.12);
        max-width: 1200px;
    }}
}}
@media print {{
    body {{ background: white; padding: 0; }}
    .report {{ box-shadow: none; padding: 0; max-width: none; }}
    .no-print {{ display: none !important; }}
}}
</style>
</head>
<body>
<div class="report">
    <div class="report-header">
        <div class="report-title">{title}</div>
        <div class="report-subtitle">الشبكة - نظام إدارة الأجهزة والأعطال</div>
    </div>
    <div class="info-grid">
        <div class="info-card">
            <span class="info-label">فترة التقرير</span>
            <span class="info-value">{period}</span>
        </div>
        <div class="info-card">
            <span class="info-label">ملخص الأعطال</span>
            <span class="info-value">{summary}</span>
        </div>
        <div class="info-card">
            <span class="info-label">تاريخ إنشاء التقرير</span>
            <span class="info-value">{generated_at}</span>
        </div>
    </div>
    <table>
        <colgroup>
            <col style="width: 10%;">
            <col style="width: 12%;">
            <col style="width: 13%;">
            <col style="width: 12%;">
            <col style="width: 13%;">
            <col style="width: 10%;">
            <col style="width: 14%;">
            <col style="width: 16%;">
        </colgroup>
        <thead>
            <tr>
                <th>رقم العطل</th>
                <th>التاريخ</th>
                <th>الجهاز</th>
                <th>المستخدم</th>
                <th>الحالة</th>
                <th>الأولوية</th>
                <th>القائم بالإصلاح</th>
                <th>وصف العطل</th>
            </tr>
        </thead>
        <tbody>
            {rows_html}
        </tbody>
    </table>
    <div class="footer">
        <div>الشبكة - نظام إدارة الأجهزة والأعطال</div>
        <div>تم إنشاء التقرير بتاريخ {generated_at}</div>
    </div>
</div>
</body>
</html>
"""


# =============================================================
# صفحة الأعطال
# =============================================================

class FaultsPage(QWidget):

    PAGE_SIZE = 50

    FROM_SQL = """
        FROM faults
        LEFT JOIN devices ON faults.device_id = devices.id
        LEFT JOIN buildings ON faults.building_id = buildings.id
        WHERE 1 = 1
    """

    TABLE_HEADERS = [
        "ID", "رقم العطل", "التاريخ", "الجهاز", "المبنى",
        "المستخدم", "وصف العطل", "الحالة", "الأولوية",
        "القائم بالإصلاح",
    ]

    def __init__(self, current_user=None):
        super().__init__()

        self.current_user = current_user or {}

        self._page = 0
        self._total = 0

        self.setup_ui()

        self.load_building_filter()
        self.load_technician_filter()

        # استماع لتغييرات الهيكل التنظيمي — تحديث فلتر المباني تلقائياً
        OrgStructureCache.instance().data_changed.connect(self.load_building_filter)

        self.load_faults()

    # =========================================================
    # بناء الواجهة
    # =========================================================

    def setup_ui(self):
        self.setObjectName("pageRoot")
        self.setStyleSheet(APP_STYLE)

        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(26, 22, 26, 22)
        main_layout.setSpacing(14)

        # ---------- العنوان والأزرار ----------
        title_layout = QHBoxLayout()
        title_layout.setSpacing(8)
        title_layout.setAlignment(Qt.AlignVCenter)

        title_box = QVBoxLayout()
        title_box.setSpacing(2)

        title = QLabel("سجل الأعطال")
        title.setObjectName("pageTitle")
        title_box.addWidget(title)

        subtitle = QLabel("متابعة الأعطال وحالات الإصلاح والتسليم بسهولة")
        subtitle.setObjectName("pageSubtitle")
        title_box.addWidget(subtitle)

        title_layout.addLayout(title_box)
        title_layout.addStretch()

        # ---- شريط التحديد (يظهر فقط عند اختيار صفوف) ----
        self.selection_bar = QFrame()
        self.selection_bar.setObjectName("selectionBar")
        sel_layout = QHBoxLayout(self.selection_bar)
        sel_layout.setContentsMargins(8, 2, 8, 2)
        sel_layout.setSpacing(6)

        self.selection_label = QLabel("")
        self.selection_label.setObjectName("selectionBadge")
        sel_layout.addWidget(self.selection_label)

        self.edit_button = QPushButton("✎  تعديل")
        self.edit_button.setObjectName("faultsEditBtn")
        self.edit_button.setCursor(Qt.PointingHandCursor)
        self.edit_button.setToolTip("تعديل العطل المحدد")
        self.edit_button.clicked.connect(lambda: self.edit_fault())
        sel_layout.addWidget(self.edit_button)

        self.delete_button = QPushButton("🗑  حذف")
        self.delete_button.setObjectName("faultsDeleteBtn")
        self.delete_button.setCursor(Qt.PointingHandCursor)
        self.delete_button.setToolTip("حذف العطل أو الأعطال المحددة")
        self.delete_button.clicked.connect(lambda: self.delete_fault())
        sel_layout.addWidget(self.delete_button)

        self.deselect_button = QPushButton("✕  إلغاء")
        self.deselect_button.setObjectName("selectionDeselectBtn")
        self.deselect_button.setCursor(Qt.PointingHandCursor)
        self.deselect_button.setToolTip("إلغاء التحديد")
        self.deselect_button.clicked.connect(lambda: self.table.clearSelection())
        sel_layout.addWidget(self.deselect_button)

        self.selection_bar.setVisible(False)
        title_layout.addWidget(self.selection_bar)

        # ---- الزر الرئيسي ----
        self.add_button = QPushButton(" إضافة عطل")
        self.add_button.setObjectName("faultsAddBtn")
        self.add_button.setCursor(Qt.PointingHandCursor)
        self.add_button.setToolTip("تسجيل عطل جديد")
        self.add_button.clicked.connect(self.add_fault)
        title_layout.addWidget(self.add_button)

        # ---- استيراد ----
        self.import_button = QPushButton("  استيراد Excel")
        self.import_button.setObjectName("faultsImportBtn")
        self.import_button.setCursor(Qt.PointingHandCursor)
        self.import_button.setToolTip("استيراد الأعطال دفعة واحدة من ملف Excel")
        self.import_button.clicked.connect(self.import_from_excel)
        title_layout.addWidget(self.import_button)

        # ---- قائمة خيارات ----
        self.more_button = QPushButton("خيارات ▾")
        self.more_button.setObjectName("faultsMoreBtn")
        self.more_button.setCursor(Qt.PointingHandCursor)
        self.more_button.setToolTip("خيارات إضافية للجدول والتقارير")

        more_menu = QMenu(self.more_button)

        self.act_select_all = more_menu.addAction("  تحديد الكل")
        self.act_select_all.triggered.connect(self.toggle_select_all)

        more_menu.addSeparator()

        self.act_report = more_menu.addAction("  تقرير الأعطال حسب فترة...")
        self.act_report.triggered.connect(self.show_reports)

        self.act_print_filtered = more_menu.addAction("  طباعة نتائج الفلتر")
        self.act_print_filtered.triggered.connect(self.print_filtered_report)

        self.more_button.setMenu(more_menu)
        title_layout.addWidget(self.more_button)

        main_layout.addLayout(title_layout)

        # ---------- كروت الإحصائيات ----------
        stats_layout = QHBoxLayout()
        stats_layout.setSpacing(12)

        self.fault_stat_cards = {}

        stat_items = [
            ("total", "إجمالي الأعطال", "#0891b2", True),
            ("serviced", "أجهزة تمت صيانتها", "#7c3aed", False),
            ("repaired", STATUS_REPAIRED, "#16a34a", True),
            ("delivered", STATUS_DELIVERED, "#0f766e", True),
            ("technicians", "القائمون بالإصلاح", "#d97706", False),
            ("overdue", "مفتوحة ومتأخرة", "#dc2626", False),
        ]

        for key, title_text, accent, clickable in stat_items:
            card = FaultStatCard(title_text, accent, clickable)
            stats_layout.addWidget(card, 1)
            self.fault_stat_cards[key] = card

        self.fault_stat_cards["total"].clicked.connect(
            lambda: self.status_filter.setCurrentIndex(0)
        )
        self.fault_stat_cards["repaired"].clicked.connect(
            lambda: self.toggle_status_filter(STATUS_REPAIRED)
        )
        self.fault_stat_cards["delivered"].clicked.connect(
            lambda: self.toggle_status_filter(STATUS_DELIVERED)
        )

        main_layout.addLayout(stats_layout)

        # ---------- البحث + زر الفلاتر ----------
        search_card = QFrame()
        search_card.setObjectName("toolBar")

        search_layout = QHBoxLayout(search_card)
        search_layout.setContentsMargins(14, 10, 14, 10)
        search_layout.setSpacing(10)

        search_icon = QLabel()
        search_icon.setPixmap(
            self.style().standardIcon(
                QStyle.SP_FileDialogContentsView
            ).pixmap(20, 20)
        )
        search_layout.addWidget(search_icon)

        self.search_box = QLineEdit()
        self.search_box.setObjectName("searchBox")
        self.search_box.setClearButtonEnabled(True)
        self.search_box.setPlaceholderText(
            "بحث برقم العطل أو الجهاز أو المستخدم أو المبنى أو وصف العطل...  (Ctrl+F)"
        )
        search_layout.addWidget(self.search_box, 1)

        self.advanced_button = QPushButton("فلاتر متقدمة")
        self.advanced_button.setObjectName("secondaryButton")
        self.advanced_button.setCheckable(True)
        self.advanced_button.clicked.connect(self.toggle_advanced_filters)
        search_layout.addWidget(self.advanced_button)

        self.clear_filters_button = QPushButton("مسح الفلاتر")
        self.clear_filters_button.setObjectName("clearFilterButton")
        self.clear_filters_button.clicked.connect(self.clear_fault_filters)
        search_layout.addWidget(self.clear_filters_button)

        main_layout.addWidget(search_card)

        # البحث بعد توقف الكتابة (بدل استعلام مع كل حرف)
        self._search_timer = QTimer(self)
        self._search_timer.setSingleShot(True)
        self._search_timer.setInterval(300)
        self._search_timer.timeout.connect(self._reset_page_and_load_faults)
        self.search_box.textChanged.connect(
            lambda *_: self._search_timer.start()
        )

        # ---------- الفلاتر المتقدمة (مطوية افتراضيًا) ----------
        self.filter_card = QFrame()
        self.filter_card.setObjectName("filterCard")

        filter_layout = QGridLayout(self.filter_card)
        filter_layout.setContentsMargins(14, 12, 14, 12)
        filter_layout.setHorizontalSpacing(12)
        filter_layout.setVerticalSpacing(10)

        def add_filter_field(row, column, label_text, widget):
            field = QWidget()
            field_layout = QVBoxLayout(field)
            field_layout.setContentsMargins(0, 0, 0, 0)
            field_layout.setSpacing(4)

            label = QLabel(label_text)
            label.setObjectName("filterLabel")
            label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)

            widget.setMinimumWidth(0)
            widget.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
            widget.setMinimumHeight(38)

            field_layout.addWidget(label)
            field_layout.addWidget(widget)
            filter_layout.addWidget(field, row, column)

        self.status_filter = QComboBox()
        self.status_filter.addItems([
            "كل الحالات",
            "جديد",
            "قيد الفحص",
            "قيد الإصلاح",
            "بانتظار قطعة غيار",
            STATUS_REPAIRED,
            STATUS_DELIVERED,
        ])

        self.priority_filter = QComboBox()
        self.priority_filter.addItems([
            "كل الأولويات", "عادي", "متوسط", "عاجل", "حرج",
        ])

        self.building_filter = QComboBox()
        self.technician_filter = QComboBox()

        self.date_from_filter = self._make_date_filter()
        self.date_to_filter = self._make_date_filter()

        for combo in (
            self.status_filter,
            self.priority_filter,
            self.building_filter,
            self.technician_filter,
        ):
            combo.currentIndexChanged.connect(
                self._reset_page_and_load_faults
            )

        for date_edit in (self.date_from_filter, self.date_to_filter):
            date_edit.dateChanged.connect(self._reset_page_and_load_faults)

        add_filter_field(0, 0, "الحالة", self.status_filter)
        add_filter_field(0, 1, "الأولوية", self.priority_filter)
        add_filter_field(0, 2, "المبنى", self.building_filter)
        add_filter_field(0, 3, "القائم بالإصلاح", self.technician_filter)
        add_filter_field(1, 0, "تاريخ الدخول: من", self.date_from_filter)
        add_filter_field(1, 1, "تاريخ الدخول: إلى", self.date_to_filter)

        self.overdue_days = QSpinBox()
        self.overdue_days.setRange(1, 3650)
        self.overdue_days.setValue(7)
        self.overdue_days.setSuffix(" يوم")
        self.overdue_days.setToolTip("العطل المفتوح يُعد متأخرًا بعد هذا العدد من الأيام")
        self.overdue_days.valueChanged.connect(self.load_faults)
        add_filter_field(1, 2, "حد التأخير", self.overdue_days)

        for column in range(4):
            filter_layout.setColumnStretch(column, 1)

        self.filter_card.setVisible(False)
        main_layout.addWidget(self.filter_card)

        # ---------- الجدول ----------
        table_card = QFrame()
        table_card.setObjectName("tableCard")

        table_layout = QVBoxLayout(table_card)
        table_layout.setContentsMargins(8, 8, 8, 8)

        self.table = QTableWidget()
        self.table.setColumnCount(len(self.TABLE_HEADERS))
        self.table.setHorizontalHeaderLabels(self.TABLE_HEADERS)
        self.table.setSelectionBehavior(QTableWidget.SelectRows)
        self.table.setSelectionMode(QTableWidget.ExtendedSelection)
        self.table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.table.setAlternatingRowColors(True)
        self.table.setShowGrid(False)
        self.table.setWordWrap(False)
        self.table.verticalHeader().setVisible(False)
        self.table.verticalHeader().setDefaultSectionSize(40)
        self.table.setColumnHidden(0, True)

        header = self.table.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.ResizeToContents)
        header.setSectionResizeMode(3, QHeaderView.Stretch)
        header.setSectionResizeMode(6, QHeaderView.Stretch)

        self.table.doubleClicked.connect(self.edit_fault)
        self.table.itemSelectionChanged.connect(self.update_action_buttons)
        self.table.setContextMenuPolicy(Qt.CustomContextMenu)
        self.table.customContextMenuRequested.connect(self.show_table_menu)

        table_layout.addWidget(self.table)

        self.empty_label = QLabel("")
        self.empty_label.setObjectName("emptyLabel")
        self.empty_label.setAlignment(Qt.AlignCenter)
        self.empty_label.setWordWrap(True)
        self.empty_label.setVisible(False)
        table_layout.addWidget(self.empty_label, 1)

        pagination_row = QHBoxLayout()

        self.stats_label = QLabel("")
        self.stats_label.setStyleSheet("color: #64748b; font-size: 11px;")
        pagination_row.addWidget(self.stats_label)

        self.notice_label = QLabel("")
        self.notice_label.setObjectName("noticeLabel")
        pagination_row.addWidget(self.notice_label)

        self._notice_timer = QTimer(self)
        self._notice_timer.setSingleShot(True)
        self._notice_timer.setInterval(4000)
        self._notice_timer.timeout.connect(
            lambda: self.notice_label.setText("")
        )

        pagination_row.addStretch()

        self.previous_button = QPushButton("◀  السابق")
        self.previous_button.setObjectName("secondaryButton")
        self.previous_button.setMinimumHeight(32)
        self.previous_button.clicked.connect(self._previous_page)

        self.page_label = QLabel("صفحة 1 / 1")
        self.page_label.setAlignment(Qt.AlignCenter)
        self.page_label.setStyleSheet("font-weight: 700; min-width: 80px;")

        self.next_button = QPushButton("التالي  ▶")
        self.next_button.setObjectName("secondaryButton")
        self.next_button.setMinimumHeight(32)
        self.next_button.clicked.connect(self._next_page)

        pagination_row.addWidget(self.previous_button)
        pagination_row.addWidget(self.page_label)
        pagination_row.addWidget(self.next_button)

        table_layout.addLayout(pagination_row)

        main_layout.addWidget(table_card, 1)

        # ---------- اختصارات ----------
        QShortcut(QKeySequence("Ctrl+F"), self).activated.connect(
            self._focus_search
        )
        select_all_shortcut = QShortcut(QKeySequence("Ctrl+A"), self.table)
        select_all_shortcut.setContext(Qt.WidgetShortcut)
        select_all_shortcut.activated.connect(self.table.selectAll)
        QShortcut(QKeySequence(Qt.Key_Escape), self.table).activated.connect(
            self.table.clearSelection
        )
        QShortcut(QKeySequence("F5"), self).activated.connect(
            self.refresh_all
        )

        self.apply_permissions()
        self.update_action_buttons()
        self._update_filter_ui()

    def _make_date_filter(self):
        edit = QDateEdit()
        edit.setCalendarPopup(True)
        edit.setDisplayFormat("dd/MM/yyyy")
        edit.setSpecialValueText("كل التواريخ")
        clear_datetime(edit)
        keep_calendar_on_today(edit)
        return edit

    def _focus_search(self):
        self.search_box.setFocus()
        self.search_box.selectAll()

    # =========================================================
    # حالة الواجهة
    # =========================================================

    def show_notice(self, message):
        """رسالة نجاح خفيفة بدل صندوق حوار يوقف المستخدم."""
        self.notice_label.setText(message)
        self._notice_timer.start()

    def toggle_select_all(self):
        """تحديد أو إلغاء تحديد كل الصفوف الظاهرة بضغطة واحدة."""
        selected = len({i.row() for i in self.table.selectedIndexes()})
        total = self.table.rowCount()

        if total == 0:
            return

        if selected >= total:
            self.table.clearSelection()
        else:
            self.table.selectAll()

    def update_action_buttons(self):
        """يظهر شريط التحديد (تعديل / حذف / إلغاء) فقط عند اختيار صفوف."""
        selected = len({i.row() for i in self.table.selectedIndexes()})
        total = self.table.rowCount()

        self.act_select_all.setText(
            "  إلغاء التحديد"
            if total > 0 and selected >= total
            else "  تحديد الكل"
        )

        can_edit = self.has_permission("faults.edit")
        can_delete = self.has_permission("faults.delete")

        if selected > 0 and (can_edit or can_delete):
            self.selection_bar.setVisible(True)
            self.selection_label.setText(f"✔  {selected} محدد")
            self.edit_button.setVisible(selected == 1 and can_edit)
            self.delete_button.setVisible(can_delete)
            self.delete_button.setText(
                f"🗑  حذف ({selected})" if selected > 1 else "🗑  حذف"
            )
        else:
            self.selection_bar.setVisible(False)

    def show_table_menu(self, pos):
        item = self.table.itemAt(pos)
        if not item:
            return
        row = item.row()
        self.table.selectRow(row)
        id_item = self.table.item(row, 0)
        if not id_item:
            return
        fault_id = int(id_item.text())
        menu = QMenu(self)

        if self.has_permission("faults.edit"):
            status_menu = menu.addMenu("تغيير الحالة بسرعة")
            for status in ("جديد", "قيد الإصلاح", STATUS_REPAIRED):
                action = status_menu.addAction(status)
                action.triggered.connect(
                    lambda checked=False, value=status, selected_id=fault_id:
                        self.quick_change_status(value, selected_id)
                )
            menu.addAction("تعديل").triggered.connect(
                lambda checked=False, selected_id=fault_id:
                    self.edit_fault(selected_id)
            )

        receipt_action = menu.addAction("طباعة سند استلام")
        status_item = self.table.item(row, 7)
        selected_status = status_item.text().strip() if status_item else ""
        receipt_allowed = selected_status in (STATUS_REPAIRED, STATUS_DELIVERED)
        receipt_action.setEnabled(receipt_allowed)
        if not receipt_allowed:
            receipt_action.setToolTip(
                "يتاح طباعة السند بعد اكتمال الإصلاح؛ ويمكن إعادة طباعته بعد التسليم."
            )
        receipt_action.triggered.connect(
            lambda checked=False, selected_id=fault_id:
                self.print_delivery_receipt(selected_id)
        )

        menu.addAction("سجل أعطال هذا الجهاز").triggered.connect(
            lambda checked=False, selected_id=fault_id:
                self.show_device_history(selected_id)
        )

        if self.has_permission("faults.delete"):
            menu.addAction("حذف").triggered.connect(
                lambda checked=False, selected_id=fault_id:
                    self.delete_fault(selected_id)
            )

        if not menu.isEmpty():
            menu.exec(self.table.viewport().mapToGlobal(pos))

    def quick_change_status(self, status, fault_id=None):
        """تحديث الحالة مباشرة من قائمة السياق مع تسجيل تاريخ الإصلاح."""
        fault_id = fault_id or self._selected_fault_id()
        if fault_id is None:
            return
        connection = get_connection()
        try:
            cursor = connection.cursor()
            if status == STATUS_REPAIRED:
                cursor.execute("UPDATE faults SET status = ?, repair_date = COALESCE(repair_date, GETDATE()), updated_at = GETDATE() WHERE id = ?", (status, fault_id))
            else:
                cursor.execute("UPDATE faults SET status = ?, updated_at = GETDATE() WHERE id = ?", (status, fault_id))
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()
        self.show_notice(f"تم تغيير الحالة إلى: {status}")
        self.load_faults()

    def show_device_history(self, fault_id=None):
        """يفتح سجل أعطال الجهاز المرتبط بالعطل المحدد."""
        fault_id = self._selected_fault_id(fault_id)
        if fault_id is None:
            return

        connection = get_connection()
        try:
            cursor = connection.cursor()
            cursor.execute("""
                SELECT faults.device_id, devices.device_name
                FROM faults
                LEFT JOIN devices ON faults.device_id = devices.id
                WHERE faults.id = ?
            """, (fault_id,))
            row = cursor.fetchone()
        finally:
            connection.close()

        if not row or not row["device_id"]:
            QMessageBox.information(
                self, "تنبيه", "هذا العطل غير مرتبط بجهاز."
            )
            return

        DeviceFaultHistoryDialog(
            row["device_id"], row["device_name"] or "", self
        ).exec()

    def _query_filtered_faults(self):
        where_sql, params = self._build_filters()
        connection = get_connection()
        try:
            cursor = connection.cursor()
            cursor.execute("SELECT faults.*, devices.device_name, buildings.name AS building_name " + self.FROM_SQL + where_sql + " ORDER BY faults.id DESC", params)
            return cursor.fetchall()
        finally:
            connection.close()

    def print_filtered_report(self):
        """طباعة كل النتائج المطابقة للفلاتر، وليس الصفحة الحالية فقط."""
        rows = self._query_filtered_faults()
        if not rows:
            QMessageBox.information(self, "لا توجد نتائج", "لا توجد أعطال مطابقة للطباعة.")
            return
        headers = ["رقم العطل", "التاريخ", "الجهاز", "المبنى", "المستخدم", "وصف العطل", "الحالة", "الأولوية", "القائم بالإصلاح"]
        body = "".join("<tr>" + "".join(f"<td>{html.escape(str(v or ''))}</td>" for v in [r["fault_number"], format_datetime(r["entry_date"]), r["device_name"], r["building_name"], r["user_name"], r["fault_description"], r["status"], r["priority"], r["technician"]]) + "</tr>" for r in rows)
        filter_summary = f"نتائج الفلاتر الحالية — عدد السجلات: {len(rows)}"
        page = """<!doctype html><html lang='ar' dir='rtl'><head><meta charset='utf-8'><title>تقرير سجل الأعطال</title><style>
@page{size:A4 landscape;margin:14mm 12mm}*{box-sizing:border-box}body{font-family:Tahoma,Arial,sans-serif;color:#172536;margin:0;font-size:11px}.masthead{display:flex;justify-content:space-between;align-items:center;border-bottom:3px solid #17324d;padding:0 0 10px;margin-bottom:14px}.identity{font-size:12px;line-height:1.9;color:#263d52}.mark{border:2px solid #a88947;color:#17324d;padding:8px 14px;font-weight:bold;text-align:center;min-width:135px}.docmeta{text-align:left;line-height:1.8;color:#475569}.title{text-align:center;color:#17324d;font-size:20px;font-weight:bold;margin:12px 0 4px}.subtitle{text-align:center;color:#526477;margin:0 0 14px}.summary{border:1px solid #cbd5e1;background:#f4f6f8;padding:8px 10px;margin-bottom:10px}.filters{color:#475569;margin-top:4px}table{width:100%;border-collapse:collapse;font-size:10px}thead{display:table-header-group}tr{break-inside:avoid;page-break-inside:avoid}th,td{border:1px solid #9aa9b7;padding:6px 5px;text-align:right;vertical-align:top}th{background:#17324d;color:#fff;font-weight:bold}tbody tr:nth-child(even){background:#f1f4f6}.footer{border-top:1px solid #aeb9c3;margin-top:14px;padding-top:7px;display:flex;justify-content:space-between;color:#586879;font-size:9px}.approval{margin-top:24px;display:flex;justify-content:space-around;color:#34495e}.approval span{min-width:150px;border-top:1px solid #758596;padding-top:6px;text-align:center}@media print{button{display:none}}
</style></head><body><header class='masthead'><div class='identity'><b>الجهة / الإدارة:</b> ........................................<br><b>الوحدة:</b> ....................................................</div><div class='mark'>إدارة الأعطال والصيانة<br>تقرير متابعة</div><div class='docmeta'><b>التاريخ:</b> """ + datetime.now().strftime('%Y-%m-%d') + "<br><b>عدد السجلات:</b> " + str(len(rows)) + "</div></header><h1 class='title'>تقرير سجل الأعطال والصيانة</h1><p class='subtitle'>""" + html.escape(filter_summary) + "</p><div class='summary'><b>معايير العرض:</b> السجلات المعروضة وفق محددات البحث والفرز المختارة في النظام.</div><table><thead><tr>" + "".join(f"<th>{h}</th>" for h in headers) + "</tr></thead><tbody>" + body + "</tbody></table><div class='approval'><span>إعداد</span><span>مراجعة</span><span>اعتماد</span></div><footer class='footer'><span>نظام الأعطال والصيانة</span><span>تاريخ الاستخراج: """ + datetime.now().strftime('%Y-%m-%d %H:%M') + "</span></footer><script>window.onload=()=>window.print()</script></body></html>"
        self._open_print_page(page, "fault_report")

    def print_delivery_receipt(self, fault_id=None):
        fault_id = fault_id or self._selected_fault_id()
        if fault_id is None:
            return
        connection = get_connection()
        try:
            cursor = connection.cursor()
            cursor.execute("SELECT faults.*, devices.device_name, devices.device_number, devices.serial_number, buildings.name AS building_name FROM faults LEFT JOIN devices ON faults.device_id=devices.id LEFT JOIN buildings ON faults.building_id=buildings.id WHERE faults.id=?", (fault_id,))
            fault = cursor.fetchone()
        finally:
            connection.close()
        if not fault:
            return
        current_status = str(fault["status"] or "").strip()
        if current_status not in (STATUS_REPAIRED, STATUS_DELIVERED):
            QMessageBox.information(
                self,
                "السند غير متاح",
                "يمكن طباعة سند التسليم بعد اكتمال الإصلاح فقط. "
                "بعد التسليم يظل السند متاحًا لإعادة الطباعة."
            )
            return
        def val(key): return html.escape(str(fault[key] or ""))
        page = f"""<!doctype html><html lang='ar' dir='rtl'><head><meta charset='utf-8'><title>سند تسليم واستلام جهاز</title><style>
@page{{size:A4 portrait;margin:9mm 11mm}}*{{box-sizing:border-box}}body{{font-family:Tahoma,Arial,sans-serif;color:#172536;margin:0;font-size:10.5px;line-height:1.42}}.masthead{{display:flex;justify-content:space-between;align-items:center;border-bottom:2px solid #17324d;padding-bottom:6px;margin-bottom:8px}}.identity{{color:#263d52;line-height:1.5}}.mark{{border:1.5px solid #a88947;color:#17324d;padding:5px 12px;text-align:center;font-weight:bold}}h1{{text-align:center;color:#17324d;font-size:18px;margin:7px 0 2px}}.subtitle{{text-align:center;color:#637386;margin:0 0 8px;font-size:10px}}.meta{{display:flex;justify-content:space-between;background:#f2f5f7;border:1px solid #cbd5e1;padding:5px 8px;margin-bottom:7px;font-size:10px}}.box{{border:1px solid #9aa9b7;padding:7px 11px;margin:7px 0}}.box h2{{font-size:11px;color:#17324d;border-bottom:1px solid #cbd5e1;padding-bottom:3px;margin:0 0 4px}}.row{{display:flex;gap:10px;margin:3px 0}}.field{{flex:1;border-bottom:1px dotted #758596;min-height:21px}}.label{{font-weight:bold;color:#34495e}}.description{{min-height:22px;max-height:31px;border:1px solid #d5dde4;padding:4px 6px;margin:3px 0;overflow:hidden;font-size:10px;line-height:1.35}}.declaration{{border:1px solid #a88947;background:#faf9f5;padding:7px 9px;margin-top:8px;font-size:10px}}.sign{{display:flex;justify-content:space-between;gap:18px;margin-top:25px}}.signbox{{width:48%;border-top:1px solid #758596;padding-top:4px;text-align:center;min-height:56px;font-size:10px;line-height:1.55}}.stamp{{position:absolute;left:11mm;bottom:17mm;width:32mm;height:21mm;border:1px dashed #aab4be;color:#778493;text-align:center;padding-top:5px;font-size:9px}}.footer{{position:fixed;bottom:-5mm;border-top:1px solid #aeb9c3;padding-top:3px;width:100%;font-size:8px;color:#586879}}@media print{{button{{display:none}}html,body{{height:279mm;max-height:279mm;overflow:hidden}}.masthead,.box,.declaration,.sign{{break-inside:avoid;page-break-inside:avoid}}.description{{display:-webkit-box;-webkit-line-clamp:2;-webkit-box-orient:vertical}}}}
</style></head><body><header class='masthead'><div class='identity'><b>الجهة / الإدارة:</b> ................................<br><b>الوحدة:</b> ..........................................</div><div class='mark'>سند عهدة جهاز<br>تسليم واستلام</div></header><h1>سند تسليم واستلام جهاز</h1><p class='subtitle'>نموذج إثبات تسليم الجهاز للمستخدم بعد أعمال الفحص أو الصيانة</p><div class='meta'><span><b>رقم العطل:</b> {val('fault_number')}</span><span><b>تاريخ تحرير السند:</b> {html.escape(datetime.now().strftime('%Y-%m-%d'))}</span><span><b>تاريخ التسليم:</b> {val('delivery_date')}</span></div><section class='box'><h2>بيانات المستلم والجهاز</h2><div class='row'><div class='field'><span class='label'>اسم المستخدم:</span> {val('user_name')}</div><div class='field'><span class='label'>اسم المستلم:</span> {val('receiver_name')}</div></div><div class='row'><div class='field'><span class='label'>نوع / اسم الجهاز:</span> {val('device_name')}</div><div class='field'><span class='label'>رقم الجهاز:</span> {val('device_number')}</div></div><div class='row'><div class='field'><span class='label'>الرقم التسلسلي:</span> {val('serial_number')}</div><div class='field'><span class='label'>المبنى / الوحدة:</span> {val('building_name')}</div></div></section><section class='box'><h2>بيانات الصيانة</h2><div><b>وصف العطل:</b><div class='description'>{val('fault_description')}</div></div><div><b>الإجراء المتخذ:</b><div class='description'>{val('repair_actions')}</div></div><div class='row'><div class='field'><span class='label'>قطع الغيار:</span> {val('spare_parts')}</div><div class='field'><span class='label'>الحالة:</span> {val('status')}</div></div></section><div class='declaration'><b>إقرار الاستلام:</b> أقر أنا الموقع أدناه باستلام الجهاز الموضحة بياناته أعلاه، وأؤكد استلامه بحالته المبينة، وأتحمل مسؤولية المحافظة عليه وفق التعليمات المعمول بها.</div><div class='sign'><div class='signbox'><b>المستلم</b><br>الاسم: ........................................<br>التوقيع: .......................................<br>التاريخ: ........................................</div><div class='signbox'><b>المسلِّم / مسؤول الصيانة</b><br>الاسم: ........................................<br>التوقيع: .......................................<br>التاريخ: ........................................</div></div><div class='stamp'>ختم الجهة</div><footer class='footer'>يحرر هذا السند من نسختين عند الحاجة، وتحفظ نسخة ضمن سجل العطل.</footer><script>window.onload=()=>window.print()</script></body></html>"""
        self._open_print_page(page, "delivery_receipt")

    @staticmethod
    def _open_print_page(content, prefix):
        path = Path(tempfile.gettempdir()) / f"{prefix}_{datetime.now():%Y%m%d_%H%M%S}.html"
        path.write_text(content, encoding="utf-8")
        webbrowser.open(path.as_uri())

    def toggle_advanced_filters(self):
        self.filter_card.setVisible(self.advanced_button.isChecked())

    def toggle_status_filter(self, status):
        """الضغط على كارت يفلتر بحالته، والضغط مرة تانية يرجّع كل الحالات."""
        if self.status_filter.currentText() == status:
            self.status_filter.setCurrentIndex(0)
        else:
            self.status_filter.setCurrentText(status)

    def _filter_date(self, edit):
        if is_unset(edit):
            return None

        return edit.date().toPython()

    def _advanced_filter_count(self):
        count = 0

        if self.status_filter.currentIndex() > 0:
            count += 1
        if self.priority_filter.currentIndex() > 0:
            count += 1
        if self.building_filter.currentData() is not None:
            count += 1
        if self.technician_filter.currentData() is not None:
            count += 1
        if not is_unset(self.date_from_filter):
            count += 1
        if not is_unset(self.date_to_filter):
            count += 1

        return count

    def _update_filter_ui(self):
        count = self._advanced_filter_count()

        self.advanced_button.setText(
            "فلاتر متقدمة" + (f"  ({count})" if count else "")
        )

        self.clear_filters_button.setEnabled(
            bool(count or self.search_box.text().strip())
        )

        current_status = self.status_filter.currentText()

        self.fault_stat_cards["total"].set_active(
            self.status_filter.currentIndex() == 0
        )
        self.fault_stat_cards["repaired"].set_active(
            current_status == STATUS_REPAIRED
        )
        self.fault_stat_cards["delivered"].set_active(
            current_status == STATUS_DELIVERED
        )

    def refresh_all(self):
        self.load_building_filter()
        self.load_technician_filter()
        self.load_faults()

    # =============================================================
    # الصلاحيات
    # =============================================================

    def has_permission(self, permission_code):
        if not self.current_user:
            return False

        if self.current_user.get("role_code") == "SUPER_ADMIN":
            return True

        permissions = self.current_user.get("permissions", {})

        return bool(permissions.get(permission_code, False))

    def apply_permissions(self):
        can_view = self.has_permission("faults.view")
        can_add = self.has_permission("faults.add")
        can_reports = self.has_permission("reports.view")

        self.add_button.setVisible(can_add)
        self.import_button.setVisible(can_add)
        self.act_report.setVisible(can_reports)
        self.act_print_filtered.setVisible(can_reports)

        # زرا التعديل والحذف يتحكم فيهما update_action_buttons
        # (حسب التحديد والصلاحية)

        self.setEnabled(can_view)

    # =============================================================
    # تحميل فلتر المباني
    # =============================================================

    def load_building_filter(self):
        """
        تحميل الوحدات التنظيمية الرئيسية (مبنى / فرع) من الكاش المركزي.

        يتحدث تلقائياً عند أي تعديل في شاشة الهيكل التنظيمي.
        """
        current_building_id = (
            self.building_filter.currentData()
            if self.building_filter.count() > 0
            else None
        )

        # جلب الوحدات النشطة من الكاش (مبان + فروع)
        buildings = OrgStructureCache.instance().get_active_by_types("مبنى", "فرع")

        self.building_filter.blockSignals(True)

        try:
            self.building_filter.clear()
            self.building_filter.addItem("كل المباني", None)

            for building in buildings:
                self.building_filter.addItem(
                    f"[{building['unit_type']}] {building['name']}",
                    building["id"]
                )

            index = (
                self.building_filter.findData(current_building_id)
                if current_building_id is not None
                else -1
            )

            self.building_filter.setCurrentIndex(index if index >= 0 else 0)

        finally:
            self.building_filter.blockSignals(False)

    # =============================================================
    # تحميل فلتر القائم بالإصلاح
    # =============================================================

    def load_technician_filter(self):
        current_value = (
            self.technician_filter.currentData()
            if self.technician_filter.count() > 0
            else None
        )

        connection = get_connection()
        cursor = connection.cursor()

        try:
            cursor.execute("""
                SELECT DISTINCT LTRIM(RTRIM(technician)) AS technician
                FROM faults
                WHERE technician IS NOT NULL
                  AND LTRIM(RTRIM(technician)) <> ''
                ORDER BY LTRIM(RTRIM(technician))
            """)

            technicians = cursor.fetchall()

        finally:
            connection.close()

        self.technician_filter.blockSignals(True)

        try:
            self.technician_filter.clear()
            self.technician_filter.addItem("كل القائمين بالإصلاح", None)

            for row in technicians:
                name = row["technician"]
                self.technician_filter.addItem(name, name)

            index = (
                self.technician_filter.findData(current_value)
                if current_value
                else -1
            )

            self.technician_filter.setCurrentIndex(index if index >= 0 else 0)

        finally:
            self.technician_filter.blockSignals(False)

    # =============================================================
    # مسح الفلاتر
    # =============================================================

    def clear_fault_filters(self):
        widgets = (
            self.search_box,
            self.status_filter,
            self.priority_filter,
            self.building_filter,
            self.technician_filter,
            self.date_from_filter,
            self.date_to_filter,
        )

        for widget in widgets:
            widget.blockSignals(True)

        self.search_box.clear()
        self.status_filter.setCurrentIndex(0)
        self.priority_filter.setCurrentIndex(0)
        self.building_filter.setCurrentIndex(0)
        self.technician_filter.setCurrentIndex(0)
        clear_datetime(self.date_from_filter)
        clear_datetime(self.date_to_filter)

        for widget in widgets:
            widget.blockSignals(False)

        self._search_timer.stop()
        self._page = 0

        self.load_faults()

    def show_reports(self):
        if not self.has_permission("reports.view"):
            QMessageBox.warning(
                self,
                "غير مسموح",
                "ليس لديك صلاحية عرض التقارير."
            )
            return

        dialog = FaultReportsDialog(self)
        dialog.exec()

    # =============================================================
    # أدوات Excel
    # =============================================================

    def _normalize_excel_header(self, value):
        if value is None:
            return ""

        value = str(value).strip()

        replacements = {
            "إسم": "اسم",
            "الإسم": "الاسم",
            "الإستلام": "الاستلام",
            "الإصلاح": "الاصلاح",
            "القطع": "قطع",
        }

        for old, new in replacements.items():
            value = value.replace(old, new)

        return " ".join(value.split()).lower()

    def _excel_cell(self, row, headers, *names):
        wanted = {self._normalize_excel_header(name) for name in names}

        for header, index in headers.items():
            if self._normalize_excel_header(header) in wanted:
                value = row[index]

                if value is None:
                    return ""

                return str(value).strip()

        return ""

    def _find_building_id(self, connection, building_name):
        if not building_name:
            return None

        cursor = connection.cursor()

        cursor.execute(
            "SELECT TOP 1 id FROM buildings WHERE name = ?",
            (building_name.strip(),)
        )

        result = cursor.fetchone()

        if result:
            return result["id"]

        cursor.execute(
            "INSERT INTO buildings (name) VALUES (?)",
            (building_name.strip(),)
        )

        cursor.execute("SELECT CAST(SCOPE_IDENTITY() AS INT) AS id")

        result = cursor.fetchone()

        return result["id"] if result else None

    def _find_device_id(self, connection, device_name, building_id=None):
        if not device_name:
            return None

        cursor = connection.cursor()

        if building_id:
            cursor.execute(
                """
                SELECT TOP 1 id
                FROM devices
                WHERE device_name = ?
                  AND is_archived = 0
                  AND building_id = ?
                """,
                (device_name.strip(), building_id)
            )

            result = cursor.fetchone()

            if result:
                return result["id"]

        cursor.execute(
            """
            SELECT TOP 1 id
            FROM devices
            WHERE device_name = ?
              AND is_archived = 0
            """,
            (device_name.strip(),)
        )

        result = cursor.fetchone()

        return result["id"] if result else None

    def _historical_fault_number(self, connection):
        cursor = connection.cursor()

        year = datetime.now().year

        cursor.execute(
            "SELECT COUNT(*) FROM faults WHERE fault_number LIKE ?",
            (f"OLD-{year}-%",)
        )

        count = cursor.fetchone()[0]
        number = count + 1

        while True:
            fault_number = f"OLD-{year}-{number:05d}"

            cursor.execute(
                "SELECT TOP 1 1 FROM faults WHERE fault_number = ?",
                (fault_number,)
            )

            if not cursor.fetchone():
                return fault_number

            number += 1

    def _parse_import_date(self, value):
        if value is None or str(value).strip() == "":
            return datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        if hasattr(value, "strftime"):
            return value.strftime("%Y-%m-%d %H:%M:%S")

        raw = str(value).strip()

        for fmt in (
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%d",
            "%d/%m/%Y",
            "%d-%m-%Y",
            "%d/%m/%y",
            "%d-%m-%y",
        ):
            try:
                return datetime.strptime(raw, fmt).strftime(
                    "%Y-%m-%d %H:%M:%S"
                )
            except ValueError:
                continue

        return raw

    # =============================================================
    # استيراد Excel
    # =============================================================

    def import_from_excel(self):
        if not self.has_permission("faults.add"):
            QMessageBox.warning(
                self,
                "غير مسموح",
                "ليس لديك صلاحية استيراد الأعطال من Excel."
            )
            return

        try:
            from openpyxl import load_workbook
        except ImportError:
            QMessageBox.critical(
                self,
                "مكتبة ناقصة",
                "مكتبة openpyxl غير مثبتة.\n\n"
                "شغّل:\n"
                "pip install openpyxl"
            )
            return

        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "اختيار ملف أعطال Excel",
            "",
            "Excel Files (*.xlsx *.xlsm)"
        )

        if not file_path:
            return

        try:
            workbook = load_workbook(
                file_path,
                data_only=True,
                read_only=True
            )

            sheet = workbook.active

            rows = list(sheet.iter_rows(values_only=True))

            if not rows:
                workbook.close()

                QMessageBox.warning(
                    self,
                    "ملف فارغ",
                    "ملف Excel لا يحتوي على بيانات."
                )
                return

            raw_headers = list(rows[0])

            headers = {
                header: index
                for index, header in enumerate(raw_headers)
                if header is not None
            }

            required_headers = {
                "المبنى/الفرع",
                "إسم الجهاز",
                "مواصفات الجهاز",
                "إسم المستخدم",
                "إسم المستلم",
                "توصيف العطل",
                "إجراءات الإصلاح وقطع الغيار",
                "القائم بالإصلاح",
                "تاريخ الإستلام",
                "ملاحظات",
            }

            normalized_headers = {
                self._normalize_excel_header(h)
                for h in raw_headers
                if h is not None
            }

            missing = [
                h
                for h in required_headers
                if self._normalize_excel_header(h) not in normalized_headers
            ]

            if missing:
                workbook.close()

                QMessageBox.warning(
                    self,
                    "أعمدة غير مكتملة",
                    "ملف Excel لا يطابق أعمدة الأعطال المطلوبة.\n\n"
                    "الأعمدة الناقصة:\n- "
                    + "\n- ".join(missing)
                )
                return

            data_rows = [
                row
                for row in rows[1:]
                if any(
                    value is not None and str(value).strip()
                    for value in row
                )
            ]

            if not data_rows:
                workbook.close()

                QMessageBox.information(
                    self,
                    "لا توجد بيانات",
                    "لا توجد صفوف أعطال قابلة للاستيراد."
                )
                return

            workbook.close()

            answer = QMessageBox.question(
                self,
                "تأكيد الاستيراد",
                f"تم العثور على {len(data_rows)} عطل في الملف.\n\n"
                "سيتم إدخالها كأعطال تاريخية داخل النظام.\n"
                "الأعطال التي لا تحتوي على اسم جهاز ستدخل بدون ربط بجهاز، "
                "حتى لا يتم اختراع بيانات غير موجودة في الملف.\n\n"
                "هل تريد المتابعة؟",
                QMessageBox.Yes | QMessageBox.No,
                QMessageBox.No,
            )

            if answer != QMessageBox.Yes:
                return

            progress = QProgressDialog(
                "جاري استيراد الأعطال...",
                "إلغاء",
                0,
                len(data_rows),
                self
            )

            progress.setWindowTitle("استيراد الأعطال")
            progress.setWindowModality(Qt.WindowModal)
            progress.setMinimumDuration(0)
            progress.show()

            connection = get_connection()
            cursor = connection.cursor()

            imported = 0
            skipped = 0
            errors = []

            try:
                for row_number, row in enumerate(data_rows, start=2):
                    if progress.wasCanceled():
                        connection.rollback()

                        QMessageBox.information(
                            self,
                            "تم الإلغاء",
                            "تم إلغاء الاستيراد.\n"
                            "لم يتم حفظ أي صف."
                        )
                        return

                    progress.setValue(row_number - 2)

                    building_name = self._excel_cell(
                        row, headers, "المبنى/الفرع"
                    )
                    device_name = self._excel_cell(
                        row, headers, "إسم الجهاز"
                    )
                    user_name = self._excel_cell(
                        row, headers, "إسم المستخدم"
                    )
                    receiver_name = self._excel_cell(
                        row, headers, "إسم المستلم"
                    )
                    fault_description = self._excel_cell(
                        row, headers, "توصيف العطل"
                    )
                    repair_actions = self._excel_cell(
                        row, headers, "إجراءات الإصلاح وقطع الغيار"
                    )
                    technician = self._excel_cell(
                        row, headers, "القائم بالإصلاح"
                    )
                    received_date_raw = self._excel_cell(
                        row, headers, "تاريخ الإستلام"
                    )
                    notes = self._excel_cell(row, headers, "ملاحظات")

                    if (
                        not building_name
                        and not fault_description
                        and not user_name
                    ):
                        skipped += 1
                        continue

                    building_id = self._find_building_id(
                        connection, building_name
                    )

                    device_id = self._find_device_id(
                        connection, device_name, building_id
                    )

                    entry_date = self._parse_import_date(received_date_raw)

                    status = STATUS_REPAIRED if repair_actions else "جديد"

                    fault_number = self._historical_fault_number(connection)

                    cursor.execute(
                        """
                        INSERT INTO faults (
                            fault_number,
                            entry_date,
                            building_id,
                            device_id,
                            user_name,
                            receiver_name,
                            fault_description,
                            repair_actions,
                            spare_parts,
                            technician,
                            priority,
                            status,
                            repair_date,
                            delivery_date,
                            notes
                        )
                        VALUES (
                            ?, ?, ?, ?, ?, ?,
                            ?, ?, ?, ?, ?, ?,
                            ?, ?, ?
                        )
                        """,
                        (
                            fault_number,
                            entry_date,
                            building_id,
                            device_id,
                            user_name,
                            receiver_name,
                            fault_description,
                            repair_actions,
                            "",
                            technician,
                            "عادي",
                            status,
                            entry_date if status == STATUS_REPAIRED else None,
                            entry_date if status == STATUS_REPAIRED else None,
                            (
                                "استيراد تاريخي من Excel"
                                + (f" | {notes}" if notes else "")
                            ),
                        )
                    )

                    imported += 1

                connection.commit()

            except Exception:
                connection.rollback()
                raise

            finally:
                connection.close()

            progress.setValue(len(data_rows))

            write_audit_log(
                user=self.current_user,
                action="استيراد",
                module_name="الأعطال",
                description=(
                    f"تم استيراد الأعطال من ملف Excel: "
                    f"{Path(file_path).name}"
                    f"\nالأعطال المضافة: {imported}"
                    f"\nالصفوف المتجاهلة: {skipped}"
                    f"\nإجمالي الصفوف: {len(data_rows)}"
                ),
                new_values={
                    "file": Path(file_path).name,
                    "total_rows": len(data_rows),
                    "imported_faults": imported,
                    "skipped_rows": skipped,
                },
            )

            self.refresh_all()

            message = (
                f"تم استيراد {imported} عطل بنجاح.\n"
                f"تم تجاهل {skipped} صف فارغ."
            )

            if errors:
                message += (
                    "\n\nصفوف بها مشاكل:\n"
                    + "\n".join(errors[:10])
                )

            QMessageBox.information(self, "تم الاستيراد", message)

        except Exception as exc:
            QMessageBox.critical(
                self,
                "خطأ في الاستيراد",
                f"حدث خطأ أثناء استيراد ملف Excel:\n\n{exc}"
            )

    # =============================================================
    # تحميل الأعطال + الفلاتر
    # =============================================================

    def _reset_page_and_load_faults(self, *_args):
        self._page = 0
        self.load_faults()

    def _previous_page(self, _checked=False):
        if self._page > 0:
            self._page -= 1
            self.load_faults()

    def _next_page(self, _checked=False):
        total_pages = max(
            1,
            (self._total + self.PAGE_SIZE - 1) // self.PAGE_SIZE
        )

        if self._page < total_pages - 1:
            self._page += 1
            self.load_faults()

    def _build_filters(self, include_status=True):
        """يبني شرط WHERE مشترك بين الجدول والإحصائيات."""
        clauses = []
        params = []

        search = self.search_box.text().strip()

        if search:
            clauses.append("""
                (
                    faults.fault_number LIKE ?
                    OR devices.device_name LIKE ?
                    OR devices.printer LIKE ?
                    OR faults.user_name LIKE ?
                    OR faults.receiver_name LIKE ?
                    OR faults.technician LIKE ?
                    OR buildings.name LIKE ?
                    OR faults.fault_description LIKE ?
                )
            """)
            params.extend([f"%{search}%"] * 8)

        if include_status and self.status_filter.currentIndex() > 0:
            clauses.append("faults.status = ?")
            params.append(self.status_filter.currentText())

        if self.priority_filter.currentIndex() > 0:
            clauses.append("faults.priority = ?")
            params.append(self.priority_filter.currentText())

        building_id = self.building_filter.currentData()

        if building_id:
            clauses.append("faults.building_id = ?")
            params.append(building_id)

        technician = self.technician_filter.currentData()

        if technician:
            clauses.append("LTRIM(RTRIM(faults.technician)) = ?")
            params.append(technician)

        date_from = self._filter_date(self.date_from_filter)
        date_to = self._filter_date(self.date_to_filter)

        # لو المستخدم عكس التاريخين نبدّلهم بدل ما النتيجة تطلع فاضية
        if date_from and date_to and date_from > date_to:
            date_from, date_to = date_to, date_from

        if date_from:
            clauses.append("TRY_CONVERT(date, faults.entry_date) >= ?")
            params.append(date_from)

        if date_to:
            clauses.append("TRY_CONVERT(date, faults.entry_date) <= ?")
            params.append(date_to)

        where_sql = "".join(f" AND {clause}" for clause in clauses)

        return where_sql, params

    @staticmethod
    def _paint_badge(item, colors):
        foreground, background = colors

        item.setForeground(QColor(foreground))

        if background:
            item.setBackground(QColor(background))

        font = item.font()
        font.setBold(True)
        item.setFont(font)

    def load_faults(self):
        where_sql, params = self._build_filters()

        connection = get_connection()
        cursor = connection.cursor()

        try:
            cursor.execute(
                "SELECT COUNT(*) " + self.FROM_SQL + where_sql,
                params
            )

            self._total = cursor.fetchone()[0] or 0

            total_pages = max(
                1,
                (self._total + self.PAGE_SIZE - 1) // self.PAGE_SIZE
            )

            self._page = min(self._page, total_pages - 1)

            cursor.execute(
                """
                SELECT
                    faults.*,
                    devices.device_name,
                    buildings.name AS building_name
                """
                + self.FROM_SQL
                + where_sql
                + """
                ORDER BY faults.id DESC
                OFFSET ? ROWS FETCH NEXT ? ROWS ONLY
                """,
                params + [self._page * self.PAGE_SIZE, self.PAGE_SIZE]
            )

            visible_faults = cursor.fetchall()

        finally:
            connection.close()

        self.table.setUpdatesEnabled(False)
        self.table.setRowCount(0)

        for fault in visible_faults:
            row = self.table.rowCount()
            self.table.insertRow(row)

            description = (
                (fault["fault_description"] or "").replace("\n", " ").strip()
            )

            short_description = (
                description
                if len(description) <= 70
                else description[:70] + "…"
            )

            values = [
                fault["id"],
                fault["fault_number"],
                format_datetime(fault["entry_date"]),
                fault["device_name"] or "",
                fault["building_name"] or "",
                fault["user_name"] or "",
                short_description,
                fault["status"] or "",
                fault["priority"] or "",
                fault["technician"] or "",
            ]

            for column, value in enumerate(values):
                item = QTableWidgetItem(str(value or ""))

                if column == 6:
                    item.setTextAlignment(Qt.AlignCenter)
                    item.setToolTip(description)
                else:
                    item.setTextAlignment(Qt.AlignCenter)

                if column == 7 and value in STATUS_COLORS:
                    self._paint_badge(item, STATUS_COLORS[value])

                if column == 8 and value in PRIORITY_COLORS:
                    self._paint_badge(item, PRIORITY_COLORS[value])

                self.table.setItem(row, column, item)

        self.table.setUpdatesEnabled(True)

        has_rows = bool(visible_faults)

        self.table.setVisible(has_rows)
        self.act_select_all.setEnabled(has_rows)
        self.empty_label.setVisible(not has_rows)

        if not has_rows:
            if self._total == 0 and (
                self.search_box.text().strip()
                or self._advanced_filter_count()
            ):
                self.empty_label.setText(
                    "لا توجد أعطال مطابقة للبحث أو الفلاتر الحالية.\n"
                    "جرّب تغيير البحث أو اضغط (مسح الفلاتر)."
                )
            else:
                self.empty_label.setText(
                    "لا توجد أعطال مسجلة بعد.\n"
                    "اضغط (إضافة عطل) لتسجيل أول عطل."
                )

        self.page_label.setText(f"صفحة {self._page + 1} / {total_pages}")
        self.previous_button.setEnabled(self._page > 0)
        self.next_button.setEnabled(self._page < total_pages - 1)

        if self._total:
            first = self._page * self.PAGE_SIZE + 1
            last = first + len(visible_faults) - 1

            self.stats_label.setText(
                f"عرض {first}–{last} من {self._total} عطل"
            )
        else:
            self.stats_label.setText("")

        self.update_action_buttons()
        self._update_filter_ui()
        self.load_fault_statistics()

    def load_fault_statistics(self):
        """إحصائيات الكروت: تتأثر بالبحث والفلاتر ما عدا فلتر الحالة (عشان الكروت تفضل ثابتة)."""
        where_sql, params = self._build_filters(include_status=False)

        sql = (
            """
            SELECT
                COUNT(*) AS total,
                COUNT(DISTINCT CASE
                    WHEN faults.status IN (?, ?) THEN faults.device_id
                END) AS serviced,
                SUM(CASE WHEN faults.status = ? THEN 1 ELSE 0 END) AS repaired,
                SUM(CASE WHEN faults.status = ? THEN 1 ELSE 0 END) AS delivered,
                COUNT(DISTINCT NULLIF(
                    LTRIM(RTRIM(faults.technician)), ''
                )) AS technicians,
                SUM(CASE WHEN faults.status NOT IN (?, ?) AND TRY_CONVERT(date, faults.entry_date) < DATEADD(day, -?, CAST(GETDATE() AS date)) THEN 1 ELSE 0 END) AS overdue
            """
            + self.FROM_SQL
            + where_sql
        )

        all_params = [
            STATUS_REPAIRED,
            STATUS_DELIVERED,
            STATUS_REPAIRED,
            STATUS_DELIVERED,
            STATUS_REPAIRED,
            STATUS_DELIVERED,
            self.overdue_days.value(),
        ] + params

        connection = get_connection()
        cursor = connection.cursor()

        try:
            cursor.execute(sql, all_params)
            stats = cursor.fetchone()
        finally:
            connection.close()

        self.fault_stat_cards["total"].set_value(stats["total"] or 0)
        self.fault_stat_cards["serviced"].set_value(stats["serviced"] or 0)
        self.fault_stat_cards["repaired"].set_value(stats["repaired"] or 0)
        self.fault_stat_cards["delivered"].set_value(stats["delivered"] or 0)
        self.fault_stat_cards["technicians"].set_value(
            stats["technicians"] or 0
        )
        self.fault_stat_cards["overdue"].set_value(stats["overdue"] or 0)

    def _selected_fault_id(self, fault_id=None):
        # QPushButton.clicked يمرر bool؛ لا نعامله كرقم عطل.
        if isinstance(fault_id, bool):
            fault_id = None
        elif fault_id is not None and hasattr(fault_id, "row"):
            # بعض إشارات Qt تمرر QModelIndex بدل معرّف السجل.
            row = fault_id.row()
            item = self.table.item(row, 0)
            return int(item.text()) if item is not None else None
        elif fault_id is not None:
            return int(fault_id)
        selected = self.table.selectedItems()

        if not selected:
            QMessageBox.information(self, "تنبيه", "اختر عطلًا أولًا.")
            return None

        row = selected[0].row()

        return int(self.table.item(row, 0).text())

    @staticmethod
    def _safe_close(connection):
        try:
            connection.rollback()
        except Exception:
            pass

        try:
            connection.close()
        except Exception:
            pass

    # =============================================================
    # إضافة عطل
    # =============================================================

    def add_fault(self):
        if not self.has_permission("faults.add"):
            QMessageBox.warning(
                self,
                "غير مسموح",
                "ليس لديك صلاحية إضافة الأعطال."
            )
            return

        dialog = FaultDialog(self)
        dialog.fault_number.setText(dialog.generate_fault_number())

        if dialog.exec() != QDialog.Accepted:
            return

        data = dialog.get_data()

        connection = get_connection()
        try:
            cursor = connection.cursor()
            cursor.execute(
                "SELECT COUNT(*) FROM faults WHERE device_id = ?",
                (data["device_id"],)
            )
            previous_count = cursor.fetchone()[0] or 0
        finally:
            connection.close()

        if previous_count:
            answer = QMessageBox.question(
                self,
                "جهاز متكرر العطل",
                f"تنبيه: لهذا الجهاز {previous_count} عطل سابق. هل تريد تسجيل العطل الجديد؟",
                QMessageBox.Yes | QMessageBox.No,
                QMessageBox.Yes,
            )
            if answer != QMessageBox.Yes:
                return

        connection = get_connection()
        cursor = connection.cursor()

        if data["spare_part_id"]:
            cursor.execute(
                """
                SELECT
                    id,
                    part_name,
                    quantity
                FROM spare_parts
                WHERE id = ?
                  AND is_archived = 0
                """,
                (data["spare_part_id"],)
            )

            part = cursor.fetchone()

            if not part:
                connection.close()

                QMessageBox.warning(
                    self,
                    "تنبيه",
                    "قطعة الغيار المختارة غير موجودة."
                )
                return

            if part["quantity"] < data["spare_part_qty"]:
                connection.close()

                QMessageBox.warning(
                    self,
                    "الكمية غير كافية",
                    f"القطعة: {part['part_name']}\n"
                    f"المتاح: {part['quantity']}\n"
                    f"المطلوب: {data['spare_part_qty']}"
                )
                return

        cursor.execute(
            "SELECT building_id FROM devices WHERE id = ?",
            (data["device_id"],)
        )

        result = cursor.fetchone()

        building_id = result["building_id"] if result else None

        cursor.execute(
            """
            INSERT INTO faults (
                fault_number,
                entry_date,
                building_id,
                device_id,
                user_name,
                receiver_name,
                fault_description,
                repair_actions,
                spare_parts,
                technician,
                priority,
                status,
                repair_date,
                delivery_date
            )
            VALUES (
                ?, ?, ?, ?, ?, ?,
                ?, ?, ?, ?, ?, ?,
                ?, ?
            )
            """,
            (
                data["fault_number"],
                data["entry_date"],
                building_id,
                data["device_id"],
                data["user_name"],
                data["receiver_name"],
                data["fault_description"],
                data["repair_actions"],
                data["spare_parts"],
                data["technician"],
                data["priority"],
                data["status"],
                data["repair_date"],
                data["delivery_date"],
            )
        )

        cursor.execute(
            """
            SELECT TOP 1 id
            FROM faults
            WHERE fault_number = ?
            ORDER BY id DESC
            """,
            (data["fault_number"],)
        )

        inserted_fault = cursor.fetchone()

        new_fault_id = inserted_fault["id"] if inserted_fault else None

        if data["spare_part_id"]:
            cursor.execute(
                """
                UPDATE spare_parts
                SET
                    quantity = quantity - ?,
                    updated_at = CURRENT_TIMESTAMP
                WHERE id = ?
                  AND quantity >= ?
                """,
                (
                    data["spare_part_qty"],
                    data["spare_part_id"],
                    data["spare_part_qty"]
                )
            )

            if cursor.rowcount != 1:
                connection.rollback()
                connection.close()

                QMessageBox.warning(
                    self,
                    "تنبيه",
                    "تعذر خصم قطعة الغيار من المخزون."
                )
                return

            cursor.execute(
                """
                INSERT INTO spare_part_movements
                (
                    spare_part_id,
                    movement_type,
                    quantity,
                    notes
                )
                VALUES (
                    ?,
                    'صرف',
                    ?,
                    ?
                )
                """,
                (
                    data["spare_part_id"],
                    data["spare_part_qty"],
                    f"صرف تلقائي لعطل {data['fault_number']}"
                )
            )

        connection.commit()
        connection.close()

        audit_new_values = dict(data)
        audit_new_values["building_id"] = building_id
        audit_new_values["id"] = new_fault_id

        write_audit_log(
            user=self.current_user,
            action="إضافة",
            module_name="الأعطال",
            record_id=new_fault_id,
            description=f"تم تسجيل العطل رقم {data['fault_number']}",
            new_values=audit_new_values,
        )

        self.refresh_all()

        self.show_notice(f"تم تسجيل العطل {data['fault_number']}")

    # =============================================================
    # تعديل عطل
    # =============================================================

    def edit_fault(self, fault_id=None):
        if not self.has_permission("faults.edit"):
            QMessageBox.warning(
                self,
                "غير مسموح",
                "ليس لديك صلاحية تعديل الأعطال."
            )
            return

        fault_id = self._selected_fault_id(fault_id)

        if fault_id is None:
            return

        connection = get_connection()
        cursor = connection.cursor()

        cursor.execute(
            "SELECT * FROM faults WHERE id = ?",
            (fault_id,)
        )

        old_fault = cursor.fetchone()

        old_values = fault_audit_data(old_fault)

        connection.close()

        if not old_fault:
            QMessageBox.warning(self, "تنبيه", "العطل غير موجود.")
            return

        dialog = FaultDialog(self, fault=old_fault)

        if dialog.exec() != QDialog.Accepted:
            return

        data = dialog.get_data()

        connection = get_connection()
        cursor = connection.cursor()

        try:
            cursor.execute(
                """
                SELECT TOP 1 id,
                    spare_part_id,
                    quantity
                FROM spare_part_movements
                WHERE movement_type = 'صرف'
                  AND notes = ?
                ORDER BY id DESC
                """,
                (f"صرف تلقائي لعطل {old_fault['fault_number']}",)
            )

            old_movement = cursor.fetchone()

            old_part_id = (
                old_movement["spare_part_id"] if old_movement else None
            )

            old_qty = int(old_movement["quantity"]) if old_movement else 0

            new_spare_id = data.get("spare_part_id")

            new_qty = int(data.get("spare_part_qty") or 0)

            old_has_part = old_part_id is not None and old_qty > 0

            if old_has_part:
                cursor.execute(
                    """
                    UPDATE spare_parts
                    SET
                        quantity = quantity + ?,
                        updated_at = CURRENT_TIMESTAMP
                    WHERE id = ?
                    """,
                    (old_qty, old_part_id)
                )

                if cursor.rowcount != 1:
                    connection.rollback()
                    connection.close()

                    QMessageBox.warning(
                        self,
                        "تنبيه",
                        "تعذر إرجاع قطعة الغيار القديمة للمخزن."
                    )
                    return

                cursor.execute(
                    "DELETE FROM spare_part_movements WHERE id = ?",
                    (old_movement["id"],)
                )

            if new_spare_id and new_qty > 0:
                cursor.execute(
                    """
                    SELECT
                        id,
                        part_name,
                        quantity
                    FROM spare_parts
                    WHERE id = ?
                      AND is_archived = 0
                    """,
                    (new_spare_id,)
                )

                new_part = cursor.fetchone()

                if not new_part:
                    connection.rollback()
                    connection.close()

                    QMessageBox.warning(
                        self,
                        "تنبيه",
                        "قطعة الغيار المختارة غير موجودة."
                    )
                    return

                if int(new_part["quantity"]) < new_qty:
                    connection.rollback()
                    connection.close()

                    QMessageBox.warning(
                        self,
                        "رصيد غير كافٍ",
                        f"القطعة «{new_part['part_name']}» "
                        f"المتاح منها {new_part['quantity']} فقط."
                    )
                    return

                cursor.execute(
                    """
                    UPDATE spare_parts
                    SET
                        quantity = quantity - ?,
                        updated_at = CURRENT_TIMESTAMP
                    WHERE id = ?
                      AND quantity >= ?
                    """,
                    (new_qty, new_spare_id, new_qty)
                )

                if cursor.rowcount != 1:
                    connection.rollback()
                    connection.close()

                    QMessageBox.warning(
                        self,
                        "تنبيه",
                        "تعذر خصم قطعة الغيار الجديدة."
                    )
                    return

                cursor.execute(
                    """
                    INSERT INTO spare_part_movements
                    (
                        spare_part_id,
                        movement_type,
                        quantity,
                        notes
                    )
                    VALUES (
                        ?,
                        'صرف',
                        ?,
                        ?
                    )
                    """,
                    (
                        new_spare_id,
                        new_qty,
                        f"صرف تلقائي لعطل {old_fault['fault_number']}"
                    )
                )

            cursor.execute(
                """
                UPDATE faults SET
                    entry_date = ?,
                    building_id = ?,
                    device_id = ?,
                    user_name = ?,
                    receiver_name = ?,
                    fault_description = ?,
                    repair_actions = ?,
                    spare_parts = ?,
                    technician = ?,
                    priority = ?,
                    status = ?,
                    repair_date = ?,
                    delivery_date = ?,
                    notes = ?,
                    updated_at = CURRENT_TIMESTAMP
                WHERE id = ?
                """,
                (
                    data["entry_date"],
                    data["building_id"],
                    data["device_id"],
                    data["user_name"],
                    data["receiver_name"],
                    data["fault_description"],
                    data["repair_actions"],
                    data["spare_parts"],
                    data["technician"],
                    data["priority"],
                    data["status"],
                    data["repair_date"],
                    data["delivery_date"],
                    data.get("notes", ""),
                    fault_id,
                )
            )

            connection.commit()
            connection.close()

            audit_new_values = dict(data)

            audit_new_values.update({
                "id": fault_id,
                "fault_number": old_fault["fault_number"],
                "building_id": data["building_id"],
                "linked_spare_part_movement": {
                    "spare_part_id": new_spare_id,
                    "quantity": new_qty if new_spare_id else 0,
                },
            })

            old_values["linked_spare_part_movement"] = {
                "spare_part_id": old_part_id,
                "quantity": old_qty,
            }

            write_audit_log(
                user=self.current_user,
                action="تعديل",
                module_name="الأعطال",
                record_id=fault_id,
                description=(
                    f"تم تعديل بيانات العطل رقم "
                    f"{old_fault['fault_number']}"
                ),
                old_values=old_values,
                new_values=audit_new_values,
            )

            self.refresh_all()

            self.show_notice(
                f"تم تعديل العطل {old_fault['fault_number']} "
                "وتحديث المخزون"
            )

        except Exception as exc:
            self._safe_close(connection)

            QMessageBox.critical(
                self,
                "خطأ",
                f"حدث خطأ أثناء تعديل العطل:\n{exc}"
            )

    # =============================================================
    # حذف عطل
    # =============================================================

    def delete_fault(self, fault_id=None):
        if not self.has_permission("faults.delete"):
            QMessageBox.warning(
                self,
                "غير مسموح",
                "ليس لديك صلاحية حذف الأعطال."
            )
            return

        # clicked(bool) لا يحمل معرّفًا، لذلك نستخدم التحديد الحالي.
        if isinstance(fault_id, bool):
            fault_id = None
        elif fault_id is not None and hasattr(fault_id, "row"):
            # دعم تمرير QModelIndex من أي إشارة/قائمة Qt.
            row = fault_id.row()
            id_item = self.table.item(row, 0)
            number_item = self.table.item(row, 1)
            if id_item is not None:
                fault_id = int(id_item.text())
                context_fault_number = number_item.text().strip() if number_item else None
            else:
                fault_id = None
                context_fault_number = None
        else:
            context_fault_number = None

        # قائمة كليك يمين تمرر معرّفًا محددًا؛ زر الشريط يستخدم كل الصفوف المختارة.
        if fault_id is not None:
            fault_refs = [(int(fault_id), context_fault_number)]
        else:
            selected_rows = sorted({item.row() for item in self.table.selectedItems()})
            fault_refs = []
            seen_fault_numbers = set()
            for row in selected_rows:
                id_item = self.table.item(row, 0)
                number_item = self.table.item(row, 1)
                if id_item is None:
                    continue
                fault_number = number_item.text().strip() if number_item else ""
                if fault_number and fault_number in seen_fault_numbers:
                    continue
                if fault_number:
                    seen_fault_numbers.add(fault_number)
                try:
                    row_id = int(id_item.text())
                except (TypeError, ValueError):
                    row_id = None
                if row_id is not None:
                    fault_refs.append((row_id, fault_number or None))

        if not fault_refs:
            QMessageBox.information(self, "تنبيه", "حدد عطلًا واحدًا على الأقل.")
            return

        connection = get_connection()
        cursor = connection.cursor()
        faults = []
        missing_ids = []
        try:
            for selected_id, fault_number in fault_refs:
                cursor.execute("SELECT * FROM faults WHERE id = ?", (selected_id,))
                fault = cursor.fetchone()
                if fault_number and (
                    not fault or str(fault["fault_number"]).strip() != fault_number
                ):
                    # لا نعتمد ID وحده: نتحقق من رقم العطل الظاهر حتى لا نحذف سجلًا آخر.
                    cursor.execute(
                        "SELECT * FROM faults WHERE fault_number = ?",
                        (fault_number,)
                    )
                    fault = cursor.fetchone()
                if fault:
                    if not any(int(existing["id"]) == int(fault["id"]) for existing in faults):
                        faults.append(fault)
                else:
                    missing_ids.append(fault_number or selected_id)
        except Exception as exc:
            self._safe_close(connection)
            QMessageBox.critical(self, "خطأ", f"تعذر قراءة الأعطال المحددة:\n{exc}")
            return

        if not faults:
            diagnostic = []
            try:
                cursor.execute(
                    "SELECT DB_NAME() AS database_name, @@SERVERNAME AS server_name"
                )
                database_info = cursor.fetchone()
                cursor.execute("SELECT COUNT(*) AS total FROM faults")
                total_in_database = cursor.fetchone()["total"]
                cursor.execute(
                    "SELECT TOP 5 id, fault_number FROM faults ORDER BY id DESC"
                )
                database_samples = cursor.fetchall()
                diagnostic.append(
                    f"الاتصال الحالي: {database_info['server_name']} / {database_info['database_name']}"
                )
                diagnostic.append(
                    f"عدد الأعطال في هذا الاتصال: {total_in_database}"
                )
                diagnostic.append(
                    "أمثلة من القاعدة: "
                    + "، ".join(
                        f"{row['id']} / {row['fault_number']}"
                        for row in database_samples
                    )
                )
            except Exception as exc:
                diagnostic.append(f"تعذر قراءة معلومات الاتصال: {exc}")
            diagnostic.append(
                "أول صفوف محددة من الجدول: "
                + "، ".join(f"{row_id} / {number}" for row_id, number in fault_refs[:5])
            )
            connection.close()
            QMessageBox.warning(
                self,
                "لا توجد سجلات للحذف",
                "لم أجد تطابقًا بين الصفوف المحددة والقاعدة المتصلة. لم يُحذف أي سجل.\n\n"
                + "\n".join(diagnostic)
            )
            self.load_faults()
            return

        listed_numbers = "، ".join(str(f["fault_number"]) for f in faults[:5])
        if len(faults) > 5:
            listed_numbers += f"، وآخرون ({len(faults) - 5})"
        prompt_title = "تأكيد حذف العطل" if len(faults) == 1 else "تأكيد الحذف الجماعي"
        prompt = (
            f"هل أنت متأكد من حذف {len(faults)} عطل؟\n\n{listed_numbers}\n\n"
            "إذا صُرفت قطع غيار لهذه الأعطال، ستُعاد الكميات للمخزون تلقائيًا."
        )
        if missing_ids:
            prompt += (
                f"\n\nملاحظة: سيتم تخطي {len(missing_ids)} صف لم يعد موجودًا "
                "في قاعدة البيانات، وحذف السجلات الموجودة فقط."
            )
        answer = QMessageBox.question(
            self, prompt_title, prompt,
            QMessageBox.Yes | QMessageBox.No, QMessageBox.No
        )
        if answer != QMessageBox.Yes:
            connection.close()
            return

        audit_records = []
        restored_count = 0
        try:
            for fault in faults:
                current_id = int(fault["id"])
                old_values = fault_audit_data(fault)
                cursor.execute(
                    """SELECT TOP 1 id, spare_part_id, quantity
                       FROM spare_part_movements
                       WHERE movement_type = 'صرف' AND notes = ?
                       ORDER BY id DESC""",
                    (f"صرف تلقائي لعطل {fault['fault_number']}",)
                )
                movement = cursor.fetchone()
                old_values["linked_spare_part_movement"] = (
                    {"spare_part_id": movement["spare_part_id"],
                     "quantity": movement["quantity"]}
                    if movement else None
                )
                if movement:
                    cursor.execute(
                        """UPDATE spare_parts
                           SET quantity = quantity + ?, updated_at = CURRENT_TIMESTAMP
                           WHERE id = ?""",
                        (movement["quantity"], movement["spare_part_id"])
                    )
                    if cursor.rowcount != 1:
                        raise RuntimeError(
                            f"تعذر إرجاع قطعة الغيار للعطل {fault['fault_number']}؛ أُلغي الحذف بالكامل."
                        )
                    cursor.execute(
                        "DELETE FROM spare_part_movements WHERE id = ?",
                        (movement["id"],)
                    )
                    restored_count += 1

                cursor.execute("DELETE FROM fault_updates WHERE fault_id = ?", (current_id,))
                cursor.execute("DELETE FROM faults WHERE id = ?", (current_id,))
                if cursor.rowcount != 1:
                    raise RuntimeError(
                        f"تعذر حذف العطل {fault['fault_number']}؛ أُلغي الحذف بالكامل."
                    )
                audit_records.append((fault, old_values))

            connection.commit()
            connection.close()

            for fault, old_values in audit_records:
                write_audit_log(
                    user=self.current_user,
                    action="حذف",
                    module_name="الأعطال",
                    record_id=fault["id"],
                    description=f"تم حذف العطل رقم {fault['fault_number']}",
                    old_values=old_values,
                    new_values={
                        "deleted": True,
                        "stock_returned": bool(old_values.get("linked_spare_part_movement")),
                    },
                )

            self.refresh_all()
            message = f"تم حذف {len(audit_records)} عطل بنجاح"
            if restored_count:
                message += f"، وإرجاع قطع الغيار لـ {restored_count} عطل"
            if missing_ids:
                message += f"، وتخطي {len(missing_ids)} سجل غير موجود"
            self.show_notice(message)

        except Exception as exc:
            self._safe_close(connection)
            QMessageBox.critical(
                self,
                "تعذر إتمام الحذف",
                f"لم تُحذف الأعطال المحددة، وتم التراجع عن العملية.\n\n{exc}"
            )