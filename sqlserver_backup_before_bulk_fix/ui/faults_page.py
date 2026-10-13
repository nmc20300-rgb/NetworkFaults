from datetime import datetime
import html
import tempfile
import webbrowser
from pathlib import Path

from PySide6.QtCore import Qt
from PySide6.QtWidgets import QCompleter
from PySide6.QtWidgets import QFileDialog, QProgressDialog
from PySide6.QtWidgets import (
    QApplication, QWidget, QVBoxLayout, QHBoxLayout, QGridLayout,
    QLabel, QPushButton, QLineEdit,
    QTableWidget, QTableWidgetItem, QMessageBox, QDialog,
    QComboBox, QTextEdit, QDialogButtonBox, QHeaderView, QFrame, QStyle,
    QDateTimeEdit, QSpinBox, QCalendarWidget, QScrollArea
)

from database import get_connection


APP_STYLE = """
QWidget {
    font-family: "Segoe UI", "Tahoma";
    font-size: 12px;
    color: #334155;
}

#pageRoot {
    background: #f4f7fb;
}

#pageTitle {
    color: #17324d;
    font-size: 26px;
    font-weight: 800;
}

#pageSubtitle {
    color: #718096;
    font-size: 12px;
}

#toolBar {
    background: white;
    border: 1px solid #e3eaf2;
    border-radius: 12px;
}

#filterCard {
    background: white;
    border: 1px solid #e3eaf2;
    border-radius: 12px;
}

#filterLabel {
    color: #64748b;
    font-weight: 700;
}

#searchBox {
    background: #f8fafc;
    border: 1px solid #d9e2ec;
    border-radius: 9px;
    padding: 9px 12px;
    min-height: 22px;
}

#searchBox:focus {
    border: 1px solid #5b8bb0;
    background: white;
}

#primaryButton {
    background: #17324d;
    color: white;
    border: none;
    border-radius: 9px;
    padding: 9px 15px;
    font-weight: 700;
}

#primaryButton:hover {
    background: #244d70;
}

#secondaryButton {
    background: #eef3f7;
    color: #29455e;
    border: 1px solid #d9e2ec;
    border-radius: 9px;
    padding: 9px 15px;
    font-weight: 700;
}

#secondaryButton:hover {
    background: #e3ebf2;
}

#clearFilterButton {
    background: #fff7ed;
    color: #9a3412;
    border: 1px solid #fed7aa;
    border-radius: 9px;
    padding: 8px 14px;
    font-weight: 700;
}

#clearFilterButton:hover {
    background: #ffedd5;
}

QComboBox {
    background: #fbfdff;
    border: 1px solid #d5dfe9;
    border-radius: 8px;
    padding: 8px 10px;
    min-width: 145px;
}

QComboBox:focus {
    border: 1px solid #5b8bb0;
    background: white;
}

#tableCard {
    background: white;
    border: 1px solid #e3eaf2;
    border-radius: 14px;
}

QTableWidget {
    background: white;
    border: none;
    gridline-color: #edf1f5;
    color: #34495e;
    selection-background-color: #e8f1f8;
    selection-color: #17324d;
    alternate-background-color: #fafcfe;
}

QHeaderView::section {
    background: #f0f4f8;
    color: #4a5d70;
    border: none;
    border-bottom: 1px solid #dfe7ef;
    padding: 10px;
    font-weight: 800;
}

#dialogRoot {
    background: #f5f8fb;
}

#dialogHeader {
    background: #17324d;
    color: white;
    border-radius: 12px;
}

#dialogHeaderIcon {
    color: white;
    font-size: 28px;
    font-weight: 900;
}

#dialogTitle {
    color: white;
    font-size: 19px;
    font-weight: 800;
}

#dialogHint {
    color: #dbe7f1;
    font-size: 11px;
}

#formLabel {
    color: #4b6074;
    font-weight: 700;
    padding: 2px 3px;
}

QLineEdit, QComboBox, QDateTimeEdit, QTextEdit {
    background: #fbfdff;
    border: 1px solid #d5dfe9;
    border-radius: 8px;
    padding: 8px 10px;
}

QLineEdit:focus, QComboBox:focus, QDateTimeEdit:focus, QTextEdit:focus {
    border: 1px solid #5b8bb0;
    background: white;
}

QTextEdit {
    min-height: 55px;
}

QDialogButtonBox QPushButton {
    min-width: 95px;
    padding: 8px 14px;
    border-radius: 8px;
    border: 1px solid #d5dfe9;
    background: white;
}

QDialogButtonBox QPushButton:hover {
    background: #eef3f7;
}
"""


class FaultDialog(QDialog):
    def __init__(self, parent=None, fault=None):
        super().__init__(parent)

        self.fault = fault

        self.setWindowTitle(
            "تعديل عطل" if fault else "إضافة عطل"
        )
        screen = self.screen() or QApplication.primaryScreen()
        available = screen.availableGeometry() if screen else None
        if available:
            width = max(320, min(840, int(available.width() * 0.86)))
            height = max(380, min(640, int(available.height() * 0.82)))
        else:
            width, height = 800, 620
        self.resize(width, height)
        self.setMinimumSize(min(620, width), min(480, height))
        if available:
            self.move(available.center() - self.rect().center())

        self.setObjectName("dialogRoot")
        self.setStyleSheet(APP_STYLE)

        root = QVBoxLayout(self)

        root.setContentsMargins(14, 14, 14, 12)

        root.setSpacing(12)

        header = QFrame()
        header.setObjectName("dialogHeader")

        hl = QHBoxLayout(header)

        hl.setContentsMargins(
            16, 13, 16, 13
        )

        icon = QLabel("⚠")
        icon.setObjectName("dialogHeaderIcon")
        icon.setAlignment(Qt.AlignCenter)
        icon.setFixedSize(42, 42)

        hl.addWidget(icon)

        text = QVBoxLayout()

        title = QLabel(
            "تسجيل عطل جديد"
            if not fault
            else "تعديل بيانات العطل"
        )

        title.setObjectName("dialogTitle")

        hint = QLabel(
            "سجل بيانات العطل ومتابعة الإصلاح من مكان واحد"
        )

        hint.setObjectName("dialogHint")

        text.addWidget(title)
        text.addWidget(hint)

        hl.addLayout(text)
        hl.addStretch()

        root.addWidget(header)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.NoFrame)
        scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        form_widget = QWidget()
        form = QGridLayout(form_widget)
        form.setContentsMargins(6, 8, 6, 8)
        form.setHorizontalSpacing(12)
        form.setVerticalSpacing(8)
        form.setColumnMinimumWidth(0, 86)
        form.setColumnMinimumWidth(2, 86)
        form.setColumnStretch(1, 1)
        form.setColumnStretch(3, 1)

        self.fault_number = QLineEdit()
        self.fault_number.setReadOnly(True)

        self.entry_date = QDateTimeEdit()
        self.entry_date.setCalendarPopup(True)
        self.entry_date.setDateTime(datetime.now())
        self.entry_date.setDisplayFormat("yyyy-MM-dd HH:mm")

        self.device = QComboBox()
        self.device.setEditable(True)
        self.device.setInsertPolicy(QComboBox.NoInsert)

        self.device.setPlaceholderText(
            "اكتب اسم الجهاز للبحث أو اختر من القائمة..."
        )

        self.device.setCompleter(QCompleter())

        self.device.completer().setCaseSensitivity(
            Qt.CaseInsensitive
        )

        self.device.completer().setFilterMode(
            Qt.MatchContains
        )

        self.device.completer().setCompletionMode(
            QCompleter.PopupCompletion
        )

        self.device.lineEdit().editingFinished.connect(
            self.select_typed_device
        )

        self.device.currentIndexChanged.connect(
            self.device_changed
        )

        self.device_specs = QTextEdit()
        self.device_specs.setReadOnly(True)

        self.device_specs.setFixedHeight(72)

        self.device_specs.setVerticalScrollBarPolicy(
            Qt.ScrollBarAsNeeded
        )

        self.device_specs.setHorizontalScrollBarPolicy(
            Qt.ScrollBarAlwaysOff
        )

        self.building = QLineEdit()
        self.building.setReadOnly(True)

        self.user_name = QLineEdit()
        self.receiver_name = QLineEdit()

        self.fault_description = QTextEdit()
        self.repair_actions = QTextEdit()
        self.fault_description.setFixedHeight(74)
        self.repair_actions.setFixedHeight(68)

        self.spare_part = QComboBox()

        self.spare_part.addItem(
            "بدون قطعة غيار",
            None
        )

        self.spare_part_qty = QSpinBox()
        self.spare_part_qty.setRange(1, 999999)
        self.spare_part_qty.setValue(1)
        self.spare_part_qty.setFixedWidth(110)

        spare_part_box = QHBoxLayout()

        spare_part_box.setContentsMargins(
            0, 0, 0, 0
        )

        spare_part_box.setSpacing(8)

        spare_part_box.addWidget(
            self.spare_part,
            1
        )

        spare_part_box.addWidget(
            QLabel("الكمية:")
        )

        spare_part_box.addWidget(
            self.spare_part_qty
        )

        self.spare_part_container = QWidget()

        self.spare_part_container.setLayout(
            spare_part_box
        )

        self.technician = QLineEdit()

        self.priority = QComboBox()

        self.priority.addItems([
            "عادي",
            "متوسط",
            "عاجل",
            "حرج"
        ])

        self.status = QComboBox()

        self.status.addItems([
            "جديد",
            "قيد الفحص",
            "قيد الإصلاح",
            "بانتظار قطعة غيار",
            "تم الإصلاح",
            "تم التسليم"
        ])

        self.repair_date = QDateTimeEdit()
        self.repair_date.setCalendarPopup(True)
        self.repair_date.setDisplayFormat("yyyy-MM-dd HH:mm")

        self.repair_date.setSpecialValueText(
            "غير محدد"
        )

        self.delivery_date = QDateTimeEdit()
        self.delivery_date.setCalendarPopup(True)
        self.delivery_date.setDisplayFormat("yyyy-MM-dd HH:mm")

        self.delivery_date.setSpecialValueText(
            "غير محدد"
        )

        def add_pair(row, right_label, right_widget, left_label, left_widget):
            for column, label_text, widget in (
                (0, right_label, right_widget),
                (2, left_label, left_widget),
            ):
                label = QLabel(label_text + ":")
                label.setObjectName("formLabel")
                label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
                form.addWidget(label, row, column)
                form.addWidget(widget, row, column + 1)

        def add_full(row, label_text, widget):
            label = QLabel(label_text + ":")
            label.setObjectName("formLabel")
            label.setAlignment(Qt.AlignRight | Qt.AlignTop)
            form.addWidget(label, row, 0)
            form.addWidget(widget, row, 1, 1, 3)

        add_pair(0, "رقم العطل", self.fault_number,
                 "تاريخ الدخول", self.entry_date)
        add_pair(1, "الجهاز *", self.device,
                 "المبنى / الفرع", self.building)
        add_full(2, "مواصفات الجهاز", self.device_specs)
        add_pair(3, "اسم المستخدم", self.user_name,
                 "اسم المستلم", self.receiver_name)
        add_full(4, "توصيف العطل *", self.fault_description)
        add_full(5, "إجراءات الإصلاح", self.repair_actions)
        add_full(6, "قطع الغيار المستخدمة", self.spare_part_container)
        add_pair(7, "القائم بالإصلاح", self.technician,
                 "الأولوية", self.priority)
        add_pair(8, "حالة العطل", self.status,
                 "تاريخ الإصلاح", self.repair_date)
        add_full(9, "تاريخ التسليم", self.delivery_date)

        scroll.setWidget(form_widget)
        root.addWidget(scroll, 1)

        buttons = QDialogButtonBox(
            QDialogButtonBox.Save |
            QDialogButtonBox.Cancel
        )

        buttons.accepted.connect(
            self.validate_and_accept
        )

        buttons.rejected.connect(
            self.reject
        )

        root.addWidget(buttons)

        self.load_devices()
        self.load_spare_parts()

        if fault:
            self.load_fault()
        else:
            self.spare_part_qty.setEnabled(True)

    def select_typed_device(self):
        typed = self.device.currentText().strip()

        if not typed:
            return

        for i in range(
            self.device.count()
        ):
            if self.device.itemText(i).strip() == typed:
                self.device.setCurrentIndex(i)
                return

    def load_devices(self):
        connection = get_connection()
        cursor = connection.cursor()

        cursor.execute("""
            SELECT
                devices.id,
                devices.device_name,
                devices.device_number,
                devices.user_name,
                devices.processor,
                devices.ram,
                devices.hard_disk,
                devices.serial_number,
                devices.monitor,
                devices.operating_system,
                devices.printer,
                buildings.name AS building_name
            FROM devices
            LEFT JOIN buildings
                ON devices.building_id = buildings.id
            WHERE devices.is_archived = 0
            ORDER BY devices.device_name
        """)

        devices = cursor.fetchall()

        connection.close()

        self.device.clear()

        self.device.addItem(
            "اختر الجهاز...",
            None
        )

        for device in devices:
            self.device.addItem(
                device["device_name"],
                device["id"]
            )

        self.device.completer().setModel(
            self.device.model()
        )

    def load_spare_parts(self):
        connection = get_connection()
        cursor = connection.cursor()

        cursor.execute("""
            SELECT
                id,
                part_name,
                part_number,
                quantity
            FROM spare_parts
            WHERE is_archived = 0
            ORDER BY part_name
        """)

        parts = cursor.fetchall()

        connection.close()

        self.spare_part.clear()

        self.spare_part.addItem(
            "بدون قطعة غيار",
            None
        )

        for part in parts:
            label = (
                f"{part['part_name']} | "
                f"المتاح: {part['quantity']}"
            )

            if part["part_number"]:
                label += f" | {part['part_number']}"

            self.spare_part.addItem(
                label,
                part["id"]
            )

    def device_changed(self):
        device_id = self.device.currentData()

        if not device_id:
            self.building.clear()
            self.user_name.clear()
            self.device_specs.clear()
            return

        connection = get_connection()
        cursor = connection.cursor()

        cursor.execute("""
            SELECT
                devices.*,
                buildings.name AS building_name
            FROM devices
            LEFT JOIN buildings
                ON devices.building_id = buildings.id
            WHERE devices.id = ?
        """, (device_id,))

        device = cursor.fetchone()

        connection.close()

        if not device:
            return

        self.building.setText(
            device["building_name"] or ""
        )

        self.user_name.setText(
            device["user_name"] or ""
        )

        specs = (
            f"النوع: {device['device_type'] or '-'}\n"
            f"المعالج: {device['processor'] or '-'}\n"
            f"RAM: {device['ram'] or '-'}\n"
            f"الهارد: {device['hard_disk'] or '-'}\n"
            f"Serial: {device['serial_number'] or '-'}\n"
            f"الشاشة: {device['monitor'] or '-'}\n"
            f"النظام: {device['operating_system'] or '-'}\n"
            f"الطابعة: {device['printer'] or '-'}"
        )

        self.device_specs.setPlainText(
            specs
        )

    def generate_fault_number(self):
        year = datetime.now().year

        connection = get_connection()
        cursor = connection.cursor()

        cursor.execute(
            "SELECT COUNT(*) FROM faults "
            "WHERE fault_number LIKE ?",
            (f"FLT-{year}-%",)
        )

        count = cursor.fetchone()[0]

        connection.close()

        return f"FLT-{year}-{count + 1:04d}"

    def validate_and_accept(self):
        if not self.device.currentData():
            QMessageBox.warning(
                self,
                "تنبيه",
                "من فضلك اختر الجهاز."
            )

            self.device.setFocus()
            return

        if not self.fault_description.toPlainText().strip():
            QMessageBox.warning(
                self,
                "تنبيه",
                "من فضلك اكتب توصيف العطل."
            )

            self.fault_description.setFocus()
            return

        self.accept()

    def get_data(self):
        building_id = None
        device_id = self.device.currentData()

        if device_id:
            connection = get_connection()
            cursor = connection.cursor()

            cursor.execute(
                "SELECT building_id FROM devices WHERE id = ?",
                (device_id,)
            )

            result = cursor.fetchone()

            connection.close()

            if result:
                building_id = result["building_id"]

        return {
            "fault_number": self.fault_number.text(),

            "entry_date": self.entry_date.dateTime().toString(
                "yyyy-MM-dd HH:mm:ss"
            ),

            "building_id": building_id,
            "device_id": device_id,

            "user_name": self.user_name.text().strip(),

            "receiver_name": self.receiver_name.text().strip(),

            "fault_description":
                self.fault_description.toPlainText().strip(),

            "repair_actions":
                self.repair_actions.toPlainText().strip(),

            "spare_part_id":
                self.spare_part.currentData(),

            "spare_part_qty":
                self.spare_part_qty.value()
                if self.spare_part.currentData()
                else 0,

            "spare_parts": (
                f"{self.spare_part.currentText().split(' | المتاح:')[0]} "
                f"× {self.spare_part_qty.value()}"
                if self.spare_part.currentData()
                else ""
            ),

            "technician":
                self.technician.text().strip(),

            "priority":
                self.priority.currentText(),

            "status":
                self.status.currentText(),

            "repair_date":
                self.repair_date.dateTime().toString(
                    "yyyy-MM-dd HH:mm:ss"
                ),

            "delivery_date":
                self.delivery_date.dateTime().toString(
                    "yyyy-MM-dd HH:mm:ss"
                ),

            "notes":
                self.fault["notes"]
                if self.fault and "notes" in self.fault.keys()
                else "",
        }

    def load_fault(self):
        self.fault_number.setText(
            self.fault["fault_number"]
        )

        self.entry_date.setDateTime(
            datetime.strptime(
                self.fault["entry_date"],
                "%Y-%m-%d %H:%M:%S"
            )
        )

        index = self.device.findData(
            self.fault["device_id"]
        )

        if index >= 0:
            self.device.setCurrentIndex(index)

        self.user_name.setText(
            self.fault["user_name"] or ""
        )

        self.receiver_name.setText(
            self.fault["receiver_name"] or ""
        )

        self.fault_description.setPlainText(
            self.fault["fault_description"] or ""
        )

        self.repair_actions.setPlainText(
            self.fault["repair_actions"] or ""
        )

        saved_parts = self.fault["spare_parts"] or ""

        self.spare_part.setEnabled(True)
        self.spare_part_qty.setEnabled(True)

        if saved_parts:
            import re

            match = re.search(
                r"(.+?)\s*[×xX]\s*(\d+)\s*$",
                saved_parts
            )

            if match:
                part_name = match.group(1).strip()
                qty = int(match.group(2))

                self.spare_part_qty.setValue(
                    max(1, qty)
                )

                for i in range(
                    self.spare_part.count()
                ):
                    label = self.spare_part.itemText(i)

                    if (
                        label.split(" | المتاح:")[0].strip()
                        == part_name
                    ):
                        self.spare_part.setCurrentIndex(i)
                        break
            else:
                self.spare_part.setCurrentIndex(0)

        self.technician.setText(
            self.fault["technician"] or ""
        )

        priority_index = self.priority.findText(
            self.fault["priority"]
        )

        if priority_index >= 0:
            self.priority.setCurrentIndex(
                priority_index
            )

        status_index = self.status.findText(
            self.fault["status"]
        )

        if status_index >= 0:
            self.status.setCurrentIndex(
                status_index
            )

        try:
            if self.fault["repair_date"]:
                self.repair_date.setDateTime(
                    datetime.strptime(
                        self.fault["repair_date"],
                        "%Y-%m-%d %H:%M:%S"
                    )
                )

            if self.fault["delivery_date"]:
                self.delivery_date.setDateTime(
                    datetime.strptime(
                        self.fault["delivery_date"],
                        "%Y-%m-%d %H:%M:%S"
                    )
                )

        except (ValueError, TypeError):
            pass


class FaultReportsDialog(QDialog):
    """
    تقرير موقف الأعطال اليومي أو الأسبوعي
    مع معاينة وطباعة من خلال المتصفح.
    """

    def __init__(self, parent=None):
        super().__init__(parent)

        self.setWindowTitle(
            "تقارير موقف الأعطال"
        )

        self.resize(1100, 700)

        self.setLayoutDirection(
            Qt.RightToLeft
        )

        root = QVBoxLayout(self)

        root.setContentsMargins(
            18, 18, 18, 18
        )

        root.setSpacing(12)

        title = QLabel(
            "تقارير موقف الأعطال"
        )

        title.setStyleSheet(
            "font-size:22px;font-weight:800;color:#17324d;"
        )

        root.addWidget(title)

        subtitle = QLabel(
            "اختر نوع التقرير والتاريخ لعرض موقف الأعطال ثم معاينته وطباعته."
        )

        subtitle.setStyleSheet(
            "color:#718096;"
        )

        root.addWidget(subtitle)

        controls = QHBoxLayout()

        controls.setSpacing(10)

        controls.addWidget(
            QLabel("نوع التقرير:")
        )

        self.report_type = QComboBox()

        self.report_type.addItems([
            "تقرير يومي",
            "تقرير أسبوعي"
        ])

        self.report_type.currentIndexChanged.connect(
            self.load_report
        )

        controls.addWidget(
            self.report_type
        )

        controls.addWidget(
            QLabel("التاريخ:")
        )

        self.calendar = QCalendarWidget()

        self.calendar.setGridVisible(True)

        self.calendar.setSelectedDate(
            __import__('PySide6').QtCore.QDate.currentDate()
        )

        self.calendar.clicked.connect(
            self.load_report
        )

        controls.addWidget(
            self.calendar
        )

        root.addLayout(controls)

        self.period_label = QLabel()

        self.period_label.setStyleSheet(
            "font-weight:700;color:#334155;padding:5px;"
        )

        root.addWidget(
            self.period_label
        )

        self.summary = QLabel()

        self.summary.setStyleSheet(
            "background:#f4f7fb;"
            "border:1px solid #e3eaf2;"
            "border-radius:8px;"
            "padding:10px;"
        )

        root.addWidget(
            self.summary
        )

        self.table = QTableWidget()

        self.table.setColumnCount(8)

        self.table.setHorizontalHeaderLabels([
            "رقم العطل",
            "التاريخ",
            "الجهاز",
            "المستخدم",
            "الحالة",
            "الأولوية",
            "القائم بالإصلاح",
            "وصف العطل"
        ])

        self.table.setSelectionBehavior(
            QTableWidget.SelectRows
        )

        self.table.setEditTriggers(
            QTableWidget.NoEditTriggers
        )

        self.table.verticalHeader().setVisible(
            False
        )

        self.table.horizontalHeader().setSectionResizeMode(
            QHeaderView.Stretch
        )

        self.table.setAlternatingRowColors(True)

        root.addWidget(
            self.table,
            1
        )

        buttons = QHBoxLayout()

        self.refresh_button = QPushButton(
            " تحديث"
        )

        self.print_button = QPushButton(
            " معاينة وطباعة"
        )

        self.close_button = QPushButton(
            "إغلاق"
        )

        self.refresh_button.clicked.connect(
            self.load_report
        )

        self.print_button.clicked.connect(
            self.print_report
        )

        self.close_button.clicked.connect(
            self.accept
        )

        buttons.addStretch()

        buttons.addWidget(
            self.refresh_button
        )

        buttons.addWidget(
            self.print_button
        )

        buttons.addWidget(
            self.close_button
        )

        root.addLayout(buttons)

        self.load_report()

    def date_range(self):
        date = self.calendar.selectedDate()

        start = date
        end = date

        if self.report_type.currentIndex() == 1:
            end = date.addDays(6)

        return (
            start.toString("yyyy-MM-dd"),
            end.toString("yyyy-MM-dd")
        )

    def load_report(self):
        start, end = self.date_range()

        self.period_label.setText(
            f"الفترة: {start}"
            if start == end
            else f"الفترة: من {start} إلى {end}"
        )

        con = get_connection()
        cur = con.cursor()

        cur.execute("""
            SELECT
                faults.fault_number,
                faults.entry_date,
                devices.device_name,
                faults.user_name,
                faults.status,
                faults.priority,
                faults.technician,
                faults.fault_description
            FROM faults
            LEFT JOIN devices
                ON faults.device_id = devices.id
            WHERE date(faults.entry_date) BETWEEN ? AND ?
            ORDER BY
                datetime(faults.entry_date) DESC,
                faults.id DESC
        """, (start, end))

        rows = cur.fetchall()

        con.close()

        self.table.setRowCount(0)

        for fault in rows:
            row = self.table.rowCount()

            self.table.insertRow(row)

            values = [
                fault["fault_number"],
                fault["entry_date"],
                fault["device_name"] or "",
                fault["user_name"] or "",
                fault["status"] or "",
                fault["priority"] or "",
                fault["technician"] or "",
                fault["fault_description"] or ""
            ]

            for col, value in enumerate(values):
                item = QTableWidgetItem(
                    str(value)
                )

                item.setTextAlignment(
                    Qt.AlignCenter
                )

                self.table.setItem(
                    row,
                    col,
                    item
                )

        counts = {}

        for fault in rows:
            status = fault["status"] or "غير محدد"

            counts[status] = (
                counts.get(status, 0) + 1
            )

        details = "  |  ".join(
            f"{k}: {v}"
            for k, v in counts.items()
        )

        self.summary.setText(
            f"إجمالي الأعطال: {len(rows)}"
            +
            (
                f"  |  {details}"
                if details
                else "  |  لا توجد أعطال خلال الفترة المحددة."
            )
        )

    def print_report(self):
        """
        تجهيز التقرير كملف HTML وفتحه في المتصفح
        لمعاينته ثم طباعته على ورق A4.
        """

        try:
            report_html = self.build_print_html()

            reports_dir = (
                Path(tempfile.gettempdir())
                / "NetworkFaultsReports"
            )

            reports_dir.mkdir(
                parents=True,
                exist_ok=True
            )

            file_name = (
                "fault_report_"
                f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.html"
            )

            file_path = reports_dir / file_name

            file_path.write_text(
                report_html,
                encoding="utf-8"
            )

            # فتح التقرير في المتصفح الافتراضي
            webbrowser.open(
                file_path.resolve().as_uri()
            )

        except Exception as exc:
            QMessageBox.critical(
                self,
                "خطأ في التقرير",
                "حدث خطأ أثناء تجهيز التقرير للطباعة:\n\n"
                f"{exc}"
            )

    def build_print_html(self):
        """
        بناء تقرير HTML كامل ومنسق
        للطباعة على A4 أفقي.
        """

        start, end = self.date_range()

        title = html.escape(
            "تقرير موقف الأعطال"
        )

        period = html.escape(
            self.period_label.text()
        )

        summary = html.escape(
            self.summary.text()
        )

        rows_html = ""

        for row in range(
            self.table.rowCount()
        ):

            cells = []

            for column in range(
                self.table.columnCount()
            ):

                item = self.table.item(
                    row,
                    column
                )

                value = (
                    item.text()
                    if item
                    else ""
                )

                cells.append(
                    html.escape(
                        str(value)
                    )
                )

            rows_html += f"""
                <tr>
                    <td>{cells[0]}</td>
                    <td>{cells[1]}</td>
                    <td>{cells[2]}</td>
                    <td>{cells[3]}</td>
                    <td>{cells[4]}</td>
                    <td>{cells[5]}</td>
                    <td>{cells[6]}</td>
                    <td class="description">
                        {cells[7]}
                    </td>
                </tr>
            """

        if not rows_html:
            rows_html = """
                <tr>
                    <td
                        colspan="8"
                        class="empty-row"
                    >
                        لا توجد أعطال خلال الفترة المحددة.
                    </td>
                </tr>
            """

        generated_at = datetime.now().strftime(
            "%Y-%m-%d %H:%M"
        )

        generated_at = html.escape(
            generated_at
        )

        return f"""<!DOCTYPE html>
<html lang="ar" dir="rtl">

<head>

<meta charset="UTF-8">

<meta
    name="viewport"
    content="width=device-width, initial-scale=1.0"
>

<title>{title}</title>

<style>

@page {{
    size: A4 landscape;
    margin: 10mm;
}}

* {{
    box-sizing: border-box;
}}

html,
body {{
    margin: 0;
    padding: 0;
}}

body {{
    direction: rtl;
    background: white;
    color: #1e293b;

    font-family:
        "Segoe UI",
        Tahoma,
        Arial,
        sans-serif;

    font-size: 11px;
    line-height: 1.5;
}}

.report {{
    width: 100%;
    max-width: 100%;
    margin: 0 auto;
}}

.report-header {{
    border-bottom: 3px solid #17324d;
    padding-bottom: 10px;
    margin-bottom: 12px;
}}

.report-title {{
    font-size: 24px;
    font-weight: 800;
    color: #17324d;
    margin-bottom: 4px;
}}

.report-subtitle {{
    color: #64748b;
    font-size: 11px;
}}

.info-grid {{
    display: grid;
    grid-template-columns:
        1fr 1fr 1fr;

    gap: 8px;
    margin-bottom: 12px;
}}

.info-card {{
    border: 1px solid #dbe3ec;
    background: #f8fafc;
    border-radius: 6px;
    padding: 8px 10px;
}}

.info-label {{
    display: block;
    color: #64748b;
    font-size: 9px;
    margin-bottom: 2px;
}}

.info-value {{
    color: #17324d;
    font-size: 11px;
    font-weight: 700;
}}

table {{
    width: 100%;
    border-collapse: collapse;
    table-layout: fixed;
    margin-top: 8px;
}}

thead {{
    display: table-header-group;
}}

tfoot {{
    display: table-footer-group;
}}

tr {{
    page-break-inside: avoid;
    break-inside: avoid;
}}

th,
td {{
    border: 1px solid #cbd5e1;
    padding: 6px 5px;
    text-align: center;
    vertical-align: middle;

    overflow-wrap: anywhere;
    word-break: break-word;
}}

th {{
    background: #17324d;
    color: white;
    font-weight: 800;
    font-size: 10px;
}}

td {{
    font-size: 9.5px;
    background: white;
}}

tbody tr:nth-child(even) td {{
    background: #f8fafc;
}}

.description {{
    text-align: right;
    line-height: 1.45;
}}

.empty-row {{
    text-align: center;
    padding: 25px;
    color: #64748b;
    font-size: 12px;
}}

.footer {{
    margin-top: 12px;
    padding-top: 8px;

    border-top:
        1px solid #dbe3ec;

    display: flex;
    justify-content: space-between;

    color: #64748b;
    font-size: 9px;
}}

@media screen {{

    body {{
        background: #e5e7eb;
        padding: 20px;
    }}

    .report {{
        background: white;
        padding: 18px;

        box-shadow:
            0 3px 15px
            rgba(0, 0, 0, 0.12);

        max-width: 1200px;
    }}
}}

@media print {{

    body {{
        background: white;
        padding: 0;
    }}

    .report {{
        box-shadow: none;
        padding: 0;
        max-width: none;
    }}

    .no-print {{
        display: none !important;
    }}

}}

</style>

</head>

<body>

<div class="report">

    <div class="report-header">

        <div class="report-title">
            {title}
        </div>

        <div class="report-subtitle">
            الشبكة - نظام إدارة الأجهزة والأعطال
        </div>

    </div>


    <div class="info-grid">

        <div class="info-card">

            <span class="info-label">
                فترة التقرير
            </span>

            <span class="info-value">
                {period}
            </span>

        </div>


        <div class="info-card">

            <span class="info-label">
                ملخص الأعطال
            </span>

            <span class="info-value">
                {summary}
            </span>

        </div>


        <div class="info-card">

            <span class="info-label">
                تاريخ إنشاء التقرير
            </span>

            <span class="info-value">
                {generated_at}
            </span>

        </div>

    </div>


    <table>

        <colgroup>

            <col style="width: 10%;">
            <col style="width: 12%;">
            <col style="width: 13%;">
            <col style="width: 12%;">
            <col style="width: 13%;">
            <col style="width: 10%;">
            <col style="width: 14%;">
            <col style="width: 16%;">

        </colgroup>

        <thead>

            <tr>

                <th>
                    رقم العطل
                </th>

                <th>
                    التاريخ
                </th>

                <th>
                    الجهاز
                </th>

                <th>
                    المستخدم
                </th>

                <th>
                    الحالة
                </th>

                <th>
                    الأولوية
                </th>

                <th>
                    القائم بالإصلاح
                </th>

                <th>
                    وصف العطل
                </th>

            </tr>

        </thead>


        <tbody>

            {rows_html}

        </tbody>

    </table>


    <div class="footer">

        <div>
            الشبكة - نظام إدارة الأجهزة والأعطال
        </div>

        <div>
            تم إنشاء التقرير بتاريخ {generated_at}
        </div>

    </div>

</div>

</body>

</html>
"""



class FaultsPage(QWidget):
    def __init__(self):
        super().__init__()

        self.setup_ui()
        self.load_building_filter()
        self.load_faults()

    def setup_ui(self):
        self.setObjectName(
            "pageRoot"
        )

        self.setStyleSheet(
            APP_STYLE
        )

        main_layout = QVBoxLayout(self)

        main_layout.setContentsMargins(
            26, 22, 26, 22
        )

        main_layout.setSpacing(14)

        # =========================================================
        # العنوان والأزرار
        # =========================================================

        title_layout = QHBoxLayout()

        title_box = QVBoxLayout()
        title_box.setSpacing(2)

        title = QLabel(
            "سجل الأعطال"
        )

        title.setObjectName(
            "pageTitle"
        )

        title_box.addWidget(
            title
        )

        subtitle = QLabel(
            "متابعة الأعطال وحالات الإصلاح والتسليم بسهولة"
        )

        subtitle.setObjectName(
            "pageSubtitle"
        )

        title_box.addWidget(
            subtitle
        )

        title_layout.addLayout(
            title_box
        )

        title_layout.addStretch()

        self.add_button = QPushButton(
            "⚠  إضافة عطل"
        )

        self.add_button.setObjectName(
            "primaryButton"
        )

        self.add_button.clicked.connect(
            self.add_fault
        )

        self.edit_button = QPushButton(
            "  تعديل"
        )

        self.edit_button.setObjectName(
            "secondaryButton"
        )

        self.edit_button.setIcon(
            self.style().standardIcon(
                QStyle.SP_FileDialogDetailedView
            )
        )

        self.edit_button.clicked.connect(
            self.edit_fault
        )

        self.report_button = QPushButton(
            "📋  تقارير الموقف"
        )

        self.report_button.setObjectName(
            "secondaryButton"
        )

        self.report_button.clicked.connect(
            self.show_reports
        )

        self.import_button = QPushButton(
            "📥  استيراد Excel"
        )

        self.import_button.setObjectName(
            "secondaryButton"
        )

        self.import_button.clicked.connect(
            self.import_from_excel
        )

        self.delete_button = QPushButton(
            "  حذف"
        )

        self.delete_button.setObjectName(
            "secondaryButton"
        )

        self.delete_button.setIcon(
            self.style().standardIcon(
                QStyle.SP_TrashIcon
            )
        )

        self.delete_button.clicked.connect(
            self.delete_fault
        )

        title_layout.addWidget(
            self.add_button
        )

        title_layout.addWidget(
            self.edit_button
        )

        title_layout.addWidget(
            self.delete_button
        )

        title_layout.addWidget(
            self.report_button
        )

        title_layout.addWidget(
            self.import_button
        )

        main_layout.addLayout(
            title_layout
        )

        # =========================================================
        # البحث
        # =========================================================

        search_card = QFrame()

        search_card.setObjectName(
            "toolBar"
        )

        search_layout = QHBoxLayout(
            search_card
        )

        search_layout.setContentsMargins(
            14, 10, 14, 10
        )

        search_icon = QLabel()

        search_icon.setPixmap(
            self.style().standardIcon(
                QStyle.SP_FileDialogContentsView
            ).pixmap(20, 20)
        )

        search_layout.addWidget(
            search_icon
        )

        self.search_box = QLineEdit()

        self.search_box.setObjectName(
            "searchBox"
        )

        self.search_box.setPlaceholderText(
            "بحث برقم العطل أو الجهاز أو المستخدم أو المستلم أو المبنى..."
        )

        self.search_box.textChanged.connect(
            self.load_faults
        )

        search_layout.addWidget(
            self.search_box,
            1
        )

        main_layout.addWidget(
            search_card
        )

        # =========================================================
        # فلاتر الأعطال
        # =========================================================

        filter_card = QFrame()

        filter_card.setObjectName(
            "filterCard"
        )

        filter_layout = QHBoxLayout(
            filter_card
        )

        filter_layout.setContentsMargins(
            14, 10, 14, 10
        )

        filter_layout.setSpacing(10)

        status_label = QLabel(
            "الحالة:"
        )

        status_label.setObjectName(
            "filterLabel"
        )

        filter_layout.addWidget(
            status_label
        )

        self.status_filter = QComboBox()

        self.status_filter.addItems([
            "كل الحالات",
            "جديد",
            "قيد الفحص",
            "قيد الإصلاح",
            "بانتظار قطعة غيار",
            "تم الإصلاح",
            "تم التسليم"
        ])

        self.status_filter.currentIndexChanged.connect(
            self.load_faults
        )

        filter_layout.addWidget(
            self.status_filter
        )

        priority_label = QLabel(
            "الأولوية:"
        )

        priority_label.setObjectName(
            "filterLabel"
        )

        filter_layout.addWidget(
            priority_label
        )

        self.priority_filter = QComboBox()

        self.priority_filter.addItems([
            "كل الأولويات",
            "عادي",
            "متوسط",
            "عاجل",
            "حرج"
        ])

        self.priority_filter.currentIndexChanged.connect(
            self.load_faults
        )

        filter_layout.addWidget(
            self.priority_filter
        )

        building_label = QLabel(
            "المبنى:"
        )

        building_label.setObjectName(
            "filterLabel"
        )

        filter_layout.addWidget(
            building_label
        )

        self.building_filter = QComboBox()

        self.building_filter.setMinimumWidth(
            180
        )

        self.building_filter.currentIndexChanged.connect(
            self.load_faults
        )

        filter_layout.addWidget(
            self.building_filter
        )

        self.clear_filters_button = QPushButton(
            "✕  مسح الفلاتر"
        )

        self.clear_filters_button.setObjectName(
            "clearFilterButton"
        )

        self.clear_filters_button.clicked.connect(
            self.clear_fault_filters
        )

        filter_layout.addWidget(
            self.clear_filters_button
        )

        filter_layout.addStretch()

        main_layout.addWidget(
            filter_card
        )

        # =========================================================
        # جدول الأعطال
        # =========================================================

        table_card = QFrame()

        table_card.setObjectName(
            "tableCard"
        )

        table_layout = QVBoxLayout(
            table_card
        )

        table_layout.setContentsMargins(
            8, 8, 8, 8
        )

        self.table = QTableWidget()

        self.table.setColumnCount(
            8
        )

        self.table.setHorizontalHeaderLabels([
            "ID",
            "رقم العطل",
            "التاريخ",
            "الجهاز",
            "المستخدم",
            "الحالة",
            "الأولوية",
            "القائم بالإصلاح"
        ])

        self.table.setSelectionBehavior(
            QTableWidget.SelectRows
        )

        self.table.setSelectionMode(
            QTableWidget.SingleSelection
        )

        self.table.setEditTriggers(
            QTableWidget.NoEditTriggers
        )

        self.table.verticalHeader().setVisible(
            False
        )

        self.table.setColumnHidden(
            0,
            True
        )

        self.table.horizontalHeader().setSectionResizeMode(
            QHeaderView.Stretch
        )

        self.table.setAlternatingRowColors(
            True
        )

        self.table.doubleClicked.connect(
            self.edit_fault
        )

        table_layout.addWidget(
            self.table
        )

        main_layout.addWidget(
            table_card,
            1
        )

    # =============================================================
    # تحميل فلتر المباني
    # =============================================================

    def load_building_filter(self):
        connection = get_connection()
        cursor = connection.cursor()

        cursor.execute("""
            SELECT id, name
            FROM buildings
            ORDER BY name
        """)

        buildings = cursor.fetchall()

        connection.close()

        self.building_filter.blockSignals(
            True
        )

        self.building_filter.clear()

        self.building_filter.addItem(
            "كل المباني",
            None
        )

        for building in buildings:
            self.building_filter.addItem(
                building["name"],
                building["id"]
            )

        self.building_filter.blockSignals(
            False
        )

    # =============================================================
    # مسح الفلاتر
    # =============================================================

    def clear_fault_filters(self):
        self.search_box.blockSignals(
            True
        )

        self.status_filter.blockSignals(
            True
        )

        self.priority_filter.blockSignals(
            True
        )

        self.building_filter.blockSignals(
            True
        )

        self.search_box.clear()

        self.status_filter.setCurrentIndex(
            0
        )

        self.priority_filter.setCurrentIndex(
            0
        )

        self.building_filter.setCurrentIndex(
            0
        )

        self.search_box.blockSignals(
            False
        )

        self.status_filter.blockSignals(
            False
        )

        self.priority_filter.blockSignals(
            False
        )

        self.building_filter.blockSignals(
            False
        )

        self.load_faults()

    def show_reports(self):
        dialog = FaultReportsDialog(
            self
        )

        dialog.exec()

    # =============================================================
    # أدوات Excel
    # =============================================================

    def _normalize_excel_header(self, value):
        if value is None:
            return ""

        value = str(value).strip()

        replacements = {
            "إسم": "اسم",
            "الإسم": "الاسم",
            "الإستلام": "الاستلام",
            "الإصلاح": "الاصلاح",
            "القطع": "قطع",
        }

        for old, new in replacements.items():
            value = value.replace(
                old,
                new
            )

        return " ".join(
            value.split()
        ).lower()

    def _excel_cell(
        self,
        row,
        headers,
        *names
    ):
        wanted = {
            self._normalize_excel_header(name)
            for name in names
        }

        for header, index in headers.items():
            if (
                self._normalize_excel_header(header)
                in wanted
            ):
                value = row[index]

                if value is None:
                    return ""

                return str(value).strip()

        return ""

    def _find_building_id(
        self,
        connection,
        building_name
    ):
        if not building_name:
            return None

        cursor = connection.cursor()

        cursor.execute(
            "SELECT id FROM buildings "
            "WHERE name = ? LIMIT 1",
            (building_name.strip(),)
        )

        result = cursor.fetchone()

        if result:
            return result["id"]

        cursor.execute(
            "INSERT INTO buildings (name) VALUES (?)",
            (building_name.strip(),)
        )

        return cursor.lastrowid

    def _find_device_id(
        self,
        connection,
        device_name,
        building_id=None
    ):
        if not device_name:
            return None

        cursor = connection.cursor()

        if building_id:
            cursor.execute(
                """
                SELECT id
                FROM devices
                WHERE device_name = ?
                  AND is_archived = 0
                  AND building_id = ?
                LIMIT 1
                """,
                (
                    device_name.strip(),
                    building_id
                )
            )

            result = cursor.fetchone()

            if result:
                return result["id"]

        cursor.execute(
            """
            SELECT id
            FROM devices
            WHERE device_name = ?
              AND is_archived = 0
            LIMIT 1
            """,
            (device_name.strip(),)
        )

        result = cursor.fetchone()

        return (
            result["id"]
            if result
            else None
        )

    def _historical_fault_number(
        self,
        connection
    ):
        cursor = connection.cursor()

        year = datetime.now().year

        cursor.execute(
            "SELECT COUNT(*) FROM faults "
            "WHERE fault_number LIKE ?",
            (f"OLD-{year}-%",)
        )

        count = cursor.fetchone()[0]
        number = count + 1

        while True:
            fault_number = (
                f"OLD-{year}-{number:05d}"
            )

            cursor.execute(
                "SELECT 1 FROM faults "
                "WHERE fault_number = ? LIMIT 1",
                (fault_number,)
            )

            if not cursor.fetchone():
                return fault_number

            number += 1

    def _parse_import_date(
        self,
        value
    ):
        """
        يحافظ على تاريخ Excel إن كان تاريخًا حقيقيًا،
        وإلا يحاول قراءة النص.
        """

        if (
            value is None
            or str(value).strip() == ""
        ):
            return datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            )

        if hasattr(value, "strftime"):
            return value.strftime(
                "%Y-%m-%d %H:%M:%S"
            )

        raw = str(value).strip()

        for fmt in (
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%d",
            "%d/%m/%Y",
            "%d-%m-%Y",
            "%d/%m/%y",
            "%d-%m-%y",
        ):
            try:
                return datetime.strptime(
                    raw,
                    fmt
                ).strftime(
                    "%Y-%m-%d %H:%M:%S"
                )

            except ValueError:
                continue

        return raw

    # =============================================================
    # استيراد Excel
    # =============================================================

    def import_from_excel(self):
        """
        استيراد الأعطال التاريخية من Excel
        وفق أعمدة النظام الحالية.
        """

        try:
            from openpyxl import load_workbook

        except ImportError:
            QMessageBox.critical(
                self,
                "مكتبة ناقصة",
                "مكتبة openpyxl غير مثبتة.\n\n"
                "شغّل:\n"
                "pip install openpyxl"
            )

            return

        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "اختيار ملف أعطال Excel",
            "",
            "Excel Files (*.xlsx *.xlsm)"
        )

        if not file_path:
            return

        try:
            workbook = load_workbook(
                file_path,
                data_only=True,
                read_only=True
            )

            sheet = workbook.active

            rows = list(
                sheet.iter_rows(
                    values_only=True
                )
            )

            if not rows:
                workbook.close()

                QMessageBox.warning(
                    self,
                    "ملف فارغ",
                    "ملف Excel لا يحتوي على بيانات."
                )

                return

            raw_headers = list(
                rows[0]
            )

            headers = {
                header: index
                for index, header in enumerate(
                    raw_headers
                )
                if header is not None
            }

            required_headers = {
                "المبنى/الفرع",
                "إسم الجهاز",
                "مواصفات الجهاز",
                "إسم المستخدم",
                "إسم المستلم",
                "توصيف العطل",
                "إجراءات الإصلاح وقطع الغيار",
                "القائم بالإصلاح",
                "تاريخ الإستلام",
                "ملاحظات",
            }

            normalized_headers = {
                self._normalize_excel_header(h)
                for h in raw_headers
                if h is not None
            }

            missing = [
                h
                for h in required_headers
                if self._normalize_excel_header(h)
                not in normalized_headers
            ]

            if missing:
                workbook.close()

                QMessageBox.warning(
                    self,
                    "أعمدة غير مكتملة",
                    "ملف Excel لا يطابق أعمدة الأعطال المطلوبة.\n\n"
                    "الأعمدة الناقصة:\n- "
                    + "\n- ".join(missing)
                )

                return

            data_rows = [
                row
                for row in rows[1:]
                if any(
                    value is not None
                    and str(value).strip()
                    for value in row
                )
            ]

            if not data_rows:
                workbook.close()

                QMessageBox.information(
                    self,
                    "لا توجد بيانات",
                    "لا توجد صفوف أعطال قابلة للاستيراد."
                )

                return

            workbook.close()

            answer = QMessageBox.question(
                self,
                "تأكيد الاستيراد",
                f"تم العثور على {len(data_rows)} عطل في الملف.\n\n"
                "سيتم إدخالها كأعطال تاريخية داخل النظام.\n"
                "الأعطال التي لا تحتوي على اسم جهاز ستدخل بدون ربط بجهاز، "
                "حتى لا يتم اختراع بيانات غير موجودة في الملف.\n\n"
                "هل تريد المتابعة؟",
                QMessageBox.Yes |
                QMessageBox.No,
                QMessageBox.No,
            )

            if answer != QMessageBox.Yes:
                return

            progress = QProgressDialog(
                "جاري استيراد الأعطال...",
                "إلغاء",
                0,
                len(data_rows),
                self
            )

            progress.setWindowTitle(
                "استيراد الأعطال"
            )

            progress.setWindowModality(
                Qt.WindowModal
            )

            progress.setMinimumDuration(
                0
            )

            progress.show()

            connection = get_connection()
            cursor = connection.cursor()

            imported = 0
            skipped = 0
            errors = []

            try:
                for row_number, row in enumerate(
                    data_rows,
                    start=2
                ):
                    if progress.wasCanceled():
                        connection.rollback()

                        QMessageBox.information(
                            self,
                            "تم الإلغاء",
                            "تم إلغاء الاستيراد.\n"
                            "لم يتم حفظ أي صف."
                        )

                        return

                    progress.setValue(
                        row_number - 2
                    )

                    building_name = self._excel_cell(
                        row,
                        headers,
                        "المبنى/الفرع"
                    )

                    device_name = self._excel_cell(
                        row,
                        headers,
                        "إسم الجهاز"
                    )

                    device_specs = self._excel_cell(
                        row,
                        headers,
                        "مواصفات الجهاز"
                    )

                    user_name = self._excel_cell(
                        row,
                        headers,
                        "إسم المستخدم"
                    )

                    receiver_name = self._excel_cell(
                        row,
                        headers,
                        "إسم المستلم"
                    )

                    fault_description = self._excel_cell(
                        row,
                        headers,
                        "توصيف العطل"
                    )

                    repair_actions = self._excel_cell(
                        row,
                        headers,
                        "إجراءات الإصلاح وقطع الغيار"
                    )

                    technician = self._excel_cell(
                        row,
                        headers,
                        "القائم بالإصلاح"
                    )

                    received_date_raw = self._excel_cell(
                        row,
                        headers,
                        "تاريخ الإستلام"
                    )

                    notes = self._excel_cell(
                        row,
                        headers,
                        "ملاحظات"
                    )

                    if (
                        not building_name
                        and not fault_description
                        and not user_name
                    ):
                        skipped += 1
                        continue

                    building_id = self._find_building_id(
                        connection,
                        building_name
                    )

                    device_id = self._find_device_id(
                        connection,
                        device_name,
                        building_id
                    )

                    entry_date = self._parse_import_date(
                        received_date_raw
                    )

                    status = (
                        "تم الإصلاح"
                        if repair_actions
                        else "جديد"
                    )

                    fault_number = (
                        self._historical_fault_number(
                            connection
                        )
                    )

                    cursor.execute(
                        """
                        INSERT INTO faults (
                            fault_number,
                            entry_date,
                            building_id,
                            device_id,
                            user_name,
                            receiver_name,
                            fault_description,
                            repair_actions,
                            spare_parts,
                            technician,
                            priority,
                            status,
                            repair_date,
                            delivery_date,
                            notes
                        )
                        VALUES (
                            ?, ?, ?, ?, ?, ?,
                            ?, ?, ?, ?, ?, ?,
                            ?, ?, ?
                        )
                        """,
                        (
                            fault_number,
                            entry_date,
                            building_id,
                            device_id,
                            user_name,
                            receiver_name,
                            fault_description,
                            repair_actions,
                            "",
                            technician,
                            "عادي",
                            status,
                            entry_date
                            if status == "تم الإصلاح"
                            else None,
                            entry_date
                            if status == "تم الإصلاح"
                            else None,
                            (
                                "استيراد تاريخي من Excel"
                                +
                                (
                                    f" | {notes}"
                                    if notes
                                    else ""
                                )
                            ),
                        )
                    )

                    imported += 1

                connection.commit()

            except Exception:
                connection.rollback()
                raise

            finally:
                connection.close()

            progress.setValue(
                len(data_rows)
            )

            self.load_building_filter()
            self.load_faults()

            message = (
                f"تم استيراد {imported} عطل بنجاح.\n"
                f"تم تجاهل {skipped} صف فارغ."
            )

            if errors:
                message += (
                    "\n\nصفوف بها مشاكل:\n"
                    + "\n".join(errors[:10])
                )

            QMessageBox.information(
                self,
                "تم الاستيراد",
                message
            )

        except Exception as exc:
            QMessageBox.critical(
                self,
                "خطأ في الاستيراد",
                f"حدث خطأ أثناء استيراد ملف Excel:\n\n{exc}"
            )

    # =============================================================
    # تحميل الأعطال + الفلاتر
    # =============================================================

    def load_faults(self):
        search = (
            self.search_box.text().strip()
            if hasattr(self, "search_box")
            else ""
        )

        status = (
            self.status_filter.currentText()
            if hasattr(self, "status_filter")
            else "كل الحالات"
        )

        priority = (
            self.priority_filter.currentText()
            if hasattr(self, "priority_filter")
            else "كل الأولويات"
        )

        building_id = (
            self.building_filter.currentData()
            if hasattr(self, "building_filter")
            else None
        )

        connection = get_connection()
        cursor = connection.cursor()

        query = """
            SELECT
                faults.*,
                devices.device_name,
                buildings.name AS building_name
            FROM faults
            LEFT JOIN devices
                ON faults.device_id = devices.id
            LEFT JOIN buildings
                ON faults.building_id = buildings.id
            WHERE 1 = 1
        """

        params = []

        # ---------------------------------------------------------
        # البحث
        # ---------------------------------------------------------

        if search:
            query += """
                AND (
                    faults.fault_number LIKE ?
                    OR devices.device_name LIKE ?
                    OR faults.user_name LIKE ?
                    OR faults.receiver_name LIKE ?
                    OR faults.technician LIKE ?
                    OR buildings.name LIKE ?
                )
            """

            keyword = f"%{search}%"

            params.extend([
                keyword,
                keyword,
                keyword,
                keyword,
                keyword,
                keyword
            ])

        # ---------------------------------------------------------
        # فلتر الحالة
        # ---------------------------------------------------------

        if status != "كل الحالات":
            query += """
                AND faults.status = ?
            """

            params.append(
                status
            )

        # ---------------------------------------------------------
        # فلتر الأولوية
        # ---------------------------------------------------------

        if priority != "كل الأولويات":
            query += """
                AND faults.priority = ?
            """

            params.append(
                priority
            )

        # ---------------------------------------------------------
        # فلتر المبنى
        # ---------------------------------------------------------

        if building_id:
            query += """
                AND faults.building_id = ?
            """

            params.append(
                building_id
            )

        query += """
            ORDER BY faults.id DESC
        """

        cursor.execute(
            query,
            params
        )

        faults = cursor.fetchall()

        connection.close()

        self.table.setRowCount(
            0
        )

        for fault in faults:
            row = self.table.rowCount()

            self.table.insertRow(
                row
            )

            values = [
                fault["id"],
                fault["fault_number"],
                fault["entry_date"],
                fault["device_name"] or "",
                fault["user_name"] or "",
                fault["status"] or "",
                fault["priority"] or "",
                fault["technician"] or ""
            ]

            for column, value in enumerate(
                values
            ):
                item = QTableWidgetItem(
                    str(value or "")
                )

                item.setTextAlignment(
                    Qt.AlignCenter
                )

                self.table.setItem(
                    row,
                    column,
                    item
                )

    # =============================================================
    # إضافة عطل
    # =============================================================

    def add_fault(self):
        dialog = FaultDialog(
            self
        )

        dialog.fault_number.setText(
            dialog.generate_fault_number()
        )

        if dialog.exec() != QDialog.Accepted:
            return

        data = dialog.get_data()

        connection = get_connection()
        cursor = connection.cursor()

        # خصم قطعة الغيار تلقائيًا عند تسجيل العطل
        if data["spare_part_id"]:
            cursor.execute(
                """
                SELECT
                    id,
                    part_name,
                    quantity
                FROM spare_parts
                WHERE id = ?
                  AND is_archived = 0
                """,
                (data["spare_part_id"],)
            )

            part = cursor.fetchone()

            if not part:
                connection.close()

                QMessageBox.warning(
                    self,
                    "تنبيه",
                    "قطعة الغيار المختارة غير موجودة."
                )

                return

            if (
                part["quantity"]
                < data["spare_part_qty"]
            ):
                connection.close()

                QMessageBox.warning(
                    self,
                    "الكمية غير كافية",
                    f"القطعة: {part['part_name']}\n"
                    f"المتاح: {part['quantity']}\n"
                    f"المطلوب: {data['spare_part_qty']}"
                )

                return

        cursor.execute(
            "SELECT building_id FROM devices "
            "WHERE id = ?",
            (data["device_id"],)
        )

        result = cursor.fetchone()

        building_id = (
            result["building_id"]
            if result
            else None
        )

        cursor.execute(
            """
            INSERT INTO faults (
                fault_number,
                entry_date,
                building_id,
                device_id,
                user_name,
                receiver_name,
                fault_description,
                repair_actions,
                spare_parts,
                technician,
                priority,
                status,
                repair_date,
                delivery_date
            )
            VALUES (
                ?, ?, ?, ?, ?, ?,
                ?, ?, ?, ?, ?, ?,
                ?, ?
            )
            """,
            (
                data["fault_number"],
                data["entry_date"],
                building_id,
                data["device_id"],
                data["user_name"],
                data["receiver_name"],
                data["fault_description"],
                data["repair_actions"],
                data["spare_parts"],
                data["technician"],
                data["priority"],
                data["status"],
                data["repair_date"],
                data["delivery_date"],
            )
        )

        # خصم المخزون وتسجيل حركة الصرف
        if data["spare_part_id"]:
            cursor.execute(
                """
                UPDATE spare_parts
                SET
                    quantity = quantity - ?,
                    updated_at = CURRENT_TIMESTAMP
                WHERE id = ?
                  AND quantity >= ?
                """,
                (
                    data["spare_part_qty"],
                    data["spare_part_id"],
                    data["spare_part_qty"]
                )
            )

            if cursor.rowcount != 1:
                connection.rollback()
                connection.close()

                QMessageBox.warning(
                    self,
                    "تنبيه",
                    "تعذر خصم قطعة الغيار من المخزون."
                )

                return

            cursor.execute(
                """
                INSERT INTO spare_part_movements
                    (
                        spare_part_id,
                        movement_type,
                        quantity,
                        notes
                    )
                VALUES (
                    ?,
                    'صرف',
                    ?,
                    ?
                )
                """,
                (
                    data["spare_part_id"],
                    data["spare_part_qty"],
                    f"صرف تلقائي لعطل "
                    f"{data['fault_number']}"
                )
            )

        connection.commit()
        connection.close()

        self.load_faults()

    # =============================================================
    # تعديل عطل
    # =============================================================

    def edit_fault(self):
        selected = self.table.selectedItems()

        if not selected:
            QMessageBox.information(
                self,
                "تنبيه",
                "اختر عطلًا أولًا."
            )

            return

        row = selected[0].row()

        fault_id = int(
            self.table.item(
                row,
                0
            ).text()
        )

        connection = get_connection()
        cursor = connection.cursor()

        cursor.execute(
            "SELECT * FROM faults WHERE id = ?",
            (fault_id,)
        )

        old_fault = cursor.fetchone()

        connection.close()

        if not old_fault:
            QMessageBox.warning(
                self,
                "تنبيه",
                "العطل غير موجود."
            )

            return

        dialog = FaultDialog(
            self,
            fault=old_fault
        )

        if dialog.exec() != QDialog.Accepted:
            return

        data = dialog.get_data()

        connection = get_connection()
        cursor = connection.cursor()

        try:
            cursor.execute(
                """
                SELECT
                    id,
                    spare_part_id,
                    quantity
                FROM spare_part_movements
                WHERE movement_type = 'صرف'
                  AND notes = ?
                ORDER BY id DESC
                LIMIT 1
                """,
                (
                    f"صرف تلقائي لعطل "
                    f"{old_fault['fault_number']}",
                )
            )

            old_movement = cursor.fetchone()

            old_part_id = (
                old_movement["spare_part_id"]
                if old_movement
                else None
            )

            old_qty = (
                int(old_movement["quantity"])
                if old_movement
                else 0
            )

            new_spare_id = data.get(
                "spare_part_id"
            )

            new_qty = int(
                data.get("spare_part_qty")
                or 0
            )

            old_has_part = (
                old_part_id is not None
                and old_qty > 0
            )

            # إرجاع القطعة القديمة
            if old_has_part:
                cursor.execute(
                    """
                    UPDATE spare_parts
                    SET
                        quantity = quantity + ?,
                        updated_at = CURRENT_TIMESTAMP
                    WHERE id = ?
                    """,
                    (
                        old_qty,
                        old_part_id
                    )
                )

                if cursor.rowcount != 1:
                    connection.rollback()
                    connection.close()

                    QMessageBox.warning(
                        self,
                        "تنبيه",
                        "تعذر إرجاع قطعة الغيار القديمة للمخزن."
                    )

                    return

                cursor.execute(
                    "DELETE FROM spare_part_movements "
                    "WHERE id = ?",
                    (old_movement["id"],)
                )

            # خصم القطعة الجديدة
            if new_spare_id and new_qty > 0:
                cursor.execute(
                    """
                    SELECT
                        id,
                        part_name,
                        quantity
                    FROM spare_parts
                    WHERE id = ?
                      AND is_archived = 0
                    """,
                    (new_spare_id,)
                )

                new_part = cursor.fetchone()

                if not new_part:
                    connection.rollback()
                    connection.close()

                    QMessageBox.warning(
                        self,
                        "تنبيه",
                        "قطعة الغيار المختارة غير موجودة."
                    )

                    return

                if (
                    int(new_part["quantity"])
                    < new_qty
                ):
                    connection.rollback()
                    connection.close()

                    QMessageBox.warning(
                        self,
                        "رصيد غير كافٍ",
                        f"القطعة «{new_part['part_name']}» "
                        f"المتاح منها "
                        f"{new_part['quantity']} فقط."
                    )

                    return

                cursor.execute(
                    """
                    UPDATE spare_parts
                    SET
                        quantity = quantity - ?,
                        updated_at = CURRENT_TIMESTAMP
                    WHERE id = ?
                      AND quantity >= ?
                    """,
                    (
                        new_qty,
                        new_spare_id,
                        new_qty
                    )
                )

                if cursor.rowcount != 1:
                    connection.rollback()
                    connection.close()

                    QMessageBox.warning(
                        self,
                        "تنبيه",
                        "تعذر خصم قطعة الغيار الجديدة."
                    )

                    return

                cursor.execute(
                    """
                    INSERT INTO spare_part_movements
                    (
                        spare_part_id,
                        movement_type,
                        quantity,
                        notes
                    )
                    VALUES (
                        ?,
                        'صرف',
                        ?,
                        ?
                    )
                    """,
                    (
                        new_spare_id,
                        new_qty,
                        f"صرف تلقائي لعطل "
                        f"{old_fault['fault_number']}"
                    )
                )

            # تحديث بيانات العطل
            cursor.execute(
                """
                UPDATE faults SET
                    entry_date = ?,
                    building_id = ?,
                    device_id = ?,
                    user_name = ?,
                    receiver_name = ?,
                    fault_description = ?,
                    repair_actions = ?,
                    spare_parts = ?,
                    technician = ?,
                    priority = ?,
                    status = ?,
                    repair_date = ?,
                    delivery_date = ?,
                    notes = ?,
                    updated_at = CURRENT_TIMESTAMP
                WHERE id = ?
                """,
                (
                    data["entry_date"],
                    data["building_id"],
                    data["device_id"],
                    data["user_name"],
                    data["receiver_name"],
                    data["fault_description"],
                    data["repair_actions"],
                    data["spare_parts"],
                    data["technician"],
                    data["priority"],
                    data["status"],
                    data["repair_date"],
                    data["delivery_date"],
                    data.get("notes", ""),
                    fault_id,
                )
            )

            connection.commit()
            connection.close()

            self.load_faults()

            QMessageBox.information(
                self,
                "تم",
                "تم تعديل العطل وتحديث المخزون "
                "وحركات قطع الغيار."
            )

        except Exception as exc:
            connection.rollback()
            connection.close()

            QMessageBox.critical(
                self,
                "خطأ",
                f"حدث خطأ أثناء تعديل العطل:\n{exc}"
            )

    # =============================================================
    # حذف عطل
    # =============================================================

    def delete_fault(self):
        selected = self.table.selectedItems()

        if not selected:
            QMessageBox.information(
                self,
                "تنبيه",
                "اختر عطلًا أولًا."
            )

            return

        row = selected[0].row()

        fault_id = int(
            self.table.item(
                row,
                0
            ).text()
        )

        connection = get_connection()
        cursor = connection.cursor()

        cursor.execute(
            """
            SELECT
                id,
                fault_number
            FROM faults
            WHERE id = ?
            """,
            (fault_id,)
        )

        fault = cursor.fetchone()

        if not fault:
            connection.close()

            QMessageBox.warning(
                self,
                "تنبيه",
                "العطل غير موجود."
            )

            return

        answer = QMessageBox.question(
            self,
            "تأكيد حذف العطل",
            f"هل أنت متأكد من حذف العطل:\n\n"
            f"{fault['fault_number']}؟\n\n"
            "إذا كان قد تم صرف قطعة غيار له من المخزون، "
            "سيتم إرجاع الكمية للمخزن تلقائيًا.",
            QMessageBox.Yes |
            QMessageBox.No,
            QMessageBox.No
        )

        if answer != QMessageBox.Yes:
            connection.close()
            return

        try:
            cursor.execute(
                """
                SELECT
                    id,
                    spare_part_id,
                    quantity
                FROM spare_part_movements
                WHERE movement_type = 'صرف'
                  AND notes = ?
                ORDER BY id DESC
                LIMIT 1
                """,
                (
                    f"صرف تلقائي لعطل "
                    f"{fault['fault_number']}"
                )
            )

            movement = cursor.fetchone()

            if movement:
                cursor.execute(
                    """
                    UPDATE spare_parts
                    SET
                        quantity = quantity + ?,
                        updated_at = CURRENT_TIMESTAMP
                    WHERE id = ?
                    """,
                    (
                        movement["quantity"],
                        movement["spare_part_id"]
                    )
                )

                if cursor.rowcount != 1:
                    connection.rollback()
                    connection.close()

                    QMessageBox.warning(
                        self,
                        "تنبيه",
                        "تعذر إرجاع قطعة الغيار للمخزن، "
                        "لذلك لم يتم حذف العطل."
                    )

                    return

                cursor.execute(
                    "DELETE FROM spare_part_movements "
                    "WHERE id = ?",
                    (movement["id"],)
                )

            # حذف سجل المتابعات المرتبط بالعطل
            cursor.execute(
                "DELETE FROM fault_updates "
                "WHERE fault_id = ?",
                (fault_id,)
            )

            cursor.execute(
                "DELETE FROM faults WHERE id = ?",
                (fault_id,)
            )

            if cursor.rowcount != 1:
                connection.rollback()
                connection.close()

                QMessageBox.warning(
                    self,
                    "تنبيه",
                    "تعذر حذف العطل."
                )

                return

            connection.commit()
            connection.close()

            self.load_faults()

            QMessageBox.information(
                self,
                "تم",
                "تم حذف العطل وإرجاع قطعة الغيار "
                "للمخزن إن وُجدت."
            )

        except Exception as exc:
            connection.rollback()
            connection.close()

            QMessageBox.critical(
                self,
                "خطأ",
                f"حدث خطأ أثناء حذف العطل:\n{exc}"
            )
